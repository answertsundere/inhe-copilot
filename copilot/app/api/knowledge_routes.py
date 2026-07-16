"""
知识库后台 API 路由
权限角色：operator / supervisor / admin
所有状态变更通过 KnowledgeLifecycleService 完成。
"""

import os
import tempfile
from flask import Blueprint, request, jsonify, current_app, send_file

from app.repositories.knowledge_entry_repository import KnowledgeEntryRepository
from app.repositories.knowledge_version_repository import KnowledgeVersionRepository
from app.repositories.knowledge_chunk_repository import KnowledgeChunkRepository
from app.services.knowledge_index_service import KnowledgeIndexService
from app.services.knowledge_import_service import KnowledgeImportService
from app.services.knowledge_quality_service import KnowledgeQualityService
from app.models.knowledge_base import KnowledgeEntry, KnowledgeChunk
from app.db import SessionLocal
from app.api.admin_auth import current_role, current_user_name, has_any_role

knowledge_bp = Blueprint("knowledge", __name__)


# ============ 权限检查 ============

def _get_user_role():
    return current_role()


def _get_user():
    return current_user_name()


def _require_role(*roles):
    if not has_any_role(*roles):
        return jsonify({"error": "authorization_denied"}), 403
    return None


# ============ 知识条目 CRUD ============

@knowledge_bp.route("/api/knowledge/entries", methods=["GET"])
def api_list_entries():
    """知识列表"""
    source_type = request.args.get("source_type", "")
    intent = request.args.get("intent", "")
    status = request.args.get("status", "")
    risk_level = request.args.get("risk_level", "")
    search_query = request.args.get("q", "")
    limit = request.args.get("limit", 50, type=int)
    offset = request.args.get("offset", 0, type=int)
    # 扩展筛选
    fact_type = request.args.get("fact_type", "")
    product_id = request.args.get("product_id", "")
    sku_id = request.args.get("sku_id", "")
    business_key = request.args.get("business_key", "")
    batch_id = request.args.get("batch_id", "")
    index_status = request.args.get("index_status", "")
    human_review_required = request.args.get("human_review_required", "")
    auto_reply_allowed = request.args.get("auto_reply_allowed", "")
    sort_by = request.args.get("sort_by", "updated_at")
    sort_order = request.args.get("sort_order", "desc")

    # 限制 limit 上限
    if limit > 200:
        limit = 200

    items, total = KnowledgeEntryRepository.list_entries(
        source_type=source_type,
        intent=intent,
        status=status,
        risk_level=risk_level,
        search_query=search_query,
        limit=limit,
        offset=offset,
        fact_type=fact_type,
        product_id=product_id,
        sku_id=sku_id,
        business_key=business_key,
        batch_id=batch_id,
        index_status=index_status,
        human_review_required=human_review_required,
        auto_reply_allowed=auto_reply_allowed,
        sort_by=sort_by,
        sort_order=sort_order,
    )
    return jsonify({
        "items": [i.to_dict(include_content=False) for i in items],
        "total": total,
        "limit": limit,
        "offset": offset,
    })


@knowledge_bp.route("/api/knowledge/entries", methods=["POST"])
def api_create_entry():
    """新增知识"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    data = request.get_json(force=True) or {}
    user = _get_user()

    result = KnowledgeLifecycleService.create_draft(
        source_type=data.get("source_type", ""),
        title=data.get("title", ""),
        content=data.get("content", ""),
        intent=data.get("intent", "general"),
        sub_intent=data.get("sub_intent", ""),
        category=data.get("category", ""),
        product_scope=data.get("product_scope", []),
        sku_scope=data.get("sku_scope", []),
        platform_scope=data.get("platform_scope", []),
        risk_level=data.get("risk_level", "low"),
        auto_reply_allowed=data.get("auto_reply_allowed", True),
        human_review_required=data.get("human_review_required", False),
        condition_text=data.get("condition_text", ""),
        forbidden_usage=data.get("forbidden_usage", ""),
        created_by=user,
    )
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result), 201


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>", methods=["GET"])
def api_get_entry(entry_id):
    """知识详情"""
    entry = KnowledgeEntryRepository.get_by_id(entry_id)
    if not entry:
        return jsonify({"error": "知识不存在"}), 404
    return jsonify(entry.to_dict(include_content=True))


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>", methods=["PUT"])
def api_update_entry(entry_id):
    """编辑知识"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    data = request.get_json(force=True) or {}
    user = _get_user()

    result = KnowledgeLifecycleService.update_draft(entry_id, data, updated_by=user)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>", methods=["DELETE"])
def api_archive_entry(entry_id):
    """归档知识"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    err = _require_role("operator", "supervisor", "admin")
    if err:
        return err
    user = _get_user()

    result = KnowledgeLifecycleService.archive(entry_id, user)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


# ============ 状态流转 ============

@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/submit-review", methods=["POST"])
def api_submit_review(entry_id):
    """提交审核"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    user = _get_user()
    result = KnowledgeLifecycleService.submit_review(entry_id, user)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/review", methods=["POST"])
def api_review_entry(entry_id):
    """审核知识"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    err = _require_role("supervisor", "admin")
    if err:
        return err

    data = request.get_json(force=True) or {}
    user = _get_user()
    approved = data.get("approved", True)
    reason = data.get("reason", "")

    if approved:
        result = KnowledgeLifecycleService.approve(entry_id, user)
    else:
        result = KnowledgeLifecycleService.reject(entry_id, user, reason)

    if "error" in result:
        # 质量检查未通过返回 409
        if "blocking_issues" in result:
            return jsonify(result), 409
        return jsonify(result), 400
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/publish", methods=["POST"])
def api_publish_entry(entry_id):
    """发布知识（与审核通过使用同一质量门槛）"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    err = _require_role("supervisor", "admin")
    if err:
        return err
    user = _get_user()

    result = KnowledgeLifecycleService.publish(entry_id, user)
    if "error" in result:
        if "blocking_issues" in result:
            return jsonify(result), 409
        return jsonify(result), 400
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/rollback", methods=["POST"])
def api_rollback_entry(entry_id):
    """回滚版本"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    err = _require_role("supervisor", "admin")
    if err:
        return err

    data = request.get_json(force=True) or {}
    user = _get_user()
    version = data.get("version") or data.get("target_version")
    if not version:
        return jsonify({"error": "缺少 version 参数"}), 400

    result = KnowledgeLifecycleService.rollback_to_draft(entry_id, version, user)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


# ============ 版本与审计 ============

@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/versions", methods=["GET"])
def api_list_versions(entry_id):
    """版本历史"""
    versions = KnowledgeVersionRepository.get_versions(entry_id)
    return jsonify({"items": [v.to_dict() for v in versions]})


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/audit-log", methods=["GET"])
def api_audit_log(entry_id):
    """审计日志"""
    from app.db import SessionLocal
    from app.models.knowledge_base import KnowledgeAuditLog
    db = SessionLocal()
    try:
        logs = (
            db.query(KnowledgeAuditLog)
            .filter(KnowledgeAuditLog.entry_id == entry_id)
            .order_by(KnowledgeAuditLog.created_at.desc())
            .all()
        )
        return jsonify({"items": [l.to_dict() for l in logs]})
    finally:
        db.close()


# ============ 统计 ============

@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/stats", methods=["GET"])
def api_entry_stats(entry_id):
    """命中次数和采纳率"""
    stats = KnowledgeEntryRepository.get_stats(entry_id)
    return jsonify(stats)


# ============ 批量导入 ============

@knowledge_bp.route("/api/knowledge/import-preview", methods=["POST"])
def api_import_preview():
    """导入预览"""
    if "file" not in request.files:
        return jsonify({"error": "未上传文件"}), 400

    file = request.files["file"]
    source_type_override = request.form.get("source_type", "")

    suffix = os.path.splitext(file.filename)[1].lower()
    if suffix not in (".xlsx", ".xls", ".csv"):
        return jsonify({"error": "仅支持 Excel/CSV 文件"}), 400

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        result = KnowledgeImportService.parse_excel_preview(tmp_path, source_type_override)
        return jsonify(result)
    finally:
        os.unlink(tmp_path)


@knowledge_bp.route("/api/knowledge/import", methods=["POST"])
def api_import_execute():
    """执行导入"""
    err = _require_role("operator", "supervisor", "admin")
    if err:
        return err

    data = request.get_json(force=True) or {}
    preview = data.get("preview", [])
    batch_id = data.get("batch_id", "")
    allow_duplicates = data.get("allow_duplicates", False)
    user = _get_user()

    if not preview:
        return jsonify({"error": "缺少预览数据"}), 400

    result = KnowledgeImportService.import_from_preview(preview, user, batch_id, allow_duplicates=allow_duplicates)
    return jsonify(result)


# ============ RAG 检索（内部用） ============

@knowledge_bp.route("/api/knowledge/retrieve", methods=["POST"])
def api_retrieve():
    """检索 published 知识分片"""
    data = request.get_json(force=True) or {}
    query = data.get("query", "")
    source_types = data.get("source_types", [])
    intent = data.get("intent", "")
    product_scope = data.get("product_scope", [])
    sku_scope = data.get("sku_scope", [])
    top_k = data.get("top_k", 5)
    min_score = data.get("min_score", 0.1)

    results = KnowledgeChunkRepository.search_chunks(
        query=query,
        source_types=source_types if source_types else None,
        intent=intent,
        product_scope=product_scope if product_scope else None,
        sku_scope=sku_scope if sku_scope else None,
        top_k=top_k,
        min_score=min_score,
    )
    return jsonify({"results": results, "count": len(results)})


# ============ 保留原有数据质量 API ============

@knowledge_bp.route("/api/knowledge/quality", methods=["GET"])
def api_knowledge_quality():
    """全局数据质量摘要"""
    svc = current_app.config.get("data_quality_service")
    if svc is None:
        from app.main import get_data_quality_service
        svc = get_data_quality_service()
    return jsonify(svc.get_quality_summary())


@knowledge_bp.route("/api/knowledge/missing-info", methods=["GET"])
def api_knowledge_missing_info():
    """缺失信息任务列表"""
    limit = request.args.get("limit", 100, type=int)
    svc = current_app.config.get("data_quality_service")
    if svc is None:
        from app.main import get_data_quality_service
        svc = get_data_quality_service()
    records = svc.get_missing_info(limit=limit)
    return jsonify({"records": records, "count": len(records)})


@knowledge_bp.route("/api/knowledge/products/<i_id>/quality", methods=["GET"])
def api_product_quality(i_id):
    """单个商品数据质量"""
    svc = current_app.config.get("data_quality_service")
    if svc is None:
        from app.main import get_data_quality_service
        svc = get_data_quality_service()
    result = svc.get_product_quality(i_id)
    return jsonify(result)


# ============ 知识库质量报告 ============

@knowledge_bp.route("/api/knowledge/quality/report", methods=["GET"])
def api_knowledge_quality_report():
    """知识库质量检查报告"""
    report = KnowledgeQualityService.generate_report()
    return jsonify(report)


@knowledge_bp.route("/api/knowledge/quality/product-facts-incomplete", methods=["GET"])
def api_product_facts_incomplete():
    """product_facts 缺字段列表（运营补全用）"""
    items = KnowledgeQualityService.get_product_facts_incomplete()
    return jsonify({"items": items, "count": len(items)})


@knowledge_bp.route("/api/knowledge/quality/product-facts-export", methods=["GET"])
def api_product_facts_export():
    """导出 product_facts 缺字段 Excel"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    items = KnowledgeQualityService.get_product_facts_incomplete()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "product_facts待补全"
    headers = ["entry_id", "title", "sku_code", "product_scope", "missing_fields", "missing_scope", "current_content", "suggested_fix_template", "status"]
    ws.append(headers)
    for h in ws[1]:
        h.font = Font(bold=True)
        h.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

    for item in items:
        ws.append([
            item["entry_id"],
            item["title"],
            ", ".join(item.get("sku_scope", [])),
            ", ".join(item.get("product_scope", [])),
            ", ".join(item.get("missing_fields", [])),
            "是" if item.get("missing_scope") else "否",
            item.get("current_content", ""),
            item.get("suggested_fix_template", ""),
            item.get("status", ""),
        ])

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return send_file(path, as_attachment=True, download_name="product_facts_incomplete.xlsx")


@knowledge_bp.route("/api/knowledge/quality/risky-claims", methods=["GET"])
def api_risky_claims():
    """List risky customer-facing convenience claims for knowledge cleanup."""
    result = KnowledgeQualityService.get_risky_convenience_claims()
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/quality/risky-claims-export", methods=["GET"])
def api_risky_claims_export():
    """Export risky customer-facing convenience claims."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    result = KnowledgeQualityService.get_risky_convenience_claims()
    items = result.get("items", [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "risky_claims"
    headers = [
        "table", "id", "source_type", "title", "status", "risk_level",
        "fact_type", "auto_reply_allowed", "human_review_required",
        "source_sheet", "row_number", "matched_phrases", "suggested_action",
    ]
    ws.append(headers)
    for h in ws[1]:
        h.font = Font(bold=True)
        h.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for item in items:
        ws.append([
            item.get("table", ""),
            item.get("id", ""),
            item.get("source_type", ""),
            item.get("title", ""),
            item.get("status", ""),
            item.get("risk_level", ""),
            item.get("fact_type", ""),
            item.get("auto_reply_allowed", ""),
            item.get("human_review_required", ""),
            item.get("source_sheet", ""),
            item.get("row_number", ""),
            "; ".join(item.get("matched_phrases", [])),
            item.get("suggested_action", ""),
        ])

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return send_file(path, as_attachment=True, download_name="risky_claims_cleanup.xlsx")


@knowledge_bp.route("/api/knowledge/quality/duplicate-review", methods=["GET"])
def api_duplicate_review():
    """duplicate_candidate / possible_variant 审核列表"""
    result = KnowledgeQualityService.get_duplicate_review_list()
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/quality/duplicate-action", methods=["POST"])
def api_duplicate_action():
    """对 duplicate/variant 执行操作"""
    data = request.get_json(force=True) or {}
    action = data.get("action", "")
    entry_ids = data.get("entry_ids", [])
    user = _get_user()

    if not entry_ids:
        return jsonify({"error": "缺少 entry_ids"}), 400

    results = []
    for eid in entry_ids:
        entry = KnowledgeEntryRepository.get_by_id(eid)
        if not entry:
            results.append({"entry_id": eid, "status": "not_found"})
            continue

        if action == "archive":
            from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService
            r = KnowledgeLifecycleService.archive(eid, user)
            results.append({"entry_id": eid, "status": r.get("status", "archived") if "error" not in r else "error", "result": r})
        elif action == "submit_review":
            from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService
            r = KnowledgeLifecycleService.submit_review(eid, user)
            results.append({"entry_id": eid, "status": "submitted" if "error" not in r else "error", "result": r})
        elif action == "mark_as_duplicate":
            entry = KnowledgeEntryRepository.get_by_id(eid)
            if entry:
                entry.condition_text = (entry.condition_text or "") + "\n[系统标记：与条目重复]"
                db = _get_db()
                db.add(entry)
                db.commit()
                db.close()
            results.append({"entry_id": eid, "status": "marked_duplicate"})
        elif action == "keep_separate":
            results.append({"entry_id": eid, "status": "kept"})
        else:
            results.append({"entry_id": eid, "status": "unknown_action"})

    return jsonify({"results": results})


@knowledge_bp.route("/api/knowledge/quality/publish-candidates", methods=["GET"])
def api_publish_candidates():
    """建议优先发布的黄金知识集候选"""
    candidates = KnowledgeQualityService.get_publish_candidates()
    return jsonify(candidates)


@knowledge_bp.route("/api/knowledge/entries/batch-submit-review", methods=["POST"])
def api_batch_submit_review():
    """批量提交审核"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    data = request.get_json(force=True) or {}
    entry_ids = data.get("entry_ids", [])
    user = _get_user()

    results = []
    for eid in entry_ids:
        r = KnowledgeLifecycleService.submit_review(eid, user)
        if "error" not in r:
            results.append({"entry_id": eid, "status": "submitted"})
        else:
            results.append({"entry_id": eid, "status": "error", "reason": r.get("error", "")})

    return jsonify({"results": results, "success": sum(1 for r in results if r["status"] == "submitted")})


# ============ 知识验证与索引 ============

@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/validate", methods=["POST"])
def api_validate_entry(entry_id):
    """验证知识条目的发布条件。"""
    entry = KnowledgeEntryRepository.get_by_id(entry_id)
    if not entry:
        return jsonify({"error": "知识不存在"}), 404

    from app.services.knowledge_quality_gate import validate_for_publish
    result = validate_for_publish(entry)
    result["entry_id"] = entry_id
    result["status"] = entry.status
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/index-status", methods=["GET"])
def api_index_status(entry_id):
    """获取条目的索引状态。"""
    from app.services.knowledge_index_pipeline import KnowledgeIndexPipeline
    result = KnowledgeIndexPipeline.get_index_status(entry_id)
    if "error" in result:
        return jsonify(result), 404
    return jsonify(result)


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/reindex", methods=["POST"])
def api_reindex_entry(entry_id):
    """重建单个条目的索引。仅允许 published 条目。"""
    from app.services.knowledge_index_pipeline import KnowledgeIndexPipeline
    try:
        result = KnowledgeIndexPipeline.rebuild_entry(entry_id)
        if result.get("rejected"):
            return jsonify(result), 400
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@knowledge_bp.route("/api/knowledge/entries/<int:entry_id>/revision", methods=["POST"])
def api_create_revision(entry_id):
    """创建修订版。"""
    from app.services.knowledge_lifecycle_service import KnowledgeLifecycleService

    err = _require_role("supervisor", "admin")
    if err:
        return err
    user = _get_user()

    result = KnowledgeLifecycleService.create_revision(entry_id, user)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result), 201


# ============ Summary ============

@knowledge_bp.route("/api/knowledge/summary", methods=["GET"])
def api_knowledge_summary():
    """RAG 知识库统计摘要"""
    db = SessionLocal()
    try:
        from sqlalchemy import func
        status_counts = dict(
            db.query(KnowledgeEntry.status, func.count(KnowledgeEntry.id))
            .group_by(KnowledgeEntry.status).all()
        )
        index_counts = dict(
            db.query(KnowledgeEntry.index_status, func.count(KnowledgeEntry.id))
            .group_by(KnowledgeEntry.index_status).all()
        )
        chunk_count = db.query(KnowledgeChunk).count()
        return jsonify({
            "total_entries": sum(status_counts.values()),
            "draft": status_counts.get("draft", 0),
            "pending_review": status_counts.get("pending_review", 0),
            "published": status_counts.get("published", 0),
            "archived": status_counts.get("archived", 0),
            "index_pending": index_counts.get("pending", 0),
            "index_ready": index_counts.get("ready", 0),
            "index_failed": index_counts.get("failed", 0),
            "chunk_count": chunk_count,
            "published_ready_count": sum(
                1 for e in db.query(KnowledgeEntry).filter(
                    KnowledgeEntry.status == "published",
                    KnowledgeEntry.index_status == "ready",
                ).all()
            ),
        })
    finally:
        db.close()


def _get_db():
    from app.db import SessionLocal
    return SessionLocal()
