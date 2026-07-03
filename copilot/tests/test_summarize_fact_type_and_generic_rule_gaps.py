import json

from openpyxl import load_workbook

from scripts.summarize_fact_type_and_generic_rule_gaps import (
    build_workbook,
    classify_fact_type_gap,
    classify_generic_rule_gap,
    summarize_gaps,
    _write_excel,
    _write_json,
)


def test_classifies_fact_type_missing_into_generic_semantic_clusters():
    installation = classify_fact_type_gap("这个护栏能补一面吗")
    assert installation["cluster"] == "配件/结构适配"
    assert installation["suggested_fact_type"] == "accessory_compatibility"
    assert installation["code_fix_recommended"] is True

    social = classify_fact_type_gap("好的")
    assert social["cluster"] == "社交/确认/等待"
    assert social["actionable"] is False
    assert social["should_score"] is False

    child_safety = classify_fact_type_gap("这个适合2周岁宝宝吗")
    assert child_safety["suggested_fact_type"] == "age_range"
    assert child_safety["needs_rag"] is True

    dimension = classify_fact_type_gap("沙发高40")
    assert dimension["cluster"] == "尺寸/空间/摆放"
    assert dimension["suggested_fact_type"] == "dimensions"

    link = classify_fact_type_gap("发我下链接")
    assert link["cluster"] == "商品链接/下单协助"
    assert link["needs_rag"] is False

    company_context = classify_fact_type_gap("账户号码：123456789 公司名称：测试公司")
    assert company_context["cluster"] == "发票/企业信息上下文"
    assert company_context["should_score"] is False

    structure = classify_fact_type_gap("如果再想改成独立的爬梯行吗")
    assert structure["cluster"] == "配件/结构适配"
    assert structure["suggested_fact_type"] == "accessory_compatibility"


def test_classifies_generic_rule_missing_without_product_fact_promises():
    promo = classify_generic_rule_gap({"query_fact_type": "promotion_policy", "buyer_message_preview": "有什么优惠吗"})
    assert promo["cluster"] == "优惠/活动/议价规则"
    assert promo["generic_rule_recommended"] is True
    assert "不能承诺具体优惠" in promo["reason"]

    other = classify_generic_rule_gap({"query_fact_type": "unknown", "buyer_message_preview": "这个"})
    assert other["generic_rule_recommended"] is False


def test_summarizes_gap_json_and_writes_chinese_excel(tmp_path):
    input_path = tmp_path / "evidence.json"
    payload = {
        "zero_evidence_records": [
            {
                "primary_reason": "query_fact_type_missing",
                "buyer_message_preview": "安装视频发我",
                "query_fact_type": "unknown",
                "turn_uid": "turn-1",
            },
            {
                "primary_reason": "query_fact_type_missing",
                "buyer_message_preview": "哦",
                "query_fact_type": "unknown",
                "turn_uid": "turn-2",
            },
            {
                "primary_reason": "generic_rule_missing",
                "buyer_message_preview": "多买能便宜吗",
                "query_fact_type": "promotion_policy",
                "turn_uid": "turn-3",
            },
        ]
    }
    input_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    report = summarize_gaps(str(input_path))

    assert report["summary"]["query_fact_type_missing_count"] == 2
    assert report["summary"]["generic_rule_missing_count"] == 1
    assert report["summary"]["fact_type_gap_clusters"]["安装/配件/结构"] == 1
    assert report["summary"]["fact_type_gap_clusters"]["社交/确认/等待"] == 1
    assert report["summary"]["generic_rule_gap_clusters"]["优惠/活动/议价规则"] == 1
    assert any(row["fix_type"] == "generic_rule_contract" for row in report["recommended_fix_candidates"])
    assert all("[ADDRESS_REDACTED]" not in row["recommendation"] for row in report["recommended_fix_candidates"])

    workbook = build_workbook(report)
    assert "FactType 缺口样本" in workbook.sheetnames
    assert "通用规则缺口样本" in workbook.sheetnames
    assert workbook["FactType 缺口样本"]["A1"].value == "语义簇"

    excel_path = tmp_path / "缺口.xlsx"
    json_path = tmp_path / "缺口.json"
    _write_excel(str(excel_path), report)
    _write_json(str(json_path), report)
    loaded = load_workbook(excel_path)
    assert loaded["通用规则缺口样本"]["A1"].value == "规则簇"
    assert json.loads(json_path.read_text(encoding="utf-8"))["summary"]["generic_rule_missing_count"] == 1
