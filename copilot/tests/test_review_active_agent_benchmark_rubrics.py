from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.db as db_module
from app.db import Base
from app.models.eval_tables import AgentBenchmarkScenario
from scripts.review_active_agent_benchmark_rubrics import review_active_rubrics


def _patch_test_db(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    Base.metadata.create_all(bind=engine, tables=[AgentBenchmarkScenario.__table__])
    return session_factory


def _add_active_installation(
    session_factory,
    *,
    scenario_uid: str = "active_install_1",
    expected: dict | None = None,
    metadata_extra: dict | None = None,
):
    db = session_factory()
    try:
        row = AgentBenchmarkScenario(
            scenario_uid=scenario_uid,
            source_type="manual",
            source_uid=scenario_uid,
            status="active",
            title="Installation benchmark",
            scenario_type="installation",
            created_by="test",
        )
        row.set_sidecar_context({"product_title": "Test shelf", "sku_code": "SKU-INSTALL"})
        row.set_conversation_turns([
            {"turn_uid": "turn_1", "speaker": "buyer", "text": "Do you have installation material?"}
        ])
        row.set_expected_reply(expected or {
            "expected_reply": "Please hand off for installation material.",
            "key_points": ["转人工"],
            "forbidden_claims": [],
            "needs_review": False,
            "quality": "valid",
            "auto_send_allowed": False,
            "must_handoff": True,
        })
        metadata = {
            "query_fact_type": "installation",
            "expected_reply_quality": "valid",
            "needs_expected_reply_review": False,
        }
        if metadata_extra:
            metadata.update(metadata_extra)
        row.set_metadata(metadata)
        db.add(row)
        db.commit()
    finally:
        db.close()


def _scenario(session_factory, scenario_uid: str):
    db = session_factory()
    try:
        return db.query(AgentBenchmarkScenario).filter_by(scenario_uid=scenario_uid).one()
    finally:
        db.close()


def test_active_rubric_review_dry_run_does_not_write(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_active_installation(
        session_factory,
        metadata_extra={
            "recommended_assets": [
                {
                    "asset_type": "manual",
                    "review_status": "approved",
                    "usable": True,
                    "auto_send_level": "auto",
                }
            ]
        },
    )

    result = review_active_rubrics(apply=False, db_factory=session_factory)

    assert result["dry_run"] is True
    assert result["changed_count"] == 1
    assert result["items"][0]["new_auto_send_allowed"] is True
    assert result["items"][0]["new_must_handoff"] is False
    row = _scenario(session_factory, "active_install_1")
    assert row.get_expected_reply()["auto_send_allowed"] is False
    assert row.get_expected_reply()["must_handoff"] is True


def test_active_rubric_review_apply_updates_existing_active(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_active_installation(
        session_factory,
        metadata_extra={
            "recommended_assets": [
                {
                    "asset_type": "manual",
                    "review_status": "approved",
                    "usable": True,
                    "auto_send_level": "auto",
                }
            ]
        },
    )

    result = review_active_rubrics(apply=True, db_factory=session_factory)

    assert result["reviewed_count"] == 1
    row = _scenario(session_factory, "active_install_1")
    expected = row.get_expected_reply()
    assert row.status == "active"
    assert expected["auto_send_allowed"] is True
    assert expected["must_handoff"] is False
    assert "安装图或说明书" in expected["key_points"]


def test_active_rubric_review_rejects_plain_product_photo_as_install_material(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_active_installation(
        session_factory,
        expected={
            "expected_reply": "Send image.",
            "key_points": ["安装图或说明书"],
            "forbidden_claims": [],
            "needs_review": False,
            "quality": "valid",
            "auto_send_allowed": True,
            "must_handoff": False,
        },
        metadata_extra={
            "recommended_assets": [
                {
                    "asset_type": "sku_image",
                    "media_role": "product_photo",
                    "review_status": "approved",
                    "usable": True,
                    "auto_send_level": "auto",
                }
            ]
        },
    )

    result = review_active_rubrics(apply=False, db_factory=session_factory)

    assert result["changed_count"] == 1
    assert result["items"][0]["new_auto_send_allowed"] is False
    assert result["items"][0]["new_must_handoff"] is True
    assert "转人工" in result["items"][0]["new_key_points"]


def test_active_rubric_review_excel_uses_chinese_headers(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_active_installation(session_factory)
    excel_output = tmp_path / "active_rubrics.xlsx"

    review_active_rubrics(apply=False, excel_output=str(excel_output), db_factory=session_factory)

    workbook = load_workbook(excel_output)
    assert workbook.sheetnames == ["重审结果"]
    assert [cell.value for cell in workbook["重审结果"][1]][:5] == [
        "场景 UID",
        "标题",
        "场景类型",
        "问题类型",
        "是否变化",
    ]
