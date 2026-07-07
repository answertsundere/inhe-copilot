import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models.eval_tables import EvalRun, EvalTrace
from app.models.kb_tables import KBGenericServiceRule, KBMediaAsset, KBProduct
from scripts.diagnose_evidence_chain_for_replay import (
    build_workbook,
    diagnose_evidence_chain,
    _write_excel,
    _write_json,
)


def _session_factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'diag.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    return sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)


def _trace(run_uid, turn_uid, *, message, qft, raw=None, sidecar=True):
    trace = EvalTrace(
        run_uid=run_uid,
        case_uid=f"case-{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message=message,
        agent_reply="",
        query_fact_type=qft,
        passed=False,
    )
    sidecar_context = {
        "product_title": "测试书架",
        "sku_code": "SKU-1",
        "i_id": "IID-1",
        "order_id": "ORDER-HASH",
        "sidecar_context_quality": "complete",
        "sidecar_context_sources": ["sidecar_product_title", "sidecar_sku_code", "sidecar_i_id"],
    }
    turn_understanding = {
        "query_fact_type": qft,
        "effective_query_fact_type": qft,
        "sidecar_context_quality": "complete" if sidecar else "missing",
    }
    trace.set_turn_understanding(turn_understanding)
    payload = raw or {}
    if sidecar:
        payload["sidecar_context"] = sidecar_context
        payload["sidecar_context_quality"] = "complete"
    trace.set_raw_response(payload)
    trace.set_answer_trace({"query_fact_type": qft, "sidecar_context": sidecar_context if sidecar else {}})
    trace.set_product_identity({"sidecar_context": sidecar_context if sidecar else {}})
    trace.set_selected_evidence([])
    trace.set_failure_labels(["rag_miss"])
    return trace


def test_diagnoses_zero_evidence_layers_and_writes_chinese_workbook(tmp_path, monkeypatch):
    monkeypatch.delenv("COPILOT_EMBEDDING_ENABLED", raising=False)
    monkeypatch.delenv("COPILOT_EMBEDDING_API_BASE", raising=False)
    monkeypatch.delenv("COPILOT_EMBEDDING_API_KEY", raising=False)
    monkeypatch.delenv("COPILOT_EMBEDDING_MODEL", raising=False)
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-evidence"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=2))
    product = KBProduct(i_id="IID-1", product_name="测试书架", status="active")
    product.set_specs({"材质": "铁", "尺寸": "有尺寸字段"})
    db.add(product)
    db.flush()
    db.add(
        KBMediaAsset(
            product_id=product.id,
            i_id="IID-1",
            sku_code="SKU-1",
            product_name="测试书架",
            asset_type="product_photo",
            asset_title="普通商品图",
            asset_url="https://example.invalid/a.jpg",
            status="approved",
            audit_status="approved",
            usable_for_agent=1,
        )
    )
    db.add(KBGenericServiceRule(rule_key="promo", title="优惠核对", fact_type="promotion_policy", status="active"))
    pack = {
        "product_identity_resolution": {
            "status": "resolved",
            "resolved_product_id": product.id,
            "i_id": "IID-1",
            "sku_code": "SKU-1",
            "identity_confidence": 1.0,
            "match_reason": "sku_exact",
        },
        "resolved_product_identity": {"product_id": product.id, "i_id": "IID-1", "sku": "SKU-1", "product_name": "测试书架"},
        "product_structured_facts": [],
        "product_media_assets": [],
        "product_scoped_chunks": [],
        "generic_fallback_rules": [],
        "missing_required_evidence": ["installation_video"],
    }
    raw = {"evidence_debug": {"product_context_pack_summary": {"evidence_pack": pack}}}
    db.add(_trace(run_uid, "turn-install", message="有安装视频吗", qft="installation", raw=raw))
    db.add(_trace(run_uid, "turn-missing-sidecar", message="这个能用吗", qft="dimensions", sidecar=False))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["total_traces"] == 2
    assert report["summary"]["zero_evidence_trace_count"] == 2
    reasons = report["summary"]["by_primary_reason"]
    assert reasons["evidence_role_mismatch"] == 1
    assert reasons["sidecar_missing_or_not_passed"] == 1
    assert report["summary"]["embedding_config_status"]["missing_env_keys"] == [
        "COPILOT_EMBEDDING_ENABLED",
        "COPILOT_EMBEDDING_API_BASE",
        "COPILOT_EMBEDDING_API_KEY",
        "COPILOT_EMBEDDING_MODEL",
    ]

    workbook = build_workbook(report)
    assert "Zero Evidence 明细" in workbook.sheetnames
    assert "商品字段缺失" in workbook.sheetnames
    assert "素材缺失或不可用" in workbook.sheetnames
    assert "RAG Embedding 问题" in workbook.sheetnames
    assert workbook["Zero Evidence 明细"]["A1"].value == "回放批次"

    xlsx_path = tmp_path / "诊断.xlsx"
    json_path = tmp_path / "诊断.json"
    _write_excel(str(xlsx_path), report)
    _write_json(str(json_path), report)
    loaded = load_workbook(xlsx_path)
    assert loaded["总览"]["A1"].value == "指标"
    assert json.loads(Path(json_path).read_text(encoding="utf-8"))["run_uid"] == run_uid


def test_latest_completed_real_run_is_used(tmp_path):
    Session = _session_factory(tmp_path)
    db = Session()
    older = EvalRun(run_uid="older", source_type="real_conversation", status="completed")
    newer = EvalRun(run_uid="newer", source_type="real_conversation", status="completed")
    db.add_all([older, newer])
    db.flush()
    older.id = 1
    newer.id = 2
    db.add(_trace("newer", "turn-1", message="优惠吗", qft="promotion_policy"))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(latest=True, db_factory=Session)

    assert report["run_uid"] == "newer"
    assert report["summary"]["total_traces"] == 1


def test_direct_answer_evidence_summary_is_counted_as_selected_evidence(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://example.invalid")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "configured-for-test")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "text-embedding-v3")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-direct-evidence"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    product = KBProduct(i_id="IID-1", product_name="测试书架", status="active")
    product.set_specs({"installation": "贴纸贴在对应标记位置"})
    db.add(product)
    db.flush()
    raw = {
        "evidence_debug": {
            "knowledge_evidence_summary": [
                {
                    "source_type": "product_facts",
                    "query_fact_type": "installation",
                    "evidence_fact_type": "installation",
                    "gate_status": "allowed",
                    "direct_answer_allowed": True,
                    "evidence_allowed_for_exact_answer": True,
                    "chunk_preview": "安装说明：贴纸贴在对应标记位置。",
                }
            ],
            "product_context_pack_summary": {
                "evidence_pack": {
                    "resolved_product_identity": {"product_id": product.id, "i_id": "IID-1", "sku": "SKU-1"},
                    "identity_confidence": 1.0,
                    "product_structured_facts": [
                        {
                            "evidence_id": "kbproduct:1:installation",
                            "fact_type": "installation",
                            "direct_answer_allowed": True,
                            "preview": "贴纸贴在对应标记位置。",
                        }
                    ],
                }
            },
        }
    }
    db.add(_trace(run_uid, "turn-direct", message="贴纸贴哪", qft="installation", raw=raw))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["total_traces"] == 1
    assert report["summary"]["zero_evidence_trace_count"] == 0
    assert report["records"][0]["selected_evidence_count"] >= 1
    assert report["summary"]["embedding_config_status"]["status"] == "configured"


def test_direct_pack_evidence_is_counted_as_selected_not_zero_evidence(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://example.invalid")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "configured-for-test")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "text-embedding-v3")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-pack-not-consumed"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    product = KBProduct(i_id="IID-1", product_name="娴嬭瘯涔︽灦", status="active")
    db.add(product)
    db.flush()
    db.add(KBGenericServiceRule(rule_key="after-sales", title="售后核对", fact_type="aftersales_policy", status="active"))
    raw = {
        "evidence_debug": {
            "product_context_pack_summary": {
                "evidence_pack": {
                    "resolved_product_identity": {"product_id": product.id, "i_id": "IID-1", "sku": "SKU-1"},
                    "identity_confidence": 1.0,
                    "product_structured_facts": [
                        {
                            "fact_type": "aftersales_policy",
                            "direct_answer_allowed": True,
                            "preview": "售后处理说明。",
                        }
                    ],
                    "product_media_assets": [],
                    "product_scoped_chunks": [],
                    "generic_fallback_rules": [{"rule_key": "after-sales", "fact_type": "aftersales_policy"}],
                }
            }
        }
    }
    db.add(_trace(run_uid, "turn-pack", message="少件了怎么办", qft="aftersales_policy", raw=raw))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["zero_evidence_trace_count"] == 0
    row = report["records"][0]
    assert row["selected_evidence_count"] == 1
    assert row["pack_structured_facts_count"] == 1
    assert row["pack_scoped_chunks_count"] == 0
    assert row["pack_generic_rules_count"] == 1


def test_generic_fallback_only_is_not_pack_not_consumed(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://example.invalid")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "configured-for-test")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "text-embedding-v3")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-generic-only-pack"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    product = KBProduct(i_id="IID-1", product_name="测试书架", status="active")
    db.add(product)
    db.flush()
    raw = {
        "evidence_debug": {
            "product_context_pack_summary": {
                "evidence_pack": {
                    "resolved_product_identity": {"product_id": product.id, "i_id": "IID-1", "sku": "SKU-1"},
                    "identity_confidence": 1.0,
                    "product_structured_facts": [],
                    "product_media_assets": [],
                    "product_scoped_chunks": [],
                    "generic_fallback_rules": [{"rule_key": "after-sales", "fact_type": "aftersales_policy"}],
                    "missing_required_evidence": [{"evidence_type": "product_fact", "fact_type": "aftersales_policy"}],
                }
            }
        }
    }
    db.add(_trace(run_uid, "turn-generic-only", message="少件了怎么办", qft="aftersales_policy", raw=raw))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["zero_evidence_trace_count"] == 1
    assert report["summary"]["by_primary_reason"] == {"true_knowledge_gap": 1}


def test_unknown_unscored_turn_is_context_or_noise_not_fact_type_missing(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://example.invalid")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "configured-for-test")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "text-embedding-v3")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-context-or-noise"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    trace = _trace(
        run_uid,
        "turn-noise",
        message="好的",
        qft="",
        raw={
            "quality_bucket": {
                "quality_bucket": "unscored_or_noise",
                "should_count_in_quality_rate": False,
            }
        },
    )
    trace.set_turn_understanding({
        "query_fact_type": "",
        "turn_actionability": "acknowledgement",
        "should_score": False,
        "skip_reason": "non_actionable_acknowledgement",
    })
    trace.set_failure_labels([])
    db.add(trace)
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["zero_evidence_trace_count"] == 1
    assert report["summary"]["by_primary_reason"] == {"context_or_noise": 1}
    assert "query_fact_type_missing" not in report["summary"]["by_primary_reason"]


def test_rag_candidates_without_selected_summary_are_classified(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://example.invalid")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "configured-for-test")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "text-embedding-v3")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-rag-summary-missing"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    product = KBProduct(i_id="IID-1", product_name="娴嬭瘯涔︽灦", status="active")
    product.set_specs({"material": "PP"})
    db.add(product)
    db.flush()
    raw = {"evidence_debug": {"evidence_gate_summary": {"candidate_count": 3, "filtered_count": 2}}}
    db.add(_trace(run_uid, "turn-rag-candidates", message="什么材质", qft="material", raw=raw))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["zero_evidence_trace_count"] == 1
    assert report["summary"]["by_primary_reason"] == {"selected_evidence_summary_missing": 1}
    assert report["zero_evidence_records"][0]["rag_candidate_count"] == 3


def test_empty_pack_without_other_breakpoint_is_classified(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://example.invalid")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "configured-for-test")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "text-embedding-v3")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-empty-pack"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    db.add(KBProduct(i_id="IID-1", product_name="娴嬭瘯涔︽灦", status="active"))
    db.add(_trace(run_uid, "turn-empty-pack", message="这个还有别的吗", qft="product_question"))
    db.commit()
    db.close()

    report = diagnose_evidence_chain(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["zero_evidence_trace_count"] == 1
    assert report["summary"]["by_primary_reason"] == {"product_first_pack_empty": 1}
