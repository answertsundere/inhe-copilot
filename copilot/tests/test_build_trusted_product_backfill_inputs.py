from __future__ import annotations

import json

from openpyxl import load_workbook


def test_build_trusted_inputs_from_json_and_preserves_internal_ids(tmp_path):
    from scripts.build_trusted_product_backfill_inputs import run_build

    source = tmp_path / "trusted.json"
    source.write_text(json.dumps([
        {
            "source": "trusted_product_sheet",
            "platform_item_id": "123456789",
            "product_url": "https://item.example.com/item?id=123456789&token=secret",
            "platform_product_title": "只做参考的标题",
            "i_id": "YH40K01",
            "sku_code": "YH40K01B01S01",
            "material": "PP",
            "gross_weight": "2.6kg",
            "asset_url": "https://cdn.example.com/install.mp4?Signature=secret",
            "asset_title": "安装视频",
            "mime_type": "video/mp4",
        }
    ], ensure_ascii=False), encoding="utf-8")

    result = run_build(inputs=[str(source)], output_dir=str(tmp_path), date="20260630")

    assert result["trusted_identity_rows"] == 1
    assert result["trusted_structured_rows"] == 1
    assert result["trusted_media_rows"] == 1
    identity = load_workbook(result["outputs"]["identity"], data_only=True).active
    headers = [cell.value for cell in identity[1]]
    values = [cell.value for cell in identity[2]]
    row = dict(zip(headers, values))
    assert row["i_id"] == "YH40K01"
    assert row["sku_code"] == "YH40K01B01S01"
    assert row["confirmation_status"] == "confirmed"
    assert "token=secret" not in (row["product_url"] or "")
    media = load_workbook(result["outputs"]["media"], data_only=True).active
    media_row = dict(zip([cell.value for cell in media[1]], [cell.value for cell in media[2]]))
    assert media_row["asset_url"] == "https://cdn.example.com/install.mp4"


def test_build_trusted_inputs_excludes_untrusted_sources(tmp_path):
    from scripts.build_trusted_product_backfill_inputs import run_build

    source = tmp_path / "untrusted.json"
    source.write_text(json.dumps([
        {
            "source": "history_chat",
            "i_id": "YH40K01",
            "sku_code": "YH40K01B01S01",
            "asset_url": "https://cdn.example.com/install.jpg",
            "asset_title": "聊天里的安装图",
        }
    ], ensure_ascii=False), encoding="utf-8")

    result = run_build(inputs=[str(source)], output_dir=str(tmp_path), date="20260630")

    assert result["trusted_identity_rows"] == 0
    assert result["trusted_structured_rows"] == 0
    assert result["trusted_media_rows"] == 0
    assert result["manual_review_rows"] == 1


def test_title_only_row_does_not_generate_identity_or_internal_ids(tmp_path):
    from scripts.build_trusted_product_backfill_inputs import run_build

    source = tmp_path / "title_only.json"
    source.write_text(json.dumps([
        {
            "source": "trusted_product_sheet",
            "platform_product_title": "看起来很像某个商品的标题",
            "product_url": "https://item.example.com/item?id=1",
        }
    ], ensure_ascii=False), encoding="utf-8")

    result = run_build(inputs=[str(source)], output_dir=str(tmp_path), date="20260630")

    assert result["trusted_identity_rows"] == 0
    identity = load_workbook(result["outputs"]["identity"], data_only=True).active
    assert identity.max_row == 1


def test_build_trusted_inputs_from_excel_with_english_headers(tmp_path):
    from openpyxl import Workbook
    from scripts.build_trusted_product_backfill_inputs import run_build

    source = tmp_path / "trusted.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["source", "i_id", "sku_code", "material", "asset_url", "asset_title"])
    ws.append(["trusted_product_sheet", "YH41K01", "YH41K01B01S01", "ABS", "https://cdn.example.com/size.jpg?token=1", "尺寸图"])
    wb.save(source)

    result = run_build(inputs=[str(source)], output_dir=str(tmp_path), date="20260630")

    assert result["trusted_structured_rows"] == 1
    assert result["trusted_media_rows"] == 1
