from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.db as db_module
from app.db import Base
from app.models.eval_tables import AgentBenchmarkScenario
from app.services.agent_benchmark_runner_service import AgentBenchmarkRunnerService
from scripts.run_agent_benchmark import run_benchmark_report


def _patch_test_db(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    Base.metadata.create_all(bind=engine, tables=[AgentBenchmarkScenario.__table__])
    return session_factory


def _add_scenario(
    session_factory,
    *,
    scenario_uid: str = "bench_runner_1",
    expected: dict | None = None,
    scenario_type: str = "presales",
    status: str = "active",
):
    expected = expected or {
        "expected_reply": "The material is reinforced board.",
        "key_points": ["reinforced board"],
        "forbidden_claims": ["zero formaldehyde"],
        "must_handoff": False,
        "auto_send_allowed": True,
    }
    db = session_factory()
    try:
        row = AgentBenchmarkScenario(
            scenario_uid=scenario_uid,
            source_type="manual",
            source_uid=scenario_uid,
            status=status,
            title="Material question",
            scenario_type=scenario_type,
        )
        row.set_sidecar_context({
            "product_title": "Kids storage cabinet",
            "sku_code": "SKU-BENCH",
            "i_id": "IID-BENCH",
            "order_id": "ORDER-BENCH",
        })
        row.set_conversation_turns([
            {"turn_uid": f"{scenario_uid}_turn_1", "speaker": "buyer", "text": "Is this material safe?"}
        ])
        row.set_expected_reply(expected)
        row.set_rubric({"fact_correctness": True})
        db.add(row)
        db.commit()
    finally:
        db.close()


def test_runner_injects_sidecar_context_into_agent_payload(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    payloads = []

    def fake_agent(payload):
        payloads.append(payload)
        return {
            "can_send": True,
            "requires_human_review": False,
            "sendable_reply": "The material is reinforced board. Please confirm against the product page.",
            "answer_trace": {"query_fact_type": "material"},
        }

    result = AgentBenchmarkRunnerService(agent_callable=fake_agent).run_scenarios(db_factory=session_factory)

    assert result["passed"] == 1
    assert result["total_scenarios"] == 1
    assert result["passed_count"] == 1
    assert result["by_scenario_type"]["presales"]["passed"] == 1
    assert result["by_query_fact_type"]["material"]["passed"] == 1
    assert result["by_reply_status"]["can_send"] == 1
    item = result["per_scenario_result"][0]
    assert item["title"] == "Material question"
    assert item["conversation_turns"]
    assert item["expected_reply"]["expected_reply"]
    assert item["can_send"] is True
    assert item["reply_status"] == "can_send"
    assert payloads[0]["product_name"] == "Kids storage cabinet"
    assert payloads[0]["sku_code"] == "SKU-BENCH"
    assert payloads[0]["copilot_context"]["sidecar_context"]["i_id"] == "IID-BENCH"


def test_runner_injects_reviewed_query_fact_type_contract(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory, scenario_type="installation")
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_runner_1").one()
        metadata = row.get_metadata()
        metadata["query_fact_type"] = "installation"
        row.set_metadata(metadata)
        db.commit()
    finally:
        db.close()
    payloads = []

    def fake_agent(payload):
        payloads.append(payload)
        return {
            "can_send": False,
            "requires_human_review": True,
            "draft_reply": "Please confirm installation material with a human agent.",
            "answer_trace": {"query_fact_type": "installation"},
        }

    AgentBenchmarkRunnerService(agent_callable=fake_agent).run_scenarios(db_factory=session_factory)

    turn_contract = payloads[0]["copilot_context"]["turn_understanding"]
    assert payloads[0]["copilot_context"]["benchmark_query_fact_type"] == "installation"
    assert turn_contract["query_fact_type"] == "installation"
    assert turn_contract["expected_query_fact_type"] == "installation"
    assert "expected_reply" not in payloads[0]["copilot_context"]


def test_runner_scores_reviewed_source_turn_not_later_buyer_turn(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory, scenario_type="installation")
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_runner_1").one()
        row.set_metadata({"source_turn_uid": "target_turn", "query_fact_type": "installation"})
        row.set_conversation_turns([
            {"turn_uid": "target_turn", "speaker": "buyer", "text": "Need installation material."},
            {"turn_uid": "csr_1", "speaker": "service", "text": "I will check it."},
            {"turn_uid": "later_turn", "speaker": "buyer", "text": "Different later question."},
        ])
        db.commit()
    finally:
        db.close()
    payloads = []

    def fake_agent(payload):
        payloads.append(payload)
        return {
            "can_send": False,
            "requires_human_review": True,
            "draft_reply": "Please confirm the installation material with a human agent.",
            "answer_trace": {"query_fact_type": "installation"},
        }

    result = AgentBenchmarkRunnerService(agent_callable=fake_agent).run_scenarios(db_factory=session_factory)

    assert result["total_scenarios"] == 1
    assert len(payloads) == 1
    assert payloads[0]["message"] == "Need installation material."
    assert payloads[0]["copilot_context"]["conversation_history"] == []
    assert "responses" not in result["per_scenario_result"][0]
    traced = AgentBenchmarkRunnerService(agent_callable=fake_agent).run_scenarios(
        db_factory=session_factory,
        include_full_trace=True,
    )
    assert traced["per_scenario_result"][0]["responses"][0]["turn_uid"] == "target_turn"


def test_runner_empty_when_no_active_scenarios(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory, status="candidate")

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {}).run_scenarios(db_factory=session_factory)

    assert result["total_scenarios"] == 0
    assert result["passed_count"] == 0
    assert result["failed_count"] == 0
    assert result["pass_rate"] == 0
    assert result["per_scenario_result"] == []


def test_missing_key_points_fail(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "Please confirm against the product page.",
        "answer_trace": {"query_fact_type": "material"},
    }).run_scenarios(db_factory=session_factory)

    assert result["failed"] == 1
    item = result["per_scenario_result"][0]
    assert item["failure_reasons"] == ["missing_key_point"]
    assert item["missing_key_points"] == ["reinforced board"]


def test_aftersales_key_points_accept_generic_aliases(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="aftersales",
        expected={
            "expected_reply": "Please comfort the customer and hand off.",
            "key_points": ["先别着急", "订单信息", "实物照片", "转人工"],
            "forbidden_claims": ["直接退款"],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "亲，我先帮您核实当前订单。麻烦把问题位置拍照发来，我这边人工核实后给您处理方案。",
        "answer_trace": {"query_fact_type": "aftersales"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["failure_reasons"] == []
    assert item["missing_key_points"] == []


def test_current_product_key_point_accepts_current_style_alias(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="promotion",
        expected={
            "expected_reply": "Verify the current product promotion.",
            "key_points": ["\u6309\u5f53\u524d\u5546\u54c1"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u6211\u5e2e\u60a8\u770b\u4e00\u4e0b\u5f53\u524d\u8fd9\u6b3e\u80fd\u7528\u7684\u4f18\u60e0\uff0c\u5177\u4f53\u4ee5\u4e0b\u5355\u9875\u9762\u663e\u793a\u4e3a\u51c6\u3002",
        "answer_trace": {"query_fact_type": "promotion"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["missing_key_points"] == []


def test_current_product_key_point_accepts_your_style_alias(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="installation",
        expected={
            "expected_reply": "Verify against this product.",
            "key_points": ["\u6309\u5f53\u524d\u8fd9\u6b3e\u5546\u54c1"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u8fd9\u4e2a\u9700\u8981\u6309\u60a8\u8fd9\u6b3e\u7684\u7ed3\u6784\u548c\u914d\u4ef6\u89c4\u683c\u6838\u5bf9\uff0c\u6211\u8fd9\u8fb9\u8f6c\u4eba\u5de5\u786e\u8ba4\u3002",
        "answer_trace": {"query_fact_type": "installation"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["missing_key_points"] == []


def test_current_product_key_point_rejects_page_only_reply(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="promotion",
        expected={
            "expected_reply": "Verify the current product promotion.",
            "key_points": ["\u6309\u5f53\u524d\u5546\u54c1"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u4f18\u60e0\u4e00\u822c\u4ee5\u4e0b\u5355\u9875\u9762\u663e\u793a\u4e3a\u51c6\uff0c\u6211\u8fd9\u8fb9\u53ef\u4ee5\u5e2e\u60a8\u6838\u5bf9\u3002",
        "answer_trace": {"query_fact_type": "promotion"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["failed"] == 1
    assert item["missing_key_points"] == ["\u6309\u5f53\u524d\u5546\u54c1"]


def test_promotion_discount_promise_key_point_accepts_page_rule_boundary(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="promotion",
        expected={
            "expected_reply": "Do not promise extra discount.",
            "key_points": ["\u4e0d\u80fd\u76f4\u63a5\u627f\u8bfa\u989d\u5916\u964d\u4ef7"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u5177\u4f53\u4ee5\u60a8\u4e0b\u5355\u9875\u9762\u663e\u793a\u4e3a\u51c6\uff0c\u6211\u8fd9\u8fb9\u53ef\u4ee5\u5e2e\u60a8\u6838\u5bf9\u6d3b\u52a8\u89c4\u5219\u3002",
        "answer_trace": {"query_fact_type": "promotion"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["missing_key_points"] == []


def test_promotion_discount_promise_key_point_rejects_promised_discount(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="promotion",
        expected={
            "expected_reply": "Do not promise extra discount.",
            "key_points": ["\u4e0d\u80fd\u76f4\u63a5\u627f\u8bfa\u989d\u5916\u964d\u4ef7"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u6211\u53ef\u4ee5\u76f4\u63a5\u7ed9\u60a8\u989d\u5916\u4fbf\u5b9c\u5341\u5143\u3002",
        "answer_trace": {"query_fact_type": "promotion"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["failed"] == 1
    assert item["missing_key_points"] == ["\u4e0d\u80fd\u76f4\u63a5\u627f\u8bfa\u989d\u5916\u964d\u4ef7"]


def test_installation_video_promise_key_point_accepts_no_sendable_video_boundary(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="installation",
        expected={
            "expected_reply": "Do not promise installation video.",
            "key_points": ["\u4e0d\u76f4\u63a5\u627f\u8bfa\u6709\u5b89\u88c5\u89c6\u9891"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u8fd9\u6b3e\u76ee\u524d\u6682\u65f6\u6ca1\u6709\u53ef\u76f4\u63a5\u53d1\u9001\u7684\u5b89\u88c5\u89c6\u9891\uff0c\u6211\u5148\u5e2e\u60a8\u6838\u5bf9\u5b89\u88c5\u8d44\u6599\u3002",
        "answer_trace": {"query_fact_type": "installation"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["missing_key_points"] == []


def test_installation_video_negative_key_point_accepts_no_video_promise(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="installation",
        expected={
            "expected_reply": "Do not promise installation video.",
            "key_points": ["\u4e0d\u76f4\u63a5\u627f\u8bfa\u6709\u5b89\u88c5\u89c6\u9891"],
            "forbidden_claims": ["\u4e00\u5b9a\u6709\u5b89\u88c5\u89c6\u9891"],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": False,
        "requires_human_review": True,
        "draft_reply": "\u4eb2\uff0c\u8fd9\u4e2a\u9700\u8981\u6309\u60a8\u8fd9\u6b3e\u7684\u7ed3\u6784\u548c\u914d\u4ef6\u89c4\u683c\u6838\u5bf9\uff0c\u6211\u8fd9\u8fb9\u8f6c\u4eba\u5de5\u786e\u8ba4\u662f\u5426\u9002\u914d\u3002",
        "answer_trace": {"query_fact_type": "installation"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["missing_key_points"] == []


def test_installation_material_key_point_accepts_diagram_and_manual_aliases(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        scenario_type="installation",
        expected={
            "expected_reply": "Send the approved installation diagram or manual when available.",
            "key_points": ["\u5b89\u88c5\u56fe\u6216\u8bf4\u660e\u4e66"],
            "forbidden_claims": [],
            "must_handoff": False,
            "auto_send_allowed": True,
        },
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "\u6211\u5148\u628a\u5b89\u88c5\u793a\u610f\u56fe/\u8bf4\u660e\u4e66\u53d1\u60a8\u53c2\u8003\uff0c\u60a8\u53ef\u4ee5\u6309\u56fe\u7eb8\u4e0a\u7684\u6b65\u9aa4\u5b89\u88c5\u3002",
        "answer_trace": {"query_fact_type": "installation"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["passed"] == 1
    assert item["missing_key_points"] == []


def test_forbidden_claims_fail(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "The material is reinforced board and has zero formaldehyde.",
        "answer_trace": {"query_fact_type": "material"},
    }).run_scenarios(db_factory=session_factory)

    item = result["per_scenario_result"][0]
    assert result["failed"] == 1
    assert "forbidden_claim_present" in item["failure_reasons"]
    assert item["forbidden_claims_hit"] == ["zero formaldehyde"]


def test_must_handoff_and_auto_send_allowed_contracts(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(
        session_factory,
        expected={
            "expected_reply": "Please ask a human agent to confirm the discount.",
            "key_points": ["human agent"],
            "forbidden_claims": [],
            "must_handoff": True,
            "auto_send_allowed": False,
        },
        scenario_type="promotion",
    )

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "Please ask a human agent to confirm the discount.",
        "answer_trace": {"query_fact_type": "promotion_policy"},
    }).run_scenarios(db_factory=session_factory)

    reasons = set(result["per_scenario_result"][0]["failure_reasons"])
    assert result["failed"] == 1
    assert "auto_send_not_allowed" in reasons
    assert "handoff_mismatch" in reasons


def test_fact_type_mismatch_fails(monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory, scenario_type="installation")

    result = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "The material is reinforced board.",
        "answer_trace": {"query_fact_type": "material"},
    }).run_scenarios(db_factory=session_factory)

    assert result["failed"] == 1
    assert "fact_type_mismatch" in result["per_scenario_result"][0]["failure_reasons"]


def test_run_benchmark_report_writes_json_and_excel(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    json_output = tmp_path / "benchmark.json"
    excel_output = tmp_path / "benchmark.xlsx"

    result, exit_code = run_benchmark_report(
        json_output=str(json_output),
        excel_output=str(excel_output),
        fail_on_failure=True,
        db_factory=session_factory,
        agent_callable=lambda _payload: {
            "can_send": True,
            "requires_human_review": False,
            "sendable_reply": "Please confirm against the product page.",
            "answer_trace": {"query_fact_type": "material"},
        },
    )

    assert exit_code == 1
    assert result["failed_count"] == 1
    assert json_output.exists()
    workbook = load_workbook(excel_output)
    assert set(workbook.sheetnames) == {"汇总", "按场景", "按问题类型", "失败原因", "逐条结果"}
    assert [cell.value for cell in workbook["逐条结果"][1]][:5] == ["场景 UID", "标题", "场景类型", "问题类型", "是否通过"]
    assert "responses" not in json_output.read_text(encoding="utf-8")
