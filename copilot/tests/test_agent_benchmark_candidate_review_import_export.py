from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.db as db_module
from app.db import Base
from app.models.eval_tables import AgentBenchmarkScenario
from app.services.agent_benchmark_runner_service import AgentBenchmarkRunnerService
from scripts.export_agent_benchmark_candidates import export_candidates
from scripts.import_reviewed_agent_benchmark_candidates import import_reviewed_candidates


def _patch_test_db(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    Base.metadata.create_all(bind=engine, tables=[AgentBenchmarkScenario.__table__])
    return session_factory


def _add_scenario(
    session_factory,
    *,
    scenario_uid: str = "bench_review_1",
    expected_reply: str = "",
    expected_quality: str = "low_quality",
    needs_review: bool = True,
    with_sidecar: bool = True,
    turns: list[dict] | None = None,
):
    db = session_factory()
    try:
        row = AgentBenchmarkScenario(
            scenario_uid=scenario_uid,
            source_type="real_conversation",
            source_uid=f"source_{scenario_uid}",
            status="candidate",
            title="benchmark candidate",
            scenario_type="presales",
            created_by="test",
        )
        row.set_sidecar_context({"product_title": "Kids storage cabinet", "sku_code": "SKU-1"} if with_sidecar else {})
        row.set_conversation_turns(
            turns
            if turns is not None
            else [
                {"speaker": "buyer", "text": "Is the material safe?"},
                {"speaker": "csr", "text": "Please wait."},
                {"speaker": "buyer", "text": "Will it get damp?"},
            ]
        )
        row.set_expected_reply({
            "expected_reply": expected_reply,
            "key_points": [],
            "forbidden_claims": [],
            "needs_review": needs_review,
            "auto_send_allowed": False,
            "must_handoff": True,
        })
        row.set_metadata({
            "expected_reply_quality": expected_quality,
            "expected_reply_block_reason": "welcome_reference_reply" if expected_quality == "low_quality" else "",
            "needs_expected_reply_review": needs_review,
            "original_cs_reply": "Welcome to our shop.",
            "query_fact_type": "material",
        })
        db.add(row)
        db.commit()
    finally:
        db.close()


def _write_review_sheet(path, rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Agent benchmark review"
    sheet.append(headers or [
        "scenario_uid",
        "当前标准答案",
        "关键点",
        "禁止话术",
        "必须转人工",
        "允许自动发送",
        "审核人",
        "审核备注",
    ])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_export_keeps_low_quality_candidate_blocked(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory, expected_reply="", expected_quality="low_quality")
    output = tmp_path / "candidates.xlsx"

    result = export_candidates(str(output), db_factory=session_factory)

    assert result["exported_count"] == 1
    workbook = load_workbook(output)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    row = {headers[idx]: sheet[2][idx].value for idx in range(len(headers))}
    assert "建议标准答案草稿" in headers
    assert row["质量状态"] == "low_quality"
    assert row["阻断原因"] == "welcome_reference_reply"
    assert "Will it get damp?" in row["客户完整对话"]
    assert row["当前标准答案"] is None


def test_import_dry_run_does_not_write_review(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [[
        "bench_review_1",
        "This material should be checked against the product record before replying.",
        "product record",
        "absolutely safe",
        "true",
        "false",
        "lead",
        "reviewed",
    ]])

    result = import_reviewed_candidates(str(review_file), apply=False, db_factory=session_factory)

    assert result["updated_count"] == 1
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_review_1").one()
        assert row.get_expected_reply()["needs_review"] is True
        assert row.get_metadata()["expected_reply_quality"] == "low_quality"
    finally:
        db.close()


def test_import_apply_writes_reviewed_expected_reply(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [[
        "bench_review_1",
        "This material should be checked against the product record before replying.",
        "product record",
        "absolutely safe",
        "true",
        "false",
        "lead",
        "reviewed",
    ]])

    result = import_reviewed_candidates(str(review_file), apply=True, db_factory=session_factory)

    assert result["updated_count"] == 1
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_review_1").one()
        expected = row.get_expected_reply()
        assert expected["needs_review"] is False
        assert expected["key_points"] == ["product record"]
        assert expected["forbidden_claims"] == ["absolutely safe"]
        assert row.get_metadata()["expected_reply_quality"] == "valid"
    finally:
        db.close()


def test_import_uses_suggested_draft_when_current_expected_reply_is_empty(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(
        review_file,
        [[
            "bench_review_1",
            "",
            "Use product record",
            "do not promise exact safety",
            "true",
            "false",
            "lead",
            "",
            "Please check the current product record before replying to the customer.",
        ]],
        headers=[
            "scenario_uid",
            "当前标准答案",
            "关键点",
            "禁止话术",
            "必须转人工",
            "允许自动发送",
            "审核人",
            "审核备注",
            "建议标准答案草稿",
        ],
    )

    result = import_reviewed_candidates(str(review_file), apply=True, db_factory=session_factory)

    assert result["updated_count"] == 1
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_review_1").one()
        assert row.get_expected_reply()["expected_reply"] == "Please check the current product record before replying to the customer."
        assert row.get_expected_reply()["needs_review"] is False
    finally:
        db.close()


def test_import_skips_missing_or_unknown_scenario_uid(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [
        ["", "Valid reviewed reply", "", "", "false", "false", "lead", ""],
        ["unknown_uid", "Valid reviewed reply", "", "", "false", "false", "lead", ""],
    ])

    result = import_reviewed_candidates(str(review_file), apply=True, db_factory=session_factory)

    assert result["skipped_count"] == 2
    assert {item["reason"] for item in result["skipped"]} == {"missing_scenario_uid", "scenario_not_found"}


def test_reviewer_required_for_import(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [["bench_review_1", "Check the product record before replying.", "", "", "false", "false", "", ""]])

    result = import_reviewed_candidates(str(review_file), apply=True, db_factory=session_factory)

    assert result["skipped_count"] == 1
    assert result["skipped"][0]["reason"] == "reviewer_required"


def test_dry_run_promote_active_requires_sidecar_context(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory, scenario_uid="bench_no_sidecar", with_sidecar=False)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [[
        "bench_no_sidecar",
        "This reply has enough business detail for the reviewed benchmark answer.",
        "",
        "",
        "false",
        "false",
        "lead",
        "",
    ]])

    result = import_reviewed_candidates(str(review_file), apply=False, promote_active=True, db_factory=session_factory)

    assert result["dry_run"] is True
    assert result["updated_count"] == 1
    assert result["would_update_count"] == 1
    assert result["promoted_count"] == 0
    assert result["would_promote_count"] == 0
    assert result["promoted_scenario_uids"] == []
    assert result["errors"] == [{"scenario_uid": "bench_no_sidecar", "reason": "sidecar_context_required"}]
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_no_sidecar").one()
        assert row.status == "candidate"
        assert row.get_expected_reply()["needs_review"] is True
        assert row.get_metadata()["expected_reply_quality"] == "low_quality"
    finally:
        db.close()


def test_dry_run_promote_active_reports_would_promote_for_eligible_scenario(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [[
        "bench_review_1",
        "This reply has enough business detail for the reviewed benchmark answer.",
        "",
        "",
        "false",
        "false",
        "lead",
        "",
    ]])

    result = import_reviewed_candidates(str(review_file), apply=False, promote_active=True, db_factory=session_factory)

    assert result["updated_count"] == 1
    assert result["would_update_count"] == 1
    assert result["promoted_count"] == 0
    assert result["would_promote_count"] == 1
    assert result["promoted_scenario_uids"] == []
    assert result["would_promote_scenario_uids"] == ["bench_review_1"]
    db = session_factory()
    try:
        row = db.query(AgentBenchmarkScenario).filter_by(scenario_uid="bench_review_1").one()
        assert row.status == "candidate"
        assert row.get_expected_reply()["needs_review"] is True
        assert row.get_metadata()["expected_reply_quality"] == "low_quality"
    finally:
        db.close()


def test_dry_run_promote_active_does_not_promote_low_quality_expected_reply(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [["bench_review_1", "ok", "", "", "false", "false", "lead", ""]])

    result = import_reviewed_candidates(str(review_file), apply=False, promote_active=True, db_factory=session_factory)

    assert result["updated_count"] == 0
    assert result["promoted_count"] == 0
    assert result["would_promote_count"] == 0
    assert result["skipped_count"] == 1
    assert result["skipped"][0]["scenario_uid"] == "bench_review_1"


def test_promote_active_only_for_eligible_reviewed_scenario(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    _add_scenario(session_factory, scenario_uid="bench_no_sidecar", with_sidecar=False)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [
        ["bench_review_1", "Check this material against the product record before replying.", "", "", "false", "false", "lead", ""],
        ["bench_no_sidecar", "Check this material against the product record before replying.", "", "", "false", "false", "lead", ""],
    ])

    result = import_reviewed_candidates(str(review_file), apply=True, promote_active=True, db_factory=session_factory)

    assert result["promoted_scenario_uids"] == ["bench_review_1"]
    assert result["error_count"] == 1
    assert result["errors"][0]["reason"] == "sidecar_context_required"


def test_imported_key_points_and_forbidden_claims_are_used_by_runner(tmp_path, monkeypatch):
    session_factory = _patch_test_db(monkeypatch)
    _add_scenario(session_factory)
    review_file = tmp_path / "review.xlsx"
    _write_review_sheet(review_file, [[
        "bench_review_1",
        "This material should be checked against the product record before replying.",
        "product record",
        "absolutely safe",
        "false",
        "true",
        "lead",
        "",
    ]])
    import_reviewed_candidates(str(review_file), apply=True, promote_active=True, db_factory=session_factory)

    missing_key_point = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "Please ask support.",
        "answer_trace": {"query_fact_type": "material"},
    }).run_scenarios(db_factory=session_factory)
    assert "missing_key_point" in missing_key_point["per_scenario_result"][0]["failure_reasons"]

    forbidden = AgentBenchmarkRunnerService(agent_callable=lambda _payload: {
        "can_send": True,
        "requires_human_review": False,
        "sendable_reply": "This material should be checked against the product record and is absolutely safe.",
        "answer_trace": {"query_fact_type": "material"},
    }).run_scenarios(db_factory=session_factory)
    assert "forbidden_claim_present" in forbidden["per_scenario_result"][0]["failure_reasons"]
