from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


@pytest.fixture()
def candidate_db(monkeypatch):
    import app.db as db_module
    from app.models.eval_tables import EvalRun, EvalTrace
    from app.models.kb_tables import KBProduct, ProductIdentityMapping

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db_module, "engine", engine)
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    db_module.Base.metadata.create_all(
        bind=engine,
        tables=[EvalRun.__table__, EvalTrace.__table__, KBProduct.__table__, ProductIdentityMapping.__table__],
    )
    return session_factory


def _add_run(db, run_uid="run_candidates"):
    from app.models.eval_tables import EvalRun

    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed"))
    db.commit()


def _add_product(db, *, i_id="YH71K01", sku="YH71K01B01S01", name="trusted product"):
    from app.models.kb_tables import KBProduct

    product = KBProduct(i_id=i_id, product_name=name, status="published")
    product.set_sku_list([{"sku_code": sku}])
    db.add(product)
    db.commit()
    return product


def _add_trace(db, *, item_hash="", product_url="", title="", i_id="", sku="", turn_uid="turn-1"):
    from app.models.eval_tables import EvalTrace

    identity = {
        "item_id_hash": item_hash,
        "product_url": product_url,
        "product_title": title,
        "i_id": i_id,
        "sku_code": sku,
    }
    trace = EvalTrace(
        run_uid="run_candidates",
        case_uid=f"case-{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message="sanitized buyer message",
        passed=False,
    )
    trace.set_product_identity({"real_context_product_identity": identity})
    trace.set_raw_response({"real_context_product_identity": identity})
    db.add(trace)
    db.commit()


def test_exact_sku_and_platform_item_generates_auto_candidate(candidate_db):
    from scripts.build_product_identity_mapping_candidates import build_candidates

    db = candidate_db()
    try:
        _add_run(db)
        _add_product(db, i_id="YH71K01", sku="YH71K01B01S01")
        _add_trace(
            db,
            item_hash="hash-strong",
            product_url="https://item.example.com/item.htm?id=1&token=secret",
            sku="YH71K01B01S01",
        )
    finally:
        db.close()

    result = build_candidates(run_uid="run_candidates", include_kb_embedded=False, db_factory=candidate_db)

    assert result["auto_candidate_count"] == 1
    row = result["auto_candidates"][0]
    assert row["i_id"] == "YH71K01"
    assert row["sku_code"] == "YH71K01B01S01"
    assert row["product_url"] == "https://item.example.com/item.htm"
    assert row["match_method"] == "exact_sku_in_context"
    assert row["confirmation_status"] == "confirmed"


def test_exact_i_id_and_product_url_generates_auto_candidate(candidate_db):
    from scripts.build_product_identity_mapping_candidates import build_candidates

    db = candidate_db()
    try:
        _add_run(db)
        _add_product(db, i_id="YH72K01", sku="YH72K01B01S01")
        _add_trace(db, item_hash="hash-iid", product_url="https://item.example.com/item.htm?id=2", i_id="YH72K01")
    finally:
        db.close()

    result = build_candidates(run_uid="run_candidates", include_kb_embedded=False, db_factory=candidate_db)

    assert result["auto_candidate_count"] == 1
    assert result["auto_candidates"][0]["i_id"] == "YH72K01"
    assert result["auto_candidates"][0]["match_method"] == "exact_i_id_in_context"


def test_title_only_goes_to_manual_review(candidate_db):
    from scripts.build_product_identity_mapping_candidates import build_candidates

    db = candidate_db()
    try:
        _add_run(db)
        _add_product(db, i_id="YH73K01", sku="YH73K01B01S01", name="same title")
        _add_trace(db, title="same title")
    finally:
        db.close()

    result = build_candidates(run_uid="run_candidates", include_kb_embedded=False, db_factory=candidate_db)

    assert result["auto_candidate_count"] == 0
    assert result["manual_candidate_count"] == 1
    assert result["manual_candidates"][0]["conflict_reason"] == "title_only_not_auto_mapped"


def test_ambiguous_candidate_goes_to_manual_review(candidate_db):
    from scripts.build_product_identity_mapping_candidates import build_candidates

    db = candidate_db()
    try:
        _add_run(db)
        _add_product(db, i_id="YH74K01", sku="YH74K01B01S01")
        _add_product(db, i_id="YH74K02", sku="YH74K02B01S01")
        _add_trace(db, item_hash="hash-ambiguous", sku="YH74K01B01S01", turn_uid="a")
        _add_trace(db, item_hash="hash-ambiguous", sku="YH74K02B01S01", turn_uid="b")
    finally:
        db.close()

    result = build_candidates(run_uid="run_candidates", include_kb_embedded=False, db_factory=candidate_db)

    assert result["auto_candidate_count"] == 0
    assert result["manual_candidate_count"] == 2
    assert {row["conflict_reason"] for row in result["manual_candidates"]} == {"ambiguous"}


def test_write_outputs_preserves_internal_uid_and_manual_headers(candidate_db, tmp_path):
    from scripts.build_product_identity_mapping_candidates import build_candidates, write_outputs

    db = candidate_db()
    try:
        _add_run(db)
        _add_product(db, i_id="YH75K01", sku="YH75K01B01S01")
        _add_trace(db, item_hash="hash-write", product_url="https://item.example.com/item.htm?id=3&Signature=secret", sku="YH75K01B01S01")
    finally:
        db.close()

    result = build_candidates(run_uid="run_candidates", include_kb_embedded=False, db_factory=candidate_db)
    outputs = write_outputs(
        result,
        auto_output=str(tmp_path / "auto.xlsx"),
        manual_output=str(tmp_path / "manual.xlsx"),
    )

    auto_sheet = load_workbook(outputs["auto_output"], data_only=True).active
    headers = [cell.value for cell in auto_sheet[1]]
    row = dict(zip(headers, [cell.value for cell in auto_sheet[2]]))
    assert row["i_id"] == "YH75K01"
    assert "Signature" not in row["product_url"]
    manual_sheet = load_workbook(outputs["manual_output"], data_only=True).active
    assert "人工确认状态" in [cell.value for cell in manual_sheet[1]]
