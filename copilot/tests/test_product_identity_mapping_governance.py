from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.services.eval_sanitizer_service import hash_sensitive


@pytest.fixture()
def mapping_db(monkeypatch):
    import app.db as db_module
    from app.models.eval_tables import EvalRun, EvalTrace
    from app.models.kb_tables import KBProduct, ProductIdentityMapping

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db_module, "engine", engine)
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    db_module.Base.metadata.create_all(
        bind=engine,
        tables=[
            EvalRun.__table__,
            EvalTrace.__table__,
            KBProduct.__table__,
            ProductIdentityMapping.__table__,
        ],
    )
    return session_factory


def _add_unresolved_trace(db, *, run_uid: str, turn_uid: str, item_hash: str, buyer_message: str = "phone 13812345678 order 123456789012345"):
    from app.models.eval_tables import EvalRun, EvalTrace

    if not db.query(EvalRun).filter(EvalRun.run_uid == run_uid).first():
        db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed"))
    trace = EvalTrace(
        run_uid=run_uid,
        case_uid=f"case_{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message=buyer_message,
        agent_reply="",
        passed=False,
    )
    trace.set_raw_response({
        "real_context_product_identity": {
            "item_id_hash": item_hash,
            "product_url": "https://item.taobao.com/item.htm?id=[LONG_ID_REDACTED:abc123]",
            "product_title": "Moon storage cabinet",
            "order_product_title": "",
        },
        "evidence_debug": {
            "product_context_pack_summary": {
                "evidence_pack": {
                    "product_identity_resolution": {
                        "status": "not_found",
                        "unresolved_reason": "no_matching_product_found",
                        "ambiguous_candidates": [],
                    }
                }
            }
        },
    })
    db.add(trace)
    db.commit()


def _add_product(db, *, i_id: str = "YH92K01", sku_code: str = "YH92K01B01S01"):
    from app.models.kb_tables import KBProduct

    product = KBProduct(
        i_id=i_id,
        product_name="Moon storage cabinet",
        status="published",
        sku_list_json=json.dumps([{"sku_code": sku_code}], ensure_ascii=False),
    )
    product.set_specs({"material": "PP"})
    db.add(product)
    db.commit()
    return product


def _write_mapping_workbook(path, *, item_hash: str, i_id: str = "YH92K01", sku_code: str = "YH92K01B01S01", status: str = "已确认"):
    wb = Workbook()
    ws = wb.active
    ws.title = "商品身份映射缺口"
    ws.append([
        "source_run_uid",
        "case_uid",
        "turn_uid",
        "platform_item_id",
        "platform_item_id_hash",
        "product_url",
        "platform_product_title",
        "order_product_title",
        "建议内部 i_id",
        "建议 SKU",
        "建议商品标题",
        "人工确认状态",
        "处理人",
        "备注",
    ])
    ws.append([
        "run_map_1",
        "case_1",
        "turn_1",
        "",
        item_hash,
        "https://item.taobao.com/item.htm?id=[LONG_ID_REDACTED:abc123]",
        "Moon storage cabinet",
        "",
        i_id,
        sku_code,
        "Moon storage cabinet",
        status,
        "ops",
        "phone 13812345678 order 123456789012345",
    ])
    wb.save(path)


def test_export_unresolved_identity_mapping_gaps_are_deduped_and_sanitized(mapping_db, tmp_path):
    from scripts.export_product_identity_mapping_gaps import HEADERS, run_export

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        _add_unresolved_trace(db, run_uid="run_map_1", turn_uid="turn_1", item_hash=item_hash)
        _add_unresolved_trace(db, run_uid="run_map_1", turn_uid="turn_2", item_hash=item_hash)
    finally:
        db.close()

    output = tmp_path / "identity_gaps.xlsx"
    result = run_export(run_uid="run_map_1", output=str(output), db_factory=mapping_db)

    assert result["count"] == 1
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["商品身份映射缺口"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == HEADERS
    values = "\n".join(str(cell.value or "") for row in sheet.iter_rows(values_only=False) for cell in row)
    assert "13812345678" not in values
    assert "123456789012345" not in values
    assert item_hash in values


def test_import_identity_mapping_dry_run_does_not_write(mapping_db, tmp_path):
    from app.models.kb_tables import ProductIdentityMapping
    from scripts.import_product_identity_mappings import run_import

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        _add_product(db)
    finally:
        db.close()
    workbook = tmp_path / "filled.xlsx"
    _write_mapping_workbook(workbook, item_hash=item_hash)

    result = run_import(str(workbook), apply=False, operator="ops", db_factory=mapping_db)

    assert result["matched_count"] == 1
    db = mapping_db()
    try:
        assert db.query(ProductIdentityMapping).count() == 0
    finally:
        db.close()


def test_import_identity_mapping_apply_writes_mapping(mapping_db, tmp_path):
    from app.models.kb_tables import ProductIdentityMapping
    from scripts.import_product_identity_mappings import run_import

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        _add_product(db)
    finally:
        db.close()
    workbook = tmp_path / "filled.xlsx"
    _write_mapping_workbook(workbook, item_hash=item_hash)

    result = run_import(str(workbook), apply=True, operator="ops", db_factory=mapping_db)

    assert result["matched_count"] == 1
    db = mapping_db()
    try:
        mapping = db.query(ProductIdentityMapping).one()
        assert mapping.platform_item_id_hash == item_hash
        assert mapping.i_id == "YH92K01"
        assert mapping.status == "active"
        metadata = mapping.get_metadata()
        assert metadata["operator"] == "ops"
        assert "123456789012345" not in json.dumps(metadata, ensure_ascii=False)
    finally:
        db.close()


def test_import_identity_mapping_skips_missing_product(mapping_db, tmp_path):
    from app.models.kb_tables import ProductIdentityMapping
    from scripts.import_product_identity_mappings import run_import

    workbook = tmp_path / "filled.xlsx"
    _write_mapping_workbook(workbook, item_hash=hash_sensitive("456789012345"), i_id="YH_DOES_NOT_EXIST")

    result = run_import(str(workbook), apply=True, operator="ops", db_factory=mapping_db)

    assert result["matched_count"] == 0
    assert result["skipped_reasons"][0]["reason"] == "kb_product_not_found"
    db = mapping_db()
    try:
        assert db.query(ProductIdentityMapping).count() == 0
    finally:
        db.close()


def test_import_identity_mapping_skips_ambiguous_existing_mapping(mapping_db, tmp_path):
    from app.models.kb_tables import ProductIdentityMapping
    from scripts.import_product_identity_mappings import run_import

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        first = _add_product(db, i_id="YH92K01", sku_code="YH92K01B01S01")
        _add_product(db, i_id="YH92K02", sku_code="YH92K02B01S01")
        db.add(ProductIdentityMapping(
            mapping_uid="pim_existing",
            platform_item_id_hash=item_hash,
            kb_product_id=first.id,
            i_id=first.i_id,
            status="active",
        ))
        db.commit()
    finally:
        db.close()
    workbook = tmp_path / "filled.xlsx"
    _write_mapping_workbook(workbook, item_hash=item_hash, i_id="YH92K02", sku_code="YH92K02B01S01")

    result = run_import(str(workbook), apply=True, operator="ops", db_factory=mapping_db)

    assert result["matched_count"] == 0
    assert result["skipped_reasons"][0]["reason"] == "ambiguous_existing_mapping"


def test_resolver_uses_imported_identity_mapping(mapping_db):
    from app.models.kb_tables import ProductIdentityMapping
    from app.services.product_identity_resolver import ProductIdentityResolver

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        product = _add_product(db)
        db.add(ProductIdentityMapping(
            mapping_uid="pim_hash",
            platform_item_id_hash=item_hash,
            kb_product_id=product.id,
            i_id=product.i_id,
            sku_code="YH92K01B01S01",
            status="active",
            confidence=1.0,
        ))
        db.commit()
    finally:
        db.close()

    result = ProductIdentityResolver().resolve(platform_product_id_hash=item_hash)

    assert result["status"] == "resolved"
    assert result["i_id"] == "YH92K01"
    assert result["source"] == "product_identity_mapping"


def test_resolver_ambiguous_identity_mapping_does_not_lock_product(mapping_db):
    from app.models.kb_tables import ProductIdentityMapping
    from app.services.product_identity_resolver import ProductIdentityResolver

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        first = _add_product(db, i_id="YH92K01", sku_code="YH92K01B01S01")
        second = _add_product(db, i_id="YH92K02", sku_code="YH92K02B01S01")
        db.add(ProductIdentityMapping(
            mapping_uid="pim_first",
            platform_item_id_hash=item_hash,
            kb_product_id=first.id,
            i_id=first.i_id,
            status="active",
        ))
        db.add(ProductIdentityMapping(
            mapping_uid="pim_second",
            platform_item_id_hash=item_hash,
            kb_product_id=second.id,
            i_id=second.i_id,
            status="active",
        ))
        db.commit()
    finally:
        db.close()

    result = ProductIdentityResolver().resolve(platform_product_id_hash=item_hash)

    assert result["status"] == "ambiguous"
    assert result["i_id"] == ""
    assert len(result["ambiguous_candidates"]) == 2


def test_product_context_pack_uses_identity_mapping_for_structured_facts(mapping_db):
    from app.models.kb_tables import ProductIdentityMapping
    from app.services.product_context_pack_service import build_product_context_pack

    item_hash = hash_sensitive("456789012345")
    db = mapping_db()
    try:
        product = _add_product(db)
        db.add(ProductIdentityMapping(
            mapping_uid="pim_pack",
            platform_item_id_hash=item_hash,
            kb_product_id=product.id,
            i_id=product.i_id,
            sku_code="YH92K01B01S01",
            status="active",
            confidence=1.0,
        ))
        db.commit()
    finally:
        db.close()

    pack = build_product_context_pack(
        {
            "copilot_context": {
                "real_context": {
                    "product": {
                        "item_id_hash": item_hash,
                        "product_url": "https://item.taobao.com/item.htm?id=[LONG_ID_REDACTED:abc123]",
                    }
                }
            },
        },
        query="what material",
        allowed_source_types=["product_facts"],
        query_fact_type="material",
    )

    product_first = pack["product_first_evidence_pack"]
    assert product_first["resolved_product_identity"]["i_id"] == "YH92K01"
    assert product_first["product_structured_facts"]
    assert product_first["product_structured_facts"][0]["fact_type"] == "material"
