from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


@pytest.fixture()
def coverage_db(monkeypatch):
    import app.db as db_module
    from app.models.eval_tables import (
        AIProvisionalKnowledge,
        EvalRun,
        EvalTrace,
        KnowledgeGapDraft,
        KnowledgeGapTask,
    )

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db_module, "engine", engine)
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    db_module.Base.metadata.create_all(
        bind=engine,
        tables=[
            AIProvisionalKnowledge.__table__,
            EvalRun.__table__,
            EvalTrace.__table__,
            KnowledgeGapDraft.__table__,
            KnowledgeGapTask.__table__,
        ],
    )
    return session_factory


def _seed_coverage_rows(db):
    from app.models.eval_tables import AIProvisionalKnowledge, EvalRun, EvalTrace, KnowledgeGapTask

    run_uid = "run_cov_1"
    db.add(EvalRun(
        run_uid=run_uid,
        source_type="real_conversation",
        status="completed",
        total_cases=1,
        total_turns=1,
    ))

    task = KnowledgeGapTask(
        task_uid="kgap_cov_1",
        gap_type="product_field_gap",
        product_title="测试柜",
        item_id="",
        sku_code="",
        query_fact_type="material",
        failure_type="rag_miss",
        suggested_fix_area="knowledge_ops",
        suggested_owner="knowledge_ops",
        missing_evidence_type="product_material",
        status="open",
        sample_count=1,
        risk_level="medium",
    )
    task.set_metadata({
        "source_run_uid": run_uid,
        "gap_category": "product_field_gap",
        "required_evidence_type": "product_material",
        "target_system": "kb_product_field",
        "current_blocker": "missing_structured_material",
    })
    task.set_related_turn_uids(["turn_cov_1"])
    task.set_latest_buyer_questions(["手机号 13812345678，订单 123456789012345，材质是什么？"])
    task.set_latest_agent_replies(["需要核对资料。"])
    db.add(task)

    draft = AIProvisionalKnowledge(
        draft_uid="aipk_keep_internal_uid",
        source_run_uid=run_uid,
        case_uid="case_cov_1",
        turn_uid="turn_cov_1",
        task_uid=task.task_uid,
        query_fact_type="material",
        field_name="material",
        provisional_answer="手机号 13812345678，订单 123456789012345，建议人工填写材质。",
        confidence="low",
        verification_status="pending_review",
        usable_for_eval=False,
        usable_for_auto_send=False,
    )
    draft.set_metadata({
        "identity_status": "unresolved",
        "usable_for_eval_blocked_reason": "missing_identity",
    })
    db.add(draft)

    trace = EvalTrace(
        run_uid=run_uid,
        case_uid="case_cov_1",
        turn_uid="turn_cov_1",
        turn_index=1,
        buyer_message=(
            "手机号 13812345678 订单 123456789012345 "
            "https://oss.example.com/a.jpg?Signature=secret data:image/png;base64,"
            + "A" * 180
        ),
        agent_reply="需要人工核对。",
        query_fact_type="material",
        passed=False,
    )
    trace.set_raw_response({
        "real_context_product_identity": {
            "item_id_hash": "hash_platform_item",
            "product_url": "https://oss.example.com/a.jpg?Signature=secret",
            "product_title": "平台标题仅供人工参考",
        },
        "evidence_debug": {
            "product_context_pack_summary": {
                "evidence_pack": {
                    "product_identity_resolution": {
                        "status": "not_found",
                        "unresolved_reason": "no_matching_product_found",
                        "ambiguous_candidates": [],
                    }
                }
            }
        },
    })
    db.add(trace)
    db.commit()
    return run_uid


def _workbook_text(workbook) -> str:
    parts = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            parts.extend(str(value or "") for value in row)
    return "\n".join(parts)


def test_coverage_gap_export_has_chinese_sheets_and_sanitized_rows(coverage_db, tmp_path):
    from scripts.export_ai_provisional_coverage_gaps import run_export

    db = coverage_db()
    try:
        run_uid = _seed_coverage_rows(db)
    finally:
        db.close()

    output = tmp_path / "coverage_gaps.xlsx"
    result = run_export(run_uid=run_uid, output=str(output), db_factory=coverage_db)

    assert result["ok"] is True
    assert result["draft_identity_gap_count"] == 1
    assert result["mapping_gap_count"] == 1
    assert result["knowledge_gap_count"] == 1

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["说明", "草稿身份缺口", "平台商品映射缺口", "知识内容缺口"]
    assert [cell.value for cell in workbook["草稿身份缺口"][1]][:6] == [
        "草稿 UID",
        "关联任务 UID",
        "来源批次",
        "案例 ID",
        "轮次 ID",
        "问题类型",
    ]
    assert [cell.value for cell in workbook["平台商品映射缺口"][1]][:6] == [
        "回放批次",
        "案例 ID",
        "轮次 ID",
        "平台商品 ID",
        "平台商品 ID Hash",
        "商品链接",
    ]

    text = _workbook_text(workbook)
    assert "aipk_keep_internal_uid" in text
    assert "kgap_cov_1" in text
    assert "hash_platform_item" in text
    assert "missing_identity" in text
    assert "否" in text
    assert "13812345678" not in text
    assert "123456789012345" not in text
    assert "Signature=secret" not in text
    assert "data:image/png;base64" not in text
