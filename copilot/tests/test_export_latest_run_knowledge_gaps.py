import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models.eval_tables import EvalRun, EvalTrace, KnowledgeGapTask
from scripts.export_latest_run_knowledge_gaps import (
    SHEET_ALL,
    SHEET_HIGH_RISK,
    SHEET_MEDIA,
    SHEET_OVERVIEW,
    SHEET_POLICY,
    SHEET_PRODUCT,
    SHEET_ROUTING,
    TASK_HEADERS,
    export_latest_run_knowledge_gaps,
)


def _session_factory():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    return sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)


def _add_run_trace(db, run_uid: str, turn_uid: str) -> None:
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed"))
    db.add(EvalTrace(
        run_uid=run_uid,
        case_uid=f"case_{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message="buyer",
        agent_reply="agent",
        passed=False,
    ))


def _add_task(
    db,
    *,
    task_uid: str,
    run_uid: str,
    turn_uid: str,
    gap_type: str,
    evidence_type: str,
    target_system: str,
    owner: str,
    query_fact_type: str,
) -> None:
    task = KnowledgeGapTask(
        task_uid=task_uid,
        gap_type=gap_type,
        product_title="\u513f\u7ae5\u6536\u7eb3\u67dc",
        item_id="ITEM-1234567890",
        sku_code="SKU-TEST",
        query_fact_type=query_fact_type,
        failure_type="rag_miss",
        suggested_fix_area="knowledge_rag",
        suggested_owner=owner,
        missing_evidence_type=evidence_type,
        risk_level="medium",
        sample_count=2,
        priority="high",
        status="open",
        summary="missing evidence",
    )
    task.set_related_turn_uids([turn_uid])
    task.set_latest_buyer_questions([
        "\u4e70\u5bb6\u624b\u673a13812345678\uff0c\u8ba2\u5355123456789012345\uff0c\u95ee\u8d44\u6599",
    ])
    task.set_latest_agent_replies([
        "agent https://demo.oss-cn/a.jpg?Signature=secret&Expires=999",
    ])
    task.set_metadata({
        "source_run_uid": run_uid,
        "gap_category": gap_type,
        "required_evidence_type": evidence_type,
        "target_system": target_system,
        "recommended_action": "fill_product_field",
        "current_blocker": "\u5f53\u524d\u6ca1\u6709\u5df2\u5ba1\u6838\u8bc1\u636e",
    })
    db.add(task)


def _seed_tasks(session_factory):
    db = session_factory()
    try:
        _add_run_trace(db, "run_export_a", "turn_media")
        _add_run_trace(db, "run_export_b", "turn_other")
        _add_task(
            db,
            task_uid="kgap_media",
            run_uid="run_export_a",
            turn_uid="turn_media",
            gap_type="media_asset_gap",
            evidence_type="installation_video",
            target_system="kb_media_asset",
            owner="media_ops",
            query_fact_type="installation",
        )
        _add_task(
            db,
            task_uid="kgap_product",
            run_uid="run_export_a",
            turn_uid="turn_media",
            gap_type="product_field_gap",
            evidence_type="product_material",
            target_system="product_profile",
            owner="knowledge_ops",
            query_fact_type="material",
        )
        _add_task(
            db,
            task_uid="kgap_policy",
            run_uid="run_export_a",
            turn_uid="turn_media",
            gap_type="promotion_policy_gap",
            evidence_type="promotion_rule",
            target_system="activity_rules",
            owner="customer_service_lead",
            query_fact_type="promotion_policy",
        )
        _add_task(
            db,
            task_uid="kgap_other_run",
            run_uid="run_export_b",
            turn_uid="turn_other",
            gap_type="product_field_gap",
            evidence_type="load_capacity",
            target_system="product_profile",
            owner="knowledge_ops",
            query_fact_type="load_capacity",
        )
        db.commit()
    finally:
        db.close()


def test_latest_run_knowledge_gap_export_script_help_runs_from_project_root():
    project_root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [sys.executable, "scripts/export_latest_run_knowledge_gaps.py", "--help"],
        cwd=project_root,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0
    assert "--run-uid" in result.stdout
    assert "--gap-category" in result.stdout


def test_latest_run_knowledge_gap_export_filters_run_and_writes_sheets(monkeypatch, tmp_path):
    session_factory = _session_factory()
    monkeypatch.setattr("scripts.export_latest_run_knowledge_gaps.SessionLocal", session_factory)
    monkeypatch.setattr("scripts.export_latest_run_knowledge_gaps.init_db", lambda: None)
    _seed_tasks(session_factory)

    output = tmp_path / "gap_export.xlsx"
    result = export_latest_run_knowledge_gaps(run_uid="run_export_a", output=str(output))

    assert result["count"] == 3
    assert result["by_gap_category"] == {
        "media_asset_gap": 1,
        "product_field_gap": 1,
        "promotion_policy_gap": 1,
    }
    workbook = load_workbook(output)
    assert {
        SHEET_OVERVIEW,
        SHEET_ALL,
        SHEET_MEDIA,
        SHEET_POLICY,
        SHEET_PRODUCT,
        SHEET_ROUTING,
        SHEET_HIGH_RISK,
    }.issubset(set(workbook.sheetnames))
    overview_values = [cell.value for row in workbook[SHEET_OVERVIEW].iter_rows() for cell in row]
    assert "\u7f3a\u53e3\u4efb\u52a1\u6570" in overview_values
    all_sheet = workbook[SHEET_ALL]
    assert [cell.value for cell in all_sheet[1]] == TASK_HEADERS
    assert "\u9700\u8981\u8865\u4ec0\u4e48" in TASK_HEADERS
    assert "\u662f\u5426\u53ef\u81ea\u52a8\u53d1\u9001" in TASK_HEADERS
    assert all_sheet.max_row == 4
    media_sheet = workbook[SHEET_MEDIA]
    assert media_sheet.max_row == 2
    assert media_sheet["B2"].value == "\u7d20\u6750\u7f3a\u53e3"
    policy_sheet = workbook[SHEET_POLICY]
    assert policy_sheet.max_row == 2
    product_sheet = workbook[SHEET_PRODUCT]
    assert product_sheet.max_row == 2
    routing_sheet = workbook[SHEET_ROUTING]
    assert routing_sheet.max_row == 1
    high_risk_sheet = workbook[SHEET_HIGH_RISK]
    assert high_risk_sheet.max_row == 2


def test_latest_run_knowledge_gap_export_is_sanitized(monkeypatch, tmp_path):
    session_factory = _session_factory()
    monkeypatch.setattr("scripts.export_latest_run_knowledge_gaps.SessionLocal", session_factory)
    monkeypatch.setattr("scripts.export_latest_run_knowledge_gaps.init_db", lambda: None)
    _seed_tasks(session_factory)

    output = tmp_path / "gap_export.xlsx"
    export_latest_run_knowledge_gaps(run_uid="run_export_a", output=str(output))

    workbook = load_workbook(output)
    values = "\n".join(str(cell.value or "") for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    assert "13812345678" not in values
    assert "123456789012345" not in values
    assert "Signature=secret" not in values
    assert "ITEM-1234567890" not in values
    assert "\u5f53\u524d\u6ca1\u6709\u5df2\u5ba1\u6838\u8bc1\u636e" in values
