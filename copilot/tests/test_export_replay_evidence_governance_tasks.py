import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from scripts.export_replay_evidence_governance_tasks import export_governance_tasks


def _write_diagnosis(path: Path) -> None:
    payload = {
        "run_uid": "run_governance_test",
        "summary": {"run_uid": "run_governance_test", "total_traces": 7},
        "zero_evidence_records": [
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_1",
                "turn_uid": "turn_1",
                "buyer_message_preview": "尺寸是多少",
                "query_fact_type": "dimensions",
                "selected_evidence_count": 0,
                "primary_reason": "structured_field_missing",
                "secondary_reasons": [],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
                "missing_structured_fields": ["dimensions"],
            },
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_2",
                "turn_uid": "turn_2",
                "buyer_message_preview": "能不能放下",
                "query_fact_type": "space_fit",
                "selected_evidence_count": 0,
                "primary_reason": "structured_field_missing",
                "secondary_reasons": [],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
                "missing_structured_fields": ["dimensions"],
            },
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_3",
                "turn_uid": "turn_3",
                "buyer_message_preview": "有安装视频吗",
                "query_fact_type": "installation",
                "selected_evidence_count": 0,
                "primary_reason": "media_missing",
                "secondary_reasons": [],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
            },
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_4",
                "turn_uid": "turn_4",
                "buyer_message_preview": "这个商品图能当安装图吗",
                "query_fact_type": "installation",
                "selected_evidence_count": 0,
                "primary_reason": "evidence_role_mismatch",
                "secondary_reasons": ["embedding_not_configured"],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
            },
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_5",
                "turn_uid": "turn_5",
                "buyer_message_preview": "买两个有优惠吗",
                "query_fact_type": "promotion_policy",
                "selected_evidence_count": 0,
                "primary_reason": "generic_rule_missing",
                "secondary_reasons": [],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
            },
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_6",
                "turn_uid": "turn_6",
                "buyer_message_preview": "宝宝能用吗",
                "query_fact_type": "age_range",
                "selected_evidence_count": 0,
                "primary_reason": "structured_field_missing",
                "secondary_reasons": [],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
                "missing_structured_fields": ["age_range"],
            },
            {
                "run_uid": "run_governance_test",
                "case_uid": "case_7",
                "turn_uid": "turn_7",
                "buyer_message_preview": "查不到证据",
                "query_fact_type": "unknown",
                "selected_evidence_count": 0,
                "primary_reason": "embedding_not_configured",
                "secondary_reasons": [],
                "has_sidecar_product_context": True,
                "kb_product_found": True,
                "kb_product_i_id": "I001",
                "kb_product_name": "测试商品",
            },
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def test_export_replay_evidence_governance_tasks_writes_chinese_workbook(tmp_path):
    diagnosis = tmp_path / "diagnosis.json"
    output_json = tmp_path / "governance.json"
    output_xlsx = tmp_path / "governance.xlsx"
    _write_diagnosis(diagnosis)

    report = export_governance_tasks(
        diagnosis_json=str(diagnosis),
        json_output=str(output_json),
        excel_output=str(output_xlsx),
    )

    assert report["summary"]["governance_task_count"] >= 6
    assert report["summary"]["auto_backfill_candidate_count"] == 2
    assert report["summary"]["manual_review_required_count"] == report["summary"]["governance_task_count"]
    assert report["summary"]["by_gap_category"]["商品字段缺口"] == 3
    assert report["summary"]["by_gap_category"]["素材缺口"] == 1
    assert report["summary"]["by_gap_category"]["素材Role待确认"] == 1
    assert report["summary"]["by_gap_category"]["活动优惠规则缺口"] == 1

    workbook = load_workbook(output_xlsx)
    assert {
        "说明",
        "总览",
        "P0优先处理",
        "商品字段缺口",
        "素材缺口",
        "素材Role待确认",
        "活动优惠规则缺口",
        "Embedding和RAG配置问题",
        "可自动回填候选",
        "必须人工确认",
    }.issubset(set(workbook.sheetnames))

    headers = [cell.value for cell in workbook["商品字段缺口"][1]]
    assert headers[:5] == ["优先级", "缺口类型", "问题类型 query_fact_type", "问题类型中文名", "买家问题"]
    assert "是否允许自动发送" in headers
    assert "建议负责人" in headers

    auto_rows = list(workbook["可自动回填候选"].iter_rows(min_row=2, values_only=True))
    assert len(auto_rows) == 2
    assert {row[1] for row in auto_rows} == {"商品字段缺口"}
    assert all(row[13] == "是" for row in auto_rows)
    assert all(row[17] == "否" for row in auto_rows)

    manual_rows = list(workbook["必须人工确认"].iter_rows(min_row=2, values_only=True))
    assert manual_rows
    assert all(row[14] == "是" for row in manual_rows)
    assert all(row[17] == "否" for row in manual_rows)


def test_export_replay_evidence_governance_tasks_help_runs_from_project_root():
    result = subprocess.run(
        [sys.executable, "scripts/export_replay_evidence_governance_tasks.py", "--help"],
        cwd=Path(__file__).resolve().parents[1],
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 0
    assert "--diagnosis-json" in result.stdout
