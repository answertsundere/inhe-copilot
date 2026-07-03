from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models.eval_tables import EvalRun, EvalTrace


def _session_factory():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    return sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)


def _add_run_with_traces(session_factory):
    db = session_factory()
    try:
        run = EvalRun(run_uid="natural_run", source_type="real_conversation", status="completed")
        db.add(run)

        redline = EvalTrace(
            run_uid="natural_run",
            case_uid="case_1",
            turn_uid="turn_redline",
            turn_index=1,
            buyer_message="材质安全吗",
            agent_reply="我这边不直接承诺无毒，当前知识库缺少证据，需要人工审核。",
            query_fact_type="material_safety",
            requires_human_review=True,
            passed=False,
        )
        db.add(redline)

        natural = EvalTrace(
            run_uid="natural_run",
            case_uid="case_2",
            turn_uid="turn_natural",
            turn_index=2,
            buyer_message="有优惠吗",
            agent_reply="亲，我帮您看下当前页面活动和优惠券，最终以您下单页显示为准；如果页面没显示，我再帮您核对。",
            query_fact_type="promotion_policy",
            requires_human_review=True,
            passed=False,
        )
        db.add(natural)

        contract = EvalTrace(
            run_uid="natural_run",
            case_uid="case_3",
            turn_uid="turn_contract",
            turn_index=3,
            buyer_message="适合两岁宝宝吗",
            agent_reply="亲，我帮您核对。",
            query_fact_type="age_range",
            requires_human_review=True,
            passed=False,
        )
        contract.set_raw_response({"can_send": True, "sendable_reply": "亲，我帮您核对。"})
        db.add(contract)
        db.commit()
    finally:
        db.close()


def test_diagnose_naturalness_detects_internal_language_and_contract_risk(monkeypatch):
    import scripts.diagnose_customer_reply_naturalness as diagnose

    session_factory = _session_factory()
    _add_run_with_traces(session_factory)
    monkeypatch.setattr(diagnose, "SessionLocal", session_factory)

    result = diagnose.diagnose_naturalness(run_uid="natural_run")

    assert result["ok"] is True
    assert result["summary"]["run_uid"] == "natural_run"
    assert result["summary"]["scanned_reply_count"] == 3
    assert result["summary"]["issue_counts"]["internal_policy_phrase"] >= 1
    assert result["summary"]["issue_counts"]["can_send_contract_risk"] == 1
    issue_turns = {item["turn_uid"]: item for item in result["items"]}
    assert "turn_redline" in issue_turns
    assert "turn_contract" in issue_turns
    assert "turn_natural" not in issue_turns


def test_diagnose_naturalness_excel_has_chinese_headers(monkeypatch, tmp_path):
    import scripts.diagnose_customer_reply_naturalness as diagnose

    session_factory = _session_factory()
    _add_run_with_traces(session_factory)
    monkeypatch.setattr(diagnose, "SessionLocal", session_factory)
    result = diagnose.diagnose_naturalness(run_uid="natural_run")
    output = tmp_path / "naturalness.xlsx"

    diagnose.write_excel(result, str(output))

    workbook = load_workbook(output)
    sheet = workbook["客服话术自然度诊断"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:5] == ["回放批次", "案例 ID", "轮次 ID", "买家问题", "问题类型"]
    assert "检测问题" in headers
