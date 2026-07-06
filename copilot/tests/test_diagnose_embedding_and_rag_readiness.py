import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models.eval_tables import EvalRun, EvalTrace
from scripts.diagnose_embedding_and_rag_readiness import (
    build_embedding_rag_readiness_report,
    main,
    write_excel,
    write_json,
)


def _session_factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'embedding_diag.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    return sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)


def _trace(run_uid: str, turn_uid: str, *, selected=None, raw=None):
    trace = EvalTrace(
        run_uid=run_uid,
        case_uid=f"case-{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message="有安装视频吗",
        query_fact_type="installation",
        passed=False,
    )
    trace.set_selected_evidence(selected or [])
    trace.set_raw_response(raw or {})
    trace.set_failure_labels(["rag_miss"])
    return trace


def test_embedding_rag_readiness_masks_secrets_and_writes_chinese_workbook(tmp_path, monkeypatch):
    monkeypatch.setenv("COPILOT_EMBEDDING_ENABLED", "true")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_BASE", "https://embedding.example.invalid/v1")
    monkeypatch.setenv("COPILOT_EMBEDDING_API_KEY", "secret-key-should-not-leak")
    monkeypatch.setenv("COPILOT_EMBEDDING_MODEL", "test-embedding")
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-embedding"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=2))
    db.add(_trace(run_uid, "turn-zero"))
    db.add(_trace(run_uid, "turn-timeout", selected=[{"id": "e1"}], raw={"rag_timeout": True}))
    db.commit()
    db.close()

    report = build_embedding_rag_readiness_report(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["total_traces"] == 2
    assert report["summary"]["selected_evidence_count_distribution"]["0"] == 1
    assert report["summary"]["rag_timeout_count"] == 1
    api_key_row = next(row for row in report["environment_variables"] if row["env_key"] == "COPILOT_EMBEDDING_API_KEY")
    assert api_key_row["value"] == "已配置（已隐藏）"
    assert "secret-key-should-not-leak" not in json.dumps(report, ensure_ascii=False)

    xlsx_path = tmp_path / "embedding.xlsx"
    json_path = tmp_path / "embedding.json"
    write_excel(str(xlsx_path), report)
    write_json(str(json_path), report)
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ["说明", "总览", "环境变量状态", "RAG超时统计", "最新Replay证据命中", "建议配置项"]
    assert workbook["环境变量状态"]["A1"].value == "环境变量"
    assert json.loads(Path(json_path).read_text(encoding="utf-8"))["run_uid"] == run_uid


def test_embedding_rag_readiness_counts_missing_config_as_zero_evidence(tmp_path, monkeypatch):
    for key in ["COPILOT_EMBEDDING_ENABLED", "COPILOT_EMBEDDING_API_BASE", "COPILOT_EMBEDDING_API_KEY", "COPILOT_EMBEDDING_MODEL"]:
        monkeypatch.delenv(key, raising=False)
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-missing-embedding"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed"))
    db.add(_trace(run_uid, "turn-zero-a"))
    db.add(_trace(run_uid, "turn-zero-b"))
    db.add(_trace(run_uid, "turn-selected", selected=[{"id": "e1"}]))
    db.commit()
    db.close()

    report = build_embedding_rag_readiness_report(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["embedding_enabled"] is False
    assert report["summary"]["embedding_not_configured_count"] == 2
    assert report["summary"]["fallback_retrieval_likely"] is True



def test_embedding_rag_readiness_json_only_does_not_write_excel(tmp_path, monkeypatch):
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-json-only"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed"))
    db.commit()
    db.close()

    import scripts.diagnose_embedding_and_rag_readiness as module

    monkeypatch.setattr(module, "SessionLocal", Session)
    json_path = tmp_path / "readiness.json"
    default_excel = tmp_path / "should-not-exist.xlsx"
    monkeypatch.setattr(module, "default_excel_path", lambda: str(default_excel))
    monkeypatch.setattr(
        "sys.argv",
        [
            "diagnose_embedding_and_rag_readiness.py",
            "--run-uid",
            run_uid,
            "--json-output",
            str(json_path),
        ],
    )

    assert main() == 0
    assert json_path.exists()
    assert not default_excel.exists()
