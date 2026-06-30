import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models.eval_tables import (
    EvalFailure,
    EvalReview,
    EvalRun,
    EvalTrace,
    KnowledgeGapTask,
)
from app.services.knowledge_gap_task_service import KnowledgeGapTaskService
from scripts.export_knowledge_gap_tasks import HEADERS, export_knowledge_gap_tasks


def _session_factory():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    return sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)


def _add_trace(db, turn_uid, query_fact_type, *, product_title="收纳柜", sku_code="", item_id=""):
    trace = EvalTrace(
        run_uid="kgap_run_1",
        case_uid=f"case_{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message=f"买家手机号 13812345678 询问 {query_fact_type}",
        reference_human_reply="原客服参考，订单号 123456789012345",
        agent_reply="Agent 回复 https://demo.oss-cn/a.jpg?Signature=secret&Expires=999",
        query_fact_type=query_fact_type,
        passed=False,
        requires_human_review=True,
    )
    trace.set_product_identity({
        "display_product_name": product_title,
        "sku_code": sku_code,
        "item_id": item_id,
    })
    trace.set_selected_evidence([])
    trace.set_rejected_evidence([])
    db.add(trace)
    return trace


def _seed_gap_run(session_factory):
    db = session_factory()
    try:
        db.add(EvalRun(run_uid="kgap_run_1", source_type="real_conversation", status="completed"))
        cases = [
            ("turn_fact", "material", "rag_miss", "knowledge_rag", "knowledge_ops"),
            ("turn_media", "installation", "rag_miss", "knowledge_rag", "knowledge_ops"),
            ("turn_promo", "promotion_policy", "rag_miss", "knowledge_rag", "knowledge_ops"),
            ("turn_service", "aftersales_policy", "rag_miss", "knowledge_rag", "knowledge_ops"),
            ("turn_shipping", "stock_shipping", "rag_miss", "knowledge_rag", "knowledge_ops"),
            ("turn_context", "dimensions", "context_gap", "sample_context_extraction", "data_pipeline"),
            ("turn_correct", "dimensions", "rag_miss", "knowledge_rag", "knowledge_ops"),
        ]
        for turn_uid, fact_type, failure_type, fix_area, owner in cases:
            if turn_uid == "turn_context":
                _add_trace(db, turn_uid, fact_type, product_title="", sku_code="", item_id="")
            else:
                _add_trace(db, turn_uid, fact_type, sku_code="SKU-1", item_id="ITEM-123456")
            db.add(EvalFailure(
                run_uid="kgap_run_1",
                case_uid=f"case_{turn_uid}",
                turn_uid=turn_uid,
                failure_type=failure_type,
                severity="high" if turn_uid == "turn_media" else "medium",
                suggested_fix_area=fix_area,
                suggested_owner=owner,
                message="missing evidence for 13812345678",
            ))
        db.add(EvalReview(
            run_uid="kgap_run_1",
            case_uid="case_turn_correct",
            turn_uid="turn_correct",
            decision="correct",
            reviewer="qa",
        ))
        db.commit()
    finally:
        db.close()


def test_knowledge_gap_generation_classifies_operational_gap_categories_and_skips_correct_review():
    session_factory = _session_factory()
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        result = KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")

        assert result.generated == 6
        assert result.skipped_correct == 1
        tasks = {task.gap_type: task for task in db.query(KnowledgeGapTask).all()}
        assert tasks["product_field_gap"].missing_evidence_type == "product_material"
        assert tasks["media_asset_gap"].missing_evidence_type == "installation_video"
        assert tasks["promotion_policy_gap"].missing_evidence_type == "promotion_rule"
        assert tasks["aftersales_policy_gap"].missing_evidence_type == "aftersales_rule"
        context_required = {
            task.missing_evidence_type
            for task in db.query(KnowledgeGapTask).all()
            if task.gap_type == "context_extraction_gap"
        }
        assert {"sku_context", "order_context"}.issubset(context_required)
        assert tasks["media_asset_gap"].get_metadata()["target_system"] == "kb_media_asset"
        assert tasks["media_asset_gap"].get_metadata()["current_blocker"]
        assert len(tasks["media_asset_gap"].get_metadata()["representative_samples"]) <= 3
        assert tasks["context_extraction_gap"].get_metadata()["recommended_action"] == "fix_context_extraction"
        assert not any("turn_correct" in task.get_related_turn_uids() for task in db.query(KnowledgeGapTask).all())
    finally:
        db.close()


def test_knowledge_gap_generation_updates_existing_task_instead_of_duplicating():
    session_factory = _session_factory()
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        first = KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")
        second = KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")

        assert first.generated == 6
        assert second.generated == 0
        assert second.updated == 6
        assert db.query(KnowledgeGapTask).count() == 6
    finally:
        db.close()


def test_knowledge_gap_list_filters_new_operational_fields():
    session_factory = _session_factory()
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")
        media = KnowledgeGapTaskService().list_tasks(db, filters={"gap_category": "media_asset_gap"})
        product = KnowledgeGapTaskService().list_tasks(db, filters={"required_evidence_type": "product_material"})
        target = KnowledgeGapTaskService().list_tasks(db, filters={"target_system": "context_extractor"})

        assert len(media["items"]) == 1
        assert media["items"][0]["required_evidence_type"] == "installation_video"
        assert len(product["items"]) == 1
        assert product["items"][0]["gap_category"] == "product_field_gap"
        assert len(target["items"]) == 2
        assert {item["recommended_action"] for item in target["items"]} == {"fix_context_extraction"}
    finally:
        db.close()


def test_knowledge_gap_list_filters_by_run_uid_and_falls_back_for_legacy_tasks():
    session_factory = _session_factory()
    db = session_factory()
    try:
        db.add(EvalRun(run_uid="run_a", source_type="real_conversation", status="completed"))
        db.add(EvalRun(run_uid="run_b", source_type="real_conversation", status="completed"))
        db.add(EvalRun(run_uid="run_legacy", source_type="real_conversation", status="completed"))
        for run_uid, turn_uid in [("run_a", "turn_a"), ("run_b", "turn_b"), ("run_legacy", "turn_legacy")]:
            trace = EvalTrace(
                run_uid=run_uid,
                case_uid=f"case_{turn_uid}",
                turn_uid=turn_uid,
                turn_index=1,
                buyer_message="buyer",
                agent_reply="agent",
                query_fact_type="material",
                passed=False,
            )
            db.add(trace)

        task_a = KnowledgeGapTask(
            task_uid="kgap_run_a",
            gap_type="product_field_gap",
            query_fact_type="material",
            missing_evidence_type="product_material",
            status="open",
            sample_count=1,
        )
        task_a.set_metadata({"source_run_uid": "run_a", "gap_category": "product_field_gap"})
        task_a.set_related_turn_uids(["turn_a"])
        task_b = KnowledgeGapTask(
            task_uid="kgap_run_b",
            gap_type="media_asset_gap",
            query_fact_type="installation",
            missing_evidence_type="installation_video",
            status="open",
            sample_count=1,
        )
        task_b.set_metadata({"source_run_uid": "run_b", "gap_category": "media_asset_gap"})
        task_b.set_related_turn_uids(["turn_b"])
        legacy_a = KnowledgeGapTask(
            task_uid="kgap_legacy_a",
            gap_type="context_extraction_gap",
            query_fact_type="",
            missing_evidence_type="sku_context",
            status="open",
            sample_count=1,
        )
        legacy_a.set_metadata({"gap_category": "context_extraction_gap"})
        legacy_a.set_related_turn_uids(["turn_a"])
        legacy_run = KnowledgeGapTask(
            task_uid="kgap_legacy_run",
            gap_type="context_extraction_gap",
            query_fact_type="",
            missing_evidence_type="sku_context",
            status="open",
            sample_count=1,
        )
        legacy_run.set_metadata({"gap_category": "context_extraction_gap"})
        legacy_run.set_related_turn_uids(["turn_legacy"])
        legacy_other = KnowledgeGapTask(
            task_uid="kgap_legacy_other",
            gap_type="promotion_policy_gap",
            query_fact_type="promotion",
            missing_evidence_type="promotion_rule",
            status="open",
            sample_count=1,
        )
        legacy_other.set_related_turn_uids(["turn_unknown"])
        db.add_all([task_a, task_b, legacy_a, legacy_run, legacy_other])
        db.commit()

        filtered = KnowledgeGapTaskService().list_tasks(db, filters={"run_uid": "run_a"})
        legacy_filtered = KnowledgeGapTaskService().list_tasks(db, filters={"run_uid": "run_legacy"})
        all_rows = KnowledgeGapTaskService().list_tasks(db, filters={})

        assert {item["task_uid"] for item in filtered["items"]} == {"kgap_run_a"}
        assert filtered["summary"]["run_uid"] == "run_a"
        assert filtered["summary"]["filtered_by_run_uid"] is True
        assert filtered["summary"]["total"] == 1
        assert filtered["summary"]["by_gap_category"] == {"product_field_gap": 1}
        assert {item["task_uid"] for item in legacy_filtered["items"]} == {"kgap_legacy_run"}
        assert legacy_filtered["summary"]["total"] == 1
        assert all_rows["summary"]["filtered_by_run_uid"] is False
        assert all_rows["summary"]["total"] == 5
    finally:
        db.close()


def test_knowledge_gap_triage_updates_review_metadata_and_status_history():
    session_factory = _session_factory()
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")
        task = db.query(KnowledgeGapTask).filter(KnowledgeGapTask.gap_type == "media_asset_gap").one()

        updated = KnowledgeGapTaskService().triage_task(
            db,
            task.task_uid,
            {
                "review_decision": "upload_media_asset",
                "assigned_team": "media_ops",
                "assigned_to": "media_lead",
                "priority": "high",
                "due_date": "2026-07-01",
                "review_note": "phone 13812345678 should be masked",
                "next_action": "collect approved installation media",
            },
            reviewer="supervisor_a",
        )

        assert updated["status"] == "triaged"
        assert updated["priority"] == "high"
        assert updated["suggested_owner"] == "media_lead"
        assert updated["review_decision"] == "upload_media_asset"
        assert updated["assigned_team"] == "media_ops"
        assert updated["assigned_to"] == "media_lead"
        assert "13812345678" not in updated["review_note"]
        assert updated["status_history"][-1]["status"] == "triaged"
        assert updated["status_history"][-1]["changed_by"] == "supervisor_a"
    finally:
        db.close()


def test_knowledge_gap_triage_rejects_illegal_review_decision():
    session_factory = _session_factory()
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")
        task = db.query(KnowledgeGapTask).first()

        with pytest.raises(ValueError):
            KnowledgeGapTaskService().triage_task(
                db,
                task.task_uid,
                {"review_decision": "publish_to_formal_kb"},
                reviewer="lead",
            )
    finally:
        db.close()


def test_knowledge_gap_status_update_records_history_and_resolved_pending_retest_is_not_verified():
    session_factory = _session_factory()
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")
        task = db.query(KnowledgeGapTask).first()

        updated = KnowledgeGapTaskService().update_status(
            db,
            task.task_uid,
            {
                "status": "resolved_pending_retest",
                "assigned_to": "knowledge_lead",
                "review_note": "staging evidence prepared",
            },
            changed_by="lead",
        )

        assert updated["status"] == "resolved_pending_retest"
        assert updated["status"] != "verified"
        assert updated["assigned_to"] == "knowledge_lead"
        assert updated["status_history"][-1]["status"] == "resolved_pending_retest"
        assert updated["status_history"][-1]["changed_by"] == "lead"
    finally:
        db.close()


def test_knowledge_gap_export_excel_is_sanitized(monkeypatch, tmp_path):
    import app.db as db_module

    session_factory = _session_factory()
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    monkeypatch.setattr("scripts.export_knowledge_gap_tasks.SessionLocal", session_factory)
    monkeypatch.setattr("scripts.export_knowledge_gap_tasks.init_db", lambda: None)
    _seed_gap_run(session_factory)
    db = session_factory()
    try:
        KnowledgeGapTaskService().generate_for_run(db, "kgap_run_1", created_by="lead")
    finally:
        db.close()

    output = tmp_path / "knowledge_gap.xlsx"
    result = export_knowledge_gap_tasks(output=str(output), status="open")

    assert result["count"] == 6
    workbook = load_workbook(output)
    sheet = workbook.active
    assert sheet.title == "知识缺口任务"
    assert [cell.value for cell in sheet[1]] == HEADERS
    values = "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "13812345678" not in values
    assert "123456789012345" not in values
    assert "Signature=secret" not in values
    assert "ITEM-123456" not in values
    assert "media_asset_gap" in values
    assert "installation_video" in values


def test_knowledge_gap_export_script_help_runs_from_project_root():
    project_root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [sys.executable, "scripts/export_knowledge_gap_tasks.py", "--help"],
        cwd=project_root,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0
    assert "--output" in result.stdout
    assert "--status" in result.stdout


def test_knowledge_gap_export_empty_template_has_chinese_headers(monkeypatch, tmp_path):
    import app.db as db_module

    session_factory = _session_factory()
    monkeypatch.setattr(db_module, "SessionLocal", session_factory)
    monkeypatch.setattr("scripts.export_knowledge_gap_tasks.SessionLocal", session_factory)
    monkeypatch.setattr("scripts.export_knowledge_gap_tasks.init_db", lambda: None)

    output = tmp_path / "empty_knowledge_gap.xlsx"
    result = export_knowledge_gap_tasks(output=str(output), status="__none__")

    workbook = load_workbook(output)
    sheet = workbook.active
    assert result["count"] == 0
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == HEADERS
