import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models.eval_tables import EvalRun, EvalTrace
from app.models.kb_tables import KBGenericServiceRule, KBMediaAsset, KBProduct
from scripts.diagnose_evidence_readiness import (
    build_evidence_readiness_report,
    write_excel,
    write_json,
)


def _session_factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'evidence_ready.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    return sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)


def _trace(run_uid: str, turn_uid: str, *, qft="installation"):
    trace = EvalTrace(
        run_uid=run_uid,
        case_uid=f"case-{turn_uid}",
        turn_uid=turn_uid,
        turn_index=1,
        buyer_message="有安装视频吗",
        query_fact_type=qft,
        passed=False,
    )
    trace.set_selected_evidence([])
    trace.set_failure_labels(["rag_miss"])
    trace.set_raw_response({"sidecar_context_quality": "complete"})
    trace.set_turn_understanding({"query_fact_type": qft, "effective_query_fact_type": qft, "sidecar_context_quality": "complete"})
    trace.set_answer_trace({"query_fact_type": qft})
    return trace


def test_evidence_readiness_reports_fields_media_rules_and_replay_gaps(tmp_path):
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-evidence-ready"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed", total_turns=1))
    product = KBProduct(i_id="IID-1", product_name="测试书架", status="active")
    product.set_specs({"material": "钢木", "dimensions": "100x40x80cm", "installation": "有说明书"})
    db.add(product)
    db.flush()
    approved = KBMediaAsset(
        product_id=product.id,
        i_id="IID-1",
        sku_code="SKU-1",
        product_name="测试书架",
        asset_type="install_image",
        asset_title="安装示意图",
        asset_url="https://example.invalid/install.jpg?token=secret",
        status="approved",
        audit_status="approved",
        usable_for_agent=1,
    )
    pending = KBMediaAsset(
        product_id=product.id,
        i_id="IID-1",
        sku_code="SKU-1",
        product_name="测试书架",
        asset_type="sku_image",
        asset_title="普通商品图",
        asset_url="https://example.invalid/sku.jpg",
        status="pending_review",
        usable_for_agent=0,
    )
    db.add_all([approved, pending])
    db.add(KBGenericServiceRule(rule_key="return-pickup", title="取件核对", fact_type="return_pickup", status="active"))
    db.add(_trace(run_uid, "turn-install"))
    db.commit()
    db.close()

    report = build_evidence_readiness_report(run_uid=run_uid, db_factory=Session)

    assert report["summary"]["product_count"] == 1
    assert report["summary"]["approved_media_count"] == 1
    assert report["summary"]["pending_media_count"] == 1
    assert report["summary"]["generic_rule_available_count"] == 1
    material = next(row for row in report["product_field_rows"] if row["field_name"] == "material")
    assert material["covered_product_count"] == 1
    install_media = next(row for row in report["media_rows"] if row["media_group"] == "installation_diagram")
    assert install_media["approved_count"] == 1
    assert any(row["asset_title"] == "普通商品图" for row in report["media_role_review_rows"])

    xlsx_path = tmp_path / "evidence.xlsx"
    json_path = tmp_path / "evidence.json"
    write_excel(str(xlsx_path), report)
    write_json(str(json_path), report)
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == [
        "说明",
        "总览",
        "商品字段覆盖",
        "素材覆盖",
        "素材Role待确认",
        "通用规则覆盖",
        "Replay缺失证据",
        "P0优先补齐",
        "必须人工确认",
    ]
    assert workbook["商品字段覆盖"]["A1"].value == "字段"
    assert json.loads(Path(json_path).read_text(encoding="utf-8"))["run_uid"] == run_uid


def test_evidence_readiness_marks_missing_generic_rule(tmp_path):
    Session = _session_factory(tmp_path)
    db = Session()
    run_uid = "run-no-rules"
    db.add(EvalRun(run_uid=run_uid, source_type="real_conversation", status="completed"))
    db.add(_trace(run_uid, "turn-promo", qft="promotion_policy"))
    db.commit()
    db.close()

    report = build_evidence_readiness_report(run_uid=run_uid, db_factory=Session)

    promo = next(row for row in report["generic_rule_rows"] if row["query_fact_type"] == "promotion_policy")
    assert promo["coverage_status"] == "缺规则"
