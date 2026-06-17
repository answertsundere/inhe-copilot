"""Fix QA-to-product linking using Excel's 关联商品 column with SKU-based matching."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from app.db import SessionLocal
from app.models.kb_tables import KBQA, KBProduct
from sqlalchemy import func

EXCEL_PATH = r"D:\桌面文件\客服\INHE全部商品_金牌客服问答手册_三级分类.xlsx"

def main():
    db = SessionLocal()

    # Build product name -> id lookup (FIRST match wins, not last)
    products = db.query(KBProduct).order_by(KBProduct.id).all()
    name_to_id = {}
    sku_to_id = {}
    for p in products:
        if p.product_name.strip() not in name_to_id:
            name_to_id[p.product_name.strip()] = p.id
        # Also index by i_id
        sku_to_id[p.i_id.strip()] = p.id
        # Index by SKU codes
        for sku in p.get_sku_list():
            if isinstance(sku, dict) and sku.get('sku_code'):
                sku_to_id[sku['sku_code'].strip()] = p.id

    print(f"Products: {len(products)}, name lookups: {len(name_to_id)}, sku lookups: {len(sku_to_id)}")

    # Read Excel and collect question -> product_id mapping
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    qa_sheets = [
        "金牌客服问答（增强版）", "物流发货常见问答（增强版）",
        "安装指导常见问答（增强版）", "售后退换货常见问答（增强版）",
        "客诉处理常见问答（增强版）", "发票与价保常见问答（增强版）",
        "推销与关联推荐话术（增强版）", "平台规则速查（增强版）",
        "安全与质保常见问答（增强版）", "特殊场景处理指南（增强版）",
        "客服用语规范速查（增强版）",
    ]

    excel_map = {}  # question -> product_id
    for sheet_name in qa_sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[6]:
                continue
            question = str(row[6]).strip()
            product_name = str(row[4]).strip() if row[4] else ""
            sku_raw = str(row[5]).strip() if row[5] else ""
            if not question:
                continue

            product_id = None
            # Try SKU match first (most precise)
            if sku_raw:
                for sku in sku_raw.replace("，", ",").split(","):
                    sku = sku.strip()
                    if sku in sku_to_id:
                        product_id = sku_to_id[sku]
                        break
            # Fallback to name match
            if not product_id and product_name and product_name in name_to_id:
                product_id = name_to_id[product_name]

            if product_id:
                excel_map[question] = product_id

    wb.close()
    print(f"Excel QA-product mappings: {len(excel_map)}")

    # Update QA entries
    fixed = 0
    already_ok = 0
    not_in_excel = 0

    all_qa = db.query(KBQA).all()
    for qa in all_qa:
        correct_pid = excel_map.get(qa.question)
        if correct_pid is None:
            not_in_excel += 1
            continue
        if qa.product_id != correct_pid:
            qa.product_id = correct_pid
            fixed += 1
        else:
            already_ok += 1

    db.commit()

    print(f"Fixed: {fixed}")
    print(f"Already OK: {already_ok}")
    print(f"Not in Excel: {not_in_excel}")

    # Verify
    with_pid = db.query(func.count(KBQA.id)).filter(KBQA.product_id.isnot(None)).scalar()
    total = db.query(func.count(KBQA.id)).scalar()
    print(f"\nAfter fix: {with_pid}/{total} QA have product_id")

    for name in ["三层火箭书架", "三层鳄鱼注塑书架", "四层火箭书架"]:
        p = db.query(KBProduct).filter(KBProduct.product_name == name).first()
        if p:
            cnt = db.query(func.count(KBQA.id)).filter(KBQA.product_id == p.id).scalar()
            print(f"  {name} (id={p.id}): {cnt} QA")

    db.close()

if __name__ == "__main__":
    main()
