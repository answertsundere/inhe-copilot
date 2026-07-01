"""Import reviewed Agent benchmark expected replies from Excel."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.db import init_db
from app.models.eval_tables import AgentBenchmarkScenario  # noqa: F401 - register table before init_db
from app.services.agent_benchmark_dataset_service import AgentBenchmarkDatasetService, expected_reply_quality
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text


HEADER_ALIASES = {
    "scenario_uid": "scenario_uid",
    "expected_reply": "expected_reply",
    "当前标准答案": "expected_reply",
    "当前expected_reply": "expected_reply",
    "建议标准答案草稿": "suggested_expected_reply",
    "关键点": "key_points",
    "禁止话术": "forbidden_claims",
    "必须转人工": "must_handoff",
    "允许自动发送": "auto_send_allowed",
    "审核人": "reviewer",
    "审核备注": "review_note",
}

POSITIONAL_HEADERS = {
    0: "scenario_uid",
    1: "expected_reply",
    2: "key_points",
    3: "forbidden_claims",
    4: "must_handoff",
    5: "auto_send_allowed",
    6: "reviewer",
    7: "review_note",
}


def _header_name(cell: Any, idx: int) -> str:
    text = sanitize_text(cell)
    if text in HEADER_ALIASES:
        return HEADER_ALIASES[text]
    lowered = text.lower()
    if "scenario_uid" in lowered:
        return "scenario_uid"
    if "expected_reply" in lowered:
        return "expected_reply"
    return POSITIONAL_HEADERS.get(idx, text)


def _as_bool(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = sanitize_text(value).lower()
    if text in {"1", "true", "yes", "y", "是", "对", "需要", "允许"}:
        return True
    if text in {"0", "false", "no", "n", "否", "不", "不需要", "不允许"}:
        return False
    return bool(value)


def _split_lines(value: Any) -> list[str]:
    text = sanitize_text(value)
    if not text:
        return []
    return [part.strip() for part in text.replace("；", "\n").replace(";", "\n").splitlines() if part.strip()]


def _read_rows(path: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_header_name(cell, idx) for idx, cell in enumerate(rows[0])]
    parsed = []
    for row in rows[1:]:
        item = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            item[header] = row[idx] if idx < len(row) else ""
        parsed.append(item)
    return parsed


def import_reviewed_candidates(
    input_path: str,
    apply: bool = False,
    promote_active: bool = False,
    db_factory=None,
) -> dict[str, Any]:
    service = AgentBenchmarkDatasetService()
    rows = _read_rows(input_path)
    updated: list[str] = []
    promoted: list[str] = []
    would_promote: list[str] = []
    skipped: list[dict[str, str]] = []
    errors: list[dict[str, str]] = []
    for row in rows:
        scenario_uid = sanitize_text(row.get("scenario_uid"))
        if not scenario_uid:
            skipped.append({"scenario_uid": "", "reason": "missing_scenario_uid"})
            continue
        detail = service.get_scenario_detail(scenario_uid, db_factory=db_factory)
        if not detail:
            skipped.append({"scenario_uid": scenario_uid, "reason": "scenario_not_found"})
            continue
        reviewer = sanitize_text(row.get("reviewer"))
        if not reviewer:
            skipped.append({"scenario_uid": scenario_uid, "reason": "reviewer_required"})
            continue
        expected_reply = sanitize_text(row.get("expected_reply")) or sanitize_text(row.get("suggested_expected_reply"))
        quality = expected_reply_quality(expected_reply)
        if quality["quality"] != "valid":
            skipped.append({
                "scenario_uid": scenario_uid,
                "reason": quality["block_reason"] or "expected_reply_low_quality",
            })
            continue
        payload = {
            "scenario_uid": scenario_uid,
            "reviewer": reviewer,
            "expected_reply": expected_reply,
            "key_points": _split_lines(row.get("key_points")),
            "forbidden_claims": _split_lines(row.get("forbidden_claims")),
            "auto_send_allowed": _as_bool(row.get("auto_send_allowed")),
            "must_handoff": _as_bool(row.get("must_handoff")),
            "review_note": sanitize_text(row.get("review_note")),
        }
        if not apply:
            updated.append(scenario_uid)
            if promote_active:
                validation = service.validate_active_eligibility(
                    scenario_uid,
                    reviewer=reviewer,
                    expected_reply_override=payload["expected_reply"],
                    db_factory=db_factory,
                )
                if validation.get("eligible"):
                    would_promote.append(scenario_uid)
                else:
                    errors.append({
                        "scenario_uid": scenario_uid,
                        "reason": sanitize_text(validation.get("reason")) or "active_eligibility_failed",
                    })
            continue
        try:
            service.mark_expected_reply_reviewed(
                scenario_uid,
                reviewer=reviewer,
                expected_reply=payload["expected_reply"],
                key_points=payload["key_points"],
                forbidden_claims=payload["forbidden_claims"],
                auto_send_allowed=payload["auto_send_allowed"],
                must_handoff=payload["must_handoff"],
                review_note=payload["review_note"],
                db_factory=db_factory,
            )
            updated.append(scenario_uid)
            if promote_active:
                service.promote_to_active(scenario_uid, reviewer=reviewer, db_factory=db_factory)
                promoted.append(scenario_uid)
        except ValueError as exc:
            errors.append({"scenario_uid": scenario_uid, "reason": str(exc)})
    return sanitize_obj({
        "input": input_path,
        "apply": bool(apply),
        "dry_run": not bool(apply),
        "promote_active": bool(promote_active),
        "rows_seen": len(rows),
        "updated_count": len(updated),
        "promoted_count": len(promoted),
        "would_update_count": len(updated) if not apply else 0,
        "would_promote_count": len(would_promote) if not apply else 0,
        "skipped_count": len(skipped),
        "error_count": len(errors),
        "updated_scenario_uids": updated,
        "promoted_scenario_uids": promoted,
        "would_promote_scenario_uids": would_promote,
        "skipped": skipped,
        "errors": errors,
    })


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Import reviewed Agent benchmark candidates.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--promote-active", action="store_true")
    parser.add_argument("--json-output", default="")
    args = parser.parse_args(argv)

    init_db()
    result = import_reviewed_candidates(
        input_path=args.input,
        apply=bool(args.apply),
        promote_active=bool(args.promote_active),
    )
    if args.json_output:
        target = Path(args.json_output)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(sanitize_obj(result), ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not result.get("error_count") else 1


if __name__ == "__main__":
    raise SystemExit(main())
