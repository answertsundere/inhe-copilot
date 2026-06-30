"""Build strong ProductIdentityMapping candidates from replay and trusted data.

Automatic candidates require a platform identity signal plus a unique internal
KBProduct identity. Title-only and ambiguous rows are exported for manual review.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.db import SessionLocal, init_db
from app.models.eval_tables import EvalTrace
from app.models.kb_tables import KBProduct, ProductIdentityMapping
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text
from scripts.diagnose_product_identity_mapping_coverage import _latest_run_uid, _trace_identity, _url_host


AUTO_HEADERS = [
    "platform_item_id",
    "platform_item_id_hash",
    "product_url",
    "product_url_host",
    "platform_product_title",
    "i_id",
    "sku_code",
    "kb_product_id",
    "match_method",
    "confidence",
    "source_run_uid",
    "source_case_uid",
    "source_turn_uid",
    "source",
    "confirmation_status",
]
MANUAL_HEADERS = [
    "platform_item_id",
    "platform_item_id_hash",
    "product_url",
    "product_url_host",
    "platform_product_title",
    "possible_i_id",
    "possible_sku_code",
    "possible_product_title",
    "match_method",
    "conflict_reason",
    "source_case_uid",
    "source_turn_uid",
    "人工确认状态",
    "人工填写 i_id",
    "人工填写 SKU",
    "备注",
]
TRUSTED_HEADERS = {
    "platform_item_id": {"platform_item_id", "平台商品 ID", "平台商品ID", "item_id"},
    "platform_item_id_hash": {"platform_item_id_hash", "平台商品 ID Hash", "平台商品ID Hash", "item_id_hash"},
    "product_url": {"product_url", "商品链接", "平台商品链接"},
    "platform_product_title": {"platform_product_title", "平台商品标题", "商品标题"},
    "i_id": {"i_id", "内部 i_id", "建议内部 i_id"},
    "sku_code": {"sku_code", "SKU", "建议 SKU"},
    "confirmation_status": {"confirmation_status", "人工确认状态"},
    "source": {"source", "来源", "数据来源"},
}


def _first(raw: dict[str, Any], key: str) -> str:
    for alias in TRUSTED_HEADERS.get(key, {key}):
        if alias in raw:
            return sanitize_text(str(raw.get(alias) or ""))
    return ""


def _safe_url(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        parts = urlsplit(raw)
        if parts.scheme:
            return sanitize_text(f"{parts.scheme}://{parts.netloc}{parts.path}")
    except Exception:
        pass
    return sanitize_text(raw)


def _read_trusted_rows(path: str) -> list[dict[str, str]]:
    source = Path(path)
    if not source.exists():
        return []
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        raw_rows = data if isinstance(data, list) else data.get("items", []) if isinstance(data, dict) else []
        return [_normalize_trusted_row(row) for row in raw_rows if isinstance(row, dict)]
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows: list[dict[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            headers = [sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            if not any(headers):
                continue
            for values in sheet.iter_rows(min_row=2, values_only=True):
                raw = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
                row = _normalize_trusted_row(raw)
                if any(row.values()):
                    rows.append(row)
    finally:
        workbook.close()
    return rows


def _normalize_trusted_row(raw: dict[str, Any]) -> dict[str, str]:
    row = {key: _first(raw, key) for key in TRUSTED_HEADERS}
    row["product_url"] = _safe_url(row.get("product_url", ""))
    row["product_url_host"] = _url_host(row.get("product_url", ""))
    return row


def _product_for_iid_or_sku(db, i_id: str = "", sku_code: str = "") -> tuple[KBProduct | None, str]:
    products: dict[int, KBProduct] = {}
    i_id = sanitize_text(i_id)
    sku_code = sanitize_text(sku_code)
    if i_id:
        product = db.query(KBProduct).filter(KBProduct.i_id == i_id).first()
        if product:
            products[product.id] = product
    if sku_code:
        for product in db.query(KBProduct).filter(KBProduct.status == "published", KBProduct.sku_list_json.like(f"%{sku_code}%")).limit(20).all():
            if any(sanitize_text(item.get("sku_code")) == sku_code for item in product.get_sku_list() if isinstance(item, dict)):
                products[product.id] = product
        family = _sku_family(sku_code)
        if family:
            product = db.query(KBProduct).filter(KBProduct.i_id == family).first()
            if product:
                products[product.id] = product
    if len(products) == 1:
        return next(iter(products.values())), ""
    if len(products) > 1:
        return None, "ambiguous_internal_identity"
    return None, "kb_product_not_found"


def _sku_family(value: str) -> str:
    import re

    match = re.match(r"^(YH\d+K\d+)", str(value or "").strip(), re.I)
    return match.group(1).upper() if match else ""


def _platform_key(row: dict[str, Any]) -> tuple[str, str, str]:
    return (
        sanitize_text(row.get("platform_item_id")),
        sanitize_text(row.get("platform_item_id_hash")),
        _url_host(row.get("product_url", "")),
    )


def _has_platform_identity(row: dict[str, Any]) -> bool:
    return any(_platform_key(row))


def _existing_mapping_conflict(db, row: dict[str, Any], product: KBProduct) -> str:
    filters = []
    if row.get("platform_item_id"):
        filters.append(ProductIdentityMapping.platform_item_id == row["platform_item_id"])
    if row.get("platform_item_id_hash"):
        filters.append(ProductIdentityMapping.platform_item_id_hash == row["platform_item_id_hash"])
    if row.get("product_url_host") and (row.get("platform_item_id") or row.get("platform_item_id_hash")):
        filters.append(ProductIdentityMapping.product_url_host == row["product_url_host"])
    if not filters:
        return ""
    from sqlalchemy import or_

    mappings = (
        db.query(ProductIdentityMapping)
        .filter(ProductIdentityMapping.status.in_(["active", "verified", "confirmed"]))
        .filter(or_(*filters))
        .limit(20)
        .all()
    )
    product_ids = {item.kb_product_id for item in mappings if item.kb_product_id}
    if product_ids and product_ids != {product.id}:
        return "ambiguous_existing_mapping"
    return ""


def _auto_row(row: dict[str, Any], product: KBProduct, *, method: str, source: str, confidence: float) -> dict[str, Any]:
    return sanitize_obj({
        "platform_item_id": sanitize_text(row.get("platform_item_id")),
        "platform_item_id_hash": sanitize_text(row.get("platform_item_id_hash")),
        "product_url": _safe_url(row.get("product_url", "")),
        "product_url_host": _url_host(row.get("product_url", "")),
        "platform_product_title": sanitize_text(row.get("platform_product_title")),
        "i_id": product.i_id,
        "sku_code": sanitize_text(row.get("sku_code")) or product.i_id,
        "kb_product_id": product.id,
        "match_method": method,
        "confidence": confidence,
        "source_run_uid": sanitize_text(row.get("source_run_uid")),
        "source_case_uid": sanitize_text(row.get("source_case_uid")),
        "source_turn_uid": sanitize_text(row.get("source_turn_uid")),
        "source": source,
        "confirmation_status": "confirmed",
    })


def _manual_row(row: dict[str, Any], *, method: str, reason: str, product: KBProduct | None = None) -> dict[str, Any]:
    return sanitize_obj({
        "platform_item_id": sanitize_text(row.get("platform_item_id")),
        "platform_item_id_hash": sanitize_text(row.get("platform_item_id_hash")),
        "product_url": _safe_url(row.get("product_url", "")),
        "product_url_host": _url_host(row.get("product_url", "")),
        "platform_product_title": sanitize_text(row.get("platform_product_title")),
        "possible_i_id": product.i_id if product else sanitize_text(row.get("i_id")),
        "possible_sku_code": sanitize_text(row.get("sku_code")),
        "possible_product_title": product.product_name if product else "",
        "match_method": method,
        "conflict_reason": reason,
        "source_case_uid": sanitize_text(row.get("source_case_uid")),
        "source_turn_uid": sanitize_text(row.get("source_turn_uid")),
        "人工确认状态": "",
        "人工填写 i_id": "",
        "人工填写 SKU": "",
        "备注": "",
    })


def _embedded_platform_rows(product: KBProduct) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for item in product.get_sku_list() or []:
        if not isinstance(item, dict):
            continue
        item_id = sanitize_text(item.get("platform_item_id") or item.get("item_id") or item.get("taobao_item_id"))
        item_hash = sanitize_text(item.get("platform_item_id_hash") or item.get("item_id_hash") or item.get("taobao_item_id_hash"))
        product_url = _safe_url(item.get("product_url") or item.get("item_url") or item.get("url") or "")
        if not (item_id or item_hash or product_url):
            continue
        rows.append({
            "platform_item_id": item_id,
            "platform_item_id_hash": item_hash,
            "product_url": product_url,
            "product_url_host": _url_host(product_url),
            "platform_product_title": sanitize_text(item.get("product_title") or product.product_name),
            "i_id": product.i_id,
            "sku_code": sanitize_text(item.get("sku_code")) or product.i_id,
        })
    return rows


def _dedupe(rows: list[dict[str, Any]], keys: list[str]) -> list[dict[str, Any]]:
    seen = set()
    result = []
    for row in rows:
        key = tuple(sanitize_text(row.get(item)) for item in keys)
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def _replay_rows(db, run_uid: str) -> list[dict[str, Any]]:
    traces = (
        db.query(EvalTrace)
        .filter(EvalTrace.run_uid == run_uid)
        .order_by(EvalTrace.case_uid.asc(), EvalTrace.turn_index.asc(), EvalTrace.id.asc())
        .all()
    )
    rows = []
    for trace in traces:
        identity = _trace_identity(trace)
        rows.append({
            **identity,
            "source_run_uid": trace.run_uid,
            "source_case_uid": trace.case_uid,
            "source_turn_uid": trace.turn_uid,
        })
    return rows


def build_candidates(
    *,
    run_uid: str = "",
    trusted_inputs: list[str] | None = None,
    include_kb_embedded: bool = True,
    db_factory=None,
) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    db = db_factory()
    auto: list[dict[str, Any]] = []
    manual: list[dict[str, Any]] = []
    try:
        target_run_uid = sanitize_text(run_uid) or _latest_run_uid(db)
        source_rows: list[dict[str, Any]] = []
        for row in _replay_rows(db, target_run_uid):
            source_rows.append({**row, "source": "replay_trace"})
        for input_path in trusted_inputs or []:
            for row in _read_trusted_rows(input_path):
                source_rows.append({**row, "source": sanitize_text(row.get("source")) or "trusted_product_sheet"})
        if include_kb_embedded:
            for product in db.query(KBProduct).filter(KBProduct.status == "published").all():
                for row in _embedded_platform_rows(product):
                    source_rows.append({**row, "source": "kb_product_embedded_platform_fields"})

        grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
        for row in source_rows:
            if _has_platform_identity(row):
                grouped[_platform_key(row)].append(row)
            elif row.get("platform_product_title") or row.get("order_product_title"):
                manual.append(_manual_row(row, method="title_only_not_used", reason="title_only_not_auto_mapped"))

        for rows in grouped.values():
            internal_candidates: dict[int, tuple[KBProduct, dict[str, Any]]] = {}
            for row in rows:
                product, reason = _product_for_iid_or_sku(db, row.get("i_id", ""), row.get("sku_code", ""))
                if product:
                    internal_candidates[product.id] = (product, row)
                elif row.get("i_id") or row.get("sku_code"):
                    manual.append(_manual_row(row, method="explicit_internal_identity_missing", reason=reason))
            if len(internal_candidates) == 1:
                product, row = next(iter(internal_candidates.values()))
                conflict = _existing_mapping_conflict(db, row, product)
                if conflict:
                    manual.append(_manual_row(row, method="explicit_platform_and_internal_identity", reason=conflict, product=product))
                    continue
                method = "exact_sku_in_context" if row.get("sku_code") else "exact_i_id_in_context"
                if row.get("source") == "kb_product_embedded_platform_fields":
                    method = "kb_product_embedded_platform_field"
                elif row.get("source") and row.get("source") != "replay_trace":
                    method = "trusted_source_same_row"
                auto.append(_auto_row(row, product, method=method, source=row.get("source") or "replay_trace", confidence=1.0))
            elif len(internal_candidates) > 1:
                for product, row in internal_candidates.values():
                    manual.append(_manual_row(row, method="multiple_internal_identities_for_platform_identity", reason="ambiguous", product=product))
            else:
                row = rows[0]
                reason = "hash_only_without_verified_internal_identity" if row.get("platform_item_id_hash") and not row.get("platform_item_id") else "missing_internal_identity"
                manual.append(_manual_row(row, method="platform_identity_without_internal_identity", reason=reason))

        auto = _dedupe(auto, ["platform_item_id", "platform_item_id_hash", "product_url_host", "i_id", "sku_code"])
        manual = _dedupe(manual, ["platform_item_id", "platform_item_id_hash", "product_url_host", "possible_i_id", "conflict_reason"])
        return sanitize_obj({
            "ok": True,
            "run_uid": target_run_uid,
            "source_row_count": len(source_rows),
            "auto_candidate_count": len(auto),
            "manual_candidate_count": len(manual),
            "auto_candidates": auto,
            "manual_candidates": manual,
        })
    finally:
        db.close()


def _write_workbook(path: str, headers: list[str], rows: list[dict[str, Any]], sheet_name: str) -> str:
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet_name
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([row.get(key, "") for key in headers])
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(14, min(42, len(header) + 4))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def write_outputs(result: dict[str, Any], *, auto_output: str, manual_output: str) -> dict[str, str]:
    return {
        "auto_output": _write_workbook(auto_output, AUTO_HEADERS, result.get("auto_candidates", []), "auto_candidates"),
        "manual_output": _write_workbook(manual_output, MANUAL_HEADERS, result.get("manual_candidates", []), "manual_review"),
    }


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    today = datetime.now().strftime("%Y%m%d")
    parser = argparse.ArgumentParser(description="Build product identity mapping candidates.")
    parser.add_argument("--run-uid", default="")
    parser.add_argument("--trusted-input", action="append", default=[])
    parser.add_argument("--no-kb-embedded", action="store_true")
    parser.add_argument("--auto-output", default=f"outputs/product_identity_mapping_candidates_auto_{today}.xlsx")
    parser.add_argument("--manual-output", default=str(Path.home() / "Desktop" / f"商品身份映射待人工确认_{today}.xlsx"))
    parser.add_argument("--json-output", default=f"outputs/product_identity_mapping_candidates_{today}.json")
    args = parser.parse_args()
    init_db()
    result = build_candidates(
        run_uid=args.run_uid,
        trusted_inputs=args.trusted_input,
        include_kb_embedded=not args.no_kb_embedded,
    )
    outputs = write_outputs(result, auto_output=args.auto_output, manual_output=args.manual_output)
    summary = {k: v for k, v in result.items() if k not in {"auto_candidates", "manual_candidates"}}
    summary["outputs"] = outputs
    if args.json_output:
        output = Path(args.json_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
