"""Export one replay run's knowledge gap tasks into an operator workbook."""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db import SessionLocal, init_db
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text
from app.services.knowledge_gap_task_service import KnowledgeGapTaskService


SHEET_README = "\u8bf4\u660e"
SHEET_OVERVIEW = "\u603b\u89c8"
SHEET_ALL = "\u77e5\u8bc6\u7f3a\u53e3\u4efb\u52a1"
SHEET_MEDIA = "\u7d20\u6750\u7f3a\u53e3"
SHEET_POLICY = "\u552e\u540e\u6d3b\u52a8\u89c4\u5219\u7f3a\u53e3"
SHEET_PRODUCT = "\u5546\u54c1\u5b57\u6bb5\u7f3a\u53e3"
SHEET_ROUTING = "\u4e0a\u4e0b\u6587\u6216\u8bc1\u636e\u8def\u7531\u7f3a\u53e3"
SHEET_HIGH_RISK = "\u9ad8\u98ce\u9669\u9700\u4eba\u5de5\u786e\u8ba4"

TASK_HEADERS = [
    "\u4f18\u5148\u7ea7",
    "\u7f3a\u53e3\u7c7b\u578b",
    "\u8bc1\u636e\u7c7b\u578b",
    "\u76ee\u6807\u7cfb\u7edf",
    "\u8d1f\u8d23\u4eba",
    "\u5546\u54c1\u7f16\u7801",
    "\u5546\u54c1\u540d\u79f0",
    "\u95ee\u9898\u7c7b\u578b",
    "\u95ee\u9898\u7c7b\u578b\u4e2d\u6587\u540d",
    "\u5931\u8d25\u6b21\u6570",
    "\u4e70\u5bb6\u95ee\u9898\uff08\u5df2\u8131\u654f\uff09",
    "\u5f53\u524d\u56de\u590d\uff08\u5df2\u8131\u654f\uff09",
    "\u5931\u8d25\u539f\u56e0",
    "\u5f53\u524d\u963b\u585e\u70b9",
    "\u9700\u8981\u8865\u4ec0\u4e48",
    "\u5efa\u8bae\u52a8\u4f5c",
    "\u5efa\u8bae\u586b\u5199\u5185\u5bb9",
    "\u4ee3\u8868\u6837\u672c\uff08\u6700\u591a3\u6761\uff09",
    "\u662f\u5426\u53ef\u81ea\u52a8\u53d1\u9001",
    "\u4eba\u5de5\u5904\u7406\u7ed3\u679c",
    "\u4eba\u5de5\u8865\u5145\u5185\u5bb9",
    "\u662f\u5426\u5df2\u8865\u77e5\u8bc6\u5e93",
    "\u662f\u5426\u5df2\u4e0a\u4f20\u7d20\u6750",
    "\u7d20\u6750\u94fe\u63a5\u6216\u7d20\u6750\u7f16\u53f7",
    "\u5546\u54c1\u5b57\u6bb5\u503c",
    "\u552e\u540e\u89c4\u5219\u53e3\u5f84",
    "\u6d3b\u52a8\u89c4\u5219\u53e3\u5f84",
    "\u5904\u7406\u4eba",
    "\u8bc1\u636e\u6765\u6e90/\u7d20\u6750\u94fe\u63a5",
    "\u662f\u5426\u5df2\u8865\u9f50",
    "\u5907\u6ce8",
    "\u4efb\u52a1ID",
    "\u72b6\u6001",
    "\u98ce\u9669\u7b49\u7ea7",
]

GAP_CATEGORY_LABELS = {
    "media_asset_gap": "\u7d20\u6750\u7f3a\u53e3",
    "product_field_gap": "\u5546\u54c1\u5b57\u6bb5\u7f3a\u53e3",
    "aftersales_policy_gap": "\u552e\u540e\u89c4\u5219\u7f3a\u53e3",
    "promotion_policy_gap": "\u6d3b\u52a8\u89c4\u5219\u7f3a\u53e3",
    "evidence_routing_gap": "\u8bc1\u636e\u8def\u7531\u7f3a\u53e3",
    "context_extraction_gap": "\u4e0a\u4e0b\u6587\u62bd\u53d6\u7f3a\u53e3",
}

QUERY_FACT_TYPE_LABELS = {
    "installation": "\u5b89\u88c5/\u8bf4\u660e\u4e66/\u6559\u7a0b",
    "accessory_usage": "\u914d\u4ef6\u7528\u9014",
    "accessory_availability": "\u914d\u4ef6\u8865\u914d/\u552e\u5356",
    "accessory_compatibility": "\u914d\u4ef6\u9002\u914d",
    "structure_function": "\u7ed3\u6784\u529f\u80fd",
    "dimensions": "\u5c3a\u5bf8",
    "space_fit": "\u7a7a\u95f4/\u6446\u653e\u9002\u914d",
    "placement_scene": "\u4f7f\u7528\u573a\u666f",
    "material": "\u6750\u8d28",
    "material_safety": "\u6750\u8d28\u5b89\u5168",
    "certification_report": "\u68c0\u6d4b/\u8bc1\u4e66",
    "load_capacity": "\u627f\u91cd",
    "gross_weight": "\u91cd\u91cf/\u6bdb\u91cd",
    "age_range": "\u9002\u7528\u5e74\u9f84",
    "child_suitability": "\u513f\u7ae5\u9002\u7528",
    "child_safety": "\u513f\u7ae5\u5b89\u5168",
    "promotion": "\u4f18\u60e0/\u798f\u5229",
    "promotion_policy": "\u6d3b\u52a8\u89c4\u5219",
    "price_negotiation": "\u8bae\u4ef7/\u591a\u4e70\u4f18\u60e0",
    "aftersales": "\u552e\u540e",
    "aftersales_policy": "\u552e\u540e\u653f\u7b56",
    "stock_shipping": "\u5e93\u5b58/\u53d1\u8d27",
    "logistics": "\u7269\u6d41",
}


def default_output_path() -> str:
    today = datetime.now().strftime("%Y%m%d")
    return str(Path.home() / "Desktop" / f"\u771f\u5b9e\u56de\u653e\u77e5\u8bc6\u7f3a\u53e3_{today}.xlsx")


def _normalize_optional_filter(value: str) -> str:
    text = sanitize_text(value)
    return "" if text in {"", "__all__"} else text


def _join(values: list[Any], limit: int = 3) -> str:
    return "\n".join(
        sanitize_text(str(value))
        for value in values[:limit]
        if sanitize_text(str(value))
    )


def _metadata(task: dict[str, Any]) -> dict[str, Any]:
    metadata = task.get("metadata")
    return metadata if isinstance(metadata, dict) else {}


def _gap_category(task: dict[str, Any]) -> str:
    return sanitize_text(task.get("gap_category") or task.get("gap_type") or "")


def _product_code(task: dict[str, Any]) -> str:
    sku_code = sanitize_text(task.get("sku_code") or "")
    if sku_code:
        return sku_code
    return sanitize_text(task.get("item_id_masked") or "")


def _current_blocker(task: dict[str, Any]) -> str:
    metadata = _metadata(task)
    return sanitize_text(metadata.get("current_blocker") or task.get("summary") or "")


def _query_fact_type_label(task: dict[str, Any]) -> str:
    fact_type = sanitize_text(task.get("query_fact_type") or "")
    return QUERY_FACT_TYPE_LABELS.get(fact_type, fact_type or "\u672a\u8bc6\u522b")


def _required_evidence(task: dict[str, Any]) -> str:
    metadata = _metadata(task)
    evidence = sanitize_text(
        task.get("required_evidence_type")
        or task.get("missing_evidence_type")
        or task.get("media_needed_type")
        or metadata.get("required_evidence_type")
    )
    missing_fields = metadata.get("missing_fields")
    if isinstance(missing_fields, list) and missing_fields:
        fields = [sanitize_text(str(value)) for value in missing_fields if sanitize_text(str(value))]
        if fields:
            return "\n".join(fields[:5])
    return evidence or _current_blocker(task)


def _suggested_fill_content(task: dict[str, Any]) -> str:
    category = _gap_category(task)
    evidence = _required_evidence(task)
    if category == "media_asset_gap":
        return f"\u586b\u5199\u5df2\u5ba1\u6838\u53ef\u7528\u7684\u7d20\u6750\u94fe\u63a5/\u7d20\u6750ID\uff1a{evidence}"
    if category == "product_field_gap":
        return f"\u586b\u5199\u5f53\u524d\u5546\u54c1\u7684\u7ed3\u6784\u5316\u5b57\u6bb5\u503c\uff1a{evidence}"
    if category == "promotion_policy_gap":
        return "\u586b\u5199\u5f53\u524d\u6d3b\u52a8/\u4f18\u60e0/\u8d60\u54c1\u89c4\u5219\uff0c\u6ce8\u660e\u751f\u6548\u6761\u4ef6\u548c\u65f6\u95f4"
    if category == "aftersales_policy_gap":
        return "\u586b\u5199\u9000\u6362/\u8865\u53d1/\u8fd0\u8d39/\u7834\u635f\u7b49\u552e\u540e\u5904\u7406\u53e3\u5f84"
    if category in {"context_extraction_gap", "evidence_routing_gap"}:
        return "\u8865\u5145\u4e0a\u4e0b\u6587\u5b57\u6bb5\u6765\u6e90\u6216\u4fee\u6b63\u8bc1\u636e role/\u8def\u7531\u6807\u8bb0"
    return "\u586b\u5199\u6709\u6765\u6e90\u7684\u53ef\u590d\u6838\u8bc1\u636e"


def _representative_samples(task: dict[str, Any]) -> str:
    questions = [sanitize_text(str(value)) for value in (task.get("latest_buyer_questions") or [])]
    replies = [sanitize_text(str(value)) for value in (task.get("latest_agent_replies") or [])]
    lines: list[str] = []
    for index, question in enumerate([value for value in questions if value][:3], start=1):
        reply = replies[index - 1] if index - 1 < len(replies) else ""
        lines.append(f"{index}. \u4e70\u5bb6\uff1a{question}")
        if reply:
            lines.append(f"   \u5f53\u524d\u56de\u590d\uff1a{reply}")
    return "\n".join(lines)


def _task_row(task: dict[str, Any]) -> list[Any]:
    sample_text = _representative_samples(task)
    return [
        sanitize_text(task.get("priority") or ""),
        GAP_CATEGORY_LABELS.get(_gap_category(task), _gap_category(task)),
        sanitize_text(task.get("required_evidence_type") or task.get("missing_evidence_type") or ""),
        sanitize_text(task.get("target_system") or ""),
        sanitize_text(task.get("suggested_owner") or ""),
        _product_code(task),
        sanitize_text(task.get("product_title_preview") or task.get("product_title") or ""),
        sanitize_text(task.get("query_fact_type") or ""),
        _query_fact_type_label(task),
        int(task.get("sample_count") or 0),
        _join(task.get("latest_buyer_questions") or []),
        _join(task.get("latest_agent_replies") or []),
        sanitize_text(task.get("failure_type") or ""),
        _current_blocker(task),
        _required_evidence(task),
        sanitize_text(task.get("recommended_action") or ""),
        _suggested_fill_content(task),
        sample_text,
        "\u5426\uff08\u9700 verified evidence + sendable contract \u901a\u8fc7\u540e\u518d\u8bc4\u4f30\uff09",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        str(task.get("task_uid") or "").strip(),
        sanitize_text(task.get("status") or ""),
        sanitize_text(task.get("risk_level") or ""),
    ]


def _apply_sheet_style(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in sheet.columns:
        letter = column[0].column_letter
        header = str(column[0].value or "")
        sheet.column_dimensions[letter].width = min(max(len(header) + 4, 12), 36)
    sheet.freeze_panes = "A2"


def _append_task_sheet(workbook: Workbook, title: str, tasks: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(TASK_HEADERS)
    for task in tasks:
        sheet.append(_task_row(task))
    _apply_sheet_style(sheet)


def _append_overview(workbook: Workbook, tasks: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(SHEET_OVERVIEW)
    category_counts = Counter(_gap_category(task) or "unknown" for task in tasks)
    fact_counts = Counter(sanitize_text(task.get("query_fact_type") or "unknown") for task in tasks)
    owner_counts = Counter(sanitize_text(task.get("suggested_owner") or "unknown") for task in tasks)
    priority_counts = Counter(sanitize_text(task.get("priority") or "unknown") for task in tasks)
    rows = [
        ["\u6307\u6807", "\u503c"],
        ["\u7f3a\u53e3\u4efb\u52a1\u6570", len(tasks)],
        ["\u8bf4\u660e", "\u672c\u8868\u53ea\u7528\u4e8e\u4eba\u5de5\u6cbb\u7406\u548c\u590d\u6d4b\uff0c\u4e0d\u4f1a\u81ea\u52a8\u5199\u6b63\u5f0f\u77e5\u8bc6\u5e93/\u7d20\u6750\u5e93\u3002"],
        ["", ""],
        ["\u6309\u7f3a\u53e3\u7c7b\u578b", ""],
    ]
    for key, count in sorted(category_counts.items()):
        rows.append([GAP_CATEGORY_LABELS.get(key, key), count])
    rows.append(["", ""])
    rows.append(["\u6309\u95ee\u9898\u7c7b\u578b", ""])
    for key, count in fact_counts.most_common():
        rows.append([QUERY_FACT_TYPE_LABELS.get(key, key), count])
    rows.append(["", ""])
    rows.append(["\u6309\u8d1f\u8d23\u4eba", ""])
    for key, count in sorted(owner_counts.items()):
        rows.append([key, count])
    rows.append(["", ""])
    rows.append(["\u6309\u4f18\u5148\u7ea7", ""])
    for key, count in sorted(priority_counts.items()):
        rows.append([key, count])
    for row in rows:
        sheet.append(row)
    _apply_sheet_style(sheet)


def _append_readme(workbook: Workbook, *, run_uid: str, tasks: list[dict[str, Any]]) -> None:
    sheet = workbook.active
    sheet.title = SHEET_README
    category_counts = Counter(_gap_category(task) or "unknown" for task in tasks)
    owner_counts = Counter(sanitize_text(task.get("suggested_owner") or "unknown") for task in tasks)
    rows = [
        ["run_uid", sanitize_text(run_uid)],
        ["total_gap_tasks", len(tasks)],
        ["note", "\u672c\u5de5\u4f5c\u7c3f\u53ea\u7528\u4e8e\u4eba\u5de5\u8865\u8d44\u6599/\u7d20\u6750/\u89c4\u5219\uff0c\u4e0d\u4f1a\u81ea\u52a8\u53d1\u5e03\u6b63\u5f0f\u77e5\u8bc6\u5e93\u3002"],
        ["workflow", "\u4eba\u5de5\u8865\u9f50 -> staging \u5ba1\u6838 -> \u56de\u653e\u590d\u6d4b -> \u518d\u51b3\u5b9a\u662f\u5426\u53d1\u5e03\u3002"],
        ["fill_rule", "\u53ea\u586b\u5145\u6709\u6765\u6e90\u7684\u771f\u5b9e\u8bc1\u636e\uff0c\u4e0d\u8981\u628a\u5ba2\u670d\u539f\u56de\u590d\u76f4\u63a5\u5199\u6210\u6807\u51c6\u7b54\u6848\u3002"],
        [],
        ["by_gap_category", ""],
    ]
    for key, count in sorted(category_counts.items()):
        rows.append([GAP_CATEGORY_LABELS.get(key, key), count])
    rows.extend([[], ["by_owner", ""]])
    for key, count in sorted(owner_counts.items()):
        rows.append([key, count])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 80


def fetch_tasks_for_run(
    *,
    run_uid: str,
    status: str = "",
    owner: str = "",
    gap_category: str = "",
    limit: int = 500,
) -> dict[str, Any]:
    init_db()
    db = SessionLocal()
    try:
        result = KnowledgeGapTaskService().list_tasks(
            db,
            filters={
                "run_uid": run_uid,
                "status": _normalize_optional_filter(status),
                "suggested_owner": _normalize_optional_filter(owner),
                "gap_category": _normalize_optional_filter(gap_category),
            },
            limit=limit,
        )
        return sanitize_obj(result)
    finally:
        db.close()


def export_latest_run_knowledge_gaps(
    *,
    run_uid: str,
    output: str | None = None,
    status: str = "",
    owner: str = "",
    gap_category: str = "",
) -> dict[str, Any]:
    target_run_uid = sanitize_text(run_uid)
    if not target_run_uid:
        raise ValueError("run_uid is required")
    result = fetch_tasks_for_run(
        run_uid=target_run_uid,
        status=status,
        owner=owner,
        gap_category=gap_category,
    )
    tasks = result.get("items") if isinstance(result.get("items"), list) else []
    tasks = [task for task in tasks if isinstance(task, dict)]

    workbook = Workbook()
    _append_readme(workbook, run_uid=target_run_uid, tasks=tasks)
    _append_overview(workbook, tasks)
    _append_task_sheet(workbook, SHEET_ALL, tasks)
    _append_task_sheet(workbook, SHEET_MEDIA, [task for task in tasks if _gap_category(task) == "media_asset_gap"])
    _append_task_sheet(
        workbook,
        SHEET_POLICY,
        [
            task for task in tasks
            if _gap_category(task) in {"aftersales_policy_gap", "promotion_policy_gap"}
        ],
    )
    _append_task_sheet(workbook, SHEET_PRODUCT, [task for task in tasks if _gap_category(task) == "product_field_gap"])
    _append_task_sheet(
        workbook,
        SHEET_ROUTING,
        [
            task for task in tasks
            if _gap_category(task) in {"context_extraction_gap", "evidence_routing_gap"}
        ],
    )
    _append_task_sheet(
        workbook,
        SHEET_HIGH_RISK,
        [
            task for task in tasks
            if sanitize_text(task.get("risk_level")).lower() in {"high", "critical"}
            or _gap_category(task) in {"aftersales_policy_gap", "promotion_policy_gap"}
        ],
    )

    path = Path(output or default_output_path())
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)

    category_counts = Counter(_gap_category(task) or "unknown" for task in tasks)
    owner_counts = Counter(sanitize_text(task.get("suggested_owner") or "unknown") for task in tasks)
    return {
        "output": str(path),
        "run_uid": target_run_uid,
        "count": len(tasks),
        "by_gap_category": dict(sorted(category_counts.items())),
        "by_owner": dict(sorted(owner_counts.items())),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Export one replay run's knowledge gap tasks to Excel.")
    parser.add_argument("--run-uid", required=True, help="Replay run uid to export")
    parser.add_argument("--output", default="", help="Excel output path")
    parser.add_argument("--status", default="", help="Optional task status filter; empty exports all statuses")
    parser.add_argument("--owner", default="", help="Optional suggested_owner filter")
    parser.add_argument("--gap-category", default="", help="Optional gap_category filter")
    args = parser.parse_args()
    result = export_latest_run_knowledge_gaps(
        run_uid=args.run_uid,
        output=args.output or None,
        status=args.status,
        owner=args.owner,
        gap_category=args.gap_category,
    )
    print(result)


if __name__ == "__main__":
    main()
