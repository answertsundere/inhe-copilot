"""Classify trusted product media URLs into KBMediaAsset asset types."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.db import SessionLocal, init_db
from app.models.kb_tables import KBMediaAsset, KBProduct
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text


FIELD_ALIASES = {
    "i_id": {"i_id", "内部 i_id", "商品货号"},
    "sku_code": {"sku_code", "SKU", "sku"},
    "product_name": {"product_name", "商品标题", "商品名称"},
    "asset_url": {"asset_url", "素材链接", "URL", "url"},
    "asset_title": {"asset_title", "素材标题", "标题", "文件名"},
    "source": {"source", "来源", "数据来源"},
    "mime_type": {"mime_type", "MIME", "文件类型"},
}
TRUSTED_SOURCES = {"dingtalk", "dingtalk_product_kb_sync", "trusted_product_sheet", "product_detail", "official_product_asset", ""}
UNTRUSTED_SOURCES = {"chat", "history_chat", "real_conversation", "customer_upload", "screenshot"}


def _read_rows(path: str) -> list[dict[str, str]]:
    source = Path(path)
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else data.get("items", []) if isinstance(data, dict) else []
        return [_normalize_row(row) for row in rows]
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = _select_sheet(workbook)
        headers = [sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        result = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            raw = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
            result.append(_normalize_row(raw))
        return result
    finally:
        workbook.close()


def _select_sheet(workbook):
    required = FIELD_ALIASES["asset_url"]
    for sheet in workbook.worksheets:
        headers = {sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))}
        if headers & required:
            return sheet
    return workbook[workbook.sheetnames[0]]


def _normalize_row(raw: dict[str, Any]) -> dict[str, str]:
    row: dict[str, str] = {}
    for key, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in raw:
                if key == "source":
                    row[key] = str(raw.get(alias) or "").strip()
                else:
                    row[key] = sanitize_text(str(raw.get(alias) or ""))
                break
        row.setdefault(key, "")
    return row


def _stable_url(url: str) -> str:
    try:
        parts = urlsplit(url or "")
    except Exception:
        return sanitize_text(url)
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "", ""))


def _content_hash(i_id: str, sku_code: str, asset_type: str, asset_url: str) -> str:
    raw = "|".join([sanitize_text(i_id), sanitize_text(sku_code), sanitize_text(asset_type), _stable_url(asset_url)])
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def classify_asset_type(row: dict[str, str]) -> tuple[str, list[str], str]:
    title = f"{row.get('asset_title', '')} {row.get('asset_url', '')} {row.get('mime_type', '')}".lower()
    is_video = any(marker in title for marker in (".mp4", ".mov", ".avi", "video", "视频", "抖音", "b站", "bilibili"))
    is_image = any(marker in title for marker in (".jpg", ".jpeg", ".png", ".webp", "image", "图片", "图"))
    if is_video and any(term in title for term in ("安装", "教程", "组装", "install", "assembly")):
        return "install_video", ["installation"], "video_install_keyword"
    if is_image and any(term in title for term in ("安装", "教程", "说明书", "步骤", "图纸", "install", "manual")):
        return "install_image", ["installation"], "image_install_keyword"
    if is_image and any(term in title for term in ("尺寸", "规格", "长宽高", "size", "dimension")):
        return "size_image", ["dimensions", "space_fit"], "image_size_keyword"
    if is_image and any(term in title for term in ("配件", "五金", "零件", "清单", "accessory", "parts")):
        return "accessory_image", ["accessories", "packing_list"], "image_accessory_keyword"
    if is_image and any(term in title for term in ("证书", "检测报告", "certificate", "report", "合格证")):
        return "certificate_image", ["certification_report"], "image_certificate_keyword"
    if is_image and any(term in title for term in ("售后", "退换", "aftersales")):
        return "aftersales_image", ["aftersales"], "image_aftersales_keyword"
    if is_image:
        return "sku_image", ["appearance"], "image_default_sku"
    if is_video:
        return "marketing_video", ["appearance"], "video_default_marketing"
    return "other", [], "unclassified"


def _find_product(db, i_id: str, sku_code: str) -> KBProduct | None:
    if i_id:
        product = db.query(KBProduct).filter(KBProduct.i_id == i_id).one_or_none()
        if product:
            return product
    if sku_code:
        for product in db.query(KBProduct).all():
            if any(sanitize_text(item.get("sku_code")) == sku_code for item in product.get_sku_list() if isinstance(item, dict)):
                return product
    return None


def run_classify(input_path: str, *, apply: bool = False, operator: str = "auto_media_classifier", db_factory=None) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    rows = _read_rows(input_path)
    db = db_factory()
    stats = Counter()
    skipped: list[dict[str, Any]] = []
    classified_by_type = Counter()
    try:
        for index, row in enumerate(rows, start=2):
            stats["media_checked_count"] += 1
            source = str(row.get("source") or "").strip().lower()
            if source in UNTRUSTED_SOURCES or source not in TRUSTED_SOURCES:
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": "untrusted_source"})
                continue
            url = sanitize_text(row.get("asset_url"))
            if not url:
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": "missing_asset_url"})
                continue
            product = _find_product(db, sanitize_text(row.get("i_id")), sanitize_text(row.get("sku_code")))
            if not product:
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": "kb_product_not_found"})
                continue
            asset_type, scene_tags, reason = classify_asset_type(row)
            if asset_type == "other":
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": "unclassified"})
                continue
            classified_by_type[asset_type] += 1
            chash = _content_hash(product.i_id, sanitize_text(row.get("sku_code")) or product.i_id, asset_type, url)
            existing = (
                db.query(KBMediaAsset)
                .filter(KBMediaAsset.content_hash == chash)
                .order_by(KBMediaAsset.id.asc())
                .first()
            )
            if existing:
                stats["media_updated_count"] += 1
                asset = existing
            else:
                stats["media_created_count"] += 1
                asset = KBMediaAsset(content_hash=chash, created_by=operator)
            if apply:
                asset.product_id = product.id
                asset.i_id = product.i_id
                asset.sku_code = sanitize_text(row.get("sku_code")) or product.i_id
                asset.product_name = product.product_name
                asset.asset_type = asset_type
                asset.asset_title = sanitize_text(row.get("asset_title"))[:255] or f"{product.product_name} {asset_type}"
                asset.asset_url = _stable_url(url)
                asset.source = "trusted_product_media_backfill"
                asset.source_doc_id = chash
                asset.match_confidence = 0.9
                asset.match_reason = reason
                asset.status = "pending_review"
                asset.audit_status = "unreviewed"
                asset.usable_for_agent = 0
                asset.updated_by = operator
                asset.set_scene_tags(scene_tags)
                asset.set_source_raw(sanitize_obj({
                    "source": row.get("source"),
                    "source_file": str(input_path),
                    "original_title": row.get("asset_title"),
                    "stable_url_only": True,
                }))
                db.add(asset)
        if apply:
            db.commit()
        else:
            db.rollback()
        return sanitize_obj({
            "ok": True,
            "dry_run": not apply,
            "media_checked_count": stats["media_checked_count"],
            "media_created_count": stats["media_created_count"],
            "media_updated_count": stats["media_updated_count"],
            "classified_by_type": dict(classified_by_type),
            "skipped_count": stats["skipped_count"],
            "skipped_reasons": skipped,
            "auto_approved_count": 0,
            "writes_verified_knowledge": False,
        })
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Classify trusted product media assets.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--operator", default="auto_media_classifier")
    parser.add_argument("--json-output", "--output", dest="json_output", default="")
    args = parser.parse_args()
    init_db()
    result = run_classify(args.input, apply=args.apply, operator=args.operator)
    if args.json_output:
        output = Path(args.json_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
