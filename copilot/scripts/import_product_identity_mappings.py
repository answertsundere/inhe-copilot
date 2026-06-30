"""Import manually confirmed platform-to-KBProduct identity mappings."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.db import SessionLocal, init_db
from app.models.kb_tables import KBProduct, ProductIdentityMapping
from app.services.eval_sanitizer_service import hash_sensitive, sanitize_obj, sanitize_text, stable_hash


SHEET_NAME = "商品身份映射缺口"
CONFIRMED_VALUES = {"已确认", "确认", "confirmed", "active", "verified", "yes", "y", "true"}
FIELD_ALIASES = {
    "source_run_uid": {"回放批次", "source_run_uid"},
    "case_uid": {"案例 ID", "case_uid"},
    "turn_uid": {"轮次 ID", "turn_uid"},
    "platform_item_id": {"平台商品 ID", "platform_item_id"},
    "platform_item_id_hash": {"平台商品 ID Hash", "platform_item_id_hash"},
    "product_url": {"商品链接", "product_url"},
    "platform_product_title": {"平台商品标题", "platform_product_title"},
    "order_product_title": {"订单商品标题", "order_product_title"},
    "i_id": {"建议内部 i_id", "内部 i_id", "i_id"},
    "sku_code": {"建议 SKU", "SKU", "sku_code"},
    "suggested_product_title": {"建议商品标题", "商品标题"},
    "confirmation_status": {"人工确认状态", "确认状态"},
    "operator": {"处理人", "operator"},
    "note": {"备注", "note"},
}


def _now_iso() -> str:
    return datetime.utcnow().isoformat()


def _canonical_header(value: Any) -> str:
    text = sanitize_text(str(value or "")).strip()
    for canonical, aliases in FIELD_ALIASES.items():
        if text in aliases:
            return canonical
    return ""


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return sanitize_text(str(value)).strip()


def _raw_cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_rows(input_path: str) -> list[dict[str, str]]:
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        return []
    sheet = workbook[SHEET_NAME]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return []
    headers = [_canonical_header(value) for value in header_row]
    rows: list[dict[str, str]] = []
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values: dict[str, str] = {"_row": str(index)}
        for header, value in zip(headers, row):
            if not header:
                continue
            if header in {"i_id", "sku_code", "platform_item_id_hash"}:
                values[header] = _raw_cell_text(value)
            else:
                values[header] = _cell_text(value)
        if any(value for key, value in values.items() if key != "_row"):
            rows.append(values)
    return rows


def _is_confirmed(value: str) -> bool:
    return str(value or "").strip().lower() in CONFIRMED_VALUES


def _url_host(value: str) -> str:
    try:
        return urlsplit(str(value or "")).netloc.lower()
    except Exception:
        return ""


def _mapping_key(row: dict[str, str]) -> dict[str, str]:
    platform_item_id = row.get("platform_item_id", "")
    if "REDACTED" in platform_item_id:
        platform_item_id = ""
    platform_item_id_hash = row.get("platform_item_id_hash", "")
    if not platform_item_id_hash and platform_item_id and platform_item_id.isdigit():
        platform_item_id_hash = hash_sensitive(platform_item_id)
    return {
        "platform_item_id": platform_item_id,
        "platform_item_id_hash": platform_item_id_hash,
        "product_url": row.get("product_url", ""),
        "product_url_host": _url_host(row.get("product_url", "")),
    }


def _product_has_sku(product: KBProduct, sku_code: str) -> bool:
    target = str(sku_code or "").strip().upper()
    if not target:
        return False
    if str(product.i_id or "").strip().upper() == target:
        return True
    for item in product.get_sku_list() or []:
        if isinstance(item, dict):
            values = [item.get("sku_code"), item.get("sku_id"), item.get("sku"), item.get("barcode_69")]
        else:
            values = [item]
        if any(str(value or "").strip().upper() == target for value in values):
            return True
    return False


def _find_product(db, i_id: str, sku_code: str) -> KBProduct | None:
    if i_id:
        product = db.query(KBProduct).filter(KBProduct.i_id == i_id).first()
        if product:
            return product
    if sku_code:
        rows = (
            db.query(KBProduct)
            .filter(KBProduct.status == "published")
            .filter(KBProduct.sku_list_json.like(f"%{sku_code}%"))
            .limit(20)
            .all()
        )
        matches = [product for product in rows if _product_has_sku(product, sku_code)]
        if len(matches) == 1:
            return matches[0]
    return None


def _existing_mappings(db, key: dict[str, str]) -> list[ProductIdentityMapping]:
    filters = []
    if key.get("platform_item_id_hash"):
        filters.append(ProductIdentityMapping.platform_item_id_hash == key["platform_item_id_hash"])
    if key.get("platform_item_id"):
        filters.append(ProductIdentityMapping.platform_item_id == key["platform_item_id"])
    if key.get("product_url_host") and (key.get("platform_item_id") or key.get("platform_item_id_hash")):
        filters.append(ProductIdentityMapping.product_url_host == key["product_url_host"])
    if not filters:
        return []
    from sqlalchemy import or_

    return (
        db.query(ProductIdentityMapping)
        .filter(ProductIdentityMapping.status.in_(["active", "verified", "confirmed"]))
        .filter(or_(*filters))
        .limit(20)
        .all()
    )


def _mapping_uid(key: dict[str, str], product: KBProduct, sku_code: str) -> str:
    raw = "|".join([
        key.get("platform_item_id_hash", ""),
        key.get("platform_item_id", ""),
        key.get("product_url_host", ""),
        product.i_id or "",
        sku_code or "",
    ])
    return f"pim_{stable_hash(raw, 18)}"


def run_import(
    input_path: str,
    *,
    apply: bool = False,
    source_run_uid: str = "",
    operator: str = "",
    json_output: str = "",
    db_factory=None,
) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    rows = _read_rows(input_path)
    db = db_factory()
    matched = 0
    skipped: list[dict[str, Any]] = []
    updated: list[str] = []
    changed_fields: dict[str, int] = {}
    try:
        for row in rows:
            row_no = row.get("_row", "")
            if not _is_confirmed(row.get("confirmation_status", "")):
                skipped.append({"row": row_no, "reason": "confirmation_status_not_confirmed"})
                continue
            key = _mapping_key(row)
            if not any(key.get(field) for field in ("platform_item_id_hash", "platform_item_id", "product_url")):
                skipped.append({"row": row_no, "reason": "missing_stable_platform_identity"})
                continue
            product = _find_product(db, row.get("i_id", ""), row.get("sku_code", ""))
            if not product:
                skipped.append({"row": row_no, "reason": "kb_product_not_found"})
                continue
            existing = _existing_mappings(db, key)
            product_ids = {item.kb_product_id for item in existing if item.kb_product_id}
            if product_ids and product.id not in product_ids:
                skipped.append({"row": row_no, "reason": "ambiguous_existing_mapping"})
                continue
            mapping_uid = existing[0].mapping_uid if existing else _mapping_uid(key, product, row.get("sku_code", ""))
            metadata = sanitize_obj({
                "source": "manual_identity_mapping_import",
                "operator": operator or row.get("operator", ""),
                "import_time": _now_iso(),
                "source_file": str(input_path),
                "dry_run": not apply,
                "source_run_uid": source_run_uid or row.get("source_run_uid", ""),
                "case_uid": row.get("case_uid", ""),
                "turn_uid": row.get("turn_uid", ""),
                "suggested_product_title": row.get("suggested_product_title", ""),
                "note": row.get("note", ""),
            })
            matched += 1
            if not apply:
                updated.append(mapping_uid)
                continue
            mapping = existing[0] if existing else ProductIdentityMapping(mapping_uid=mapping_uid)
            mapping.platform = "taobao" if "taobao" in row.get("product_url", "") else ""
            mapping.platform_item_id = key.get("platform_item_id", "")
            mapping.platform_item_id_hash = key.get("platform_item_id_hash", "")
            mapping.product_url_host = key.get("product_url_host", "")
            mapping.platform_product_title = row.get("platform_product_title") or row.get("order_product_title") or ""
            mapping.kb_product_id = product.id
            mapping.i_id = product.i_id
            mapping.sku_code = row.get("sku_code", "")
            mapping.confidence = 1.0
            mapping.status = "active"
            mapping.source = "manual_import"
            mapping.created_by = operator or row.get("operator", "")
            mapping.set_metadata(metadata)
            db.add(mapping)
            updated.append(mapping_uid)
            for field in ("platform_item_id_hash", "platform_item_id", "product_url_host", "i_id", "sku_code"):
                if getattr(mapping, field, ""):
                    changed_fields[field] = changed_fields.get(field, 0) + 1
        if apply:
            db.commit()
        result = sanitize_obj({
            "ok": True,
            "apply": apply,
            "dry_run": not apply,
            "matched_count": matched,
            "skipped_count": len(skipped),
            "error_count": 0,
            "updated_mapping_uids": updated,
            "skipped_reasons": skipped,
            "changed_fields_summary": changed_fields,
        })
    except Exception as exc:
        if apply:
            db.rollback()
        result = sanitize_obj({
            "ok": False,
            "apply": apply,
            "dry_run": not apply,
            "matched_count": matched,
            "skipped_count": len(skipped),
            "error_count": 1,
            "updated_mapping_uids": updated,
            "skipped_reasons": skipped,
            "error": f"{type(exc).__name__}: {exc}",
        })
    finally:
        db.close()
    if json_output:
        Path(json_output).parent.mkdir(parents=True, exist_ok=True)
        Path(json_output).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Import confirmed product identity mappings.")
    parser.add_argument("--input", required=True, help="Filled mapping xlsx path.")
    parser.add_argument("--run-uid", default="", help="Optional source replay run_uid.")
    parser.add_argument("--operator", default="", help="Operator name for audit metadata.")
    parser.add_argument("--apply", action="store_true", help="Write mappings. Omit for dry-run.")
    parser.add_argument("--json-output", default="", help="Optional JSON output path.")
    args = parser.parse_args()
    init_db()
    result = run_import(
        args.input,
        apply=args.apply,
        source_run_uid=args.run_uid,
        operator=args.operator,
        json_output=args.json_output,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
