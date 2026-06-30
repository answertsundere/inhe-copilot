"""Backfill trusted platform product identity mappings.

The script is intentionally conservative:
- it only writes when a row has an explicit platform item id/hash/url and a
  unique internal i_id or SKU match;
- product titles are exported as review context only and are never used as an
  automatic match key;
- default mode is dry-run, --apply is required to write.
"""

from __future__ import annotations

import argparse
import json
import sys
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.db import SessionLocal, init_db
from app.models.kb_tables import KBProduct, ProductIdentityMapping
from app.services.eval_sanitizer_service import hash_sensitive, sanitize_obj, sanitize_text


INPUT_ALIASES = {
    "platform_item_id": {"platform_item_id", "平台商品 ID", "平台商品ID", "商品ID"},
    "platform_item_id_hash": {"platform_item_id_hash", "平台商品 ID Hash", "平台商品ID Hash", "商品ID Hash"},
    "product_url": {"product_url", "商品链接", "平台商品链接"},
    "platform_product_title": {"platform_product_title", "平台商品标题"},
    "i_id": {"i_id", "内部 i_id", "建议内部 i_id", "内部商品ID"},
    "sku_code": {"sku_code", "SKU", "建议 SKU", "内部 SKU"},
    "confirmation_status": {"confirmation_status", "人工确认状态", "确认状态"},
    "source": {"source", "来源", "数据来源"},
}
CONFIRMED_STATUSES = {"", "confirmed", "active", "verified", "已确认", "确认", "有效"}


def _read_rows(path: str) -> list[dict[str, Any]]:
    if not sanitize_text(path):
        return []
    source = Path(path)
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else data.get("items", []) if isinstance(data, dict) else []
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = _select_sheet(workbook)
        headers = [sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            raw = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
            rows.append(_normalize_row(raw))
        return rows
    finally:
        workbook.close()


def _select_sheet(workbook):
    platform_headers = INPUT_ALIASES["platform_item_id_hash"] | INPUT_ALIASES["platform_item_id"] | INPUT_ALIASES["product_url"]
    internal_headers = INPUT_ALIASES["i_id"] | INPUT_ALIASES["sku_code"]
    fallback = None
    for sheet in workbook.worksheets:
        headers = {sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))}
        if headers & platform_headers:
            return sheet
        if fallback is None and headers & internal_headers:
            fallback = sheet
    if fallback is not None:
        return fallback
    return workbook[workbook.sheetnames[0]]


def _normalize_row(raw: dict[str, Any]) -> dict[str, str]:
    row: dict[str, str] = {}
    for key, aliases in INPUT_ALIASES.items():
        for alias in aliases:
            if alias in raw:
                row[key] = sanitize_text(str(raw.get(alias) or ""))
                break
        row.setdefault(key, "")
    return row


def _url_host(value: str) -> str:
    try:
        return sanitize_text(urlsplit(value or "").netloc)
    except Exception:
        return ""


def _product_for_iid_or_sku(db, *, i_id: str = "", sku_code: str = "") -> tuple[KBProduct | None, str]:
    candidates: dict[int, KBProduct] = {}
    if sanitize_text(i_id):
        product = db.query(KBProduct).filter(KBProduct.i_id == sanitize_text(i_id)).one_or_none()
        if product:
            candidates[product.id] = product
    if sanitize_text(sku_code):
        sku = sanitize_text(sku_code)
        for product in db.query(KBProduct).all():
            if any(sanitize_text(item.get("sku_code")) == sku for item in product.get_sku_list() if isinstance(item, dict)):
                candidates[product.id] = product
    if len(candidates) == 1:
        return next(iter(candidates.values())), ""
    if len(candidates) > 1:
        return None, "ambiguous_internal_identity"
    return None, "kb_product_not_found"


def _mapping_lookup(db, *, platform_item_id: str, platform_item_id_hash: str):
    filters = []
    if platform_item_id:
        filters.append(ProductIdentityMapping.platform_item_id == platform_item_id)
    if platform_item_id_hash:
        filters.append(ProductIdentityMapping.platform_item_id_hash == platform_item_id_hash)
    if not filters:
        return []
    from sqlalchemy import or_

    return db.query(ProductIdentityMapping).filter(or_(*filters), ProductIdentityMapping.status == "active").all()


def _row_identity(row: dict[str, str]) -> tuple[str, str]:
    platform_item_id = sanitize_text(row.get("platform_item_id"))
    platform_item_id_hash = sanitize_text(row.get("platform_item_id_hash"))
    if platform_item_id and not platform_item_id_hash:
        platform_item_id_hash = hash_sensitive(platform_item_id)
    return platform_item_id, platform_item_id_hash


def run_backfill(input_path: str, *, apply: bool = False, operator: str = "auto_backfill", db_factory=None) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    rows = _read_rows(input_path)
    db = db_factory()
    stats = Counter()
    skipped: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    written: list[str] = []
    try:
        for index, row in enumerate(rows, start=2):
            stats["checked_count"] += 1
            status = sanitize_text(row.get("confirmation_status")).lower()
            if status not in CONFIRMED_STATUSES:
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": "not_confirmed"})
                continue
            platform_item_id, platform_item_id_hash = _row_identity(row)
            if not (platform_item_id or platform_item_id_hash or sanitize_text(row.get("product_url"))):
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": "missing_platform_identity"})
                continue
            product, reason = _product_for_iid_or_sku(db, i_id=row.get("i_id", ""), sku_code=row.get("sku_code", ""))
            if not product:
                stats["skipped_count"] += 1
                skipped.append({"row": index, "reason": reason})
                continue
            existing = _mapping_lookup(db, platform_item_id=platform_item_id, platform_item_id_hash=platform_item_id_hash)
            conflicting = [item for item in existing if item.kb_product_id != product.id]
            if conflicting:
                stats["conflict_count"] += 1
                conflicts.append({"row": index, "reason": "identity_conflict", "mapping_uids": [item.mapping_uid for item in conflicting]})
                continue
            if existing:
                stats["matched_count"] += 1
                continue
            mapping_uid = f"pim_auto_{uuid.uuid4().hex[:16]}"
            stats["matched_count"] += 1
            stats["written_mapping_count"] += 1
            written.append(mapping_uid)
            if apply:
                mapping = ProductIdentityMapping(
                    mapping_uid=mapping_uid,
                    platform="auto_trusted_source",
                    platform_item_id=platform_item_id,
                    platform_item_id_hash=platform_item_id_hash,
                    product_url_host=_url_host(row.get("product_url", "")),
                    platform_product_title=sanitize_text(row.get("platform_product_title")),
                    kb_product_id=product.id,
                    i_id=product.i_id,
                    sku_code=sanitize_text(row.get("sku_code")) or product.i_id,
                    confidence=1.0,
                    status="active",
                    source=sanitize_text(row.get("source")) or "trusted_auto_backfill",
                    created_by=operator,
                )
                mapping.set_metadata(sanitize_obj({
                    "operator": operator,
                    "imported_at": datetime.utcnow().isoformat(),
                    "source_file": str(input_path),
                    "rule": "explicit_platform_identity_and_unique_internal_identity",
                    "writes_verified_knowledge": False,
                }))
                db.add(mapping)
        if apply:
            db.commit()
        else:
            db.rollback()
        return sanitize_obj({
            "ok": True,
            "dry_run": not apply,
            "matched_count": stats["matched_count"],
            "skipped_count": stats["skipped_count"],
            "conflict_count": stats["conflict_count"],
            "written_mapping_count": stats["written_mapping_count"],
            "skipped_reasons": skipped,
            "conflicts": conflicts,
            "written_mapping_uids": written,
            "writes_verified_knowledge": False,
        })
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Backfill trusted product identity mappings.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--operator", default="auto_backfill")
    parser.add_argument("--json-output", "--output", dest="json_output", default="")
    args = parser.parse_args()
    init_db()
    result = run_backfill(args.input, apply=args.apply, operator=args.operator)
    if args.json_output:
        output = Path(args.json_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
