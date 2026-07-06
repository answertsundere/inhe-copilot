"""Diagnose why real replay turns end with zero selected evidence.

This script is intentionally read-only. It inspects EvalTrace payloads and the
current KB tables, then exports a human-readable evidence-chain report.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def _load_dotenv_safely() -> None:
    try:
        from dotenv import load_dotenv

        load_dotenv(PROJECT_ROOT / ".env")
    except Exception:
        pass


_load_dotenv_safely()

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from sqlalchemy import or_  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models.eval_tables import EvalRun, EvalTrace  # noqa: E402
from app.models.kb_tables import KBGenericServiceRule, KBMediaAsset, KBProduct  # noqa: E402
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text  # noqa: E402
from app.services.media_asset_service import get_auto_send_level, get_media_purpose  # noqa: E402


PRIMARY_REASONS = {
    "sidecar_missing_or_not_passed",
    "product_identity_unresolved",
    "kb_product_missing",
    "structured_field_missing",
    "media_missing",
    "media_role_not_sendable",
    "generic_rule_missing",
    "generic_rule_blocked_by_product_fact",
    "embedding_not_configured",
    "rag_timeout",
    "retrieval_filter_too_strict",
    "evidence_role_mismatch",
    "query_fact_type_missing",
    "final_gate_blocked",
    "unknown",
}

STRUCTURED_FACT_ALIASES: dict[str, list[str]] = {
    "material": ["material", "材质"],
    "material_safety": ["material", "材质", "safety", "安全", "检测", "证书"],
    "certification_report": ["certificate", "certification", "report", "检测", "证书", "报告"],
    "dimensions": ["dimensions", "dimension", "size", "length", "width", "height", "尺寸", "长", "宽", "高"],
    "space_fit": ["dimensions", "dimension", "size", "尺寸", "长", "宽", "高", "空间"],
    "placement_scene": ["placement", "scene", "适用场景", "摆放", "阳台", "场景"],
    "load_capacity": ["load_capacity", "capacity", "bearing", "承重", "载重"],
    "gross_weight": ["gross_weight", "package_weight", "weight", "毛重", "重量"],
    "accessories": ["accessories", "parts", "配件", "清单"],
    "accessory_usage": ["accessories", "parts", "配件", "用途", "使用"],
    "accessory_availability": ["accessories", "parts", "配件", "补配", "单独"],
    "accessory_compatibility": ["accessories", "parts", "配件", "适配", "加装"],
    "installation": ["installation", "install", "manual", "guide", "安装", "说明书", "教程"],
    "age_range": ["age_range", "age", "适用年龄", "年龄", "宝宝", "儿童"],
    "child_suitability": ["age_range", "age", "适用年龄", "宝宝", "儿童"],
    "child_safety": ["safety", "安全", "宝宝", "儿童", "防夹", "防倒"],
    "structure_function": ["structure_function", "structure", "function", "结构", "功能"],
}

MEDIA_FACT_TYPES = {
    "installation",
    "installation_media_request",
    "accessory_usage",
    "accessory_availability",
    "accessory_compatibility",
    "dimensions",
    "space_fit",
}

POLICY_FACT_TYPES = {
    "promotion",
    "promotion_policy",
    "price_negotiation",
    "aftersales",
    "aftersales_policy",
    "invoice_policy",
    "stock_shipping",
    "logistics",
    "price_protection",
}

INSTALL_MEDIA_PURPOSES = {
    "install_video",
    "installation_video",
    "video",
    "install_image",
    "installation_image",
    "installation_guide",
    "manual",
    "manual_image",
    "pack_guide_image",
}

SIZE_MEDIA_PURPOSES = {
    "size_chart",
    "dimension_image",
    "space_fit_image",
}

ACCESSORY_MEDIA_PURPOSES = {
    "accessory_photo",
    "accessory_image",
    "parts_image",
}

SUMMARY_COLUMNS = [
    ("指标", "metric"),
    ("数值", "value"),
]

DETAIL_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("买家问题预览", "buyer_message_preview"),
    ("问题类型", "query_fact_type"),
    ("质量分桶", "quality_bucket"),
    ("失败标签", "failure_labels"),
    ("选中证据数", "selected_evidence_count"),
    ("主原因", "primary_reason"),
    ("次要原因", "secondary_reasons"),
    ("侧栏质量", "sidecar_context_quality"),
    ("侧栏商品", "sidecar_product_title"),
    ("侧栏 SKU", "sidecar_sku_code"),
    ("侧栏 i_id", "sidecar_i_id"),
    ("侧栏订单", "sidecar_order_id"),
    ("商品身份状态", "identity_status"),
    ("身份置信度", "identity_confidence"),
    ("匹配原因", "match_reason"),
    ("未解析原因", "unresolved_reason"),
    ("KB 商品存在", "kb_product_found"),
    ("KB 商品 i_id", "kb_product_i_id"),
    ("结构化字段覆盖", "structured_field_coverage"),
    ("缺失结构化字段", "missing_structured_fields"),
    ("素材总数", "media_total_count"),
    ("可用素材数", "media_available_count"),
    ("可发送匹配素材数", "media_sendable_match_count"),
    ("素材问题", "media_issue"),
    ("Generic 规则数", "generic_rule_count"),
    ("Product-first Pack", "has_product_first_pack"),
    ("Pack 字段数", "pack_structured_facts_count"),
    ("Pack 素材数", "pack_media_assets_count"),
    ("Pack Chunks 数", "pack_scoped_chunks_count"),
    ("Pack Generic 数", "pack_generic_rules_count"),
    ("缺失证据", "missing_required_evidence"),
    ("Block Reasons", "block_reasons"),
]

IDENTITY_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("侧栏商品", "sidecar_product_title"),
    ("侧栏 SKU", "sidecar_sku_code"),
    ("侧栏 i_id", "sidecar_i_id"),
    ("商品身份状态", "identity_status"),
    ("身份来源", "identity_sources"),
    ("置信度", "identity_confidence"),
    ("匹配原因", "match_reason"),
    ("未解析原因", "unresolved_reason"),
    ("候选商品", "ambiguous_candidates"),
    ("主原因", "primary_reason"),
]

FIELD_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("KB 商品 i_id", "kb_product_i_id"),
    ("KB 商品标题", "kb_product_name"),
    ("缺失字段", "missing_structured_fields"),
    ("已有字段覆盖", "structured_field_coverage"),
    ("建议动作", "suggested_action"),
]

MEDIA_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("KB 商品 i_id", "kb_product_i_id"),
    ("素材总数", "media_total_count"),
    ("可用素材数", "media_available_count"),
    ("可发送匹配素材数", "media_sendable_match_count"),
    ("素材详情", "media_summary"),
    ("素材问题", "media_issue"),
    ("建议动作", "suggested_action"),
]

GENERIC_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("Generic 规则数", "generic_rule_count"),
    ("Generic 规则标题", "generic_rule_titles"),
    ("主原因", "primary_reason"),
    ("建议动作", "suggested_action"),
]

RAG_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("Embedding 状态", "embedding_config_status"),
    ("缺失环境变量", "embedding_missing_env_keys"),
    ("RAG 超时", "rag_timeout"),
    ("候选证据数", "rag_candidate_count"),
    ("过滤后证据数", "rag_filtered_count"),
    ("选中证据数", "selected_evidence_count"),
    ("主原因", "primary_reason"),
]

AUTO_FIX_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("主原因", "primary_reason"),
    ("可自动修复候选", "auto_fix_candidate"),
    ("建议动作", "suggested_action"),
]

MANUAL_COLUMNS = [
    ("回放批次", "run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("主原因", "primary_reason"),
    ("需要人工确认", "manual_review_reason"),
    ("建议动作", "suggested_action"),
]


def _text(value: Any) -> str:
    return sanitize_text(value).strip()


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _as_list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _json_cell(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(sanitize_obj(value), ensure_ascii=False)


def _dig(obj: Any, *keys: str) -> Any:
    current = obj
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def _first_text(*values: Any) -> str:
    for value in values:
        text = _text(value)
        if text:
            return text
    return ""


def _first_dict(*values: Any) -> dict[str, Any]:
    for value in values:
        if isinstance(value, dict) and value:
            return value
    return {}


def _first_list(*values: Any) -> list[Any]:
    for value in values:
        if isinstance(value, list) and value:
            return value
    return []


def _latest_real_run(db) -> EvalRun | None:
    return (
        db.query(EvalRun)
        .filter(EvalRun.source_type == "real_conversation", EvalRun.status == "completed")
        .order_by(EvalRun.created_at.desc(), EvalRun.id.desc())
        .first()
    )


def _quality_bucket(trace: EvalTrace) -> str:
    quality = trace.get_quality_bucket() or {}
    return _text(quality.get("quality_bucket")) or ("auto_sendable" if trace.passed else "agent_error")


def _query_fact_type(trace: EvalTrace) -> str:
    turn = trace.get_turn_understanding() or {}
    answer = trace.get_answer_trace() or {}
    raw = trace.get_raw_response() or {}
    return (
        _first_text(
            trace.query_fact_type,
            turn.get("effective_query_fact_type"),
            turn.get("actual_query_fact_type"),
            turn.get("query_fact_type"),
            answer.get("effective_query_fact_type"),
            answer.get("actual_query_fact_type"),
            answer.get("query_fact_type"),
            raw.get("query_fact_type"),
        )
        or "unknown"
    )


def _selected_evidence(trace: EvalTrace) -> list[Any]:
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    debug = _as_dict(raw.get("evidence_debug"))
    selected = _first_list(
        trace.get_selected_evidence(),
        debug.get("selected_evidence"),
        debug.get("evidence_selected"),
        raw.get("selected_evidence"),
        answer.get("selected_evidence"),
    )
    if selected:
        return selected
    direct: list[Any] = []
    for key in ("knowledge_evidence_summary", "filtered_evidence_summary"):
        direct.extend(item for item in _as_list(debug.get(key)) if _is_direct_answer_evidence(item))
    pack = _extract_product_context_pack(raw, answer)
    for bucket in ("product_structured_facts", "product_scoped_chunks"):
        direct.extend(item for item in _as_list(pack.get(bucket)) if _is_direct_answer_evidence(item))
    return direct


def _selected_evidence_count(trace: EvalTrace) -> int:
    selected = _selected_evidence(trace)
    if selected:
        return len(selected)
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    for value in (
        raw.get("selected_evidence_count"),
        answer.get("selected_evidence_count"),
        _dig(raw, "evidence_debug", "selected_evidence_count"),
    ):
        try:
            if value is not None:
                return int(value)
        except (TypeError, ValueError):
            continue
    return 0


def _extract_product_context_pack(raw: dict[str, Any], answer: dict[str, Any]) -> dict[str, Any]:
    candidates = [
        raw.get("product_first_evidence_pack"),
        _dig(raw, "product_context_pack", "product_first_evidence_pack"),
        _dig(raw, "product_context_pack", "evidence_pack"),
        _dig(raw, "evidence_debug", "product_context_pack_summary", "product_first_evidence_pack"),
        _dig(raw, "evidence_debug", "product_context_pack_summary", "evidence_pack"),
        answer.get("product_first_evidence_pack"),
        _dig(answer, "product_context_pack", "product_first_evidence_pack"),
        _dig(answer, "product_context_pack", "evidence_pack"),
        _dig(answer, "evidence_debug", "product_context_pack_summary", "product_first_evidence_pack"),
        _dig(answer, "evidence_debug", "product_context_pack_summary", "evidence_pack"),
    ]
    return _first_dict(*candidates)


def _is_direct_answer_evidence(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    if item.get("reference_only") is True:
        return False
    if _text(item.get("gate_status")).lower() in {"blocked", "reference_only"}:
        return False
    for key in ("direct_answer_allowed", "can_direct_answer", "evidence_allowed_for_exact_answer"):
        if item.get(key) is False:
            return False
    return bool(
        item.get("chunk_text")
        or item.get("chunk_preview")
        or item.get("preview")
        or item.get("content")
        or item.get("fact")
        or item.get("evidence_id")
        or item.get("chunk_id")
        or item.get("entry_id")
    )


def _extract_sidecar(trace: EvalTrace) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    turn = trace.get_turn_understanding() or {}
    identity = trace.get_product_identity() or {}
    context_sufficiency = _as_dict(turn.get("context_sufficiency"))
    sidecar = _first_dict(
        raw.get("sidecar_context"),
        answer.get("sidecar_context"),
        _dig(raw, "copilot_context", "sidecar_context"),
        _dig(answer, "copilot_context", "sidecar_context"),
        _dig(identity, "sidecar_context"),
    )
    quality = _first_text(
        sidecar.get("sidecar_context_quality"),
        raw.get("sidecar_context_quality"),
        answer.get("sidecar_context_quality"),
        turn.get("sidecar_context_quality"),
        context_sufficiency.get("sidecar_context_quality"),
    )
    sources = _first_list(
        sidecar.get("sidecar_context_sources"),
        raw.get("sidecar_context_sources"),
        answer.get("sidecar_context_sources"),
        turn.get("sidecar_context_sources"),
        context_sufficiency.get("sidecar_context_sources"),
    )
    product_title = _first_text(
        sidecar.get("product_title"),
        sidecar.get("product_name"),
        _dig(identity, "sidecar_context", "product_title"),
        _dig(identity, "sidecar_context", "product_name"),
    )
    sku_code = _first_text(sidecar.get("sku_code"), _dig(identity, "sidecar_context", "sku_code"))
    i_id = _first_text(sidecar.get("i_id"), _dig(identity, "sidecar_context", "i_id"))
    order_id = _first_text(
        sidecar.get("order_id"),
        sidecar.get("platform_order_id"),
        _dig(identity, "sidecar_context", "order_id"),
        trace.order_identity_hash,
    )
    product_candidates = _first_list(sidecar.get("product_candidates"), _dig(identity, "sidecar_context", "product_candidates"))
    return sanitize_obj(
        {
            "product_title": product_title,
            "sku_code": sku_code,
            "i_id": i_id,
            "order_id": order_id,
            "sidecar_context_quality": quality or ("complete" if (product_title or sku_code or i_id or order_id) else "missing"),
            "sidecar_context_sources": sources,
            "product_candidates": product_candidates,
            "has_sidecar_product_context": bool(product_title or sku_code or i_id or product_candidates),
            "has_sidecar_order_context": bool(order_id),
            "in_copilot_context": bool(_dig(raw, "copilot_context", "sidecar_context") or _dig(answer, "copilot_context", "sidecar_context")),
        }
    )


def _identity_resolution(trace: EvalTrace, pack: dict[str, Any], sidecar: dict[str, Any]) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    product_identity = trace.get_product_identity() or {}
    resolution = _first_dict(
        pack.get("product_identity_resolution"),
        _dig(pack, "identity", "product_identity_resolution"),
        _dig(pack, "evidence_pack_trace", "product_identity_resolution"),
        raw.get("product_identity_resolution"),
        answer.get("product_identity_resolution"),
        _dig(raw, "product_context_pack", "product_identity_resolution"),
        _dig(answer, "product_context_pack", "product_identity_resolution"),
    )
    resolved = _first_dict(pack.get("resolved_product_identity"), _dig(pack, "identity"), product_identity)
    status = _first_text(resolution.get("status"), "resolved" if resolved.get("resolved_product_id") or resolved.get("product_id") else "")
    confidence = resolution.get("identity_confidence", resolution.get("confidence", pack.get("identity_confidence", resolved.get("identity_confidence", ""))))
    sources = _first_list(resolution.get("identity_sources"), pack.get("identity_sources"), resolved.get("identity_sources"))
    return sanitize_obj(
        {
            "resolver_executed": bool(resolution or pack),
            "status": status or "unknown",
            "resolved_product_id": resolved.get("resolved_product_id") or resolved.get("product_id") or resolution.get("resolved_product_id"),
            "i_id": _first_text(resolution.get("i_id"), resolution.get("internal_i_id"), resolved.get("i_id"), sidecar.get("i_id")),
            "sku_code": _first_text(resolution.get("sku_code"), resolution.get("sku_id"), resolved.get("sku_code"), resolved.get("sku"), sidecar.get("sku_code")),
            "display_product_name": _first_text(
                resolution.get("display_product_name"),
                resolution.get("canonical_product_name"),
                resolved.get("product_name"),
                resolved.get("display_product_name"),
                sidecar.get("product_title"),
            ),
            "identity_confidence": confidence,
            "identity_sources": sources,
            "match_reason": _first_text(resolution.get("match_reason"), resolution.get("reason"), resolved.get("match_reason")),
            "unresolved_reason": _first_text(resolution.get("unresolved_reason"), resolved.get("unresolved_reason")),
            "ambiguous_candidates": _first_list(resolution.get("ambiguous_candidates"), resolution.get("candidates")),
        }
    )


def _find_kb_product(db, identity: dict[str, Any], sidecar: dict[str, Any]) -> KBProduct | None:
    product_id = identity.get("resolved_product_id")
    if product_id:
        product = db.query(KBProduct).filter(KBProduct.id == product_id).first()
        if product:
            return product
    for value in (identity.get("i_id"), sidecar.get("i_id"), identity.get("sku_code"), sidecar.get("sku_code")):
        text = _text(value)
        if not text:
            continue
        product = db.query(KBProduct).filter(KBProduct.i_id == text).first()
        if product:
            return product
        product = db.query(KBProduct).filter(KBProduct.sku_list_json.contains(text)).first()
        if product:
            return product
    title = _first_text(identity.get("display_product_name"), sidecar.get("product_title"))
    if title:
        product = db.query(KBProduct).filter(KBProduct.product_name == title).first()
        if product:
            return product
        product = db.query(KBProduct).filter(KBProduct.product_name.contains(title[:20])).first()
        if product:
            return product
    return None


def _contains_value(obj: Any, aliases: list[str]) -> bool:
    lowered_aliases = [alias.lower() for alias in aliases]
    if isinstance(obj, dict):
        for key, value in obj.items():
            key_text = str(key).lower()
            if any(alias in key_text for alias in lowered_aliases) and value not in (None, "", [], {}):
                return True
            if _contains_value(value, aliases):
                return True
    elif isinstance(obj, list):
        return any(_contains_value(item, aliases) for item in obj)
    return False


def _structured_coverage(product: KBProduct | None) -> dict[str, bool]:
    if not product:
        return {field: False for field in STRUCTURED_FACT_ALIASES}
    profile = {
        "specs": product.get_specs(),
        "logistics": product.get_logistics(),
        "warranty": product.get_warranty(),
        "missing_fields": product.get_missing_fields(),
        "sku_list": product.get_sku_list(),
    }
    return {field: _contains_value(profile, aliases) for field, aliases in STRUCTURED_FACT_ALIASES.items()}


def _required_structured_fields(query_fact_type: str) -> list[str]:
    qft = _text(query_fact_type)
    if qft in STRUCTURED_FACT_ALIASES:
        return [qft]
    if qft in {"product_question", "unknown"}:
        return []
    return []


def _media_role_matches(query_fact_type: str, asset_type: str, purpose: str, title: str) -> bool:
    qft = _text(query_fact_type)
    role = _text(purpose or asset_type).lower()
    asset = _text(asset_type).lower()
    title_lower = _text(title).lower()
    if qft in {"installation", "installation_media_request"}:
        return role in INSTALL_MEDIA_PURPOSES or asset in INSTALL_MEDIA_PURPOSES
    if qft in {"dimensions", "space_fit"}:
        if role in SIZE_MEDIA_PURPOSES or asset in SIZE_MEDIA_PURPOSES:
            return True
        return asset in {"sku_image", "product_photo"} and any(token in title_lower for token in ["尺寸", "size", "dimension"])
    if qft in {"accessory_usage", "accessory_availability", "accessory_compatibility"}:
        return role in ACCESSORY_MEDIA_PURPOSES or asset in ACCESSORY_MEDIA_PURPOSES
    return True


def _asset_is_usable(asset: KBMediaAsset) -> bool:
    return bool(asset.status == "approved" and asset.usable_for_agent == 1 and _text(asset.refresh_status) not in {"needs_refresh", "error"})


def _collect_media(db, product: KBProduct | None, identity: dict[str, Any], sidecar: dict[str, Any], query_fact_type: str) -> dict[str, Any]:
    filters = []
    if product:
        filters.append(KBMediaAsset.product_id == product.id)
        if product.i_id:
            filters.append(KBMediaAsset.i_id == product.i_id)
        if product.product_name:
            filters.append(KBMediaAsset.product_name == product.product_name)
    for field in (identity.get("i_id"), identity.get("sku_code"), sidecar.get("i_id"), sidecar.get("sku_code")):
        text = _text(field)
        if text:
            filters.append(KBMediaAsset.i_id == text)
            filters.append(KBMediaAsset.sku_code == text)
    if not filters:
        return {"items": [], "total_count": 0, "available_count": 0, "sendable_match_count": 0, "issue": "missing_product_identity"}
    assets = db.query(KBMediaAsset).filter(or_(*filters)).order_by(KBMediaAsset.id.asc()).all()
    items: list[dict[str, Any]] = []
    for asset in assets:
        purpose = get_media_purpose(asset)
        auto_send_level = get_auto_send_level(asset)
        role_match = _media_role_matches(query_fact_type, asset.asset_type, purpose, asset.asset_title)
        usable = _asset_is_usable(asset)
        sendable_match = usable and role_match and auto_send_level == "auto" and bool(_text(asset.asset_url))
        items.append(
            sanitize_obj(
                {
                    "asset_id": asset.id,
                    "asset_type": asset.asset_type,
                    "media_purpose": purpose,
                    "asset_title": asset.asset_title,
                    "status": asset.status,
                    "audit_status": asset.audit_status,
                    "usable_for_agent": bool(asset.usable_for_agent),
                    "auto_send_level": auto_send_level,
                    "refresh_status": asset.refresh_status,
                    "has_asset_url": bool(_text(asset.asset_url)),
                    "role_matches_query": role_match,
                    "sendable_match": sendable_match,
                }
            )
        )
    available = [item for item in items if item["status"] == "approved" and item["usable_for_agent"] and item["refresh_status"] not in {"needs_refresh", "error"}]
    sendable = [item for item in items if item["sendable_match"]]
    if not items:
        issue = "media_missing"
    elif not available:
        issue = "media_not_approved_or_usable"
    elif not any(item["role_matches_query"] for item in available):
        issue = "media_role_mismatch"
    elif not sendable:
        issue = "media_not_auto_sendable"
    else:
        issue = ""
    return {
        "items": items,
        "total_count": len(items),
        "available_count": len(available),
        "sendable_match_count": len(sendable),
        "issue": issue,
    }


def _collect_generic_rules(db, query_fact_type: str) -> dict[str, Any]:
    qft = _text(query_fact_type)
    if not qft or qft == "unknown":
        return {"items": [], "count": 0}
    rules = (
        db.query(KBGenericServiceRule)
        .filter(KBGenericServiceRule.status == "active", KBGenericServiceRule.fact_type == qft)
        .order_by(KBGenericServiceRule.priority.asc(), KBGenericServiceRule.id.asc())
        .all()
    )
    return {
        "items": [
            sanitize_obj(
                {
                    "rule_key": rule.rule_key,
                    "title": rule.title,
                    "fact_type": rule.fact_type,
                    "allowed_when_product_fact_missing": bool(rule.allowed_when_product_fact_missing),
                    "auto_reply_allowed": bool(rule.auto_reply_allowed),
                    "risk_level": rule.risk_level,
                }
            )
            for rule in rules
        ],
        "count": len(rules),
    }


def _embedding_status() -> dict[str, Any]:
    enabled = _text(os.environ.get("COPILOT_EMBEDDING_ENABLED")).lower() in {"1", "true", "yes", "on"}
    required = [
        "COPILOT_EMBEDDING_ENABLED",
        "COPILOT_EMBEDDING_API_BASE",
        "COPILOT_EMBEDDING_API_KEY",
        "COPILOT_EMBEDDING_MODEL",
    ]
    missing = [key for key in required if not _text(os.environ.get(key))]
    return {
        "embedding_enabled": enabled,
        "embedding_api_base_configured": bool(_text(os.environ.get("COPILOT_EMBEDDING_API_BASE"))),
        "embedding_api_key_configured": bool(_text(os.environ.get("COPILOT_EMBEDDING_API_KEY"))),
        "embedding_model_configured": bool(_text(os.environ.get("COPILOT_EMBEDDING_MODEL"))),
        "missing_env_keys": missing,
        "status": "configured" if enabled and not missing[1:] else "not_configured",
    }


def _pack_counts(pack: dict[str, Any]) -> dict[str, Any]:
    return {
        "has_product_first_pack": bool(pack),
        "pack_structured_facts_count": len(_as_list(pack.get("product_structured_facts"))),
        "pack_media_assets_count": len(_as_list(pack.get("product_media_assets"))),
        "pack_scoped_chunks_count": len(_as_list(pack.get("product_scoped_chunks"))),
        "pack_generic_rules_count": len(_as_list(pack.get("generic_fallback_rules"))),
        "missing_required_evidence": _as_list(pack.get("missing_required_evidence")),
    }


def _trace_rag_numbers(raw: dict[str, Any]) -> dict[str, int | bool]:
    debug = _as_dict(raw.get("evidence_debug"))
    evidence_lists = [
        _as_list(debug.get("knowledge_evidence_summary")),
        _as_list(debug.get("filtered_evidence_summary")),
    ]
    summaries = [_as_dict(debug.get("evidence_gate_summary"))]
    text = json.dumps(sanitize_obj(debug), ensure_ascii=False).lower()
    timeout = "timeout" in text or "超时" in text
    candidate = max((len(items) for items in evidence_lists), default=0)
    filtered = max((sum(1 for item in items if _is_direct_answer_evidence(item)) for items in evidence_lists), default=0)
    for summary in summaries:
        for key in ("candidate_count", "candidates", "raw_count", "total"):
            try:
                candidate = max(candidate, int(summary.get(key) or 0))
            except (TypeError, ValueError):
                pass
        for key in ("filtered_count", "usable_count", "count"):
            try:
                filtered = max(filtered, int(summary.get(key) or 0))
            except (TypeError, ValueError):
                pass
    return {"rag_timeout": timeout, "rag_candidate_count": candidate, "rag_filtered_count": filtered}


def _block_reasons(trace: EvalTrace) -> list[str]:
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    audit = trace.get_final_audit() or {}
    semantic = trace.get_semantic_compiler() or {}
    reasons: list[Any] = []
    for container in (raw, answer, audit, semantic, _as_dict(raw.get("evidence_debug"))):
        reasons.extend(_as_list(container.get("block_reasons")))
        reasons.extend(_as_list(container.get("issues")))
    return [_text(item) for item in reasons if _text(item)]


def _classify_reason(
    *,
    query_fact_type: str,
    sidecar: dict[str, Any],
    identity: dict[str, Any],
    product: KBProduct | None,
    required_fields: list[str],
    missing_fields: list[str],
    media: dict[str, Any],
    generic_rules: dict[str, Any],
    embedding: dict[str, Any],
    rag: dict[str, Any],
    block_reasons: list[str],
) -> tuple[str, list[str]]:
    reasons: list[str] = []
    qft = _text(query_fact_type)
    if not qft or qft == "unknown":
        reasons.append("query_fact_type_missing")
    if sidecar.get("sidecar_context_quality") in {"missing", "insufficient"} or not sidecar.get("has_sidecar_product_context"):
        reasons.append("sidecar_missing_or_not_passed")
    if identity.get("resolver_executed") and identity.get("status") not in {"resolved", "exact", "matched"}:
        reasons.append("product_identity_unresolved")
    if qft in MEDIA_FACT_TYPES:
        if media.get("total_count", 0) == 0:
            reasons.append("media_missing")
        elif media.get("sendable_match_count", 0) == 0:
            issue = media.get("issue")
            reasons.append("evidence_role_mismatch" if issue == "media_role_mismatch" else "media_role_not_sendable")
    if not product:
        reasons.append("kb_product_missing")
    if required_fields and missing_fields:
        reasons.append("structured_field_missing")
    if qft in POLICY_FACT_TYPES and generic_rules.get("count", 0) == 0:
        reasons.append("generic_rule_missing")
    if embedding.get("status") != "configured":
        reasons.append("embedding_not_configured")
    if rag.get("rag_timeout"):
        reasons.append("rag_timeout")
    if block_reasons:
        reasons.append("final_gate_blocked")
    if not reasons:
        reasons.append("unknown")
    primary = reasons[0]
    if primary not in PRIMARY_REASONS:
        primary = "unknown"
    return primary, [reason for reason in reasons[1:] if reason != primary]


def _analyze_trace(db, trace: EvalTrace, embedding: dict[str, Any]) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    qft = _query_fact_type(trace)
    sidecar = _extract_sidecar(trace)
    pack = _extract_product_context_pack(raw, answer)
    identity = _identity_resolution(trace, pack, sidecar)
    product = _find_kb_product(db, identity, sidecar)
    coverage = _structured_coverage(product)
    required_fields = _required_structured_fields(qft)
    missing_fields = [field for field in required_fields if not coverage.get(field)]
    media = _collect_media(db, product, identity, sidecar, qft)
    generic_rules = _collect_generic_rules(db, qft)
    pack_counts = _pack_counts(pack)
    rag = _trace_rag_numbers(raw)
    block_reasons = _block_reasons(trace)
    selected_count = _selected_evidence_count(trace)
    primary, secondary = _classify_reason(
        query_fact_type=qft,
        sidecar=sidecar,
        identity=identity,
        product=product,
        required_fields=required_fields,
        missing_fields=missing_fields,
        media=media,
        generic_rules=generic_rules,
        embedding=embedding,
        rag=rag,
        block_reasons=block_reasons,
    )
    structured_field_coverage = {field: bool(value) for field, value in coverage.items() if value}
    record = {
        "run_uid": trace.run_uid,
        "case_uid": trace.case_uid,
        "turn_uid": trace.turn_uid,
        "turn_index": trace.turn_index,
        "buyer_message_preview": _text(trace.buyer_message)[:120],
        "query_fact_type": qft,
        "quality_bucket": _quality_bucket(trace),
        "failure_labels": trace.get_failure_labels() or [],
        "selected_evidence_count": selected_count,
        "primary_reason": primary,
        "secondary_reasons": secondary,
        "sidecar_context_quality": sidecar.get("sidecar_context_quality"),
        "sidecar_context_sources": sidecar.get("sidecar_context_sources"),
        "sidecar_product_title": sidecar.get("product_title"),
        "sidecar_sku_code": sidecar.get("sku_code"),
        "sidecar_i_id": sidecar.get("i_id"),
        "sidecar_order_id": sidecar.get("order_id"),
        "sidecar_in_copilot_context": sidecar.get("in_copilot_context"),
        "has_sidecar_product_context": sidecar.get("has_sidecar_product_context"),
        "has_sidecar_order_context": sidecar.get("has_sidecar_order_context"),
        "identity_resolver_executed": identity.get("resolver_executed"),
        "identity_status": identity.get("status"),
        "resolved_product_id": identity.get("resolved_product_id"),
        "identity_i_id": identity.get("i_id"),
        "identity_sku_code": identity.get("sku_code"),
        "identity_display_product_name": identity.get("display_product_name"),
        "identity_confidence": identity.get("identity_confidence"),
        "identity_sources": identity.get("identity_sources"),
        "match_reason": identity.get("match_reason"),
        "unresolved_reason": identity.get("unresolved_reason"),
        "ambiguous_candidates": identity.get("ambiguous_candidates"),
        "kb_product_found": bool(product),
        "kb_product_id": product.id if product else "",
        "kb_product_i_id": product.i_id if product else "",
        "kb_product_name": product.product_name if product else "",
        "structured_field_coverage": structured_field_coverage,
        "missing_structured_fields": missing_fields,
        "media_total_count": media.get("total_count", 0),
        "media_available_count": media.get("available_count", 0),
        "media_sendable_match_count": media.get("sendable_match_count", 0),
        "media_issue": media.get("issue", ""),
        "media_summary": media.get("items", [])[:8],
        "generic_rule_count": generic_rules.get("count", 0),
        "generic_rule_titles": [rule.get("title") or rule.get("rule_key") for rule in generic_rules.get("items", [])],
        "generic_rules": generic_rules.get("items", []),
        "embedding_config_status": embedding.get("status"),
        "embedding_missing_env_keys": embedding.get("missing_env_keys"),
        "rag_timeout": rag.get("rag_timeout"),
        "rag_candidate_count": rag.get("rag_candidate_count", 0),
        "rag_filtered_count": rag.get("rag_filtered_count", 0),
        "block_reasons": block_reasons,
        **pack_counts,
    }
    if primary in {"structured_field_missing", "media_missing", "media_role_not_sendable", "generic_rule_missing"}:
        record["auto_fix_candidate"] = "no"
        record["manual_review_reason"] = "需要运营补字段/素材/规则后复测"
    elif primary == "embedding_not_configured":
        record["auto_fix_candidate"] = "env_config"
        record["manual_review_reason"] = ""
    else:
        record["auto_fix_candidate"] = ""
        record["manual_review_reason"] = "需要工程确认证据链路"
    record["suggested_action"] = _suggested_action(record)
    return sanitize_obj(record)


def _suggested_action(row: dict[str, Any]) -> str:
    reason = row.get("primary_reason")
    if reason == "sidecar_missing_or_not_passed":
        return "补齐千牛侧栏商品名/SKU/i_id/订单号输入；不要用平台 hash 猜商品"
    if reason == "product_identity_unresolved":
        return "补商品身份映射，人工确认平台商品与内部 i_id/SKU 的对应关系"
    if reason == "kb_product_missing":
        return "检查内部商品库是否有该 i_id/SKU；缺商品主数据时先补商品档案"
    if reason == "structured_field_missing":
        return "按缺失字段补商品结构化资料；强事实字段需人工确认"
    if reason in {"media_missing", "media_role_not_sendable", "evidence_role_mismatch"}:
        return "补充或重审素材用途/角色；普通商品图不能冒充安装图、尺寸图或配件图"
    if reason == "generic_rule_missing":
        return "补活动/售后/发票/物流等通用服务规则，不能写成商品事实"
    if reason == "embedding_not_configured":
        return "配置 COPILOT_EMBEDDING_ENABLED/API_BASE/API_KEY/MODEL 后重建/重跑检索"
    if reason == "rag_timeout":
        return "检查 RAG 超时和候选过滤耗时，先确认 embedding 服务可用"
    if reason == "query_fact_type_missing":
        return "修 turn understanding/fact_type 合同，不按具体买家原话特判"
    if reason == "final_gate_blocked":
        return "查看 final gate block_reasons，确认是否缺证据或回复主题错误"
    return "需要继续人工排查该 trace 的 raw_response/evidence_debug"


def diagnose_evidence_chain(
    *,
    run_uid: str = "",
    latest: bool = False,
    limit: int = 0,
    db_factory=None,
) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    db = db_factory()
    try:
        run = None
        if run_uid:
            run = db.query(EvalRun).filter(EvalRun.run_uid == run_uid).first()
        if not run and (latest or not run_uid):
            run = _latest_real_run(db)
        if not run:
            return {"run_uid": run_uid, "summary": {"error": "run_not_found"}, "records": []}
        query = db.query(EvalTrace).filter(EvalTrace.run_uid == run.run_uid).order_by(EvalTrace.case_uid.asc(), EvalTrace.turn_index.asc(), EvalTrace.id.asc())
        if limit and limit > 0:
            query = query.limit(limit)
        traces = query.all()
        embedding = _embedding_status()
        records = [_analyze_trace(db, trace, embedding) for trace in traces]
        zero_records = [row for row in records if int(row.get("selected_evidence_count") or 0) == 0]
        total = len(records)
        summary = {
            "run_uid": run.run_uid,
            "total_traces": total,
            "zero_evidence_trace_count": len(zero_records),
            "zero_evidence_trace_rate": round(len(zero_records) / total, 4) if total else 0,
            "by_query_fact_type": dict(Counter(row.get("query_fact_type") or "unknown" for row in zero_records).most_common()),
            "by_primary_reason": dict(Counter(row.get("primary_reason") or "unknown" for row in zero_records).most_common()),
            "product_identity_resolved_count": sum(1 for row in records if row.get("identity_status") == "resolved"),
            "kb_product_found_count": sum(1 for row in records if row.get("kb_product_found")),
            "structured_field_available_count": sum(1 for row in records if row.get("structured_field_coverage")),
            "media_available_count": sum(1 for row in records if int(row.get("media_available_count") or 0) > 0),
            "media_sendable_count": sum(1 for row in records if int(row.get("media_sendable_match_count") or 0) > 0),
            "generic_rule_available_count": sum(1 for row in records if int(row.get("generic_rule_count") or 0) > 0),
            "embedding_config_status": embedding,
            "rag_timeout_count": sum(1 for row in records if row.get("rag_timeout")),
        }
        return sanitize_obj({"run_uid": run.run_uid, "summary": summary, "records": records, "zero_evidence_records": zero_records})
    finally:
        db.close()


def _rows_from_summary(summary: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for key, value in summary.items():
        rows.append({"metric": key, "value": _json_cell(value)})
    return rows


def _write_sheet(workbook: Workbook, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]):
    if title in workbook.sheetnames:
        sheet = workbook[title]
    else:
        sheet = workbook.create_sheet(title=title)
    sheet.append([label for label, _key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        sheet.append([_json_cell(row.get(key)) for _label, key in columns])
    widths = {}
    for row in sheet.iter_rows():
        for cell in row:
            widths[cell.column_letter] = min(max(widths.get(cell.column_letter, 0), len(str(cell.value or "")) + 2), 60)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"


def build_workbook(report: dict[str, Any]) -> Workbook:
    records = report.get("records") or []
    zero_records = report.get("zero_evidence_records") or []
    workbook = Workbook()
    default = workbook.active
    default.title = "说明"
    default.append(["说明", "内容"])
    default.append(["用途", "只读诊断真实回放 selected_evidence_count=0 的证据链断点；不会写库或改任务状态。"])
    default.append(["回放批次", report.get("run_uid", "")])
    default.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    for cell in default[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    _write_sheet(workbook, "总览", SUMMARY_COLUMNS, _rows_from_summary(report.get("summary") or {}))
    _write_sheet(workbook, "Zero Evidence 明细", DETAIL_COLUMNS, zero_records)
    _write_sheet(
        workbook,
        "商品身份问题",
        IDENTITY_COLUMNS,
        [row for row in zero_records if row.get("primary_reason") in {"sidecar_missing_or_not_passed", "product_identity_unresolved", "kb_product_missing"}],
    )
    _write_sheet(
        workbook,
        "商品字段缺失",
        FIELD_COLUMNS,
        [row for row in zero_records if row.get("primary_reason") == "structured_field_missing" or row.get("missing_structured_fields")],
    )
    _write_sheet(
        workbook,
        "素材缺失或不可用",
        MEDIA_COLUMNS,
        [row for row in zero_records if row.get("primary_reason") in {"media_missing", "media_role_not_sendable", "evidence_role_mismatch"} or row.get("media_issue")],
    )
    _write_sheet(
        workbook,
        "Generic 规则缺失",
        GENERIC_COLUMNS,
        [row for row in zero_records if row.get("primary_reason") in {"generic_rule_missing", "generic_rule_blocked_by_product_fact"}],
    )
    _write_sheet(
        workbook,
        "RAG Embedding 问题",
        RAG_COLUMNS,
        [row for row in zero_records if row.get("primary_reason") in {"embedding_not_configured", "rag_timeout", "retrieval_filter_too_strict"} or row.get("embedding_config_status") != "configured"],
    )
    _write_sheet(
        workbook,
        "可自动修复候选",
        AUTO_FIX_COLUMNS,
        [row for row in records if row.get("auto_fix_candidate")],
    )
    _write_sheet(
        workbook,
        "需要人工确认",
        MANUAL_COLUMNS,
        [row for row in zero_records if row.get("manual_review_reason")],
    )
    return workbook


def default_excel_path() -> str:
    return str(Path.home() / "Desktop" / f"Evidence链路诊断_{datetime.now().strftime('%Y%m%d')}.xlsx")


def _write_json(path: str, payload: dict[str, Any]) -> None:
    if not path:
        return
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(_json_report(payload), ensure_ascii=True, indent=2), encoding="utf-8")


def _json_report(payload: dict[str, Any]) -> dict[str, Any]:
    """Return a machine-friendly report without deeply nested trace payloads.

    Excel keeps the richer human-readable details. JSON is intentionally
    flattened so Python json.load and PowerShell ConvertFrom-Json can parse it
    reliably in local Windows automation.
    """
    detail_keys = [key for _label, key in DETAIL_COLUMNS]
    records = payload.get("records") or []
    zero_records = payload.get("zero_evidence_records") or []
    return sanitize_obj(
        {
            "run_uid": payload.get("run_uid", ""),
            "summary": payload.get("summary") or {},
            "records": [_compact_json_record(row, detail_keys) for row in records],
            "zero_evidence_records": [_compact_json_record(row, detail_keys) for row in zero_records],
        }
    )


def _compact_json_record(row: dict[str, Any], keys: list[str]) -> dict[str, Any]:
    compact: dict[str, Any] = {}
    for key in keys:
        value = row.get(key)
        if isinstance(value, (dict, list)):
            compact[key] = _json_cell(value)
        else:
            compact[key] = value
    return compact


def _write_excel(path: str, payload: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(payload)
    workbook.save(target)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Diagnose replay evidence chain gaps without writing database records.")
    parser.add_argument("--run-uid", default="", help="EvalRun run_uid. Omit with --latest to use latest completed real_conversation run.")
    parser.add_argument("--latest", action="store_true", help="Use latest completed real_conversation run.")
    parser.add_argument("--json-output", default="", help="Optional JSON output path.")
    parser.add_argument("--excel-output", default="", help="Optional Excel output path. Excel is only written when this is provided.")
    parser.add_argument("--limit", type=int, default=0, help="Limit traces for smoke diagnostics.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = diagnose_evidence_chain(run_uid=args.run_uid, latest=args.latest, limit=args.limit)
    excel_output = args.excel_output
    if excel_output:
        _write_excel(excel_output, report)
    if args.json_output:
        _write_json(args.json_output, report)
    summary = report.get("summary") or {}
    print(json.dumps({
        "run_uid": report.get("run_uid", ""),
        "total_traces": summary.get("total_traces", 0),
        "zero_evidence_trace_count": summary.get("zero_evidence_trace_count", 0),
        "by_primary_reason": summary.get("by_primary_reason", {}),
        "excel_output": excel_output,
        "json_output": args.json_output,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
