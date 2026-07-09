"""Diagnose sidecar-like fields in real conversation source files.

This script is read-only. It inspects source file schemas and row-level fields
so replay input gaps can be attributed to source collection instead of Agent
behavior. It does not write the eval DB or modify source files.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.services.eval_sanitizer_service import hash_sensitive, sanitize_obj, sanitize_text  # noqa: E402
from app.services.real_conversation_import_service import default_source_dir  # noqa: E402


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "product_title": (
        "sidecar_product_title",
        "sidecar_product_name",
        "product_title",
        "product_name",
        "item_title",
        "order_product_title",
        "purchased_product_title",
        "商品标题",
        "商品名称",
        "商品名",
        "宝贝标题",
        "宝贝名称",
        "当前咨询商品",
        "咨询商品",
        "订单商品标题",
    ),
    "sku_code": (
        "sidecar_sku_code",
        "sidecar_sku",
        "sku_code",
        "sku",
        "order_sku_code",
        "商家编码",
        "商品编码",
        "SKU",
        "sku编码",
        "货号",
    ),
    "i_id": (
        "sidecar_i_id",
        "i_id",
        "internal_i_id",
        "product_i_id",
        "内部i_id",
        "内部商品ID",
        "YH编码",
    ),
    "order_id": (
        "sidecar_order_id",
        "sidecar_platform_order_id",
        "order_id",
        "order_no",
        "tid",
        "platform_order_id",
        "订单号",
        "子订单号",
        "交易单号",
    ),
    "product_url": (
        "product_url",
        "item_url",
        "item_link",
        "url",
        "link",
        "商品链接",
        "宝贝链接",
    ),
    "item_hash": (
        "item_id_hash",
        "platform_item_id_hash",
        "item_hash",
        "商品IDHash",
    ),
}

ROW_LIMIT_PER_FILE = 5000
SAMPLE_LIMIT = 12
SUPPORTED_SUFFIXES = {".xlsx", ".csv", ".json", ".jsonl", ".txt"}
SKIP_DIRS = {".git", ".stfolder", "node_modules", ".pytest_cache", "logs", "log", "__pycache__"}
SKIP_SUFFIXES = {".log", ".png", ".jpg", ".jpeg", ".md", ".yml", ".yaml"}


def _norm(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().lower())


ALIAS_LOOKUP: dict[str, str] = {}
for canonical, aliases in FIELD_ALIASES.items():
    for alias in aliases:
        ALIAS_LOOKUP[_norm(alias)] = canonical


def _write_json(path: str, payload: dict[str, Any]) -> None:
    if not path:
        return
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(sanitize_obj(payload), ensure_ascii=False, indent=2), encoding="utf-8")


def _should_skip(path: Path) -> bool:
    if any(part in SKIP_DIRS for part in path.parts):
        return True
    if path.suffix.lower() in SKIP_SUFFIXES:
        return True
    name = path.name.lower()
    return any(term in name for term in ("backend", "snapshot", "quickstart", "terms", "models", "faq", "sync-conflict", "调试"))


def _is_replay_like_headers(headers: list[str]) -> bool:
    normalized = {_norm(item) for item in headers}
    has_conversation = bool(normalized & {"会话id", "conversation_id", "chat_id", "session_id"})
    has_speaker = bool(normalized & {"说话人类型", "说话类型", "speaker_type", "role", "speaker"})
    has_message = bool(normalized & {"消息内容", "聊天内容", "content", "text", "message", "msg", "原始数据"})
    return (has_conversation and has_message) or (has_speaker and has_message)


def _iter_source_files(source_dir: str) -> list[Path]:
    root = Path(source_dir)
    if not root.exists():
        return []
    files = [
        path
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES and not _should_skip(path)
    ]
    return sorted(files, key=lambda p: (0 if "聊天记录" in str(p) else 1, p.suffix.lower(), str(p)))


def _flatten_json_keys(value: Any, prefix: str = "", depth: int = 0) -> dict[str, Any]:
    if depth > 4:
        return {}
    rows: dict[str, Any] = {}
    if isinstance(value, dict):
        for key, item in value.items():
            key_text = str(key)
            path = f"{prefix}.{key_text}" if prefix else key_text
            rows[path] = item
            rows.update(_flatten_json_keys(item, path, depth + 1))
    elif isinstance(value, list):
        for item in value[:5]:
            rows.update(_flatten_json_keys(item, prefix, depth + 1))
    return rows


def _field_hits(row: dict[str, Any]) -> dict[str, str]:
    hits: dict[str, str] = {}
    for key, value in row.items():
        canonical = ALIAS_LOOKUP.get(_norm(key))
        if canonical and value not in (None, ""):
            hits.setdefault(canonical, sanitize_text(value))
    return hits


def _row_status(hits: dict[str, str]) -> str:
    has_product = bool(hits.get("product_title") or hits.get("sku_code") or hits.get("i_id"))
    has_order = bool(hits.get("order_id"))
    if has_product and has_order:
        return "complete"
    if has_product or has_order:
        return "partial"
    return "missing"


def _redacted_sample(row: dict[str, Any], hits: dict[str, str], file_path: Path) -> dict[str, Any]:
    sample = {
        "file": str(file_path),
        "fields": sorted(row.keys())[:20],
        "matched_fields": sorted(hits.keys()),
        "sidecar_status": _row_status(hits),
    }
    if hits.get("product_title"):
        sample["product_title_preview"] = sanitize_text(hits["product_title"])[:80]
    if hits.get("sku_code"):
        sample["sku_code_present"] = True
    if hits.get("i_id"):
        sample["i_id_present"] = True
    if hits.get("order_id"):
        sample["order_id_hash"] = hash_sensitive(hits["order_id"])
    if hits.get("product_url"):
        sample["product_url_present"] = True
    if hits.get("item_hash"):
        sample["item_hash_present"] = True
    return sanitize_obj(sample)


def _inspect_row(row: dict[str, Any], file_path: Path, state: dict[str, Any]) -> None:
    state["total_rows"] += 1
    hits = _field_hits(row)
    for canonical in hits:
        state["field_counts"][canonical] += 1
    status = _row_status(hits)
    state["sidecar_counts"][status] += 1
    for key in row:
        canonical = ALIAS_LOOKUP.get(_norm(key))
        if canonical:
            state["detected_aliases"][canonical].add(str(key))
    if hits and len(state["samples_with_sidecar"]) < SAMPLE_LIMIT:
        state["samples_with_sidecar"].append(_redacted_sample(row, hits, file_path))
    if not hits and len(state["samples_missing_sidecar"]) < SAMPLE_LIMIT:
        state["samples_missing_sidecar"].append(_redacted_sample(row, hits, file_path))


def _inspect_xlsx(path: Path, state: dict[str, Any]) -> bool:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        state["file_errors"].append({"file": str(path), "error": sanitize_text(str(exc))})
        return False
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            headers = [sanitize_text(item) for item in header]
            state["file_headers"][str(path)].append({"sheet": worksheet.title, "headers": headers})
            if not _is_replay_like_headers(headers):
                continue
            for index, raw_row in enumerate(rows):
                if index >= ROW_LIMIT_PER_FILE:
                    break
                row = {
                    headers[i]: raw_row[i]
                    for i in range(min(len(headers), len(raw_row)))
                    if headers[i]
                }
                if row:
                    _inspect_row(row, path, state)
    finally:
        workbook.close()
    return True


def _inspect_csv(path: Path, state: dict[str, Any]) -> bool:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            state["file_headers"][str(path)].append({"sheet": "", "headers": list(reader.fieldnames or [])})
            for index, row in enumerate(reader):
                if index >= ROW_LIMIT_PER_FILE:
                    break
                _inspect_row(dict(row), path, state)
        return True
    except Exception as exc:
        state["file_errors"].append({"file": str(path), "error": sanitize_text(str(exc))})
        return False


def _inspect_json(path: Path, state: dict[str, Any]) -> bool:
    try:
        text = path.read_text(encoding="utf-8-sig")
        payload = json.loads(text)
    except Exception as exc:
        state["file_errors"].append({"file": str(path), "error": sanitize_text(str(exc))})
        return False
    rows = payload if isinstance(payload, list) else [payload]
    for item in rows[:ROW_LIMIT_PER_FILE]:
        flattened = _flatten_json_keys(item)
        _inspect_row(flattened, path, state)
    state["file_headers"][str(path)].append({"sheet": "", "headers": sorted(_flatten_json_keys(payload).keys())[:80]})
    return True


def _inspect_jsonl(path: Path, state: dict[str, Any]) -> bool:
    readable = False
    try:
        for index, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines()):
            if index >= ROW_LIMIT_PER_FILE:
                break
            if not line.strip():
                continue
            try:
                item = json.loads(line)
            except Exception:
                continue
            readable = True
            _inspect_row(_flatten_json_keys(item), path, state)
        return readable
    except Exception as exc:
        state["file_errors"].append({"file": str(path), "error": sanitize_text(str(exc))})
        return False


def _inspect_text(path: Path, state: dict[str, Any]) -> bool:
    try:
        lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    except Exception as exc:
        state["file_errors"].append({"file": str(path), "error": sanitize_text(str(exc))})
        return False
    for line in lines[:ROW_LIMIT_PER_FILE]:
        _inspect_row({"text": line}, path, state)
    return True


def diagnose_source_sidecar_fields(
    *,
    source_dir: str,
    json_output: str = "",
    excel_output: str = "",
) -> dict[str, Any]:
    state: dict[str, Any] = {
        "total_rows": 0,
        "field_counts": Counter(),
        "sidecar_counts": Counter(),
        "detected_aliases": defaultdict(set),
        "file_headers": defaultdict(list),
        "file_errors": [],
        "samples_with_sidecar": [],
        "samples_missing_sidecar": [],
    }
    files = _iter_source_files(source_dir)
    readable = 0
    suffix_counts = Counter(path.suffix.lower() for path in files)
    for path in files:
        suffix = path.suffix.lower()
        ok = False
        if suffix == ".xlsx":
            ok = _inspect_xlsx(path, state)
        elif suffix == ".csv":
            ok = _inspect_csv(path, state)
        elif suffix == ".json":
            ok = _inspect_json(path, state)
        elif suffix == ".jsonl":
            ok = _inspect_jsonl(path, state)
        elif suffix == ".txt":
            ok = _inspect_text(path, state)
        if ok:
            readable += 1

    field_counts = dict(state["field_counts"])
    sidecar_counts = dict(state["sidecar_counts"])
    result = {
        "source_dir": str(Path(source_dir)),
        "file_count": len(files),
        "readable_file_count": readable,
        "suffix_counts": dict(suffix_counts),
        "total_rows": int(state["total_rows"]),
        "product_title_field_found_count": int(field_counts.get("product_title", 0)),
        "sku_field_found_count": int(field_counts.get("sku_code", 0)),
        "i_id_field_found_count": int(field_counts.get("i_id", 0)),
        "order_id_field_found_count": int(field_counts.get("order_id", 0)),
        "product_url_field_found_count": int(field_counts.get("product_url", 0)),
        "item_hash_field_found_count": int(field_counts.get("item_hash", 0)),
        "sidecar_complete_candidate_count": int(sidecar_counts.get("complete", 0)),
        "sidecar_partial_candidate_count": int(sidecar_counts.get("partial", 0)),
        "sidecar_missing_candidate_count": int(sidecar_counts.get("missing", 0)),
        "detected_field_aliases": {
            key: sorted(values)
            for key, values in sorted(state["detected_aliases"].items())
        },
        "file_headers": dict(list(state["file_headers"].items())[:30]),
        "file_errors": state["file_errors"][:20],
        "sample_rows_with_sidecar": state["samples_with_sidecar"],
        "sample_rows_missing_sidecar": state["samples_missing_sidecar"],
        "notes": [
            "product_url/item_hash are supplemental identity only and do not count as complete product sidecar.",
            "This script is read-only and does not import or modify replay data.",
        ],
    }
    _write_json(json_output, result)
    if excel_output:
        _write_excel(excel_output, result)
    return sanitize_obj(result)


def _write_excel(path: str, result: dict[str, Any]) -> None:
    import openpyxl

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    rows = [
        ("指标", "数值"),
        ("文件数", result.get("file_count", 0)),
        ("可读取文件数", result.get("readable_file_count", 0)),
        ("总行数", result.get("total_rows", 0)),
        ("商品标题字段命中", result.get("product_title_field_found_count", 0)),
        ("SKU字段命中", result.get("sku_field_found_count", 0)),
        ("内部i_id字段命中", result.get("i_id_field_found_count", 0)),
        ("订单号字段命中", result.get("order_id_field_found_count", 0)),
        ("商品链接字段命中", result.get("product_url_field_found_count", 0)),
        ("item hash字段命中", result.get("item_hash_field_found_count", 0)),
        ("完整sidecar候选", result.get("sidecar_complete_candidate_count", 0)),
        ("部分sidecar候选", result.get("sidecar_partial_candidate_count", 0)),
        ("缺sidecar候选", result.get("sidecar_missing_candidate_count", 0)),
    ]
    for row in rows:
        ws.append(row)
    ws_alias = wb.create_sheet("字段别名")
    ws_alias.append(["标准字段", "源字段别名"])
    for key, aliases in (result.get("detected_field_aliases") or {}).items():
        ws_alias.append([key, ", ".join(aliases)])
    ws_samples = wb.create_sheet("样例")
    ws_samples.append(["类型", "文件", "匹配字段", "sidecar状态", "商品标题预览", "SKU", "i_id", "订单hash", "商品链接", "item hash"])
    for kind, items in (("有sidecar字段", result.get("sample_rows_with_sidecar") or []), ("缺sidecar字段", result.get("sample_rows_missing_sidecar") or [])):
        for item in items:
            ws_samples.append([
                kind,
                item.get("file", ""),
                ", ".join(item.get("matched_fields") or []),
                item.get("sidecar_status", ""),
                item.get("product_title_preview", ""),
                "是" if item.get("sku_code_present") else "",
                "是" if item.get("i_id_present") else "",
                item.get("order_id_hash", ""),
                "是" if item.get("product_url_present") else "",
                "是" if item.get("item_hash_present") else "",
            ])
    wb.save(target)


def main() -> None:
    parser = argparse.ArgumentParser(description="Diagnose sidecar fields in real conversation source files.")
    parser.add_argument("--source-dir", default=default_source_dir())
    parser.add_argument("--json-output", default="")
    parser.add_argument("--excel-output", default="")
    args = parser.parse_args()
    result = diagnose_source_sidecar_fields(
        source_dir=args.source_dir,
        json_output=args.json_output,
        excel_output=args.excel_output,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
