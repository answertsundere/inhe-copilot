"""Export replay evidence governance tasks from evidence-chain diagnosis.

Read-only by design. This script turns replay evidence-chain diagnostics into
human-actionable governance rows for product fields, media roles, policies, and
RAG configuration. It does not write database records or verified knowledge.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text  # noqa: E402


TASK_COLUMNS = [
    ("优先级", "priority"),
    ("缺口类型", "gap_category"),
    ("问题类型 query_fact_type", "query_fact_type"),
    ("问题类型中文名", "query_fact_type_name"),
    ("买家问题", "buyer_message"),
    ("当前商品标题", "product_title"),
    ("SKU/i_id", "sku_i_id"),
    ("当前回复", "current_reply"),
    ("selected_evidence_count", "selected_evidence_count"),
    ("primary_reason", "primary_reason"),
    ("secondary_reasons", "secondary_reasons"),
    ("需要补什么", "required_input"),
    ("建议补充来源", "suggested_source"),
    ("是否有可信源可自动回填", "auto_backfill_candidate"),
    ("是否需要人工确认", "requires_manual_review"),
    ("是否高风险字段", "high_risk_field"),
    ("是否允许进入 eval/provisional", "allow_eval_provisional"),
    ("是否允许自动发送", "allow_auto_send"),
    ("建议负责人", "suggested_owner"),
    ("代表样本", "representative_samples"),
    ("备注", "note"),
]

SUMMARY_COLUMNS = [
    ("指标", "metric"),
    ("数值", "value"),
]

HIGH_RISK_FACT_TYPES = {
    "material_safety",
    "certification_report",
    "age_range",
    "child_safety",
    "child_suitability",
    "load_capacity",
    "promotion_policy",
    "price_negotiation",
    "aftersales_policy",
    "aftersales",
    "stock_shipping",
}

PRODUCT_FIELD_FACT_TYPES = {
    "dimensions",
    "space_fit",
    "placement_scene",
    "material",
    "material_safety",
    "certification_report",
    "load_capacity",
    "gross_weight",
    "accessories",
    "accessory_availability",
    "accessory_compatibility",
    "structure_function",
    "age_range",
    "child_safety",
    "child_suitability",
}

MEDIA_FACT_TYPES = {
    "installation",
    "installation_media_request",
    "accessory_usage",
    "accessory_availability",
    "accessory_compatibility",
    "dimensions",
    "space_fit",
}

PROMOTION_FACT_TYPES = {"promotion", "promotion_policy", "price_negotiation", "gift_policy", "price_protection"}
AFTERSALES_FACT_TYPES = {"aftersales", "aftersales_policy"}
LOGISTICS_FACT_TYPES = {"stock_shipping", "logistics", "order_status"}

FACT_TYPE_NAMES = {
    "unknown": "未识别",
    "dimensions": "尺寸",
    "space_fit": "空间适配",
    "placement_scene": "摆放场景",
    "material": "材质",
    "material_safety": "材质安全",
    "certification_report": "检测/证书",
    "load_capacity": "承重",
    "gross_weight": "毛重",
    "installation": "安装资料",
    "accessory_usage": "配件用途",
    "accessory_availability": "配件补配",
    "accessory_compatibility": "配件/结构适配",
    "structure_function": "结构功能",
    "promotion": "活动优惠",
    "promotion_policy": "活动优惠规则",
    "price_negotiation": "议价/多买优惠",
    "aftersales": "售后",
    "aftersales_policy": "售后政策",
    "stock_shipping": "物流/发货",
    "invoice_policy": "发票规则",
    "age_range": "年龄适用",
    "child_safety": "儿童安全",
    "child_suitability": "儿童适用",
}


def _text(value: Any) -> str:
    return sanitize_text(value).strip()


def _as_list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _json_cell(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(sanitize_obj(value), ensure_ascii=False)


def _safe_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _load_diagnosis(path: str) -> dict[str, Any]:
    with open(path, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if not isinstance(payload, dict):
        return {"run_uid": "", "summary": {}, "records": []}
    return payload


def _records_from_diagnosis(payload: dict[str, Any]) -> list[dict[str, Any]]:
    records = payload.get("zero_evidence_records") or payload.get("records") or []
    if not isinstance(records, list):
        return []
    return [row for row in records if isinstance(row, dict)]


def _representative_sample(row: dict[str, Any]) -> str:
    parts = []
    for key, label in (
        ("case_uid", "case"),
        ("turn_uid", "turn"),
        ("buyer_message_preview", "买家"),
    ):
        value = _text(row.get(key))
        if value:
            parts.append(f"{label}: {value}")
    return " | ".join(parts)


def _current_reply(row: dict[str, Any]) -> str:
    return _text(
        row.get("agent_reply")
        or row.get("sendable_reply")
        or row.get("draft_reply")
        or row.get("current_reply")
        or row.get("reply_preview")
    )


def _gap_category(row: dict[str, Any]) -> str:
    primary = _text(row.get("primary_reason"))
    qft = _text(row.get("query_fact_type")) or "unknown"
    if primary in {"embedding_not_configured", "rag_timeout", "retrieval_filter_too_strict"}:
        return "Embedding/RAG配置问题"
    if primary == "structured_field_missing" or qft in PRODUCT_FIELD_FACT_TYPES:
        return "商品字段缺口"
    if primary in {"media_role_not_sendable", "evidence_role_mismatch"}:
        return "素材Role待确认"
    if primary == "media_missing" or (qft in MEDIA_FACT_TYPES and primary not in {"structured_field_missing"}):
        return "素材缺口"
    if primary == "generic_rule_missing":
        if qft in PROMOTION_FACT_TYPES:
            return "活动优惠规则缺口"
        if qft in AFTERSALES_FACT_TYPES:
            return "售后政策缺口"
        if qft in LOGISTICS_FACT_TYPES:
            return "物流规则缺口"
        return "Generic规则缺口"
    if primary in {"query_fact_type_missing", "sidecar_missing_or_not_passed", "product_identity_unresolved", "kb_product_missing"}:
        return "上下文/证据路由缺口"
    if qft in PROMOTION_FACT_TYPES:
        return "活动优惠规则缺口"
    if qft in AFTERSALES_FACT_TYPES:
        return "售后政策缺口"
    if qft in LOGISTICS_FACT_TYPES:
        return "物流规则缺口"
    return "Generic规则缺口"


def _required_input(row: dict[str, Any], category: str) -> str:
    qft = _text(row.get("query_fact_type")) or "unknown"
    missing_fields = _json_cell(row.get("missing_structured_fields"))
    if category == "商品字段缺口":
        return f"补充当前商品的结构化字段：{missing_fields or FACT_TYPE_NAMES.get(qft, qft)}"
    if category == "素材缺口":
        return "补充当前商品可用素材，并标明素材用途：安装视频/安装图/说明书/尺寸图/配件图"
    if category == "素材Role待确认":
        return "确认现有素材是否真的是对应用途；普通商品图不能冒充安装图、尺寸图或配件图"
    if category == "活动优惠规则缺口":
        return "补充当前店铺活动、优惠券、满减、晒图返现、议价边界的运营口径"
    if category == "售后政策缺口":
        return "补充少件、破损、退换、补发、赔付等售后处理边界"
    if category == "物流规则缺口":
        return "补充发货、物流查询、签收未收到、运单号查询等服务规则"
    if category == "Embedding/RAG配置问题":
        return "检查 embedding/RAG 配置、索引构建、检索超时和过滤条件"
    if category == "上下文/证据路由缺口":
        return "确认侧栏商品/订单上下文、商品身份解析、fact_type 或 evidence routing 是否完整"
    return "补充对应通用客服规则，但不能写成具体商品事实"


def _suggested_source(category: str, row: dict[str, Any]) -> str:
    if category == "商品字段缺口":
        return "可信商品资料源、商品主数据、已审核商品说明"
    if category in {"素材缺口", "素材Role待确认"}:
        return "已审核素材库、商品说明书、安装图/视频源文件"
    if category == "活动优惠规则缺口":
        return "运营活动规则、店铺优惠配置、下单页活动说明"
    if category == "售后政策缺口":
        return "售后SOP、平台售后规则、主管确认口径"
    if category == "物流规则缺口":
        return "物流SOP、发货规则、平台物流后台"
    if category == "Embedding/RAG配置问题":
        return "系统配置、embedding 服务、RAG 索引任务"
    if category == "上下文/证据路由缺口":
        return "千牛侧栏、订单上下文、商品身份映射、trace/evidence debug"
    return "主管确认后的通用服务口径"


def _owner(category: str) -> str:
    if category == "商品字段缺口":
        return "knowledge_ops"
    if category in {"素材缺口", "素材Role待确认"}:
        return "media_ops"
    if category in {"活动优惠规则缺口", "售后政策缺口", "物流规则缺口", "Generic规则缺口"}:
        return "ops_policy"
    if category == "Embedding/RAG配置问题":
        return "engineering"
    return "engineering"


def _is_high_risk(row: dict[str, Any]) -> bool:
    qft = _text(row.get("query_fact_type")) or "unknown"
    fields = _json_cell(row.get("missing_structured_fields")).lower()
    if qft in HIGH_RISK_FACT_TYPES:
        return True
    high_risk_terms = ("安全", "无毒", "检测", "证书", "年龄", "宝宝", "儿童", "承重", "退款", "补发", "赔付", "优惠金额", "返现")
    return any(term in fields for term in high_risk_terms)


def _auto_backfill_candidate(row: dict[str, Any], category: str, high_risk: bool) -> bool:
    if high_risk:
        return False
    if category != "商品字段缺口":
        return False
    if not _text(row.get("kb_product_i_id") or row.get("identity_i_id") or row.get("sidecar_i_id")):
        return False
    return _text(row.get("primary_reason")) == "structured_field_missing"


def _priority(row: dict[str, Any], category: str, group_count: int, high_risk: bool) -> str:
    has_product = bool(row.get("kb_product_found")) and bool(row.get("has_sidecar_product_context"))
    primary = _text(row.get("primary_reason"))
    if (
        group_count >= 2
        and has_product
        and _safe_int(row.get("selected_evidence_count")) == 0
        and primary in {"structured_field_missing", "generic_rule_missing", "media_missing", "media_role_not_sendable", "evidence_role_mismatch"}
        and not high_risk
    ):
        return "P0"
    if category in {"素材缺口", "素材Role待确认", "活动优惠规则缺口", "售后政策缺口", "物流规则缺口"}:
        return "P1"
    if category == "Embedding/RAG配置问题":
        return "P2"
    if high_risk or primary in {"query_fact_type_missing", "sidecar_missing_or_not_passed", "product_identity_unresolved"}:
        return "P2"
    return "P1"


def _row_key(row: dict[str, Any]) -> tuple[str, str, str, str]:
    category = _gap_category(row)
    return (
        category,
        _text(row.get("query_fact_type")) or "unknown",
        _text(row.get("primary_reason")) or "unknown",
        _text(row.get("kb_product_i_id") or row.get("identity_i_id") or row.get("sidecar_i_id")),
    )


def _merge_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[_row_key(row)].append(row)
    tasks: list[dict[str, Any]] = []
    for key, items in grouped.items():
        category, qft, primary, _product_key = key
        base = items[0]
        high_risk = any(_is_high_risk(item) for item in items)
        auto_backfill = any(_auto_backfill_candidate(item, category, high_risk) for item in items)
        priority = _priority(base, category, len(items), high_risk)
        samples = []
        for item in items:
            sample = _representative_sample(item)
            if sample and sample not in samples:
                samples.append(sample)
            if len(samples) >= 3:
                break
        buyer_messages = []
        for item in items:
            message = _text(item.get("buyer_message_preview"))
            if message and message not in buyer_messages:
                buyer_messages.append(message)
            if len(buyer_messages) >= 3:
                break
        task = {
            "priority": priority,
            "gap_category": category,
            "query_fact_type": qft,
            "query_fact_type_name": FACT_TYPE_NAMES.get(qft, qft),
            "buyer_message": " / ".join(buyer_messages),
            "product_title": _text(base.get("kb_product_name") or base.get("identity_display_product_name") or base.get("sidecar_product_title")),
            "sku_i_id": _text(base.get("kb_product_i_id") or base.get("identity_i_id") or base.get("sidecar_i_id") or base.get("identity_sku_code") or base.get("sidecar_sku_code")),
            "current_reply": _current_reply(base),
            "selected_evidence_count": min(_safe_int(item.get("selected_evidence_count")) for item in items),
            "primary_reason": primary,
            "secondary_reasons": sorted({reason for item in items for reason in _as_list(item.get("secondary_reasons")) if reason}),
            "required_input": _required_input(base, category),
            "suggested_source": _suggested_source(category, base),
            "auto_backfill_candidate": "是" if auto_backfill else "否",
            "requires_manual_review": "是",
            "high_risk_field": "是" if high_risk else "否",
            "allow_eval_provisional": "是" if category != "Embedding/RAG配置问题" else "否",
            "allow_auto_send": "否",
            "suggested_owner": _owner(category),
            "representative_samples": samples,
            "note": f"聚合 {len(items)} 条 trace；本表只用于 staging/治理，不写 verified 知识。",
            "_count": len(items),
        }
        tasks.append(task)
    return sorted(tasks, key=lambda row: (row["priority"], -int(row["_count"]), row["gap_category"], row["query_fact_type"]))


def build_governance_report(diagnosis: dict[str, Any], limit: int = 0) -> dict[str, Any]:
    records = _records_from_diagnosis(diagnosis)
    actionable_records = [
        row
        for row in records
        if _safe_int(row.get("selected_evidence_count")) == 0
        or _text(row.get("primary_reason")) in {"generic_rule_missing", "structured_field_missing", "media_missing", "media_role_not_sendable", "evidence_role_mismatch"}
    ]
    tasks = _merge_rows(actionable_records)
    if limit and limit > 0:
        tasks = tasks[:limit]
    summary = {
        "run_uid": diagnosis.get("run_uid") or (diagnosis.get("summary") or {}).get("run_uid", ""),
        "source_total_records": len(records),
        "governance_task_count": len(tasks),
        "p0_count": sum(1 for row in tasks if row["priority"] == "P0"),
        "p1_count": sum(1 for row in tasks if row["priority"] == "P1"),
        "p2_count": sum(1 for row in tasks if row["priority"] == "P2"),
        "auto_backfill_candidate_count": sum(1 for row in tasks if row["auto_backfill_candidate"] == "是"),
        "manual_review_required_count": sum(1 for row in tasks if row["requires_manual_review"] == "是"),
        "by_gap_category": dict(Counter(row["gap_category"] for row in tasks).most_common()),
        "by_query_fact_type": dict(Counter(row["query_fact_type"] for row in tasks).most_common()),
        "by_primary_reason": dict(Counter(row["primary_reason"] for row in tasks).most_common()),
    }
    return sanitize_obj({"run_uid": summary["run_uid"], "summary": summary, "tasks": tasks})


def _summary_rows(summary: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"metric": key, "value": _json_cell(value)} for key, value in summary.items()]


def _write_sheet(workbook: Workbook, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append([label for label, _key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        sheet.append([_json_cell(row.get(key)) for _label, key in columns])
    for row in sheet.iter_rows():
        for cell in row:
            width = min(max(sheet.column_dimensions[cell.column_letter].width or 10, len(str(cell.value or "")) + 2), 64)
            sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"


def build_workbook(report: dict[str, Any]) -> Workbook:
    tasks = report.get("tasks") or []
    workbook = Workbook()
    readme = workbook.active
    readme.title = "说明"
    readme.append(["说明", "内容"])
    readme.append(["用途", "把真实回放 knowledge_gap/rag_miss 的证据断点整理成治理任务；只读导出，不写 DB、不写正式知识库。"])
    readme.append(["回放批次", report.get("run_uid", "")])
    readme.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    readme.append(["安全边界", "默认不允许自动发送；自动回填只表示 staging/pending 候选，必须复测和人工确认后才能进入正式知识。"])
    for cell in readme[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    _write_sheet(workbook, "总览", SUMMARY_COLUMNS, _summary_rows(report.get("summary") or {}))
    _write_sheet(workbook, "P0优先处理", TASK_COLUMNS, [row for row in tasks if row.get("priority") == "P0"])
    _write_sheet(workbook, "商品字段缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "商品字段缺口"])
    _write_sheet(workbook, "素材缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "素材缺口"])
    _write_sheet(workbook, "素材Role待确认", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "素材Role待确认"])
    _write_sheet(workbook, "活动优惠规则缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "活动优惠规则缺口"])
    _write_sheet(workbook, "售后政策缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "售后政策缺口"])
    _write_sheet(workbook, "物流规则缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "物流规则缺口"])
    _write_sheet(workbook, "Generic规则缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "Generic规则缺口"])
    _write_sheet(workbook, "上下文证据路由缺口", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "上下文/证据路由缺口"])
    _write_sheet(workbook, "Embedding和RAG配置问题", TASK_COLUMNS, [row for row in tasks if row.get("gap_category") == "Embedding/RAG配置问题"])
    _write_sheet(workbook, "可自动回填候选", TASK_COLUMNS, [row for row in tasks if row.get("auto_backfill_candidate") == "是"])
    _write_sheet(workbook, "必须人工确认", TASK_COLUMNS, [row for row in tasks if row.get("requires_manual_review") == "是"])
    return workbook


def export_governance_tasks(
    *,
    diagnosis_json: str = "",
    run_uid: str = "",
    latest: bool = False,
    json_output: str = "",
    excel_output: str = "",
    limit: int = 0,
) -> dict[str, Any]:
    if diagnosis_json:
        diagnosis = _load_diagnosis(diagnosis_json)
    else:
        from scripts.diagnose_evidence_chain_for_replay import diagnose_evidence_chain

        diagnosis = diagnose_evidence_chain(run_uid=run_uid, latest=latest or not run_uid)
    report = build_governance_report(diagnosis, limit=limit)
    if json_output:
        target = Path(json_output)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(sanitize_obj(report), ensure_ascii=True, indent=2), encoding="utf-8")
    if excel_output:
        target = Path(excel_output)
        target.parent.mkdir(parents=True, exist_ok=True)
        build_workbook(report).save(target)
    return report


def default_excel_path() -> str:
    return str(Path.home() / "Desktop" / f"真实回放证据治理任务_{datetime.now().strftime('%Y%m%d')}.xlsx")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export replay evidence governance tasks from evidence-chain diagnosis.")
    parser.add_argument("--run-uid", default="", help="EvalRun run_uid. Used when --diagnosis-json is omitted.")
    parser.add_argument("--latest", action="store_true", help="Use latest completed real_conversation run when --diagnosis-json is omitted.")
    parser.add_argument("--diagnosis-json", default="", help="Path to evidence-chain diagnosis JSON.")
    parser.add_argument("--json-output", default="", help="Optional JSON output path.")
    parser.add_argument("--excel-output", default="", help="Optional Excel output path. Defaults to Desktop.")
    parser.add_argument("--limit", type=int, default=0, help="Limit exported governance tasks after grouping.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    excel_output = args.excel_output or default_excel_path()
    report = export_governance_tasks(
        diagnosis_json=args.diagnosis_json,
        run_uid=args.run_uid,
        latest=args.latest,
        json_output=args.json_output,
        excel_output=excel_output,
        limit=args.limit,
    )
    summary = report.get("summary") or {}
    print(
        json.dumps(
            {
                "run_uid": report.get("run_uid", ""),
                "governance_task_count": summary.get("governance_task_count", 0),
                "p0_count": summary.get("p0_count", 0),
                "p1_count": summary.get("p1_count", 0),
                "p2_count": summary.get("p2_count", 0),
                "auto_backfill_candidate_count": summary.get("auto_backfill_candidate_count", 0),
                "manual_review_required_count": summary.get("manual_review_required_count", 0),
                "by_gap_category": summary.get("by_gap_category", {}),
                "excel_output": excel_output,
                "json_output": args.json_output,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
