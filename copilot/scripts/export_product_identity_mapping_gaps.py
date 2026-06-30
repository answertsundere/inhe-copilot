"""Export unresolved product identity mappings from real replay traces."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.db import SessionLocal, init_db
from app.models.eval_tables import EvalRun, EvalTrace
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text


SHEET_NAME = "商品身份映射缺口"
README_SHEET = "说明"
HEADERS = [
    "source_run_uid",
    "case_uid",
    "turn_uid",
    "platform_item_id",
    "platform_item_id_hash",
    "product_url",
    "platform_product_title",
    "order_product_title",
    "buyer_message_preview",
    "unresolved_reason",
    "candidate_count",
    "ambiguous_candidates",
    "建议内部 i_id",
    "建议 SKU",
    "建议商品标题",
    "人工确认状态",
    "处理人",
    "备注",
]


def default_output_path() -> str:
    today = datetime.now().strftime("%Y%m%d")
    return str(Path.home() / "Desktop" / f"商品身份映射缺口_{today}.xlsx")


def _latest_run_uid(db) -> str:
    run = (
        db.query(EvalRun)
        .filter(EvalRun.source_type == "real_conversation", EvalRun.status == "completed")
        .order_by(EvalRun.created_at.desc(), EvalRun.id.desc())
        .first()
    )
    return run.run_uid if run else ""


def _pack_from_trace(trace: EvalTrace) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    if not isinstance(raw, dict):
        return {}
    summary = ((raw.get("evidence_debug") or {}).get("product_context_pack_summary") or {})
    pack = summary.get("evidence_pack") or {}
    return pack if isinstance(pack, dict) else {}


def _real_identity(trace: EvalTrace) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    answer_trace = trace.get_answer_trace() or {}
    if isinstance(raw, dict) and isinstance(raw.get("real_context_product_identity"), dict):
        return raw["real_context_product_identity"]
    if isinstance(answer_trace, dict) and isinstance(answer_trace.get("real_context_product_identity"), dict):
        return answer_trace["real_context_product_identity"]
    product_identity = trace.get_product_identity() or {}
    return product_identity if isinstance(product_identity, dict) else {}


def _real_context_product(trace: EvalTrace) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    answer_trace = trace.get_answer_trace() or {}
    for container in (raw, answer_trace):
        if not isinstance(container, dict):
            continue
        real_context = container.get("real_context")
        if isinstance(real_context, dict):
            product = real_context.get("product")
            if isinstance(product, dict):
                return product
    return {}


def _first_text(*values: Any) -> str:
    for value in values:
        text = sanitize_text(str(value or "")).strip()
        if text:
            return text
    return ""


def _identity_row(trace: EvalTrace) -> dict[str, Any] | None:
    pack = _pack_from_trace(trace)
    resolution = pack.get("product_identity_resolution") or (pack.get("evidence_pack_trace") or {}).get("product_identity_resolution") or {}
    if not resolution or resolution.get("status") == "resolved":
        return None
    identity = _real_identity(trace)
    product = _real_context_product(trace)
    ambiguous = resolution.get("ambiguous_candidates") or []
    row = {
        "source_run_uid": trace.run_uid,
        "case_uid": trace.case_uid,
        "turn_uid": trace.turn_uid,
        "platform_item_id": _first_text(identity.get("item_id"), product.get("item_id")),
        "platform_item_id_hash": _first_text(identity.get("item_id_hash"), product.get("item_id_hash")),
        "product_url": _first_text(identity.get("product_url"), product.get("product_url")),
        "platform_product_title": _first_text(identity.get("product_title"), product.get("product_title")),
        "order_product_title": _first_text(identity.get("order_product_title")),
        "buyer_message_preview": sanitize_text(trace.buyer_message)[:120],
        "unresolved_reason": _first_text(resolution.get("unresolved_reason"), resolution.get("reason")),
        "candidate_count": len(ambiguous),
        "ambiguous_candidates": json.dumps(sanitize_obj(ambiguous), ensure_ascii=False),
        "建议内部 i_id": "",
        "建议 SKU": "",
        "建议商品标题": "",
        "人工确认状态": "",
        "处理人": "",
        "备注": "",
    }
    if not any(row.get(key) for key in ("platform_item_id_hash", "product_url", "platform_item_id", "platform_product_title", "order_product_title")):
        return None
    return sanitize_obj(row)


def collect_unresolved_identity_rows(run_uid: str = "", db_factory=None) -> tuple[str, list[dict[str, Any]]]:
    db_factory = db_factory or SessionLocal
    db = db_factory()
    try:
        target_run_uid = sanitize_text(run_uid) or _latest_run_uid(db)
        if not target_run_uid:
            return "", []
        traces = (
            db.query(EvalTrace)
            .filter(EvalTrace.run_uid == target_run_uid)
            .order_by(EvalTrace.case_uid.asc(), EvalTrace.turn_index.asc(), EvalTrace.id.asc())
            .all()
        )
        rows: list[dict[str, Any]] = []
        seen: set[str] = set()
        for trace in traces:
            row = _identity_row(trace)
            if not row:
                continue
            key = row.get("platform_item_id_hash") or row.get("product_url") or row.get("platform_item_id") or f"{row.get('platform_product_title')}|{row.get('order_product_title')}"
            if key in seen:
                continue
            seen.add(key)
            rows.append(row)
        return target_run_uid, rows
    finally:
        db.close()


def write_workbook(rows: list[dict[str, Any]], output: str) -> str:
    workbook = Workbook()
    ws = workbook.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        ws.append([row.get(header, "") for header in HEADERS])
    for idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(14, min(36, len(header) + 4))

    readme = workbook.create_sheet(README_SHEET)
    readme.append(["字段", "说明"])
    readme.append(["人工确认状态", "确认映射后填写：已确认 / confirmed / active。未确认不会导入。"])
    readme.append(["建议内部 i_id / 建议 SKU", "至少填写一个，导入时会校验 KBProduct 是否存在。"])
    readme.append(["安全说明", "导出内容已脱敏；如果只有 hash，不能也不应该反推原始 item_id。"])
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def run_export(run_uid: str = "", output: str = "", db_factory=None) -> dict[str, Any]:
    target_run_uid, rows = collect_unresolved_identity_rows(run_uid, db_factory=db_factory)
    output_path = output or default_output_path()
    write_workbook(rows, output_path)
    return sanitize_obj({
        "ok": True,
        "run_uid": target_run_uid,
        "count": len(rows),
        "output": output_path,
    })


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Export unresolved product identity mapping gaps.")
    parser.add_argument("--run-uid", default="", help="Replay run_uid; defaults to latest completed real_conversation run.")
    parser.add_argument("--output", default="", help="Output xlsx path.")
    args = parser.parse_args()
    init_db()
    result = run_export(run_uid=args.run_uid, output=args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
