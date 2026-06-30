"""Import manually filled knowledge gap workbook fields into staging metadata."""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.db import SessionLocal, init_db
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text
from app.services.knowledge_gap_task_service import _metadata_with_status_history


SUPPORTED_SHEETS = {
    "\u77e5\u8bc6\u7f3a\u53e3\u4efb\u52a1",
    "\u7d20\u6750\u7f3a\u53e3",
    "\u552e\u540e\u6d3b\u52a8\u89c4\u5219\u7f3a\u53e3",
    "\u5546\u54c1\u5b57\u6bb5\u7f3a\u53e3",
}

FIELD_ALIASES = {
    "task_uid": {"task_uid", "\u4efb\u52a1ID", "\u4efb\u52a1id"},
    "manual_conclusion": {"\u4eba\u5de5\u5904\u7406\u7ed3\u679c", "\u4eba\u5de5\u5904\u7406\u7ed3\u8bba"},
    "manual_content": {"\u4eba\u5de5\u8865\u5145\u5185\u5bb9", "\u8865\u5145\u5185\u5bb9"},
    "knowledge_filled": {"\u662f\u5426\u5df2\u8865\u77e5\u8bc6\u5e93"},
    "media_uploaded": {"\u662f\u5426\u5df2\u4e0a\u4f20\u7d20\u6750"},
    "media_reference": {"\u7d20\u6750\u94fe\u63a5\u6216\u7d20\u6750\u7f16\u53f7", "\u8bc1\u636e\u6765\u6e90/\u7d20\u6750\u94fe\u63a5"},
    "product_field_value": {"\u5546\u54c1\u5b57\u6bb5\u503c"},
    "aftersales_rule_text": {"\u552e\u540e\u89c4\u5219\u53e3\u5f84"},
    "promotion_rule_text": {"\u6d3b\u52a8\u89c4\u5219\u53e3\u5f84"},
    "handler": {"\u5904\u7406\u4eba"},
    "filled": {"\u662f\u5426\u5df2\u8865\u9f50"},
    "note": {"\u5907\u6ce8", "\u5904\u7406\u5907\u6ce8"},
}

WRITABLE_FIELDS = [
    "manual_conclusion",
    "manual_content",
    "knowledge_filled",
    "media_uploaded",
    "media_reference",
    "product_field_value",
    "aftersales_rule_text",
    "promotion_rule_text",
    "handler",
    "filled",
    "note",
]


def _now_iso() -> str:
    return datetime.utcnow().isoformat()


def _normalize_header(value: Any) -> str:
    return sanitize_text(str(value or "")).strip()


def _canonical_header(value: Any) -> str:
    header = _normalize_header(value)
    for canonical, aliases in FIELD_ALIASES.items():
        if header in aliases:
            return canonical
    return ""


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return sanitize_text(str(value))


def _internal_id_text(value: Any) -> str:
    return str(value or "").strip()


def _has_manual_fill(fields: dict[str, str]) -> bool:
    return any(fields.get(key) for key in WRITABLE_FIELDS)


def _merge_fields(existing: dict[str, str], incoming: dict[str, str]) -> dict[str, str]:
    merged = dict(existing)
    for key, value in incoming.items():
        if key == "task_uid":
            continue
        text = sanitize_text(value)
        if text:
            merged[key] = text
    return merged


def _read_workbook_rows(input_path: str) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    rows_by_task: dict[str, dict[str, Any]] = {}
    skipped: list[dict[str, Any]] = []
    for sheet_name in workbook.sheetnames:
        if sheet_name not in SUPPORTED_SHEETS:
            continue
        sheet = workbook[sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            continue
        headers = [_canonical_header(value) for value in header_row]
        if "task_uid" not in headers:
            skipped.append({"sheet": sheet_name, "row": 1, "reason": "missing_task_uid_header"})
            continue
        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values: dict[str, str] = {}
            for header, value in zip(headers, row):
                if header:
                    values[header] = _internal_id_text(value) if header == "task_uid" else _cell_text(value)
            task_uid = _internal_id_text(values.get("task_uid"))
            if not task_uid:
                if any(_cell_text(value) for value in row):
                    skipped.append({"sheet": sheet_name, "row": index, "reason": "missing_task_uid"})
                continue
            if not _has_manual_fill(values):
                skipped.append({"task_uid": task_uid, "sheet": sheet_name, "row": index, "reason": "no_manual_fill"})
                continue
            existing = rows_by_task.get(task_uid, {"task_uid": task_uid, "sheets": [], "rows": [], "fields": {}})
            existing["sheets"].append(sheet_name)
            existing["rows"].append(index)
            existing["fields"] = _merge_fields(existing["fields"], values)
            rows_by_task[task_uid] = existing
    return rows_by_task, skipped


def _status_after_import(current_status: str) -> str:
    return "resolved_pending_retest"


def import_filled_knowledge_gap_tasks(
    *,
    input_path: str,
    run_uid: str = "",
    apply: bool = False,
) -> dict[str, Any]:
    target_run_uid = sanitize_text(run_uid)
    rows_by_task, skipped = _read_workbook_rows(input_path)
    init_db()
    db = SessionLocal()
    matched = 0
    updated: list[str] = []
    errors: list[dict[str, Any]] = []
    changed_counter: Counter[str] = Counter()
    try:
        from app.models.eval_tables import KnowledgeGapTask

        for task_uid, row in rows_by_task.items():
            task = db.query(KnowledgeGapTask).filter(KnowledgeGapTask.task_uid == _internal_id_text(task_uid)).one_or_none()
            if task is None:
                skipped.append({"task_uid": task_uid, "reason": "task_not_found"})
                continue
            metadata = task.get_metadata()
            source_run_uid = sanitize_text(metadata.get("source_run_uid"))
            if target_run_uid and source_run_uid and source_run_uid != target_run_uid:
                skipped.append({"task_uid": task_uid, "reason": "run_uid_mismatch"})
                continue
            if sanitize_text(task.status) in {"verified", "closed", "published"}:
                skipped.append({"task_uid": task_uid, "reason": "locked_status"})
                continue
            fields = sanitize_obj(row.get("fields") or {})
            matched += 1
            for key in fields:
                if key in WRITABLE_FIELDS and sanitize_text(fields.get(key)):
                    changed_counter[key] += 1
            if not apply:
                continue
            try:
                manual_fill = metadata.get("manual_fill")
                if not isinstance(manual_fill, dict):
                    manual_fill = {}
                manual_fill.update({
                    **fields,
                    "source": "filled_knowledge_gap_excel",
                    "source_sheets": row.get("sheets") or [],
                    "source_rows": row.get("rows") or [],
                    "imported_at": _now_iso(),
                    "import_mode": "staging_review_only",
                })
                metadata["manual_fill"] = sanitize_obj(manual_fill)
                metadata["manual_fill_status"] = "ready_for_retest"
                metadata["manual_fill_source"] = "excel_import"
                metadata["manual_fill_imported_at"] = manual_fill["imported_at"]
                if target_run_uid:
                    metadata["source_run_uid"] = target_run_uid
                task.set_metadata(sanitize_obj(metadata))
                status = _status_after_import(task.status)
                metadata = _metadata_with_status_history(
                    task,
                    status=status,
                    changed_by=sanitize_text(fields.get("handler")),
                    note="manual fill imported; pending replay retest",
                )
                task.status = status
                task.set_metadata(sanitize_obj(metadata))
                updated.append(task.task_uid)
            except Exception as exc:
                errors.append({"task_uid": task_uid, "reason": sanitize_text(str(exc))})
        if apply:
            db.commit()
        else:
            db.rollback()
        result = sanitize_obj({
            "ok": not errors,
            "dry_run": not apply,
            "input": str(input_path),
            "run_uid": target_run_uid,
            "matched_count": matched,
            "skipped_count": len(skipped),
            "error_count": len(errors),
            "updated_task_uids": updated,
            "skipped_reasons": skipped,
            "errors": errors,
            "changed_fields_summary": dict(sorted(changed_counter.items())),
            "writes_formal_knowledge_base": False,
            "requires_retest_before_verified": True,
        })
        result["updated_task_uids"] = list(updated)
        return result
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import filled knowledge gap workbook into staging task metadata.")
    parser.add_argument("--input", required=True, help="Filled Excel workbook path")
    parser.add_argument("--run-uid", default="", help="Optional source run uid guard")
    parser.add_argument("--apply", action="store_true", help="Write staging/review metadata; default is dry-run")
    parser.add_argument("--json-output", default="", help="Optional sanitized JSON result path")
    args = parser.parse_args()
    result = import_filled_knowledge_gap_tasks(
        input_path=args.input,
        run_uid=args.run_uid,
        apply=bool(args.apply),
    )
    if args.json_output:
        output_path = Path(args.json_output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
