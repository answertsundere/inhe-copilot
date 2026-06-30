"""Export coverage gaps that block AI provisional/Product-first replay usage."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.db import SessionLocal, init_db
from app.models.eval_tables import AIProvisionalKnowledge, EvalRun, EvalTrace
from app.services.ai_provisional_usage_trace_service import extract_ai_provisional_usage_from_trace
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text
from app.services.knowledge_gap_task_service import KnowledgeGapTaskService


README_SHEET = "说明"
DRAFT_IDENTITY_SHEET = "草稿身份缺口"
MAPPING_GAP_SHEET = "平台商品映射缺口"
KNOWLEDGE_GAP_SHEET = "知识内容缺口"

DRAFT_IDENTITY_COLUMNS = [
    ("草稿 UID", "draft_uid"),
    ("关联任务 UID", "task_uid"),
    ("来源批次", "source_run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("问题类型", "query_fact_type"),
    ("字段名", "field_name"),
    ("当前 i_id", "i_id"),
    ("当前 SKU", "sku_code"),
    ("当前身份状态", "identity_status"),
    ("当前是否可用于 eval", "usable_for_eval_label"),
    ("阻断原因", "blocked_reason"),
    ("AI 预填答案预览", "provisional_answer_preview"),
    ("建议内部 i_id", "suggested_i_id"),
    ("建议 SKU", "suggested_sku"),
    ("人工确认状态", "confirmation_status"),
    ("处理人", "operator"),
    ("备注", "note"),
]

MAPPING_GAP_COLUMNS = [
    ("回放批次", "source_run_uid"),
    ("案例 ID", "case_uid"),
    ("轮次 ID", "turn_uid"),
    ("平台商品 ID", "platform_item_id"),
    ("平台商品 ID Hash", "platform_item_id_hash"),
    ("商品链接", "product_url"),
    ("商品链接域名", "product_url_host"),
    ("平台商品标题", "platform_product_title"),
    ("订单商品标题", "order_product_title"),
    ("买家问题预览", "buyer_message_preview"),
    ("未解析原因", "unresolved_reason"),
    ("候选数量", "candidate_count"),
    ("候选商品", "ambiguous_candidates"),
    ("建议内部 i_id", "suggested_i_id"),
    ("建议 SKU", "suggested_sku"),
    ("建议商品标题", "suggested_product_title"),
    ("人工确认状态", "confirmation_status"),
    ("处理人", "operator"),
    ("备注", "note"),
]

KNOWLEDGE_GAP_COLUMNS = [
    ("任务 UID", "task_uid"),
    ("来源批次", "source_run_uid"),
    ("缺口分类", "gap_category"),
    ("问题类型", "query_fact_type"),
    ("失败类型", "failure_type"),
    ("所需证据类型", "required_evidence_type"),
    ("目标系统", "target_system"),
    ("样本数", "sample_count"),
    ("风险等级", "risk_level"),
    ("当前阻断", "current_blocker"),
    ("已有 AI 草稿数", "provisional_draft_count"),
    ("可用 AI 草稿数", "usable_provisional_draft_count"),
    ("买家问题预览", "buyer_question_preview"),
    ("Agent 回复预览", "agent_reply_preview"),
    ("人工补充内容", "manual_content"),
    ("处理人", "operator"),
    ("备注", "note"),
]


def _latest_run_uid(db) -> str:
    run = (
        db.query(EvalRun)
        .filter(EvalRun.source_type == "real_conversation", EvalRun.status == "completed")
        .order_by(EvalRun.created_at.desc(), EvalRun.id.desc())
        .first()
    )
    return sanitize_text(run.run_uid) if run else ""


def _run_uid_from_json(path: str) -> str:
    if not sanitize_text(path):
        return ""
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        return ""
    return sanitize_text(
        data.get("run_uid")
        or (data.get("replay") or {}).get("run_uid")
        or (data.get("schedule") or {}).get("run_uid")
    )


def _json_preview(value: Any, limit: int = 240) -> str:
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, str):
        return sanitize_text(value)[:limit]
    return sanitize_text(json.dumps(sanitize_obj(value), ensure_ascii=False))[:limit]


def _first_text(*values: Any) -> str:
    for value in values:
        text = sanitize_text(str(value or "")).strip()
        if text:
            return text
    return ""


def _url_host(url: str) -> str:
    text = str(url or "")
    if not text:
        return ""
    try:
        return sanitize_text(urlsplit(text).netloc)
    except Exception:
        return ""


def _iter_dicts(value: Any):
    if isinstance(value, dict):
        yield value
        for item in value.values():
            yield from _iter_dicts(item)
    elif isinstance(value, list):
        for item in value:
            yield from _iter_dicts(item)


def _identity_from_trace(trace: EvalTrace) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    answer_trace = trace.get_answer_trace() or {}
    product_identity = trace.get_product_identity() or {}
    candidates = []
    for container in (product_identity, raw.get("real_context_product_identity"), answer_trace.get("real_context_product_identity")):
        if isinstance(container, dict):
            candidates.append(container)
    for container in (raw, answer_trace):
        if isinstance(container, dict):
            real_context = container.get("real_context")
            product = real_context.get("product") if isinstance(real_context, dict) else None
            if isinstance(product, dict):
                candidates.append(product)
    merged: dict[str, Any] = {}
    for candidate in candidates:
        for key, value in candidate.items():
            if value and key not in merged:
                merged[key] = value
    return merged


def _resolution_from_trace(trace: EvalTrace) -> dict[str, Any]:
    raw = trace.get_raw_response() or {}
    for item in _iter_dicts(raw):
        resolution = item.get("product_identity_resolution")
        if isinstance(resolution, dict) and resolution:
            return resolution
        if any(key in item for key in ("unresolved_reason", "ambiguous_candidates", "identity_confidence")):
            status = sanitize_text(item.get("status") or item.get("identity_status"))
            if status and status != "resolved":
                return item
    pack = raw.get("product_first_evidence_pack") if isinstance(raw, dict) else {}
    if isinstance(pack, dict):
        trace_data = pack.get("evidence_pack_trace") if isinstance(pack.get("evidence_pack_trace"), dict) else {}
        resolution = trace_data.get("product_identity_resolution")
        if isinstance(resolution, dict):
            return resolution
    return {}


def _identity_status(row: AIProvisionalKnowledge) -> str:
    metadata = row.get_metadata()
    status = sanitize_text(metadata.get("identity_status"))
    if status:
        return status
    trace = metadata.get("identity_binding_trace") if isinstance(metadata.get("identity_binding_trace"), dict) else {}
    status = sanitize_text(trace.get("identity_status"))
    if status:
        return status
    if sanitize_text(row.i_id) or sanitize_text(row.sku_code) or row.kb_product_id:
        return "resolved"
    return "unresolved"


def _blocked_reason(row: AIProvisionalKnowledge) -> str:
    metadata = row.get_metadata()
    trace = metadata.get("identity_binding_trace") if isinstance(metadata.get("identity_binding_trace"), dict) else {}
    return _first_text(
        metadata.get("usable_for_eval_blocked_reason"),
        trace.get("skipped_reason"),
        "missing_identity" if _identity_status(row) != "resolved" else "",
    )


def _collect_draft_identity_gaps(db) -> list[dict[str, Any]]:
    rows = db.query(AIProvisionalKnowledge).order_by(AIProvisionalKnowledge.updated_at.desc(), AIProvisionalKnowledge.id.desc()).all()
    result = []
    for row in rows:
        status = _identity_status(row)
        if bool(row.usable_for_eval) and status == "resolved":
            continue
        result.append(sanitize_obj({
            "draft_uid": row.draft_uid,
            "task_uid": row.task_uid,
            "source_run_uid": row.source_run_uid,
            "case_uid": row.case_uid,
            "turn_uid": row.turn_uid,
            "query_fact_type": row.query_fact_type,
            "field_name": row.field_name,
            "i_id": row.i_id,
            "sku_code": row.sku_code,
            "identity_status": status,
            "usable_for_eval_label": "是" if row.usable_for_eval else "否",
            "blocked_reason": _blocked_reason(row),
            "provisional_answer_preview": _json_preview(row.provisional_answer or row.provisional_value, 180),
            "suggested_i_id": "",
            "suggested_sku": "",
            "confirmation_status": "",
            "operator": "",
            "note": "",
        }))
    return result


def _collect_mapping_gaps(db, run_uid: str) -> list[dict[str, Any]]:
    traces = (
        db.query(EvalTrace)
        .filter(EvalTrace.run_uid == run_uid)
        .order_by(EvalTrace.case_uid.asc(), EvalTrace.turn_index.asc(), EvalTrace.id.asc())
        .all()
    )
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for trace in traces:
        resolution = _resolution_from_trace(trace)
        status = sanitize_text(resolution.get("status") or resolution.get("identity_status"))
        unresolved_reason = _first_text(resolution.get("unresolved_reason"), resolution.get("reason"))
        ambiguous = resolution.get("ambiguous_candidates") if isinstance(resolution.get("ambiguous_candidates"), list) else []
        if status == "resolved" and not unresolved_reason and not ambiguous:
            continue
        identity = _identity_from_trace(trace)
        row = {
            "source_run_uid": trace.run_uid,
            "case_uid": trace.case_uid,
            "turn_uid": trace.turn_uid,
            "platform_item_id": _first_text(identity.get("item_id"), identity.get("platform_item_id")),
            "platform_item_id_hash": _first_text(identity.get("item_id_hash"), identity.get("platform_item_id_hash")),
            "product_url": _first_text(identity.get("product_url"), identity.get("url")),
            "product_url_host": _url_host(_first_text(identity.get("product_url"), identity.get("url"))),
            "platform_product_title": _first_text(identity.get("product_title"), identity.get("platform_product_title")),
            "order_product_title": _first_text(identity.get("order_product_title")),
            "buyer_message_preview": sanitize_text(trace.buyer_message)[:120],
            "unresolved_reason": unresolved_reason or status or "identity_not_resolved",
            "candidate_count": len(ambiguous),
            "ambiguous_candidates": _json_preview(ambiguous),
            "suggested_i_id": "",
            "suggested_sku": "",
            "suggested_product_title": "",
            "confirmation_status": "",
            "operator": "",
            "note": "",
        }
        if not any(row.get(key) for key in ("platform_item_id_hash", "product_url", "platform_item_id", "platform_product_title", "order_product_title")):
            continue
        key = row.get("platform_item_id_hash") or row.get("product_url") or row.get("platform_item_id") or f"{row.get('platform_product_title')}|{row.get('order_product_title')}"
        if key in seen:
            continue
        seen.add(key)
        rows.append(sanitize_obj(row))
    return rows


def _collect_knowledge_content_gaps(db, run_uid: str) -> list[dict[str, Any]]:
    tasks_result = KnowledgeGapTaskService().list_tasks(db, filters={"run_uid": run_uid}, limit=500)
    tasks = tasks_result.get("items") or []
    task_uids = [sanitize_text(item.get("task_uid")) for item in tasks if sanitize_text(item.get("task_uid"))]
    drafts_by_task: dict[str, list[AIProvisionalKnowledge]] = {}
    if task_uids:
        rows = db.query(AIProvisionalKnowledge).filter(AIProvisionalKnowledge.task_uid.in_(task_uids)).all()
        for row in rows:
            drafts_by_task.setdefault(row.task_uid, []).append(row)
    result = []
    for item in tasks:
        draft_rows = drafts_by_task.get(sanitize_text(item.get("task_uid")), [])
        result.append(sanitize_obj({
            "task_uid": item.get("task_uid", ""),
            "source_run_uid": (item.get("metadata") or {}).get("source_run_uid") or run_uid,
            "gap_category": item.get("gap_category") or item.get("gap_type") or "",
            "query_fact_type": item.get("query_fact_type") or "",
            "failure_type": item.get("failure_type") or "",
            "required_evidence_type": item.get("required_evidence_type") or item.get("missing_evidence_type") or "",
            "target_system": item.get("target_system") or "",
            "sample_count": item.get("sample_count") or 0,
            "risk_level": item.get("risk_level") or "",
            "current_blocker": item.get("current_blocker") or "",
            "provisional_draft_count": len(draft_rows),
            "usable_provisional_draft_count": sum(1 for draft in draft_rows if draft.usable_for_eval),
            "buyer_question_preview": _json_preview((item.get("latest_buyer_questions") or [""])[0], 120),
            "agent_reply_preview": _json_preview((item.get("latest_agent_replies") or [""])[0], 120),
            "manual_content": "",
            "operator": "",
            "note": "",
        }))
    return result


def _usage_metrics(db, run_uid: str) -> dict[str, Any]:
    traces = db.query(EvalTrace).filter(EvalTrace.run_uid == run_uid).all()
    used_drafts: set[str] = set()
    used_turns = 0
    for trace in traces:
        usage = extract_ai_provisional_usage_from_trace(trace)
        if usage.get("provisional_used_turn"):
            used_turns += 1
        for item in usage.get("provisional_evidence") or []:
            draft_uid = sanitize_text(item.get("draft_uid") or item.get("provisional_draft_uid"))
            if draft_uid:
                used_drafts.add(draft_uid)
    return {"used_draft_count": len(used_drafts), "used_turn_count": used_turns}


def _summary_rows(
    *,
    run_uid: str,
    draft_gaps: list[dict[str, Any]],
    mapping_gaps: list[dict[str, Any]],
    knowledge_gaps: list[dict[str, Any]],
    metrics: dict[str, Any],
    db,
) -> list[list[Any]]:
    drafts = db.query(AIProvisionalKnowledge).all()
    status_counts = Counter(sanitize_text(row.verification_status) or "unknown" for row in drafts)
    skipped_reasons = Counter(row.get("blocked_reason") or "unknown" for row in draft_gaps)
    mapping_hosts = Counter(row.get("product_url_host") or "unknown" for row in mapping_gaps)
    query_gaps = Counter(row.get("query_fact_type") or "unknown" for row in knowledge_gaps)
    category_gaps = Counter(row.get("gap_category") or "unknown" for row in knowledge_gaps)
    with_identity = sum(1 for row in drafts if _identity_status(row) == "resolved")
    rows: list[list[Any]] = [
        ["字段", "值"],
        ["回放批次", run_uid],
        ["AI provisional 草稿总数", len(drafts)],
        ["usable_for_eval 草稿数", sum(1 for row in drafts if row.usable_for_eval)],
        ["usable_for_auto_send 草稿数", sum(1 for row in drafts if row.usable_for_auto_send)],
        ["已绑定身份草稿数", with_identity],
        ["本 run 使用 provisional draft 数", metrics.get("used_draft_count", 0)],
        ["本 run 使用 provisional turn 数", metrics.get("used_turn_count", 0)],
        ["草稿身份缺口数", len(draft_gaps)],
        ["平台商品映射缺口数", len(mapping_gaps)],
        ["知识内容缺口数", len(knowledge_gaps)],
        ["verification_status 分布", json.dumps(dict(status_counts), ensure_ascii=False)],
        ["Top 草稿阻断原因", json.dumps(dict(skipped_reasons.most_common(10)), ensure_ascii=False)],
        ["Top 缺映射域名", json.dumps(dict(mapping_hosts.most_common(10)), ensure_ascii=False)],
        ["Top query_fact_type 缺口", json.dumps(dict(query_gaps.most_common(10)), ensure_ascii=False)],
        ["Top gap_category 缺口", json.dumps(dict(category_gaps.most_common(10)), ensure_ascii=False)],
    ]
    return sanitize_obj(rows)


def _write_sheet(workbook: Workbook, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]):
    ws = workbook.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    headers = [label for label, _key in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        ws.append([row.get(key, "") for _label, key in columns])
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = max(14, min(38, len(str(header)) + 6))


def build_coverage_gap_workbook(run_uid: str = "", *, db_factory=None) -> tuple[str, Workbook, dict[str, Any]]:
    db_factory = db_factory or SessionLocal
    db = db_factory()
    try:
        target = sanitize_text(run_uid) or _latest_run_uid(db)
        draft_gaps = _collect_draft_identity_gaps(db)
        mapping_gaps = _collect_mapping_gaps(db, target) if target else []
        knowledge_gaps = _collect_knowledge_content_gaps(db, target) if target else []
        metrics = _usage_metrics(db, target) if target else {"used_draft_count": 0, "used_turn_count": 0}

        workbook = Workbook()
        readme = workbook.active
        readme.title = README_SHEET
        for row in _summary_rows(
            run_uid=target,
            draft_gaps=draft_gaps,
            mapping_gaps=mapping_gaps,
            knowledge_gaps=knowledge_gaps,
            metrics=metrics,
            db=db,
        ):
            readme.append(row)
        for cell in readme[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        readme.column_dimensions["A"].width = 30
        readme.column_dimensions["B"].width = 80

        _write_sheet(workbook, DRAFT_IDENTITY_SHEET, DRAFT_IDENTITY_COLUMNS, draft_gaps)
        _write_sheet(workbook, MAPPING_GAP_SHEET, MAPPING_GAP_COLUMNS, mapping_gaps)
        _write_sheet(workbook, KNOWLEDGE_GAP_SHEET, KNOWLEDGE_GAP_COLUMNS, knowledge_gaps)
        summary = sanitize_obj({
            "run_uid": target,
            "draft_identity_gap_count": len(draft_gaps),
            "mapping_gap_count": len(mapping_gaps),
            "knowledge_gap_count": len(knowledge_gaps),
            **metrics,
        })
        return target, workbook, summary
    finally:
        db.close()


def default_output_path() -> str:
    today = datetime.now().strftime("%Y%m%d")
    return str(PROJECT_ROOT / "outputs" / f"ai_provisional_coverage_gaps_{today}.xlsx")


def run_export(
    *,
    run_uid: str = "",
    run_json: str = "",
    output: str = "",
    desktop_output: str = "",
    db_factory=None,
) -> dict[str, Any]:
    target = sanitize_text(run_uid) or _run_uid_from_json(run_json)
    target, workbook, summary = build_coverage_gap_workbook(target, db_factory=db_factory)
    output_path = Path(output or default_output_path())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    desktop_path = sanitize_text(desktop_output)
    if desktop_path:
        Path(desktop_path).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(output_path, desktop_path)
    return sanitize_obj({
        "ok": True,
        **summary,
        "output": str(output_path),
        "desktop_output": desktop_path,
    })


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Export AI provisional/Product-first coverage gap workbook.")
    parser.add_argument("--run-uid", default="")
    parser.add_argument("--run-json", default="")
    parser.add_argument("--output", default="")
    parser.add_argument("--desktop-output", default="")
    args = parser.parse_args()
    init_db()
    result = run_export(
        run_uid=args.run_uid,
        run_json=args.run_json,
        output=args.output,
        desktop_output=args.desktop_output,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
