"""Backfill empty KBProduct structured fields from trusted product rows."""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.db import SessionLocal, init_db
from app.models.kb_tables import KBProduct
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text


FIELD_ALIASES = {
    "i_id": {"i_id", "内部 i_id", "商品货号", "商品ID"},
    "sku_code": {"sku_code", "SKU", "sku"},
    "product_title": {"product_title", "商品标题", "商品名称"},
    "dimensions": {"dimensions", "尺寸", "规格尺寸"},
    "material": {"material", "材质"},
    "gross_weight": {"gross_weight", "gross_weight_kg", "毛重", "毛重(kg)"},
    "load_capacity": {"load_capacity", "承重", "承重/容量"},
    "accessories": {"accessories", "配件清单", "配件"},
    "installation_method": {"installation_method", "安装方式"},
    "package_list": {"package_list", "包装清单", "打包清单"},
    "applicable_age": {"applicable_age", "适用年龄"},
    "color": {"color", "颜色"},
    "variant": {"variant", "款式", "规格"},
    "style": {"style", "风格"},
    "has_installation_video": {"has_installation_video", "是否有安装视频"},
    "has_installation_diagram": {"has_installation_diagram", "是否有安装图"},
    "certification_report": {"certification_report", "检测报告", "合格证"},
    "food_grade": {"food_grade", "食品级"},
    "safe_claim": {"safe_claim", "安全承诺"},
}

SPEC_FIELDS = {
    "dimensions": "dimensions",
    "material": "material",
    "load_capacity": "load_capacity",
    "accessories": "accessories",
    "installation_method": "installation_method",
    "package_list": "package_list",
    "applicable_age": "age_range",
    "color": "color",
    "variant": "variant",
    "style": "style",
    "has_installation_video": "has_installation_video",
    "has_installation_diagram": "has_installation_diagram",
}
LOGISTICS_FIELDS = {"gross_weight": "gross_weight_kg"}
HIGH_RISK_FIELDS = {"certification_report", "food_grade", "safe_claim"}


def _read_rows(path: str) -> list[dict[str, str]]:
    source = Path(path)
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        raw_rows = data if isinstance(data, list) else data.get("items", []) if isinstance(data, dict) else []
        return [_normalize_row(item) for item in raw_rows]
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = _select_sheet(workbook)
        headers = [sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            raw = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
            rows.append(_normalize_row(raw))
        return rows
    finally:
        workbook.close()


def _select_sheet(workbook):
    required = FIELD_ALIASES["i_id"] | FIELD_ALIASES["sku_code"]
    field_headers = set().union(*(aliases for key, aliases in FIELD_ALIASES.items() if key not in {"i_id", "sku_code"}))
    for sheet in workbook.worksheets:
        headers = {sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))}
        if (headers & required) and (headers & field_headers):
            return sheet
    return workbook[workbook.sheetnames[0]]


def _normalize_row(raw: dict[str, Any]) -> dict[str, str]:
    row: dict[str, str] = {}
    for key, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in raw:
                row[key] = sanitize_text(str(raw.get(alias) or ""))
                break
        row.setdefault(key, "")
    return row


def _find_product(db, row: dict[str, str]) -> tuple[KBProduct | None, str]:
    i_id = sanitize_text(row.get("i_id"))
    sku_code = sanitize_text(row.get("sku_code"))
    candidates: dict[int, KBProduct] = {}
    if i_id:
        product = db.query(KBProduct).filter(KBProduct.i_id == i_id).one_or_none()
        if product:
            candidates[product.id] = product
    if sku_code:
        for product in db.query(KBProduct).all():
            if any(sanitize_text(item.get("sku_code")) == sku_code for item in product.get_sku_list() if isinstance(item, dict)):
                candidates[product.id] = product
    if len(candidates) == 1:
        return next(iter(candidates.values())), ""
    if len(candidates) > 1:
        return None, "ambiguous_internal_identity"
    return None, "kb_product_not_found"


def _is_empty(value: Any) -> bool:
    text = sanitize_text(str(value or "")).strip()
    return not text or text in {"-", "--", "暂无", "无"}


def _write_conflicts(rows: list[dict[str, Any]], output: str) -> None:
    if not output:
        return
    workbook = Workbook()
    ws = workbook.active
    ws.title = "字段冲突"
    headers = ["产品 i_id", "字段", "已有值", "待导入值", "原因", "处理人", "备注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([row.get("i_id", ""), row.get("field", ""), row.get("existing_value", ""), row.get("incoming_value", ""), row.get("reason", ""), "", ""])
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def run_backfill(
    input_path: str,
    *,
    apply: bool = False,
    conflict_output: str = "",
    operator: str = "auto_structured_backfill",
    db_factory=None,
) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    rows = _read_rows(input_path)
    db = db_factory()
    stats = Counter()
    conflicts: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    try:
        for index, row in enumerate(rows, start=2):
            stats["product_checked_count"] += 1
            product, reason = _find_product(db, row)
            if not product:
                stats["missing_required_count"] += 1
                skipped.append({"row": index, "reason": reason})
                continue
            specs = product.get_specs()
            logistics = product.get_logistics()
            changed = False
            for source_key, target_key in SPEC_FIELDS.items():
                value = sanitize_text(row.get(source_key))
                if not value:
                    continue
                existing = specs.get(target_key)
                if _is_empty(existing):
                    specs[target_key] = value
                    changed = True
                    stats["field_updated_count"] += 1
                elif sanitize_text(str(existing)) != value:
                    stats["conflict_count"] += 1
                    conflicts.append({
                        "i_id": product.i_id,
                        "field": target_key,
                        "existing_value": existing,
                        "incoming_value": value,
                        "reason": "existing_verified_field_not_overwritten",
                    })
            for source_key, target_key in LOGISTICS_FIELDS.items():
                value = sanitize_text(row.get(source_key))
                if not value:
                    continue
                existing = logistics.get(target_key)
                if _is_empty(existing):
                    logistics[target_key] = value
                    changed = True
                    stats["field_updated_count"] += 1
                elif sanitize_text(str(existing)) != value:
                    stats["conflict_count"] += 1
                    conflicts.append({
                        "i_id": product.i_id,
                        "field": target_key,
                        "existing_value": existing,
                        "incoming_value": value,
                        "reason": "existing_verified_field_not_overwritten",
                    })
            for field in HIGH_RISK_FIELDS:
                if sanitize_text(row.get(field)):
                    stats["high_risk_pending_count"] += 1
                    conflicts.append({
                        "i_id": product.i_id,
                        "field": field,
                        "existing_value": "",
                        "incoming_value": sanitize_text(row.get(field)),
                        "reason": "high_risk_field_requires_manual_review",
                    })
            if changed:
                stats["product_updated_count"] += 1
                if apply:
                    specs.setdefault("trusted_auto_backfill", {})
                    specs["trusted_auto_backfill"].update({
                        "source_file": str(input_path),
                        "operator": operator,
                    })
                    product.set_specs(specs)
                    product.set_logistics(logistics)
                    product.updated_by = operator
        _write_conflicts(sanitize_obj(conflicts), conflict_output)
        if apply:
            db.commit()
        else:
            db.rollback()
        return sanitize_obj({
            "ok": True,
            "dry_run": not apply,
            "product_checked_count": stats["product_checked_count"],
            "product_updated_count": stats["product_updated_count"],
            "field_updated_count": stats["field_updated_count"],
            "conflict_count": stats["conflict_count"],
            "missing_required_count": stats["missing_required_count"],
            "high_risk_pending_count": stats["high_risk_pending_count"],
            "skipped_reasons": skipped,
            "conflicts": conflicts,
            "writes_verified_knowledge": False,
        })
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Backfill empty KBProduct structured fields from trusted data.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--conflict-output", default="")
    parser.add_argument("--operator", default="auto_structured_backfill")
    parser.add_argument("--json-output", "--output", dest="json_output", default="")
    args = parser.parse_args()
    init_db()
    result = run_backfill(args.input, apply=args.apply, conflict_output=args.conflict_output, operator=args.operator)
    if args.json_output:
        output = Path(args.json_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
