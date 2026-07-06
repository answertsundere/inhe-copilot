"""Read-only embedding and RAG readiness diagnostics.

This script does not call any embedding provider and never prints secrets. It
summarizes the current embedding configuration and recent replay evidence
selection so operators can distinguish configuration gaps from Agent logic bugs.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def _load_dotenv_safely() -> None:
    try:
        from dotenv import load_dotenv

        load_dotenv(PROJECT_ROOT / ".env")
    except Exception:
        pass


_load_dotenv_safely()

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models.eval_tables import EvalRun, EvalTrace  # noqa: E402
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text  # noqa: E402

try:  # noqa: E402
    from app.services.embedding_config_status_service import get_embedding_config_status
except Exception:  # pragma: no cover - fallback for clean trees without the optional service
    def get_embedding_config_status() -> dict[str, Any]:
        from app import config

        enabled = bool(getattr(config, "EMBEDDING_ENABLED", False))
        api_base = str(getattr(config, "EMBEDDING_API_BASE", "") or "").strip()
        api_key = str(getattr(config, "EMBEDDING_API_KEY", "") or "").strip()
        model_alias = str(getattr(config, "EMBEDDING_MODEL", "") or "").strip()
        reasons = []
        if not enabled:
            reasons.append("embedding_disabled")
        if not api_base:
            reasons.append("api_base_missing")
        if not api_key:
            reasons.append("api_key_missing")
        return {
            "embedding_enabled": enabled and bool(api_base) and bool(api_key),
            "embedding_config_enabled": enabled,
            "embedding_provider": _provider_from_base(api_base),
            "embedding_model_alias": model_alias,
            "api_base_configured": bool(api_base),
            "api_key_configured": bool(api_key),
            "last_embedding_call_status": "not_tracked",
            "fallback_reason": ",".join(reasons),
            "required_env_vars": ENV_KEYS,
        }


ENV_KEYS = [
    "COPILOT_EMBEDDING_ENABLED",
    "COPILOT_EMBEDDING_API_BASE",
    "COPILOT_EMBEDDING_API_KEY",
    "COPILOT_EMBEDDING_MODEL",
]


def _provider_from_base(api_base: str) -> str:
    text = str(api_base or "").lower()
    if not text:
        return ""
    if "dashscope" in text or "aliyun" in text or "alibabacloud" in text:
        return "dashscope"
    if "openai" in text:
        return "openai_compatible"
    return "openai_compatible"


def build_embedding_rag_readiness_report(
    *,
    run_uid: str = "",
    latest: bool = False,
    limit: int = 0,
    db_factory=None,
) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    db = db_factory()
    try:
        run = _resolve_run(db, run_uid=run_uid, latest=latest)
        traces: list[EvalTrace] = []
        if run:
            query = db.query(EvalTrace).filter(EvalTrace.run_uid == run.run_uid).order_by(EvalTrace.turn_index.asc(), EvalTrace.id.asc())
            if limit and limit > 0:
                query = query.limit(limit)
            traces = query.all()

        embedding_status = get_embedding_config_status()
        selected_counts = [_selected_evidence_count(trace) for trace in traces]
        selected_distribution = dict(Counter(_selected_bucket(count) for count in selected_counts))
        timeout_rows = [_rag_timeout_row(trace) for trace in traces if _has_rag_timeout(trace)]
        embedding_ready = bool(embedding_status.get("embedding_enabled"))
        zero_evidence_count = sum(1 for count in selected_counts if count == 0)
        embedding_not_configured_count = zero_evidence_count if not embedding_ready else 0

        evidence_rows = [
            {
                "run_uid": trace.run_uid,
                "case_uid": trace.case_uid,
                "turn_uid": trace.turn_uid,
                "query_fact_type": _query_fact_type(trace),
                "quality_bucket": _quality_bucket(trace),
                "selected_evidence_count": _selected_evidence_count(trace),
                "rag_timeout": _has_rag_timeout(trace),
                "embedding_not_configured": not embedding_ready and _selected_evidence_count(trace) == 0,
                "failure_labels": ", ".join(str(item) for item in (trace.get_failure_labels() or [])),
            }
            for trace in traces
        ]
        env_rows = [_env_row(key) for key in ENV_KEYS]
        recommendation_rows = _config_recommendations(embedding_status, timeout_rows)

        summary = {
            "run_uid": run.run_uid if run else run_uid,
            "total_traces": len(traces),
            "selected_evidence_count_distribution": selected_distribution,
            "selected_evidence_zero_count": zero_evidence_count,
            "embedding_not_configured_count": embedding_not_configured_count,
            "rag_timeout_count": len(timeout_rows),
            "embedding_enabled": embedding_ready,
            "fallback_retrieval_likely": not embedding_ready,
        }
        return sanitize_obj(
            {
                "run_uid": summary["run_uid"],
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "summary": summary,
                "embedding_config_status": _safe_embedding_status(embedding_status),
                "environment_variables": env_rows,
                "rag_timeout_rows": timeout_rows,
                "replay_evidence_rows": evidence_rows,
                "configuration_recommendations": recommendation_rows,
            }
        )
    finally:
        db.close()


def _resolve_run(db, *, run_uid: str, latest: bool):
    if run_uid:
        return db.query(EvalRun).filter(EvalRun.run_uid == run_uid).first()
    if latest or not run_uid:
        return (
            db.query(EvalRun)
            .filter(EvalRun.source_type == "real_conversation", EvalRun.status == "completed")
            .order_by(EvalRun.created_at.desc(), EvalRun.id.desc())
            .first()
        )
    return None


def _selected_evidence_count(trace: EvalTrace) -> int:
    selected = trace.get_selected_evidence() or []
    if isinstance(selected, list) and selected:
        return len(selected)
    raw = trace.get_raw_response() or {}
    answer = trace.get_answer_trace() or {}
    for value in (
        raw.get("selected_evidence_count"),
        answer.get("selected_evidence_count"),
        _dig(raw, "evidence_debug", "selected_evidence_count"),
    ):
        try:
            if value is not None:
                return int(value)
        except (TypeError, ValueError):
            continue
    return len(selected) if isinstance(selected, list) else 0


def _selected_bucket(count: int) -> str:
    if count <= 0:
        return "0"
    if count == 1:
        return "1"
    if count <= 3:
        return "2-3"
    return "4+"


def _query_fact_type(trace: EvalTrace) -> str:
    answer = trace.get_answer_trace() or {}
    tu = trace.get_turn_understanding() or {}
    return str(
        trace.query_fact_type
        or answer.get("query_fact_type")
        or tu.get("effective_query_fact_type")
        or tu.get("query_fact_type")
        or ""
    )


def _quality_bucket(trace: EvalTrace) -> str:
    bucket = trace.get_quality_bucket() or {}
    return str(bucket.get("quality_bucket") or bucket.get("bucket") or "")


def _has_rag_timeout(trace: EvalTrace) -> bool:
    return _contains_timeout(trace.get_answer_trace()) or _contains_timeout(trace.get_raw_response())


def _contains_timeout(value: Any) -> bool:
    if isinstance(value, dict):
        for key, item in value.items():
            key_text = str(key).lower()
            if "rag" in key_text and "timeout" in key_text and bool(item):
                return True
            if key_text in {"rag_timeout", "timeout"} and bool(item):
                return True
            if _contains_timeout(item):
                return True
    if isinstance(value, list):
        return any(_contains_timeout(item) for item in value)
    return False


def _rag_timeout_row(trace: EvalTrace) -> dict[str, Any]:
    return {
        "run_uid": trace.run_uid,
        "case_uid": trace.case_uid,
        "turn_uid": trace.turn_uid,
        "query_fact_type": _query_fact_type(trace),
        "selected_evidence_count": _selected_evidence_count(trace),
        "buyer_message_preview": sanitize_text(trace.buyer_message)[:80],
    }


def _env_row(key: str) -> dict[str, Any]:
    raw = str(os.environ.get(key) or "").strip()
    if key.endswith("API_KEY"):
        display = "已配置（已隐藏）" if raw else "未配置"
    elif key.endswith("API_BASE"):
        display = "已配置" if raw else "未配置"
    else:
        display = raw if raw else "未配置"
    return {"env_key": key, "configured": bool(raw), "value": display}


def _safe_embedding_status(status: dict[str, Any]) -> dict[str, Any]:
    safe = dict(status or {})
    if "api_key" in safe:
        safe["api_key"] = "已隐藏" if safe.get("api_key") else ""
    return safe


def _config_recommendations(status: dict[str, Any], timeout_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not status.get("embedding_config_enabled"):
        rows.append({"priority": "P0", "item": "启用 embedding", "action": "设置 COPILOT_EMBEDDING_ENABLED=true，重启 5011 服务后重跑 replay"})
    if not status.get("api_base_configured"):
        rows.append({"priority": "P0", "item": "配置 embedding API base", "action": "设置 COPILOT_EMBEDDING_API_BASE，不要写入日志或报告"})
    if not status.get("api_key_configured"):
        rows.append({"priority": "P0", "item": "配置 embedding API key", "action": "设置 COPILOT_EMBEDDING_API_KEY，诊断报告只显示是否配置，不显示明文"})
    if not status.get("embedding_model_alias"):
        rows.append({"priority": "P1", "item": "配置 embedding model", "action": "设置 COPILOT_EMBEDDING_MODEL，并确认当前 provider 支持"})
    if timeout_rows:
        rows.append({"priority": "P1", "item": "RAG 查询超时", "action": "确认 embedding 服务可用后，复查 RAG 候选过滤和超时阈值"})
    if not rows:
        rows.append({"priority": "P2", "item": "配置可用", "action": "继续检查 evidence role、商品字段和 verified 规则覆盖"})
    return rows


def _dig(value: Any, *keys: str) -> Any:
    current = value
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def build_workbook(report: dict[str, Any]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    _append_rows(ws, [["说明", "内容"], ["用途", "只读检查 embedding/RAG 配置和最新 replay 证据命中，不调用外部 API，不写数据库。"], ["回放批次", report.get("run_uid", "")], ["生成时间", report.get("generated_at", "")]])
    _write_sheet(wb, "总览", [("指标", "metric"), ("值", "value")], [{"metric": k, "value": _json_cell(v)} for k, v in (report.get("summary") or {}).items()])
    _write_sheet(wb, "环境变量状态", [("环境变量", "env_key"), ("是否配置", "configured"), ("值", "value")], report.get("environment_variables") or [])
    _write_sheet(wb, "RAG超时统计", [("回放批次", "run_uid"), ("案例ID", "case_uid"), ("轮次ID", "turn_uid"), ("问题类型", "query_fact_type"), ("选中证据数", "selected_evidence_count"), ("买家问题预览", "buyer_message_preview")], report.get("rag_timeout_rows") or [])
    _write_sheet(wb, "最新Replay证据命中", [("回放批次", "run_uid"), ("案例ID", "case_uid"), ("轮次ID", "turn_uid"), ("问题类型", "query_fact_type"), ("质量分桶", "quality_bucket"), ("选中证据数", "selected_evidence_count"), ("RAG超时", "rag_timeout"), ("Embedding未配置", "embedding_not_configured"), ("失败标签", "failure_labels")], report.get("replay_evidence_rows") or [])
    _write_sheet(wb, "建议配置项", [("优先级", "priority"), ("项目", "item"), ("建议动作", "action")], report.get("configuration_recommendations") or [])
    return wb


def _write_sheet(wb: Workbook, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title)
    _append_rows(ws, [[label for label, _ in columns]])
    for row in rows:
        ws.append([_json_cell(row.get(key)) for _label, key in columns])
    _format_sheet(ws)


def _append_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    _format_sheet(ws)


def _format_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            widths[cell.column_letter] = min(max(widths.get(cell.column_letter, 0), len(str(cell.value or "")) + 2), 60)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"


def _json_cell(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


def default_excel_path() -> str:
    return str(PROJECT_ROOT / "outputs" / f"embedding_rag_readiness_{datetime.now().strftime('%Y%m%d')}.xlsx")


def write_json(path: str, report: dict[str, Any]) -> None:
    if not path:
        return
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def write_excel(path: str, report: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(report).save(target)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Diagnose embedding/RAG readiness without external calls or database writes.")
    parser.add_argument("--run-uid", default="")
    parser.add_argument("--latest", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--json-output", default="")
    parser.add_argument("--excel-output", default="")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = build_embedding_rag_readiness_report(run_uid=args.run_uid, latest=args.latest, limit=args.limit)
    excel_output = args.excel_output
    if excel_output:
        write_excel(excel_output, report)
    write_json(args.json_output, report)
    summary = report.get("summary") or {}
    print(json.dumps({
        "run_uid": report.get("run_uid", ""),
        "selected_evidence_count_distribution": summary.get("selected_evidence_count_distribution", {}),
        "embedding_not_configured_count": summary.get("embedding_not_configured_count", 0),
        "rag_timeout_count": summary.get("rag_timeout_count", 0),
        "excel_output": excel_output,
        "json_output": args.json_output,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
