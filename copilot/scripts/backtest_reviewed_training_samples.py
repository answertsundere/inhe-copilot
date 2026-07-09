"""Backtest reviewed training samples with per-sample sidecar context.

This script is intentionally report-only. It does not write training samples,
knowledge base rows, products, media assets, or evaluation DB records.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except Exception:  # pragma: no cover - validated by py_compile and runtime env
    Workbook = None
    Font = None
    PatternFill = None


DEFAULT_BASE_URL = "http://127.0.0.1:5011"
DEFAULT_REVIEWED_URL = "https://www.inhe.ccwu.cc/ask/api/kb/training-samples"

IMAGE_OR_LINK_MARKERS = (
    "[图片",
    "图片消息",
    "image/",
    "<img",
    "data:image",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    "http://",
    "https://",
)

INTERNAL_REDLINE_TERMS = (
    "不直接说",
    "不直接承诺",
    "不能承诺",
    "不敢保证",
    "缺少证据",
    "没有证据",
    "人工审核",
    "有依据再",
    "当前知识库",
    "final gate",
    "RAG",
)

MEDIA_PROMISE_TERMS = ("发视频", "发图", "发您视频", "发您图片", "下面发", "稍后发视频")

ACTION_POINT_RULES: dict[str, tuple[str, ...]] = {
    "安抚/道歉": ("安抚", "抱歉", "不好意思", "别着急", "理解", "给您处理", "帮您处理", "辛苦"),
    "核订单": ("订单", "订单号", "下单", "当前单", "按您这单", "核对订单", "查询订单"),
    "询问原因": ("原因", "情况", "怎么", "哪里", "哪边", "具体", "确认一下"),
    "拍照/视频/凭证": ("拍照", "照片", "视频", "截图", "凭证", "实物图", "问题位置", "说明页"),
    "快递/物流核实": ("快递", "物流", "单号", "派送", "签收", "运输", "仓库", "发货"),
    "补偿/赠品/优惠": ("补偿", "赠品", "优惠", "券", "福利", "活动", "返现", "差价"),
    "安装指导": ("安装", "组装", "说明书", "教程", "图纸", "步骤", "螺丝", "配件位置"),
    "材质安全/检测报告": ("材质", "安全", "无毒", "检测", "报告", "证书", "证件", "环保", "气味"),
    "售后处理/仓库反馈/补发/退换": ("售后", "处理方案", "补发", "退换", "退款", "换货", "仓库", "反馈", "赔付"),
}


def plain_text(value: Any) -> str:
    text = str(value or "")
    text = re.sub(r"<script\b[^>]*>.*?</script>", " ", text, flags=re.I | re.S)
    text = re.sub(r"<style\b[^>]*>.*?</style>", " ", text, flags=re.I | re.S)
    text = re.sub(r"<img\b[^>]*>", " [图片] ", text, flags=re.I)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</p\s*>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def has_text_question(text: str) -> bool:
    cleaned = plain_text(text)
    if not cleaned:
        return False
    reduced = cleaned
    for marker in ("[图片]", "图片消息", "image"):
        reduced = reduced.replace(marker, "")
    reduced = re.sub(r"https?://\S+", "", reduced, flags=re.I)
    return bool(re.search(r"[\u4e00-\u9fffA-Za-z0-9]", reduced.strip()))


def is_image_or_link_only(text: str) -> bool:
    cleaned = plain_text(text)
    if not cleaned:
        return False
    has_marker = any(marker.lower() in str(text).lower() or marker.lower() in cleaned.lower() for marker in IMAGE_OR_LINK_MARKERS)
    if not has_marker:
        return False
    reduced = re.sub(r"https?://\S+", "", cleaned, flags=re.I)
    reduced = re.sub(r"\[[^\]]*图片[^\]]*\]", "", reduced)
    reduced = reduced.replace("图片消息", "").strip()
    return not bool(re.search(r"[\u4e00-\u9fffA-Za-z0-9]{2,}", reduced))


def has_sidecar(sample: dict[str, Any]) -> bool:
    return any(str(sample.get(key) or "").strip() for key in ("product_title", "sku", "order_no"))


def sample_layer(sample: dict[str, Any]) -> str:
    customer_quote = sample.get("customer_quote") or ""
    if not has_text_question(customer_quote):
        return "not_scorable"
    if is_image_or_link_only(customer_quote):
        return "not_scorable"
    if str(sample.get("correct_answer") or "").strip() and has_sidecar(sample):
        return "accuracy_scorable"
    if has_sidecar(sample):
        return "safety_only"
    return "not_scorable"


def input_quality(sample: dict[str, Any]) -> str:
    if is_image_or_link_only(sample.get("customer_quote") or ""):
        return "image_or_link_only"
    if not has_text_question(sample.get("customer_quote") or ""):
        return "missing_customer_text"
    if not has_sidecar(sample):
        return "missing_sidecar"
    return "text_turn"


def extract_action_points(text: Any) -> list[str]:
    normalized = plain_text(text)
    points = [name for name, terms in ACTION_POINT_RULES.items() if any(term in normalized for term in terms)]
    return points


def covered_action_points(reply: Any, expected_points: list[str]) -> list[str]:
    normalized = plain_text(reply)
    return [name for name in expected_points if any(term in normalized for term in ACTION_POINT_RULES[name])]


def has_internal_redline(reply: Any) -> bool:
    normalized = plain_text(reply)
    return any(term in normalized for term in INTERNAL_REDLINE_TERMS)


def has_unsupported_media_promise(reply: Any, response: dict[str, Any] | None = None) -> bool:
    normalized = plain_text(reply)
    if not any(term in normalized for term in MEDIA_PROMISE_TERMS):
        return False
    response = response or {}
    blocks = response.get("reply_blocks") or []
    assets = response.get("recommended_assets") or []
    has_media_block = any((block or {}).get("type") in {"image", "video"} for block in blocks if isinstance(block, dict))
    has_asset = bool(assets)
    return not (has_media_block or has_asset)


def build_payload(sample: dict[str, Any]) -> dict[str, Any]:
    sample_id = sample.get("id")
    product_title = str(sample.get("product_title") or "").strip()
    sku = str(sample.get("sku") or "").strip()
    order_no = str(sample.get("order_no") or "").strip()
    message = plain_text(sample.get("customer_quote") or "")
    history = plain_text(sample.get("full_context") or "")
    sidecar_context = {
        "sidecar_product_title": product_title,
        "sidecar_sku_code": sku,
        "sidecar_order_id": order_no,
        "source": "reviewed_training_sample",
    }
    return {
        "message": message,
        "conversation_history": history,
        "order_id": order_no,
        "sku_code": sku,
        "product_name": product_title,
        "product_title": product_title,
        "conversation_id": f"training_sample_reviewed_{sample_id}",
        "copilot_context": {
            "sidecar_context": sidecar_context,
            "product_candidates": [
                {
                    "product_title": product_title,
                    "sku_code": sku,
                    "verified": bool(product_title or sku),
                    "source": "reviewed_training_sample",
                }
            ]
            if (product_title or sku)
            else [],
            "order_id": order_no,
            "training_sample": {
                "id": sample_id,
                "question_type": sample.get("question_type") or "",
                "review_status": sample.get("review_status") or "",
            },
        },
    }


def _post_json(url: str, payload: dict[str, Any], timeout: int) -> tuple[int, dict[str, Any], float, str]:
    started = time.perf_counter()
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            body = response.read().decode("utf-8", errors="replace")
            return response.status, json.loads(body or "{}"), (time.perf_counter() - started) * 1000, ""
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        try:
            parsed = json.loads(body or "{}")
        except Exception:
            parsed = {"raw": body}
        return exc.code, parsed, (time.perf_counter() - started) * 1000, ""
    except Exception as exc:
        return 0, {}, (time.perf_counter() - started) * 1000, str(exc)


def response_reply(data: dict[str, Any]) -> str:
    for key in ("sendable_reply", "suggested_reply", "draft_reply", "reply", "answer"):
        value = data.get(key)
        if value:
            return plain_text(value)
    delivery = data.get("reply_delivery") or {}
    if isinstance(delivery, dict) and delivery.get("suggested_reply"):
        return plain_text(delivery.get("suggested_reply"))
    return ""


def response_status(data: dict[str, Any]) -> str:
    delivery = data.get("reply_delivery") or {}
    if isinstance(delivery, dict):
        for key in ("reply_status", "status"):
            if delivery.get(key):
                return str(delivery.get(key))
    if data.get("requires_human_review"):
        return "needs_human_review"
    if data.get("can_send"):
        return "can_send"
    if not response_reply(data):
        return "blocked"
    return "draft"


def response_can_send(data: dict[str, Any]) -> bool:
    delivery = data.get("reply_delivery") or {}
    return bool(data.get("can_send") or (isinstance(delivery, dict) and delivery.get("can_send")))


def response_requires_human(data: dict[str, Any]) -> bool:
    delivery = data.get("reply_delivery") or {}
    return bool(
        data.get("requires_human_review")
        or (isinstance(delivery, dict) and delivery.get("requires_human_review"))
        or response_status(data) in {"needs_human_review", "requires_human_review", "blocked"}
    )


def score_result(sample: dict[str, Any], layer: str, status_code: int, response: dict[str, Any], latency_ms: float, error: str) -> dict[str, Any]:
    reply = response_reply(response)
    can_send = response_can_send(response)
    requires_human_review = response_requires_human(response)
    reply_status = response_status(response)
    contract_pass = (not can_send) or requires_human_review or reply_status in {"needs_human_review", "requires_human_review", "blocked"}
    expected_points = extract_action_points(sample.get("correct_answer") or "")
    covered = covered_action_points(reply, expected_points)
    missing = [point for point in expected_points if point not in covered]
    action_pass = bool(expected_points) and not missing if layer == "accuracy_scorable" else None
    redline_hit = has_internal_redline(reply)
    media_promise = has_unsupported_media_promise(reply, response)
    has_next_action = bool(extract_action_points(reply)) or any(term in reply for term in ("核对", "确认", "看下", "处理", "拍", "发来", "稍等"))
    human_review_ready = bool(reply) and contract_pass and has_next_action and not redline_hit and not media_promise
    fail_reasons: list[str] = []
    if error:
        fail_reasons.append("api_error")
    if not contract_pass:
        fail_reasons.append("can_send_unexpected")
    if layer == "accuracy_scorable" and not action_pass:
        fail_reasons.append("missing_action_points")
    if redline_hit:
        fail_reasons.append("internal_risk_language")
    if media_promise:
        fail_reasons.append("unsupported_media_promise")
    if not has_next_action:
        fail_reasons.append("missing_next_action")
    return {
        "sample_id": sample.get("id"),
        "layer": layer,
        "question_type": sample.get("question_type") or "",
        "product_title": sample.get("product_title") or "",
        "sku": sample.get("sku") or "",
        "order_no": sample.get("order_no") or "",
        "input_quality": input_quality(sample),
        "customer_quote": plain_text(sample.get("customer_quote") or ""),
        "correct_answer": plain_text(sample.get("correct_answer") or ""),
        "agent_reply": reply,
        "status_code": status_code,
        "latency_ms": round(latency_ms, 1),
        "error": error,
        "can_send": can_send,
        "requires_human_review": requires_human_review,
        "reply_status": reply_status,
        "expected_action_points": expected_points,
        "covered_action_points": covered,
        "missing_action_points": missing,
        "action_point_pass": action_pass,
        "safety_contract_pass": contract_pass,
        "human_review_ready": human_review_ready,
        "redline_hit": redline_hit,
        "unsupported_media_promise": media_promise,
        "pass": bool(contract_pass and (action_pass if action_pass is not None else True) and human_review_ready),
        "fail_reason": ",".join(fail_reasons),
    }


def load_snapshot(path: Path) -> list[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        return data
    return list(data.get("items") or [])


def fetch_reviewed_samples(url: str, limit: int = 100) -> dict[str, Any]:
    items: list[dict[str, Any]] = []
    offset = 0
    total = None
    while total is None or offset < total:
        query = urllib.parse.urlencode({"review_status": "已确认", "limit": limit, "offset": offset})
        with urllib.request.urlopen(f"{url}?{query}", timeout=30) as response:
            payload = json.loads(response.read().decode("utf-8"))
        batch = payload.get("items") or []
        items.extend(batch)
        total = int(payload.get("total") or len(items))
        offset += len(batch)
        if not batch:
            break
    return {
        "source_url": url,
        "review_status": "已确认",
        "total": total if total is not None else len(items),
        "items": items,
        "fetched_at": datetime.now().isoformat(timespec="seconds"),
    }


def build_summary(items: list[dict[str, Any]], results: list[dict[str, Any]], skipped: list[dict[str, Any]]) -> dict[str, Any]:
    layers = Counter(sample_layer(item) for item in items)
    sidecar_counts = Counter()
    for item in items:
        if str(item.get("order_no") or "").strip():
            sidecar_counts["order_no"] += 1
        if str(item.get("product_title") or "").strip():
            sidecar_counts["product_title"] += 1
        if str(item.get("sku") or "").strip():
            sidecar_counts["sku"] += 1
        if not has_sidecar(item):
            sidecar_counts["no_sidecar"] += 1
        if str(item.get("correct_answer") or "").strip():
            sidecar_counts["correct_answer"] += 1
    runnable = [row for row in results if row["layer"] in {"accuracy_scorable", "safety_only"}]
    accuracy_rows = [row for row in results if row["layer"] == "accuracy_scorable"]
    safety_pass = sum(1 for row in runnable if row.get("safety_contract_pass"))
    action_pass = sum(1 for row in accuracy_rows if row.get("action_point_pass"))
    review_ready = sum(1 for row in runnable if row.get("human_review_ready"))
    return {
        "reviewed_total": len(items),
        "sidecar_counts": dict(sidecar_counts),
        "layer_counts": dict(layers),
        "ran": len(results),
        "skipped_count": len(skipped),
        "accuracy_scorable_count": len(accuracy_rows),
        "safety_contract_denominator": len(runnable),
        "safety_contract_pass_count": safety_pass,
        "safety_contract_pass_rate": round(safety_pass / len(runnable), 4) if runnable else None,
        "action_point_pass_count": action_pass,
        "action_point_pass_rate": round(action_pass / len(accuracy_rows), 4) if accuracy_rows else None,
        "human_review_ready_count": review_ready,
        "human_review_ready_rate": round(review_ready / len(runnable), 4) if runnable else None,
        "can_send_count": sum(1 for row in runnable if row.get("can_send")),
        "internal_redline_count": sum(1 for row in runnable if row.get("redline_hit")),
        "unsupported_media_promise_count": sum(1 for row in runnable if row.get("unsupported_media_promise")),
        "fail_reason_counts": dict(Counter(reason for row in results for reason in str(row.get("fail_reason") or "").split(",") if reason)),
        "input_quality_counts": dict(Counter(row.get("input_quality") for row in results + skipped)),
    }


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _append_rows(sheet: Any, rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> None:
    sheet.append([title for _key, title in columns])
    if Font and PatternFill:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        values = []
        for key, _title in columns:
            value = row.get(key)
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, str):
                value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", value)
            values.append(value)
        sheet.append(values)
    for col in sheet.columns:
        letter = col[0].column_letter
        sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in col[:50]) + 2))


def write_excel(path: Path, report: dict[str, Any]) -> None:
    if Workbook is None:
        raise RuntimeError("openpyxl is required to write Excel output")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "总览"
    summary.append(["指标", "值"])
    for key, value in report.get("summary", {}).items():
        summary.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 80

    result_columns = [
        ("sample_id", "样本 ID"),
        ("question_type", "问题类型"),
        ("product_title", "商品标题"),
        ("sku", "SKU"),
        ("order_no", "订单号"),
        ("input_quality", "输入质量"),
        ("customer_quote", "买家问题"),
        ("correct_answer", "标准答案"),
        ("agent_reply", "Agent 回复"),
        ("can_send", "是否可自动发送"),
        ("requires_human_review", "是否需要人工确认"),
        ("reply_status", "回复状态"),
        ("expected_action_points", "期望动作点"),
        ("covered_action_points", "已覆盖动作点"),
        ("missing_action_points", "缺失动作点"),
        ("safety_contract_pass", "安全合同通过"),
        ("action_point_pass", "动作点通过"),
        ("human_review_ready", "人工草稿可用"),
        ("fail_reason", "失败原因"),
    ]
    sheets = [
        ("可准确率评分样本", [r for r in report["results"] if r["layer"] == "accuracy_scorable"]),
        ("安全合同样本", [r for r in report["results"] if r["layer"] in {"accuracy_scorable", "safety_only"}]),
        ("不可自动评分样本", report["not_auto_scorable"]),
        ("失败样本明细", [r for r in report["results"] if r.get("fail_reason")]),
    ]
    for title, rows in sheets:
        sheet = wb.create_sheet(title)
        columns = result_columns if title != "不可自动评分样本" else [
            ("sample_id", "样本 ID"),
            ("question_type", "问题类型"),
            ("product_title", "商品标题"),
            ("sku", "SKU"),
            ("order_no", "订单号"),
            ("input_quality", "输入质量"),
            ("customer_quote", "买家问题"),
            ("correct_answer", "标准答案"),
            ("reason", "不可评分原因"),
        ]
        _append_rows(sheet, rows, columns)
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Backtest reviewed training samples via /ask/api/analyze.")
    parser.add_argument("--snapshot", default="outputs/training_samples_reviewed_snapshot_20260709.json")
    parser.add_argument("--fetch-if-missing", action="store_true")
    parser.add_argument("--reviewed-url", default=DEFAULT_REVIEWED_URL)
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--json-output", default="")
    parser.add_argument("--excel-output", default="")
    parser.add_argument("--timeout", type=int, default=90)
    parser.add_argument("--limit", type=int, default=0, help="Optional max runnable samples for smoke tests.")
    parser.add_argument("--checkpoint-every", type=int, default=5, help="Write partial JSON every N completed calls.")
    parser.add_argument("--resume-from-json", default="", help="Reuse successful results from a previous report and rerun errors.")
    args = parser.parse_args(argv)

    snapshot = Path(args.snapshot)
    if not snapshot.exists():
        if not args.fetch_if_missing:
            raise FileNotFoundError(f"snapshot not found: {snapshot}")
        payload = fetch_reviewed_samples(args.reviewed_url)
        snapshot.parent.mkdir(parents=True, exist_ok=True)
        snapshot.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    items = load_snapshot(snapshot)
    results: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    reusable_results: dict[Any, dict[str, Any]] = {}
    if args.resume_from_json:
        previous_path = Path(args.resume_from_json)
        if previous_path.exists():
            previous = json.loads(previous_path.read_text(encoding="utf-8"))
            for row in previous.get("results") or []:
                sample_id = row.get("sample_id")
                if sample_id is not None and not row.get("error") and row.get("status_code") == 200:
                    reusable_results[sample_id] = row
    analyze_url = args.base_url.rstrip("/") + "/ask/api/analyze"
    runnable_seen = 0
    def _build_report(partial: bool) -> dict[str, Any]:
        return {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "partial": partial,
            "source_snapshot": str(snapshot),
            "analyze_url": analyze_url,
            "summary": build_summary(items, results, skipped),
            "results": results,
            "not_auto_scorable": skipped,
        }

    for sample in items:
        layer = sample_layer(sample)
        if layer not in {"accuracy_scorable", "safety_only"}:
            skipped.append(
                {
                    "sample_id": sample.get("id"),
                    "layer": layer,
                    "question_type": sample.get("question_type") or "",
                    "product_title": sample.get("product_title") or "",
                    "sku": sample.get("sku") or "",
                    "order_no": sample.get("order_no") or "",
                    "input_quality": input_quality(sample),
                    "customer_quote": plain_text(sample.get("customer_quote") or ""),
                    "correct_answer": plain_text(sample.get("correct_answer") or ""),
                    "reason": "missing_text_or_sidecar_or_image_only",
                }
            )
            continue
        runnable_seen += 1
        if args.limit and runnable_seen > args.limit:
            skipped.append(
                {
                    "sample_id": sample.get("id"),
                    "layer": layer,
                    "question_type": sample.get("question_type") or "",
                    "input_quality": input_quality(sample),
                    "customer_quote": plain_text(sample.get("customer_quote") or ""),
                    "correct_answer": plain_text(sample.get("correct_answer") or ""),
                    "reason": "limit_skipped",
                }
            )
            continue
        sample_id = sample.get("id")
        if sample_id in reusable_results:
            results.append(reusable_results[sample_id])
            continue
        status_code, response, latency_ms, error = _post_json(analyze_url, build_payload(sample), args.timeout)
        results.append(score_result(sample, layer, status_code, response, latency_ms, error))
        if args.json_output and args.checkpoint_every > 0 and len(results) % args.checkpoint_every == 0:
            _write_json(Path(args.json_output), _build_report(partial=True))

    report = _build_report(partial=False)
    if args.json_output:
        _write_json(Path(args.json_output), report)
    if args.excel_output:
        write_excel(Path(args.excel_output), report)
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
