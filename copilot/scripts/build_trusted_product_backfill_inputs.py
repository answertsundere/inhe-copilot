"""Build standardized trusted inputs for product backfill tools.

This script normalizes trusted product data into three small Excel files used by
the auto-backfill scripts. It does not write the database.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.services.eval_sanitizer_service import hash_sensitive, sanitize_obj, sanitize_text


TRUSTED_SOURCES = {
    "dingtalk",
    "dingtalk_product_kb_sync",
    "trusted_product_sheet",
    "official_product_asset",
    "kb_export",
}
UNTRUSTED_SOURCES = {"chat", "history_chat", "real_conversation", "customer_upload", "screenshot"}

IDENTITY_HEADERS = [
    "platform_item_id",
    "platform_item_id_hash",
    "product_url",
    "platform_product_title",
    "i_id",
    "sku_code",
    "confirmation_status",
    "source",
]
STRUCTURED_HEADERS = [
    "i_id",
    "sku_code",
    "product_title",
    "dimensions",
    "material",
    "gross_weight",
    "load_capacity",
    "accessories",
    "installation_method",
    "package_list",
    "applicable_age",
    "color",
    "variant",
    "style",
    "has_installation_video",
    "has_installation_diagram",
]
MEDIA_HEADERS = ["i_id", "sku_code", "product_name", "asset_url", "asset_title", "source", "mime_type"]

ALIASES = {
    "platform_item_id": {"platform_item_id", "平台商品 ID", "平台商品ID", "商品ID", "item_id"},
    "platform_item_id_hash": {"platform_item_id_hash", "平台商品 ID Hash", "平台商品ID Hash", "item_id_hash"},
    "product_url": {"product_url", "商品链接", "平台商品链接", "url"},
    "platform_product_title": {"platform_product_title", "平台商品标题", "商品标题"},
    "i_id": {"i_id", "内部 i_id", "商品货号", "商品ID", "商品货号/证书货号"},
    "sku_code": {"sku_code", "SKU", "sku", "建议 SKU"},
    "product_title": {"product_title", "product_name", "商品标题", "商品名称"},
    "dimensions": {"dimensions", "尺寸", "规格", "规格尺寸"},
    "material": {"material", "材质"},
    "gross_weight": {"gross_weight", "gross_weight_kg", "毛重", "毛重(kg)"},
    "load_capacity": {"load_capacity", "承重", "承重/容量"},
    "accessories": {"accessories", "配件清单", "配件"},
    "installation_method": {"installation_method", "安装方式"},
    "package_list": {"package_list", "包装清单", "打包清单"},
    "applicable_age": {"applicable_age", "适用年龄", "age_range"},
    "color": {"color", "颜色"},
    "variant": {"variant", "款式", "规格"},
    "style": {"style", "风格"},
    "has_installation_video": {"has_installation_video", "是否有安装视频"},
    "has_installation_diagram": {"has_installation_diagram", "是否有安装图"},
    "asset_url": {"asset_url", "素材链接", "URL", "url"},
    "asset_title": {"asset_title", "素材标题", "文件名", "标题"},
    "source": {"source", "来源", "数据来源"},
    "mime_type": {"mime_type", "文件类型", "MIME"},
}


def _today(value: str = "") -> str:
    return sanitize_text(value) or datetime.now().strftime("%Y%m%d")


def _safe_url(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        parts = urlsplit(raw)
    except Exception:
        return sanitize_text(raw)
    if not parts.scheme:
        return sanitize_text(raw)
    return sanitize_text(urlunsplit((parts.scheme, parts.netloc, parts.path, "", "")))


def _first(raw: dict[str, Any], key: str) -> str:
    for alias in ALIASES.get(key, {key}):
        if alias in raw:
            return sanitize_text(str(raw.get(alias) or ""))
    return ""


def _first_raw(raw: dict[str, Any], key: str) -> str:
    for alias in ALIASES.get(key, {key}):
        if alias in raw:
            return str(raw.get(alias) or "")
    return ""


def _normalize_raw_row(raw: dict[str, Any], default_source: str = "") -> dict[str, str]:
    row = {key: _first(raw, key) for key in ALIASES}
    row["source"] = row.get("source") or sanitize_text(default_source)
    row["product_url"] = _safe_url(_first_raw(raw, "product_url"))
    row["asset_url"] = _safe_url(_first_raw(raw, "asset_url"))
    if row.get("platform_item_id") and not row.get("platform_item_id_hash"):
        row["platform_item_id_hash"] = hash_sensitive(row["platform_item_id"])
    return row


def _read_input_file(path: str, default_source: str = "") -> list[dict[str, str]]:
    source = Path(path)
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else data.get("items", []) if isinstance(data, dict) else []
        return [_normalize_raw_row(row, default_source=default_source) for row in rows if isinstance(row, dict)]
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows: list[dict[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            raw_headers = [sanitize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            if not any(raw_headers):
                continue
            for values in sheet.iter_rows(min_row=2, values_only=True):
                raw = {raw_headers[index]: values[index] for index in range(min(len(raw_headers), len(values)))}
                normalized = _normalize_raw_row(raw, default_source=default_source)
                if any(normalized.values()):
                    rows.append(normalized)
    finally:
        workbook.close()
    return rows


def _dingtalk_rows(args: argparse.Namespace) -> list[dict[str, str]]:
    import scripts.sync_dingtalk_product_kb as sync

    config = sync._resolve_config(args)
    token = sync.get_access_token(config["client_id"], config["client_secret"])
    product_records = sync.fetch_all_records(
        config["base_id"],
        config["product_sheet_id"],
        config["operator_id"],
        token,
        limit=args.limit_products,
    )
    sku_records = sync.fetch_all_records(
        config["base_id"],
        config["sku_sheet_id"],
        config["operator_id"],
        token,
        limit=args.limit_skus,
    )
    product_by_record_id = {sync._record_id(record): sync._product_info_from_record(record) for record in product_records if sync._record_id(record)}
    sku_infos = [sync._sku_info_from_record(record, product_by_record_id) for record in sku_records]
    product_by_iid = {row.get("i_id"): row for row in product_by_record_id.values() if row.get("i_id")}
    rows: list[dict[str, str]] = []
    for sku_record, sku_info in zip(sku_records, sku_infos):
        i_id = sanitize_text(sku_info.get("i_id"))
        product_info = product_by_iid.get(i_id, {})
        if not i_id:
            continue
        rows.append({
            "i_id": i_id,
            "sku_code": sanitize_text(sku_info.get("sku_code") or i_id),
            "product_title": sanitize_text(product_info.get("english_name") or product_info.get("english_customer_report_name") or sku_info.get("english_name") or i_id),
            "material": sanitize_text(product_info.get("material")),
            "gross_weight": sanitize_text(sku_info.get("gross_weight_kg")),
            "load_capacity": sanitize_text(product_info.get("load_capacity")),
            "applicable_age": sanitize_text(product_info.get("age_range")),
            "color": sanitize_text(sku_info.get("color")),
            "variant": sanitize_text(sku_info.get("spec")),
            "source": "dingtalk_product_kb_sync",
        })
        for media_row in sync._build_media_rows(sku_record, i_id, None, sku_info=sku_info):
            rows.append({
                "i_id": i_id,
                "sku_code": sanitize_text(sku_info.get("sku_code") or i_id),
                "product_title": sanitize_text(product_info.get("english_name") or i_id),
                "asset_url": _safe_url(media_row.get("asset_url", "")),
                "asset_title": sanitize_text(media_row.get("asset_title") or media_row.get("asset_type")),
                "source": "dingtalk_product_kb_sync",
                "mime_type": "video" if "video" in sanitize_text(media_row.get("asset_type")) else "image",
            })
    for product_info in product_by_record_id.values():
        if product_info.get("i_id"):
            rows.append({
                "i_id": sanitize_text(product_info.get("i_id")),
                "product_title": sanitize_text(product_info.get("english_name") or product_info.get("english_customer_report_name") or product_info.get("i_id")),
                "material": sanitize_text(product_info.get("material")),
                "load_capacity": sanitize_text(product_info.get("load_capacity")),
                "applicable_age": sanitize_text(product_info.get("age_range")),
                "source": "dingtalk_product_kb_sync",
            })
    return rows


def _is_trusted(row: dict[str, str]) -> bool:
    source = str(row.get("source") or "").strip().lower()
    return source in TRUSTED_SOURCES and source not in UNTRUSTED_SOURCES


def _has_any(row: dict[str, str], keys: list[str]) -> bool:
    return any(sanitize_text(row.get(key)) for key in keys)


def _dedupe(rows: list[dict[str, str]], keys: list[str]) -> list[dict[str, str]]:
    seen: set[tuple[str, ...]] = set()
    result = []
    for row in rows:
        key = tuple(sanitize_text(row.get(item)) for item in keys)
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def build_inputs(rows: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
    trusted = [row for row in rows if _is_trusted(row)]
    identity = [
        {key: row.get(key, "") for key in IDENTITY_HEADERS}
        for row in trusted
        if _has_any(row, ["platform_item_id", "platform_item_id_hash", "product_url"])
        and _has_any(row, ["i_id", "sku_code"])
    ]
    for row in identity:
        row["confirmation_status"] = row.get("confirmation_status") or "confirmed"
    structured = [
        {key: row.get(key, "") for key in STRUCTURED_HEADERS}
        for row in trusted
        if _has_any(row, ["i_id", "sku_code"])
        and _has_any(row, [key for key in STRUCTURED_HEADERS if key not in {"i_id", "sku_code", "product_title"}])
    ]
    media = [
        {key: row.get(key, "") for key in MEDIA_HEADERS}
        for row in trusted
        if _has_any(row, ["i_id", "sku_code"]) and sanitize_text(row.get("asset_url"))
    ]
    return {
        "identity": _dedupe(identity, ["platform_item_id_hash", "product_url", "i_id", "sku_code"]),
        "structured": _dedupe(structured, ["i_id", "sku_code", "material", "dimensions", "gross_weight", "load_capacity"]),
        "media": _dedupe(media, ["i_id", "sku_code", "asset_url"]),
        "manual_review": [sanitize_obj(row) for row in rows if not _is_trusted(row) or not _has_any(row, ["i_id", "sku_code"])],
    }


def _write_workbook(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "trusted_input"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([row.get(key, "") for key in headers])
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = max(14, min(36, len(header) + 4))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run_build(
    *,
    inputs: list[str] | None = None,
    output_dir: str = "outputs",
    date: str = "",
    from_dingtalk: bool = False,
    args: argparse.Namespace | None = None,
) -> dict[str, Any]:
    rows: list[dict[str, str]] = []
    for input_path in inputs or []:
        rows.extend(_read_input_file(input_path, default_source="trusted_product_sheet"))
    if from_dingtalk:
        if args is None:
            raise ValueError("args is required for DingTalk mode")
        rows.extend(_dingtalk_rows(args))
    built = build_inputs(rows)
    suffix = _today(date)
    base = Path(output_dir)
    identity_path = base / f"trusted_identity_mapping_input_{suffix}.xlsx"
    structured_path = base / f"trusted_product_structured_fields_input_{suffix}.xlsx"
    media_path = base / f"trusted_product_media_input_{suffix}.xlsx"
    _write_workbook(identity_path, IDENTITY_HEADERS, built["identity"])
    _write_workbook(structured_path, STRUCTURED_HEADERS, built["structured"])
    _write_workbook(media_path, MEDIA_HEADERS, built["media"])
    return sanitize_obj({
        "ok": True,
        "source_row_count": len(rows),
        "trusted_identity_rows": len(built["identity"]),
        "trusted_structured_rows": len(built["structured"]),
        "trusted_media_rows": len(built["media"]),
        "manual_review_rows": len(built["manual_review"]),
        "outputs": {
            "identity": str(identity_path),
            "structured": str(structured_path),
            "media": str(media_path),
        },
        "skipped_untrusted_count": sum(1 for row in rows if not _is_trusted(row)),
    })


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Build standardized trusted product backfill input workbooks.")
    parser.add_argument("--input", action="append", default=[])
    parser.add_argument("--output-dir", default="outputs")
    parser.add_argument("--date", default="")
    parser.add_argument("--from-dingtalk", action="store_true")
    parser.add_argument("--limit-products", type=int, default=0)
    parser.add_argument("--limit-skus", type=int, default=0)
    parser.add_argument("--base-id", default="")
    parser.add_argument("--product-sheet-id", default="")
    parser.add_argument("--sku-sheet-id", default="")
    parser.add_argument("--credentials-file", default="")
    parser.add_argument("--json-output", "--output", dest="json_output", default="")
    args = parser.parse_args()
    result = run_build(
        inputs=args.input,
        output_dir=args.output_dir,
        date=args.date,
        from_dingtalk=args.from_dingtalk,
        args=args,
    )
    if args.json_output:
        output = Path(args.json_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
