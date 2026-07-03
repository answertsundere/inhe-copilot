"""Diagnose customer-facing reply naturalness for real conversation replays.

This script is read-only. It checks whether customer-visible replies leak
internal policy wording or lack basic customer-service tone/actions. It does
not change replay scoring, knowledge, products, or media.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models.eval_tables import EvalRun, EvalTrace  # noqa: E402
from app.services.customer_facing_safe_handoff_service import (  # noqa: E402
    CUSTOMER_FACING_INTERNAL_REDLINE_TERMS,
)
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text  # noqa: E402


INTERNAL_POLICY_TERMS = tuple(
    sorted(
        set(CUSTOMER_FACING_INTERNAL_REDLINE_TERMS)
        | {
            "安全边界",
            "风控",
            "策略",
            "contract",
            "证据不足",
            "内部判断",
            "系统判断",
            "审核不通过",
        }
    )
)

EMPATHY_TERMS = (
    "亲",
    "您好",
    "抱歉",
    "别着急",
    "我帮您",
    "我先帮您",
    "我来帮您",
    "辛苦",
    "麻烦",
)

NEXT_ACTION_TERMS = (
    "核对",
    "核实",
    "确认",
    "处理",
    "方案",
    "发我",
    "拍",
    "截图",
    "稍等",
    "看一下",
    "一起看",
    "回复您",
    "给您",
)

AWKWARD_CUSTOMER_PHRASES = (
    "我这边不直接",
    "这边不直接",
    "有依据后",
    "有依据再",
    "不跟您绕",
    "按事实查",
    "命中",
    "可发送素材",
)

FORBIDDEN_COMMITMENT_TERMS = (
    "绝对安全",
    "完全安全",
    "保证安全",
    "保护宝宝安全",
    "适合0-6岁",
    "适合 0-6 岁",
    "无毒无味",
    "无毒",
    "有证书",
    "有检测报告",
    "一定可以",
    "肯定可以",
    "保证可以",
)

HIGH_RISK_FACT_TYPES = {
    "material",
    "material_safety",
    "certification_report",
    "age_range",
    "child_suitability",
    "child_safety",
    "load_capacity",
    "stability",
}

EXCEL_HEADERS = [
    "回放批次",
    "案例 ID",
    "轮次 ID",
    "买家问题",
    "问题类型",
    "是否可发送",
    "是否需人工核对",
    "回复状态",
    "AI 草稿/回复",
    "检测问题",
    "严重程度",
    "建议改写",
    "是否建议改代码",
    "建议修复归口",
]


def _count(values: list[str]) -> dict[str, int]:
    return dict(Counter(values).most_common())


def _latest_run(db) -> EvalRun | None:
    return (
        db.query(EvalRun)
        .filter(EvalRun.source_type == "real_conversation", EvalRun.status == "completed")
        .order_by(EvalRun.created_at.desc(), EvalRun.id.desc())
        .first()
    )


def _quality_bucket(trace: EvalTrace) -> str:
    quality = trace.get_quality_bucket() or {}
    return sanitize_text(quality.get("quality_bucket")) or ("auto_sendable" if trace.passed else "agent_error")


def _raw_can_send(trace: EvalTrace) -> bool:
    raw = trace.get_raw_response() or {}
    value = raw.get("can_send")
    if isinstance(value, bool):
        return value
    if raw.get("sendable_reply"):
        return True
    return bool(trace.agent_reply and not trace.requires_human_review and trace.passed)


def _reply_status(trace: EvalTrace) -> str:
    raw = trace.get_raw_response() or {}
    if sanitize_text(raw.get("reply_status")):
        return sanitize_text(raw.get("reply_status"))
    if trace.requires_human_review:
        return "requires_human_review"
    if _raw_can_send(trace):
        return "can_send"
    if not trace.agent_reply:
        return "empty_or_skipped"
    return "draft_or_blocked"


def _external_reply(trace: EvalTrace) -> str:
    raw = trace.get_raw_response() or {}
    candidates = [
        raw.get("sendable_reply"),
        raw.get("draft_reply"),
        raw.get("final_reply"),
        raw.get("suggested_reply"),
        trace.agent_reply,
    ]
    for value in candidates:
        text = sanitize_text(value)
        if text:
            return text
    return ""


def _expected_handoff(trace: EvalTrace) -> bool:
    bucket = _quality_bucket(trace)
    return bool(
        trace.requires_human_review
        or bucket in {"safe_handoff", "knowledge_gap"}
        or _reply_status(trace) == "requires_human_review"
    )


def _query_fact_type(trace: EvalTrace) -> str:
    answer = trace.get_answer_trace() or {}
    understanding = trace.get_turn_understanding() or {}
    return (
        sanitize_text(trace.query_fact_type)
        or sanitize_text(answer.get("query_fact_type"))
        or sanitize_text(answer.get("effective_query_fact_type"))
        or sanitize_text(understanding.get("query_fact_type"))
        or sanitize_text(understanding.get("effective_query_fact_type"))
    )


def _has_any(text: str, terms: tuple[str, ...] | list[str]) -> bool:
    return any(term and term in text for term in terms)


def _matching_terms(text: str, terms: tuple[str, ...] | list[str]) -> list[str]:
    return [term for term in terms if term and term in text]


def _detect_issues(trace: EvalTrace) -> list[dict[str, str]]:
    reply = _external_reply(trace)
    if not reply:
        return []
    fact_type = _query_fact_type(trace)
    handoff_expected = _expected_handoff(trace)
    issues: list[dict[str, str]] = []

    internal_hits = _matching_terms(reply, INTERNAL_POLICY_TERMS)
    if internal_hits:
        issues.append({
            "type": "internal_policy_phrase",
            "severity": "high",
            "detail": "、".join(internal_hits[:6]),
        })

    awkward_hits = _matching_terms(reply, AWKWARD_CUSTOMER_PHRASES)
    if awkward_hits:
        issues.append({
            "type": "awkward_customer_phrase",
            "severity": "medium",
            "detail": "、".join(awkward_hits[:6]),
        })

    if re.search(r"(为什么|原因).{0,8}(不能|不可以|无法)", reply) or "因为没有" in reply:
        issues.append({
            "type": "over_explaining_safety",
            "severity": "medium",
            "detail": "过度解释为什么不能回答",
        })

    if handoff_expected and not _has_any(reply, EMPATHY_TERMS):
        issues.append({
            "type": "missing_empathy",
            "severity": "low",
            "detail": "缺少自然承接或安抚",
        })

    if handoff_expected and not _has_any(reply, NEXT_ACTION_TERMS):
        issues.append({
            "type": "missing_next_action",
            "severity": "medium",
            "detail": "缺少下一步客服动作",
        })

    forbidden_hits = _matching_terms(reply, FORBIDDEN_COMMITMENT_TERMS)
    if forbidden_hits and (handoff_expected or fact_type in HIGH_RISK_FACT_TYPES):
        issues.append({
            "type": "forbidden_commitment",
            "severity": "high",
            "detail": "、".join(forbidden_hits[:6]),
        })

    if trace.requires_human_review and _raw_can_send(trace):
        issues.append({
            "type": "can_send_contract_risk",
            "severity": "high",
            "detail": "requires_human_review=true but can_send/sendable reply appears set",
        })

    return issues


def _suggested_rewrite(trace: EvalTrace) -> str:
    fact_type = _query_fact_type(trace)
    if fact_type in {"material", "material_safety", "certification_report", "odor"}:
        return "亲，材质、气味和检测说明我帮您按这款商品资料核对一下，避免说错。您稍等，我确认后给您准确回复。"
    if fact_type in {"age_range", "child_suitability", "child_safety"}:
        return "亲，宝宝适用和安全说明我需要按这款商品的适用年龄、材质和结构资料核对清楚。您稍等，我确认后给您准确建议。"
    if fact_type in {"load_capacity", "stability"}:
        return "亲，承重和稳定性我帮您按这款商品资料核对一下。您准备放什么物品、大概多重呢？我一起帮您看是否合适。"
    if fact_type == "installation":
        return "亲，安装资料我帮您按这款商品核对一下。您现在卡在哪一步？可以拍下当前位置，我一起帮您看。"
    if fact_type in {"promotion", "promotion_policy", "price_negotiation"}:
        return "亲，我帮您看下当前页面活动、优惠券和满减规则，最终以您下单页显示为准；如果页面没显示，我再帮您核对。"
    if fact_type in {"aftersales", "aftersales_policy"}:
        return "亲，先别着急，我先按当前订单帮您核对。麻烦发一下问题位置照片，我这边一起确认处理方案。"
    return "亲，这个我帮您按当前商品和订单信息核对一下，避免说错。您稍等，我确认后给您准确回复。"


def _row(trace: EvalTrace, issues: list[dict[str, str]]) -> dict[str, Any]:
    issue_types = [item["type"] for item in issues]
    severities = [item["severity"] for item in issues]
    severity = "high" if "high" in severities else ("medium" if "medium" in severities else ("low" if severities else "none"))
    return sanitize_obj({
        "run_uid": trace.run_uid,
        "case_uid": trace.case_uid,
        "turn_uid": trace.turn_uid,
        "buyer_message": trace.buyer_message,
        "query_fact_type": _query_fact_type(trace),
        "can_send": _raw_can_send(trace),
        "requires_human_review": bool(trace.requires_human_review),
        "reply_status": _reply_status(trace),
        "draft_reply": _external_reply(trace),
        "sendable_reply": sanitize_text((trace.get_raw_response() or {}).get("sendable_reply")),
        "detected_issues": issue_types,
        "issue_details": [item.get("detail", "") for item in issues],
        "severity": severity,
        "suggested_rewrite": _suggested_rewrite(trace) if issues else "",
        "should_fix_code": bool(issues),
        "suggested_fix_area": "customer_facing_safe_handoff" if issues else "",
    })


def diagnose_naturalness(run_uid: str = "", include_clean: bool = False, db_factory=None) -> dict[str, Any]:
    db = (db_factory or SessionLocal)()
    try:
        run = None
        if run_uid:
            run = db.query(EvalRun).filter(EvalRun.run_uid == sanitize_text(run_uid)).first()
        else:
            run = _latest_run(db)
        if run is None:
            return {"ok": False, "error": "no completed real_conversation run found"}
        traces = (
            db.query(EvalTrace)
            .filter(EvalTrace.run_uid == run.run_uid)
            .order_by(EvalTrace.turn_index.asc(), EvalTrace.id.asc())
            .all()
        )
        records = []
        issue_counts: Counter[str] = Counter()
        severity_counts: Counter[str] = Counter()
        scanned = 0
        for trace in traces:
            reply = _external_reply(trace)
            if not reply:
                continue
            scanned += 1
            issues = _detect_issues(trace)
            if issues or include_clean:
                records.append(_row(trace, issues))
            for issue in issues:
                issue_counts[issue["type"]] += 1
                severity_counts[issue["severity"]] += 1
        summary = {
            "run_uid": run.run_uid,
            "total_traces": len(traces),
            "scanned_reply_count": scanned,
            "issue_turn_count": sum(1 for record in records if record.get("detected_issues")),
            "issue_counts": dict(issue_counts.most_common()),
            "severity_counts": dict(severity_counts.most_common()),
            "requires_human_review_reply_count": sum(1 for trace in traces if trace.requires_human_review and _external_reply(trace)),
            "can_send_contract_risk_count": issue_counts.get("can_send_contract_risk", 0),
        }
        return sanitize_obj({
            "ok": True,
            "summary": summary,
            "items": records,
        })
    finally:
        db.close()


def write_excel(result: dict[str, Any], output: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客服话术自然度诊断"
    sheet.append(EXCEL_HEADERS)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for item in result.get("items") or []:
        sheet.append([
            item.get("run_uid", ""),
            item.get("case_uid", ""),
            item.get("turn_uid", ""),
            item.get("buyer_message", ""),
            item.get("query_fact_type", ""),
            item.get("can_send", False),
            item.get("requires_human_review", False),
            item.get("reply_status", ""),
            item.get("draft_reply", ""),
            "\n".join(item.get("detected_issues") or []),
            item.get("severity", ""),
            item.get("suggested_rewrite", ""),
            item.get("should_fix_code", False),
            item.get("suggested_fix_area", ""),
        ])
    widths = [18, 18, 24, 36, 18, 12, 16, 18, 50, 28, 12, 50, 16, 24]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=idx).column_letter].width = width
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


def _safe_json(value: Any) -> str:
    return json.dumps(sanitize_obj(value), ensure_ascii=False, indent=2)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Diagnose customer-facing reply naturalness for replay runs.")
    parser.add_argument("--run-uid", default="", help="Specific real_conversation run_uid. Defaults to latest completed run.")
    parser.add_argument("--json-output", default="", help="Optional JSON output path.")
    parser.add_argument("--excel-output", default="", help="Optional Excel output path.")
    parser.add_argument("--include-clean", action="store_true", help="Include replies without detected issues.")
    args = parser.parse_args(argv)

    result = diagnose_naturalness(run_uid=args.run_uid, include_clean=args.include_clean)
    if args.json_output:
        target = Path(args.json_output)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(_safe_json(result) + "\n", encoding="utf-8")
    if args.excel_output:
        write_excel(result, args.excel_output)

    summary = result.get("summary") if isinstance(result.get("summary"), dict) else {}
    print(json.dumps({
        "ok": result.get("ok"),
        "run_uid": summary.get("run_uid", ""),
        "scanned_reply_count": summary.get("scanned_reply_count", 0),
        "issue_turn_count": summary.get("issue_turn_count", 0),
        "issue_counts": summary.get("issue_counts", {}),
        "json_output": args.json_output or "",
        "excel_output": args.excel_output or "",
    }, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
