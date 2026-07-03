"""Summarize fact-type and generic-rule gaps from evidence-chain diagnostics.

Read-only by design. This script consumes the flattened JSON produced by
diagnose_evidence_chain_for_replay.py and exports grouped gap analysis for
human review before any code or data changes.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text  # noqa: E402


FACT_SAMPLE_COLUMNS = [
    ("语义簇", "cluster"),
    ("样本数", "count"),
    ("建议 fact_type", "suggested_fact_type"),
    ("是否可行动", "actionable"),
    ("是否需要 RAG", "needs_rag"),
    ("是否计分", "should_score"),
    ("建议代码修复", "code_fix_recommended"),
    ("修复理由", "reason"),
    ("代表买家问题", "sample_messages"),
    ("代表轮次", "sample_turn_uids"),
]

GENERIC_SAMPLE_COLUMNS = [
    ("规则簇", "cluster"),
    ("样本数", "count"),
    ("当前 fact_type", "query_fact_type"),
    ("建议规则类型", "suggested_rule_type"),
    ("是否建议补通用规则", "generic_rule_recommended"),
    ("是否建议代码修复", "code_fix_recommended"),
    ("原因", "reason"),
    ("代表买家问题", "sample_messages"),
    ("代表轮次", "sample_turn_uids"),
]

CANDIDATE_COLUMNS = [
    ("修复类型", "fix_type"),
    ("语义/规则簇", "cluster"),
    ("影响样本数", "count"),
    ("建议", "recommendation"),
    ("风险边界", "guardrail"),
    ("代表样本", "sample_messages"),
]

NO_CODE_COLUMNS = [
    ("原因类型", "reason_type"),
    ("样本数", "count"),
    ("说明", "description"),
    ("代表样本", "sample_messages"),
]


def _text(value: Any) -> str:
    return sanitize_text(value).strip()


def _json_cell(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(sanitize_obj(value), ensure_ascii=False)


def _unique_limited(values: list[str], limit: int = 5) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _text(value)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
        if len(result) >= limit:
            break
    return result


def _has_any(text: str, keywords: tuple[str, ...]) -> bool:
    return any(keyword in text for keyword in keywords)


def classify_fact_type_gap(message: str) -> dict[str, Any]:
    msg = _text(message)
    compact = msg.replace(" ", "")
    if not compact:
        return _cluster("context_insufficient", "", False, False, False, False, "空消息或无法读取，不应按代码硬补")
    if compact in {"哦", "嗯", "嗯嗯", "好", "好的", "可以", "行", "在吗", "在的", "人工", "谢谢", "麻烦了"}:
        return _cluster("社交/确认/等待", "", False, False, False, True, "应归为 acknowledgement/social，不应进入商品事实检索")
    if _has_any(compact, ("稍等", "等一下", "我看看", "人工", "客服")) and len(compact) <= 8:
        return _cluster("社交/确认/等待", "", False, False, False, True, "短等待/人工呼叫应单独识别，避免 unknown 污染统计")
    if _has_any(compact, ("开户", "账户", "账号", "公司", "邮箱", "qq.com", "@", "税号", "统一社会信用")):
        return _cluster("发票/企业信息上下文", "invoice_policy", False, False, False, True, "企业/账户/邮箱信息多为上下文补充，应识别为 context_update 或发票资料上下文")
    if _has_any(compact, ("链接", "给个链接", "有链接", "发链接", "发我下链接", "发我链接", "给链接", "拍哪个", "我拍")):
        return _cluster("商品链接/下单协助", "product_link", True, False, True, True, "链接/拍下协助是稳定客服动作，不应走商品事实 RAG")
    if _has_any(compact, ("发货", "物流", "快递", "到哪", "签收", "送货", "派送", "多久到", "几天到", "现货")):
        return _cluster("物流/发货/签收", "stock_shipping", True, False, True, True, "通用物流/发货意图，可走订单或发货规则边界")
    if _has_any(compact, ("订单", "下单", "拍下", "付款", "发票", "开票", "税票")):
        fact_type = "invoice_policy" if _has_any(compact, ("发票", "开票", "税票")) else "order_status"
        return _cluster("订单/发票状态", fact_type, True, False, True, True, "订单或发票类可用通用 fact_type，不依赖具体商品事实")
    if _has_any(compact, ("退货", "退款", "补发", "换货", "坏", "破", "少件", "少了", "不对", "发错", "赔", "售后", "纽扣少", "差的是")):
        return _cluster("售后/补发/退换", "aftersales", True, False, True, True, "售后动作应进入售后 fact_type，保持人工核实边界")
    if _has_any(compact, ("清洗", "清洁", "保养", "怎么洗", "能洗吗")):
        return _cluster("清洁/保养", "cleaning_care", True, True, True, True, "清洁保养有稳定 fact_type，缺资料仍应转知识缺口")
    if _has_any(compact, ("安装", "组装", "说明书", "教程", "视频", "图纸", "螺丝", "孔", "扣", "固定", "护栏", "围栏", "侧板", "配件", "补一面", "第四面", "爬梯", "改成")):
        if _has_any(compact, ("补", "第四面", "加装", "适配", "能用吗", "能不能用", "单独配", "独立", "改成", "行吗")):
            return _cluster("配件/结构适配", "accessory_compatibility", True, True, True, True, "配件补配/适配有稳定语义，但不能编造具体适配结论")
        return _cluster("安装/配件/结构", "installation", True, True, True, True, "安装资料和配件位置应进入安装/配件类 fact_type")
    if _has_any(compact, ("贺卡", "赠送", "送下", "小玩具", "礼物", "赠品")):
        return _cluster("赠品/贺卡/附加服务", "promotion_policy", True, False, True, True, "赠品/贺卡属于服务规则或活动边界，不能承诺具体赠品")
    if _has_any(compact, ("优惠", "福利", "便宜", "券", "满减", "活动", "返现", "多买", "最低价", "价格")):
        fact_type = "price_negotiation" if _has_any(compact, ("便宜", "多买", "最低价", "少点")) else "promotion_policy"
        return _cluster("优惠/活动/议价", fact_type, True, False, True, True, "促销/议价可走通用服务规则，不能承诺具体优惠")
    if _looks_like_dimension_or_space(compact):
        return _cluster("尺寸/空间/摆放", "dimensions", True, True, True, True, "尺寸/空间问题需要明确 fact_type，缺字段仍保留知识缺口")
    if _has_any(compact, ("材质", "味道", "有毒", "甲醛", "安全", "证书", "检测", "宝宝", "几岁", "周岁", "儿童")):
        if _has_any(compact, ("几岁", "周岁", "宝宝", "儿童")):
            return _cluster("儿童适用/年龄安全", "age_range", True, True, True, True, "儿童适用是高风险事实，识别后仍需证据/人工核实")
        return _cluster("材质/安全/检测", "material_safety", True, True, True, True, "材质安全和检测证明需高风险 fact_type，不可直接承诺")
    if _has_any(compact, ("哪个", "哪款", "买哪个", "怎么选", "区别", "一样吗", "有大点", "高一点", "几款", "组合也行", "一起的吗", "带盖", "配套", "想要")):
        return _cluster("选款/对比/指代不足", "comparison", True, True, True, False, "通常依赖上下文和候选商品，单靠代码规则容易误判")
    if _has_any(compact, ("可调节", "一体", "带拉链", "挂的吗", "能放", "会掉", "硬度", "承重", "压不住", "够吗")):
        return _cluster("结构功能/承重追问", "structure_function", True, True, True, True, "结构功能/承重追问可识别 fact_type，但缺事实时不能硬答")
    if _has_any(compact, ("奶瓶", "奶粉", "放几", "放多少", "容量", "几勺")):
        return _cluster("容量/收纳空间", "capacity", True, True, True, True, "容量/收纳空间是商品事实，需要结构化字段或人工补资料")
    if len(compact) <= 4:
        return _cluster("上下文不足/短句", "", False, False, False, False, "短句缺少可稳定语义，不建议代码硬归类")
    return _cluster("未归类商品/服务追问", "product_question", True, True, True, False, "需要人工复核语义簇后再决定是否补 fact_type")


def _looks_like_dimension_or_space(text: str) -> bool:
    if _has_any(text, ("尺寸", "多长", "多宽", "多高", "长度", "高度", "宽度", "厘米", "cm", "空间", "放得下", "放下", "阳台", "一米", "米", "高", "宽", "长")):
        return True
    return bool(re.search(r"\d+(\.\d+)?(cm|厘米|米|m|高|宽|长)", text, flags=re.IGNORECASE))


def _cluster(cluster: str, fact_type: str, actionable: bool, needs_rag: bool, should_score: bool, code_fix: bool, reason: str) -> dict[str, Any]:
    return {
        "cluster": cluster,
        "suggested_fact_type": fact_type,
        "actionable": actionable,
        "needs_rag": needs_rag,
        "should_score": should_score,
        "code_fix_recommended": code_fix,
        "reason": reason,
    }


def classify_generic_rule_gap(row: dict[str, Any]) -> dict[str, Any]:
    qft = _text(row.get("query_fact_type")) or "unknown"
    message = _text(row.get("buyer_message_preview"))
    if qft in {"promotion", "promotion_policy", "price_negotiation"}:
        return _generic("优惠/活动/议价规则", qft, "promotion_policy", True, False, "缺少通用优惠核对规则；只能给客服动作，不能承诺具体优惠")
    if qft in {"aftersales", "aftersales_policy"}:
        return _generic("售后处理规则", qft, "aftersales_policy", True, False, "缺少通用售后核实规则；不得承诺退款/补发/赔付")
    if qft in {"stock_shipping", "logistics", "order_status"}:
        return _generic("物流/发货规则", qft, "logistics_policy", True, False, "缺少发货/物流核对规则；订单状态仍需后台核实")
    if qft in {"invoice_policy"}:
        return _generic("发票规则", qft, "invoice_policy", True, False, "缺少发票服务规则；不涉及具体商品事实")
    if qft in {"installation"}:
        return _generic("安装无证据兜底", qft, "installation_no_evidence", True, False, "可补安装资料核对类安全兜底，但不能承诺有视频/图")
    if qft in {"material", "material_safety", "certification_report"}:
        return _generic("材质/检测安全兜底", qft, "material_safety", True, False, "高风险事实缺证据时只允许核对资料/转人工")
    if qft in {"age_range", "child_safety", "child_suitability"}:
        return _generic("儿童适用/安全兜底", qft, "child_safety", True, False, "儿童适用和安全承诺必须保守兜底")
    if qft in {"load_capacity"}:
        return _generic("承重兜底", qft, "load_capacity", True, False, "无承重事实时不能报公斤数，只能核对资料")
    if qft in {"placement_scene", "space_fit", "dimensions"}:
        return _generic("摆放/空间兜底", qft, "placement_scene", True, False, "缺尺寸或空间证据时只允许核对当前商品资料")
    if qft in {"accessory_availability", "accessory_compatibility", "structure_function"}:
        return _generic("配件/结构兜底", qft, qft, True, False, "配件补配和结构适配需要按当前款式核实")
    if _has_any(message, ("好的", "嗯", "在吗", "人工")):
        return _generic("社交/人工呼叫", qft, "human_handoff", False, False, "更适合 turn understanding 识别为不计分/人工呼叫")
    return _generic("其他通用规则缺口", qft, qft, False, False, "先人工复核，不建议直接 seed 规则")


def _generic(cluster: str, qft: str, rule_type: str, recommended: bool, code_fix: bool, reason: str) -> dict[str, Any]:
    return {
        "cluster": cluster,
        "query_fact_type": qft,
        "suggested_rule_type": rule_type,
        "generic_rule_recommended": recommended,
        "code_fix_recommended": code_fix,
        "reason": reason,
    }


def _load_records(path: str) -> list[dict[str, Any]]:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    records = data.get("zero_evidence_records") or data.get("records") or []
    if not isinstance(records, list):
        return []
    return [row for row in records if isinstance(row, dict)]


def summarize_gaps(input_path: str) -> dict[str, Any]:
    records = _load_records(input_path)
    fact_rows = [row for row in records if row.get("primary_reason") == "query_fact_type_missing"]
    generic_rows = [row for row in records if row.get("primary_reason") == "generic_rule_missing"]
    fact_groups = _group_fact_rows(fact_rows)
    generic_groups = _group_generic_rows(generic_rows)
    candidates = _build_fix_candidates(fact_groups, generic_groups)
    no_code = _build_no_code_rows(fact_groups, generic_groups)
    summary = {
        "input_path": input_path,
        "total_records": len(records),
        "query_fact_type_missing_count": len(fact_rows),
        "generic_rule_missing_count": len(generic_rows),
        "fact_type_gap_clusters": {row["cluster"]: row["count"] for row in fact_groups},
        "generic_rule_gap_clusters": {row["cluster"]: row["count"] for row in generic_groups},
        "code_fix_candidate_count": len(candidates),
        "no_code_fix_group_count": len(no_code),
    }
    return sanitize_obj(
        {
            "summary": summary,
            "fact_type_gap_samples": fact_groups,
            "generic_rule_gap_samples": generic_groups,
            "recommended_fix_candidates": candidates,
            "not_recommended_code_fixes": no_code,
        }
    )


def _group_fact_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[str, list[tuple[dict[str, Any], dict[str, Any]]]] = defaultdict(list)
    for row in rows:
        classified = classify_fact_type_gap(row.get("buyer_message_preview", ""))
        buckets[classified["cluster"]].append((row, classified))
    result = []
    for cluster, items in buckets.items():
        classified = items[0][1]
        messages = _unique_limited([item[0].get("buyer_message_preview", "") for item in items], 5)
        turns = _unique_limited([item[0].get("turn_uid", "") for item in items], 5)
        result.append(
            {
                "cluster": cluster,
                "count": len(items),
                "suggested_fact_type": classified["suggested_fact_type"],
                "actionable": classified["actionable"],
                "needs_rag": classified["needs_rag"],
                "should_score": classified["should_score"],
                "code_fix_recommended": classified["code_fix_recommended"],
                "reason": classified["reason"],
                "sample_messages": messages,
                "sample_turn_uids": turns,
            }
        )
    return sorted(result, key=lambda item: (-int(item["count"]), item["cluster"]))


def _group_generic_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[tuple[str, str], list[tuple[dict[str, Any], dict[str, Any]]]] = defaultdict(list)
    for row in rows:
        classified = classify_generic_rule_gap(row)
        buckets[(classified["cluster"], classified["query_fact_type"])].append((row, classified))
    result = []
    for (cluster, qft), items in buckets.items():
        classified = items[0][1]
        result.append(
            {
                "cluster": cluster,
                "count": len(items),
                "query_fact_type": qft,
                "suggested_rule_type": classified["suggested_rule_type"],
                "generic_rule_recommended": classified["generic_rule_recommended"],
                "code_fix_recommended": classified["code_fix_recommended"],
                "reason": classified["reason"],
                "sample_messages": _unique_limited([item[0].get("buyer_message_preview", "") for item in items], 5),
                "sample_turn_uids": _unique_limited([item[0].get("turn_uid", "") for item in items], 5),
            }
        )
    return sorted(result, key=lambda item: (-int(item["count"]), item["cluster"], item["query_fact_type"]))


def _build_fix_candidates(fact_groups: list[dict[str, Any]], generic_groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for group in fact_groups:
        if group.get("code_fix_recommended"):
            rows.append(
                {
                    "fix_type": "fact_type_contract",
                    "cluster": group["cluster"],
                    "count": group["count"],
                    "recommendation": f"补通用 fact_type: {group.get('suggested_fact_type') or 'non_actionable'}",
                    "guardrail": "只做语义/行动性识别；缺商品事实或素材仍保持 knowledge_gap/safe_handoff",
                    "sample_messages": group.get("sample_messages", []),
                }
            )
    for group in generic_groups:
        if group.get("generic_rule_recommended"):
            rows.append(
                {
                    "fix_type": "generic_rule_contract",
                    "cluster": group["cluster"],
                    "count": group["count"],
                    "recommendation": f"补充通用服务规则（{group.get('suggested_rule_type')}）",
                    "guardrail": "只允许客服动作/安全兜底，不替代具体商品事实、素材、活动政策",
                    "sample_messages": group.get("sample_messages", []),
                }
            )
    return sorted(rows, key=lambda item: (-int(item["count"]), item["fix_type"], item["cluster"]))


def _build_no_code_rows(fact_groups: list[dict[str, Any]], generic_groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for group in fact_groups:
        if not group.get("code_fix_recommended"):
            rows.append(
                {
                    "reason_type": f"FactType: {group['cluster']}",
                    "count": group["count"],
                    "description": group["reason"],
                    "sample_messages": group.get("sample_messages", []),
                }
            )
    for group in generic_groups:
        if not group.get("generic_rule_recommended"):
            rows.append(
                {
                    "reason_type": f"GenericRule: {group['cluster']}",
                    "count": group["count"],
                    "description": group["reason"],
                    "sample_messages": group.get("sample_messages", []),
                }
            )
    return sorted(rows, key=lambda item: (-int(item["count"]), item["reason_type"]))


def _write_sheet(workbook: Workbook, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append([label for label, _key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        sheet.append([_json_cell(row.get(key)) for _label, key in columns])
    for row in sheet.iter_rows():
        for cell in row:
            sheet.column_dimensions[cell.column_letter].width = min(max(sheet.column_dimensions[cell.column_letter].width or 10, len(str(cell.value or "")) + 2), 70)
    sheet.freeze_panes = "A2"


def build_workbook(report: dict[str, Any]) -> Workbook:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "说明"
    readme.append(["说明", "内容"])
    readme.append(["用途", "汇总 evidence chain 中 query_fact_type_missing 和 generic_rule_missing 的代表语义簇；只读，不写库。"])
    readme.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    for cell in readme[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    summary_rows = [{"metric": key, "value": _json_cell(value)} for key, value in (report.get("summary") or {}).items()]
    _write_sheet(workbook, "总览", [("指标", "metric"), ("数值", "value")], summary_rows)
    _write_sheet(workbook, "FactType 缺口样本", FACT_SAMPLE_COLUMNS, report.get("fact_type_gap_samples") or [])
    _write_sheet(workbook, "通用规则缺口样本", GENERIC_SAMPLE_COLUMNS, report.get("generic_rule_gap_samples") or [])
    _write_sheet(workbook, "建议修复候选", CANDIDATE_COLUMNS, report.get("recommended_fix_candidates") or [])
    _write_sheet(workbook, "不建议代码修复", NO_CODE_COLUMNS, report.get("not_recommended_code_fixes") or [])
    return workbook


def _write_json(path: str, report: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(sanitize_obj(report), ensure_ascii=True, indent=2), encoding="utf-8")


def _write_excel(path: str, report: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(report).save(target)


def default_excel_path() -> str:
    return str(Path.home() / "Desktop" / f"FactType和通用规则缺口_{datetime.now().strftime('%Y%m%d')}.xlsx")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Summarize fact-type and generic-rule gaps from evidence-chain diagnosis JSON.")
    parser.add_argument("--input", required=True, help="Path to evidence_chain diagnosis JSON.")
    parser.add_argument("--json-output", default="", help="Optional JSON output path.")
    parser.add_argument("--excel-output", default="", help="Optional Excel output path.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = summarize_gaps(args.input)
    excel_output = args.excel_output or default_excel_path()
    _write_excel(excel_output, report)
    if args.json_output:
        _write_json(args.json_output, report)
    print(
        json.dumps(
            {
                "query_fact_type_missing_count": report["summary"]["query_fact_type_missing_count"],
                "generic_rule_missing_count": report["summary"]["generic_rule_missing_count"],
                "fact_type_gap_clusters": report["summary"]["fact_type_gap_clusters"],
                "generic_rule_gap_clusters": report["summary"]["generic_rule_gap_clusters"],
                "excel_output": excel_output,
                "json_output": args.json_output,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
