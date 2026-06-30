"""Import staging-only AI provisional knowledge drafts from Excel."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from app.db import init_db
from app.services.ai_provisional_knowledge_service import AIProvisionalKnowledgeService
from app.services.eval_sanitizer_service import sanitize_text


FIELD_ALIASES = {
    "draft_uid": {"draft_uid", "草稿ID"},
    "source_run_uid": {"source_run_uid", "回放批次"},
    "case_uid": {"case_uid", "案例ID"},
    "turn_uid": {"turn_uid", "轮次ID"},
    "task_uid": {"task_uid", "知识缺口任务ID"},
    "i_id": {"i_id", "内部i_id"},
    "sku_code": {"sku_code", "SKU"},
    "query_fact_type": {"query_fact_type", "问题类型"},
    "field_name": {"field_name", "字段名"},
    "provisional_value": {"provisional_value", "预填字段值"},
    "provisional_answer": {"provisional_answer", "预填回复"},
    "confidence": {"confidence", "置信度"},
    "verification_status": {"verification_status", "审核状态"},
    "usable_for_eval": {"usable_for_eval", "可用于评测"},
    "note": {"note", "备注"},
}


def _canonical(value: Any) -> str:
    text = sanitize_text(value)
    for key, aliases in FIELD_ALIASES.items():
        if text in aliases:
            return key
    return ""


def read_rows(input_path: str) -> list[dict[str, str]]:
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    sheet = workbook["AI预填知识"] if "AI预填知识" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return []
    headers = [_canonical(value) for value in header_row]
    rows: list[dict[str, str]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item: dict[str, str] = {}
        for header, value in zip(headers, row):
            if header:
                item[header] = sanitize_text(value)
        if any(item.values()):
            rows.append(item)
    return rows


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Import AI provisional knowledge into staging table.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--apply", action="store_true", help="Persist drafts; default is dry-run.")
    parser.add_argument("--operator", default="codex")
    parser.add_argument("--json-output", default="")
    args = parser.parse_args()
    init_db()
    rows = read_rows(args.input)
    result = AIProvisionalKnowledgeService().import_rows(rows, apply=bool(args.apply), operator=args.operator)
    if args.json_output:
        output = Path(args.json_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
