"""Read-only evidence readiness diagnostics for Product-first/RAG replay.

The report answers a narrow question: do we already have verified structured
fields, usable media roles, or active generic rules that could support recent
replay questions? It does not write product data, media status, or knowledge.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models.eval_tables import EvalRun, EvalTrace  # noqa: E402
from app.models.kb_tables import KBGenericServiceRule, KBMediaAsset, KBProduct  # noqa: E402
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text  # noqa: E402
from app.services.media_asset_service import get_auto_send_level, get_media_purpose  # noqa: E402
from scripts.diagnose_evidence_chain_for_replay import diagnose_evidence_chain  # noqa: E402


FIELD_ALIASES = {
    "material": ["material", "材质"],
    "dimensions": ["dimensions", "dimension", "size", "尺寸", "长", "宽", "高"],
    "load_capacity": ["load_capacity", "capacity", "bearing", "承重", "载重"],
    "gross_weight": ["gross_weight", "package_weight", "weight", "毛重", "重量"],
    "accessories": ["accessories", "parts", "配件", "清单"],
    "installation": ["installation", "install", "manual", "guide", "安装", "说明书", "教程"],
    "age_range": ["age_range", "age", "适用年龄", "年龄", "宝宝", "儿童"],
    "structure_function": ["structure_function", "structure", "function", "结构", "功能"],
}

MEDIA_GROUPS = {
    "installation_video": {"installation_video", "install_video", "video"},
    "installation_diagram": {"installation_diagram", "install_image", "installation_image", "installation_guide", "pack_guide_image"},
    "manual": {"manual", "manual_image", "instruction", "instructions"},
    "size_chart": {"size_chart", "dimension_image", "space_fit_image"},
    "accessory_image": {"accessory_photo", "accessory_image", "parts_image"},
    "product_photo": {"product_photo", "sku_image", "image"},
}

RULE_FACT_TYPES = [
    "promotion_policy",
    "stock_shipping",
    "aftersales",
    "return_pickup",
    "order_assistance",
    "material_safety",
    "age_range",
    "load_capacity",
    "placement_scene",
]


def build_evidence_readiness_report(
    *,
    run_uid: str = "",
    latest: bool = False,
    limit: int = 0,
    db_factory=None,
) -> dict[str, Any]:
    db_factory = db_factory or SessionLocal
    db = db_factory()
    try:
        run = _resolve_run(db, run_uid=run_uid, latest=latest)
        traces = _load_traces(db, run.run_uid if run else "", limit=limit)
        products = db.query(KBProduct).all()
        media_assets = db.query(KBMediaAsset).all()
        rules = db.query(KBGenericServiceRule).all()

        field_rows = _product_field_rows(products)
        media_rows, media_role_rows = _media_rows(media_assets)
        rule_rows = _rule_rows(rules)
        chain_report = diagnose_evidence_chain(run_uid=run.run_uid if run else run_uid, limit=limit, db_factory=db_factory) if (run or run_uid) else {"records": [], "zero_evidence_records": [], "summary": {}}
        replay_rows = _replay_missing_rows(chain_report)
        p0_rows = _p0_rows(replay_rows)
        manual_rows = _manual_rows(media_role_rows, replay_rows)

        available_verified_field_count = sum(int(row["covered_product_count"]) for row in field_rows)
        approved_media_count = sum(1 for asset in media_assets if asset.status == "approved")
        pending_media_count = sum(1 for asset in media_assets if asset.status != "approved")
        generic_rule_available_count = sum(1 for rule in rules if rule.status == "active")
        selected_zero_count = sum(1 for trace in traces if _selected_evidence_count(trace) == 0)
        evidence_not_used_count = _evidence_not_used_count(chain_report)
        summary = {
            "run_uid": run.run_uid if run else run_uid,
            "product_count": len(products),
            "trace_count": len(traces),
            "selected_evidence_zero_count": selected_zero_count,
            "available_verified_field_count": available_verified_field_count,
            "approved_media_count": approved_media_count,
            "pending_media_count": pending_media_count,
            "generic_rule_available_count": generic_rule_available_count,
            "evidence_not_used_count": evidence_not_used_count,
            "query_fact_type_distribution": dict(Counter(_query_fact_type(trace) or "unknown" for trace in traces).most_common()),
        }
        return sanitize_obj(
            {
                "run_uid": summary["run_uid"],
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "summary": summary,
                "product_field_rows": field_rows,
                "media_rows": media_rows,
                "media_role_review_rows": media_role_rows,
                "generic_rule_rows": rule_rows,
                "replay_missing_evidence_rows": replay_rows,
                "p0_priority_rows": p0_rows,
                "manual_review_rows": manual_rows,
            }
        )
    finally:
        db.close()


def _resolve_run(db, *, run_uid: str, latest: bool):
    if run_uid:
        return db.query(EvalRun).filter(EvalRun.run_uid == run_uid).first()
    if latest or not run_uid:
        return (
            db.query(EvalRun)
            .filter(EvalRun.source_type == "real_conversation", EvalRun.status == "completed")
            .order_by(EvalRun.created_at.desc(), EvalRun.id.desc())
            .first()
        )
    return None


def _load_traces(db, run_uid: str, *, limit: int) -> list[EvalTrace]:
    if not run_uid:
        return []
    query = db.query(EvalTrace).filter(EvalTrace.run_uid == run_uid).order_by(EvalTrace.turn_index.asc(), EvalTrace.id.asc())
    if limit and limit > 0:
        query = query.limit(limit)
    return query.all()


def _product_field_rows(products: list[KBProduct]) -> list[dict[str, Any]]:
    rows = []
    total = len(products)
    for field, aliases in FIELD_ALIASES.items():
        covered = [product for product in products if _product_has_field(product, aliases)]
        rows.append(
            {
                "field_name": field,
                "covered_product_count": len(covered),
                "total_product_count": total,
                "coverage_rate": round(len(covered) / total, 4) if total else 0,
                "sample_products": "；".join((product.product_name or product.i_id or "") for product in covered[:5]),
            }
        )
    return rows


def _product_has_field(product: KBProduct, aliases: list[str]) -> bool:
    payloads = [product.get_specs(), product.get_logistics(), product.get_warranty()]
    for payload in payloads:
        if _dict_has_alias(payload, aliases):
            return True
    return False


def _dict_has_alias(value: Any, aliases: list[str]) -> bool:
    if isinstance(value, dict):
        for key, item in value.items():
            key_text = str(key).lower()
            if any(alias.lower() in key_text for alias in aliases) and str(item or "").strip():
                return True
            if _dict_has_alias(item, aliases):
                return True
    if isinstance(value, list):
        return any(_dict_has_alias(item, aliases) for item in value)
    return False


def _media_rows(media_assets: list[KBMediaAsset]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    grouped: dict[str, list[KBMediaAsset]] = defaultdict(list)
    role_review_rows: list[dict[str, Any]] = []
    for asset in media_assets:
        group = _media_group(asset)
        grouped[group].append(asset)
        purpose = _safe_media_purpose(asset)
        auto_level = _safe_auto_send_level(asset)
        if asset.status != "approved" or group == "unknown" or auto_level != "auto":
            role_review_rows.append(
                {
                    "asset_title": sanitize_text(asset.asset_title)[:120],
                    "asset_url": _safe_url(asset.asset_url),
                    "asset_type": asset.asset_type,
                    "media_role": purpose,
                    "status": asset.status,
                    "usable_for_agent": bool(asset.usable_for_agent),
                    "auto_send_level": auto_level,
                    "suggested_role": _suggest_media_role(asset),
                    "reason": _media_review_reason(asset, group, auto_level),
                    "high_risk": group in {"unknown", "product_photo"} or asset.status != "approved",
                    "auto_send_default": "否",
                    "owner": "素材运营/主管确认",
                }
            )
    rows: list[dict[str, Any]] = []
    for group in list(MEDIA_GROUPS.keys()) + ["unknown"]:
        assets = grouped.get(group, [])
        rows.append(
            {
                "media_group": group,
                "total_count": len(assets),
                "approved_count": sum(1 for asset in assets if asset.status == "approved"),
                "pending_review_count": sum(1 for asset in assets if asset.status != "approved"),
                "usable_for_agent_count": sum(1 for asset in assets if asset.usable_for_agent == 1),
                "auto_send_count": sum(1 for asset in assets if _safe_auto_send_level(asset) == "auto"),
            }
        )
    return rows, role_review_rows


def _media_group(asset: KBMediaAsset) -> str:
    tokens = {str(asset.asset_type or "").strip(), _safe_media_purpose(asset)}
    for group, aliases in MEDIA_GROUPS.items():
        if tokens & aliases:
            return group
    return "unknown"


def _safe_media_purpose(asset: KBMediaAsset) -> str:
    try:
        return str(get_media_purpose(asset) or "")
    except Exception:
        return ""


def _safe_auto_send_level(asset: KBMediaAsset) -> str:
    try:
        return str(get_auto_send_level(asset) or "")
    except Exception:
        return "disabled"


def _safe_url(url: str) -> str:
    text = str(url or "")
    return text.split("?", 1)[0][:160]


def _suggest_media_role(asset: KBMediaAsset) -> str:
    title = f"{asset.asset_title} {asset.asset_type} {_safe_media_purpose(asset)}".lower()
    if any(term in title for term in ["install", "安装", "说明书", "manual"]):
        return "installation_diagram/manual"
    if any(term in title for term in ["size", "dimension", "尺寸"]):
        return "size_chart"
    if any(term in title for term in ["accessory", "parts", "配件"]):
        return "accessory_image"
    return "需要人工确认"


def _media_review_reason(asset: KBMediaAsset, group: str, auto_level: str) -> str:
    if asset.status != "approved":
        return "素材未审核通过，不能当可发送证据"
    if asset.usable_for_agent != 1:
        return "素材未标记可供 Agent 使用"
    if group == "unknown":
        return "素材 role 无法归类"
    if auto_level != "auto":
        return "素材不满足自动发送等级"
    return "建议人工复核 role 是否准确"


def _rule_rows(rules: list[KBGenericServiceRule]) -> list[dict[str, Any]]:
    rows = []
    by_fact = defaultdict(list)
    for rule in rules:
        if rule.status == "active":
            by_fact[str(rule.fact_type or "")].append(rule)
    for fact_type in RULE_FACT_TYPES:
        matched = by_fact.get(fact_type, [])
        rows.append(
            {
                "query_fact_type": fact_type,
                "active_rule_count": len(matched),
                "rule_titles": "；".join(rule.title or rule.rule_key for rule in matched[:5]),
                "coverage_status": "已覆盖" if matched else "缺规则",
            }
        )
    return rows


def _replay_missing_rows(chain_report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for row in chain_report.get("zero_evidence_records") or []:
        rows.append(
            {
                "run_uid": row.get("run_uid"),
                "case_uid": row.get("case_uid"),
                "turn_uid": row.get("turn_uid"),
                "query_fact_type": row.get("query_fact_type"),
                "quality_bucket": row.get("quality_bucket"),
                "primary_reason": row.get("primary_reason"),
                "missing_evidence": row.get("missing_required_evidence") or row.get("missing_structured_fields") or row.get("media_issue"),
                "selected_evidence_count": row.get("selected_evidence_count"),
                "buyer_message_preview": row.get("buyer_message_preview"),
                "suggested_action": row.get("suggested_action") or _suggest_action(row.get("primary_reason")),
            }
        )
    return rows


def _p0_rows(replay_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for row in replay_rows:
        reason = str(row.get("primary_reason") or "")
        if reason in {"structured_field_missing", "media_missing", "media_role_not_sendable", "generic_rule_missing", "evidence_role_mismatch"}:
            copied = dict(row)
            copied["priority"] = "P0"
            copied["why"] = "问题明确且已有 replay 需求，只差可信字段/素材/规则证据"
            result.append(copied)
    return result


def _manual_rows(media_role_rows: list[dict[str, Any]], replay_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [
        {
            "item_type": "素材Role",
            "item": row.get("asset_title"),
            "reason": row.get("reason"),
            "owner": row.get("owner"),
            "auto_send_default": "否",
        }
        for row in media_role_rows[:200]
    ]
    for row in replay_rows:
        if row.get("primary_reason") in {"media_role_not_sendable", "evidence_role_mismatch", "product_identity_unresolved"}:
            rows.append(
                {
                    "item_type": "Replay缺失证据",
                    "item": row.get("turn_uid"),
                    "reason": row.get("primary_reason"),
                    "owner": "知识/素材运营确认",
                    "auto_send_default": "否",
                }
            )
    return rows


def _selected_evidence_count(trace: EvalTrace) -> int:
    selected = trace.get_selected_evidence() or []
    if isinstance(selected, list) and selected:
        return len(selected)
    raw = trace.get_raw_response() or {}
    value = raw.get("selected_evidence_count")
    try:
        return int(value) if value is not None else 0
    except (TypeError, ValueError):
        return 0


def _query_fact_type(trace: EvalTrace) -> str:
    answer = trace.get_answer_trace() or {}
    tu = trace.get_turn_understanding() or {}
    return str(trace.query_fact_type or answer.get("query_fact_type") or tu.get("effective_query_fact_type") or tu.get("query_fact_type") or "")


def _evidence_not_used_count(chain_report: dict[str, Any]) -> int:
    count = 0
    for row in chain_report.get("records") or []:
        if int(row.get("selected_evidence_count") or 0) != 0:
            continue
        has_available = bool(row.get("structured_field_coverage")) or int(row.get("media_available_count") or 0) > 0 or int(row.get("generic_rule_count") or 0) > 0
        if has_available:
            count += 1
    return count


def _suggest_action(reason: str | None) -> str:
    if reason == "embedding_not_configured":
        return "配置 embedding 环境变量并重启服务后重跑 replay"
    if reason == "structured_field_missing":
        return "补齐商品结构化字段，人工确认后再进入 verified"
    if reason in {"media_missing", "media_role_not_sendable", "evidence_role_mismatch"}:
        return "上传或确认 approved/usable 且 role 匹配的素材"
    if reason == "generic_rule_missing":
        return "补充 verified 通用服务规则，不替代商品事实"
    if reason == "product_identity_unresolved":
        return "补商品身份映射或侧栏商品信息"
    return "人工复核证据链断点"


def build_workbook(report: dict[str, Any]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    _append_rows(ws, [["说明", "内容"], ["用途", "只读诊断商品字段、素材 role、通用规则和 replay 缺失证据；不写数据库，不改素材状态。"], ["回放批次", report.get("run_uid", "")], ["生成时间", report.get("generated_at", "")]])
    _write_sheet(wb, "总览", [("指标", "metric"), ("值", "value")], [{"metric": k, "value": _json_cell(v)} for k, v in (report.get("summary") or {}).items()])
    _write_sheet(wb, "商品字段覆盖", [("字段", "field_name"), ("已覆盖商品数", "covered_product_count"), ("商品总数", "total_product_count"), ("覆盖率", "coverage_rate"), ("示例商品", "sample_products")], report.get("product_field_rows") or [])
    _write_sheet(wb, "素材覆盖", [("素材类型", "media_group"), ("总数", "total_count"), ("已审核", "approved_count"), ("待审核", "pending_review_count"), ("Agent可用", "usable_for_agent_count"), ("可自动发送", "auto_send_count")], report.get("media_rows") or [])
    _write_sheet(wb, "素材Role待确认", [("素材标题", "asset_title"), ("素材URL", "asset_url"), ("当前asset_type", "asset_type"), ("当前media_role", "media_role"), ("状态", "status"), ("Agent可用", "usable_for_agent"), ("自动发送等级", "auto_send_level"), ("建议Role", "suggested_role"), ("原因", "reason"), ("高风险", "high_risk"), ("默认可自动发送", "auto_send_default"), ("确认人", "owner")], report.get("media_role_review_rows") or [])
    _write_sheet(wb, "通用规则覆盖", [("问题类型", "query_fact_type"), ("Active规则数", "active_rule_count"), ("规则标题", "rule_titles"), ("覆盖状态", "coverage_status")], report.get("generic_rule_rows") or [])
    _write_sheet(wb, "Replay缺失证据", [("回放批次", "run_uid"), ("案例ID", "case_uid"), ("轮次ID", "turn_uid"), ("问题类型", "query_fact_type"), ("质量分桶", "quality_bucket"), ("主原因", "primary_reason"), ("缺失证据", "missing_evidence"), ("选中证据数", "selected_evidence_count"), ("买家问题预览", "buyer_message_preview"), ("建议动作", "suggested_action")], report.get("replay_missing_evidence_rows") or [])
    _write_sheet(wb, "P0优先补齐", [("优先级", "priority"), ("回放批次", "run_uid"), ("案例ID", "case_uid"), ("轮次ID", "turn_uid"), ("问题类型", "query_fact_type"), ("主原因", "primary_reason"), ("缺失证据", "missing_evidence"), ("为什么优先", "why"), ("建议动作", "suggested_action")], report.get("p0_priority_rows") or [])
    _write_sheet(wb, "必须人工确认", [("类型", "item_type"), ("项目", "item"), ("原因", "reason"), ("负责人", "owner"), ("默认可自动发送", "auto_send_default")], report.get("manual_review_rows") or [])
    return wb


def _write_sheet(wb: Workbook, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title)
    _append_rows(ws, [[label for label, _ in columns]])
    for row in rows:
        ws.append([_json_cell(row.get(key)) for _label, key in columns])
    _format_sheet(ws)


def _append_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    _format_sheet(ws)


def _format_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            widths[cell.column_letter] = min(max(widths.get(cell.column_letter, 0), len(str(cell.value or "")) + 2), 60)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"


def _json_cell(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


def default_excel_path() -> str:
    return str(Path.home() / "Desktop" / f"证据就绪度诊断_{datetime.now().strftime('%Y%m%d')}.xlsx")


def write_json(path: str, report: dict[str, Any]) -> None:
    if not path:
        return
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def write_excel(path: str, report: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(report).save(target)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Diagnose evidence readiness without database writes.")
    parser.add_argument("--run-uid", default="")
    parser.add_argument("--latest", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--json-output", default="")
    parser.add_argument("--excel-output", default="")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = build_evidence_readiness_report(run_uid=args.run_uid, latest=args.latest, limit=args.limit)
    excel_output = args.excel_output or default_excel_path()
    write_excel(excel_output, report)
    write_json(args.json_output, report)
    summary = report.get("summary") or {}
    print(json.dumps({
        "run_uid": report.get("run_uid", ""),
        "available_verified_field_count": summary.get("available_verified_field_count", 0),
        "approved_media_count": summary.get("approved_media_count", 0),
        "pending_media_count": summary.get("pending_media_count", 0),
        "generic_rule_available_count": summary.get("generic_rule_available_count", 0),
        "evidence_not_used_count": summary.get("evidence_not_used_count", 0),
        "excel_output": excel_output,
        "json_output": args.json_output,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
