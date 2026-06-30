"""Export manual review workbook for unresolved auto-backfill items."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text


SHEETS = {
    "身份冲突": ["来源", "行号", "原因", "内部 ID", "候选/冲突信息", "处理人", "备注"],
    "标题弱匹配": ["来源", "行号", "原因", "平台标题", "订单标题", "处理人", "备注"],
    "字段冲突": ["来源", "产品 i_id", "字段", "已有值", "待导入值", "原因", "处理人", "备注"],
    "素材不确定": ["来源", "行号", "原因", "素材链接", "素材标题", "处理人", "备注"],
    "高风险字段待确认": ["来源", "产品 i_id", "字段", "待确认值", "原因", "处理人", "备注"],
}


def default_output_path() -> str:
    today = datetime.now().strftime("%Y%m%d")
    return str(Path.home() / "Desktop" / f"自动回填后仍需人工确认_{today}.xlsx")


def _load_json(path: str) -> dict[str, Any]:
    if not sanitize_text(path) or not Path(path).exists():
        return {}
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _safe_url(value: str) -> str:
    text = sanitize_text(value)
    try:
        parts = urlsplit(text)
    except Exception:
        return text
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "", "")) if parts.scheme else text


def _append_rows_from_reports(workbook: Workbook, reports: list[tuple[str, dict[str, Any]]]) -> None:
    for source_name, report in reports:
        for item in report.get("conflicts") or []:
            reason = sanitize_text(item.get("reason"))
            if reason == "identity_conflict":
                workbook["身份冲突"].append([
                    source_name,
                    item.get("row", ""),
                    reason,
                    item.get("i_id", ""),
                    json.dumps(sanitize_obj(item), ensure_ascii=False),
                    "",
                    "",
                ])
            elif reason == "high_risk_field_requires_manual_review":
                workbook["高风险字段待确认"].append([
                    source_name,
                    item.get("i_id", ""),
                    item.get("field", ""),
                    item.get("incoming_value", ""),
                    reason,
                    "",
                    "",
                ])
            else:
                workbook["字段冲突"].append([
                    source_name,
                    item.get("i_id", ""),
                    item.get("field", ""),
                    item.get("existing_value", ""),
                    item.get("incoming_value", ""),
                    reason,
                    "",
                    "",
                ])
        for item in report.get("skipped_reasons") or []:
            reason = sanitize_text(item.get("reason"))
            if reason in {"missing_platform_identity", "kb_product_not_found", "ambiguous_internal_identity"}:
                workbook["标题弱匹配"].append([source_name, item.get("row", ""), reason, "", "", "", ""])
            elif reason in {"untrusted_source", "unclassified", "missing_asset_url", "kb_product_not_found"}:
                workbook["素材不确定"].append([
                    source_name,
                    item.get("row", ""),
                    reason,
                    _safe_url(sanitize_text(item.get("asset_url"))),
                    sanitize_text(item.get("asset_title")),
                    "",
                    "",
                ])


def run_export(
    *,
    output: str = "",
    identity_report: str = "",
    field_report: str = "",
    media_report: str = "",
) -> dict[str, Any]:
    workbook = Workbook()
    first = workbook.active
    first.title = "身份冲突"
    for name, headers in SHEETS.items():
        ws = first if name == "身份冲突" else workbook.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for index, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = max(14, min(40, len(header) + 6))
    reports = [
        ("identity", _load_json(identity_report)),
        ("fields", _load_json(field_report)),
        ("media", _load_json(media_report)),
    ]
    _append_rows_from_reports(workbook, reports)
    output_path = Path(output or default_output_path())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return sanitize_obj({
        "ok": True,
        "output": str(output_path),
        "sheet_counts": {name: workbook[name].max_row - 1 for name in SHEETS},
    })


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Export unresolved auto-backfill items for manual review.")
    parser.add_argument("--output", default="")
    parser.add_argument("--identity-report", default="")
    parser.add_argument("--field-report", default="")
    parser.add_argument("--media-report", default="")
    args = parser.parse_args()
    result = run_export(
        output=args.output,
        identity_report=args.identity_report,
        field_report=args.field_report,
        media_report=args.media_report,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
