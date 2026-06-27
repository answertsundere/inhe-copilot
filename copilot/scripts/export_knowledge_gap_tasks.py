"""Export knowledge gap tasks to an Excel review workbook."""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.db import SessionLocal, init_db
from app.models.eval_tables import KnowledgeGapTask
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text
from app.services.knowledge_gap_task_service import filter_tasks_by_run


HEADERS = [
    "任务ID",
    "缺口类型",
    "需要证据类型",
    "补充目标系统",
    "推荐动作",
    "缺失字段",
    "商品标题预览",
    "商品ID(脱敏)",
    "SKU",
    "问题类型",
    "失败类型",
    "建议归口",
    "负责人",
    "样本数",
    "买家问题示例",
    "Agent当前回复示例",
    "原客服回复参考",
    "风险等级",
    "状态",
    "优先级",
    "当前上下文摘要",
    "备注",
]


def default_output_path() -> str:
    today = datetime.now().strftime("%Y%m%d")
    return str(Path("D:/") / f"知识缺口任务_{today}.xlsx")


def _joined(values: list[str]) -> str:
    return "\n".join(sanitize_text(value) for value in values[:5] if sanitize_text(value))


def _metadata(task: KnowledgeGapTask) -> dict[str, Any]:
    value = task.get_metadata()
    return value if isinstance(value, dict) else {}


def _task_dict(task: KnowledgeGapTask) -> dict[str, Any]:
    return sanitize_obj(task.to_dict())


def _context_summary_text(task: KnowledgeGapTask) -> str:
    metadata = _metadata(task)
    summary = metadata.get("current_context_summary") or {}
    if not isinstance(summary, dict):
        return ""
    parts = [
        f"product={bool(summary.get('has_product_context'))}",
        f"order={bool(summary.get('has_order_context'))}",
        f"media={bool(summary.get('has_media_context'))}",
        f"selected={int(summary.get('selected_evidence_count') or 0)}",
        f"sendable_media={int(summary.get('sendable_media_asset_count') or 0)}",
    ]
    if summary.get("conversation_media_rejected_reason"):
        parts.append(f"media_rejected={sanitize_text(summary.get('conversation_media_rejected_reason'))}")
    return "; ".join(parts)


def export_knowledge_gap_tasks(output: str | None = None, status: str = "open", run_uid: str = "") -> dict:
    init_db()
    path = Path(output or default_output_path())
    path.parent.mkdir(parents=True, exist_ok=True)

    db = SessionLocal()
    try:
        query = db.query(KnowledgeGapTask).order_by(KnowledgeGapTask.updated_at.desc(), KnowledgeGapTask.id.desc())
        if status:
            query = query.filter(KnowledgeGapTask.status == sanitize_text(status))
        tasks = query.all()
        target_run_uid = sanitize_text(run_uid)
        if target_run_uid:
            tasks = filter_tasks_by_run(db, tasks, target_run_uid)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "知识缺口任务"
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")

        for task in tasks:
            row = _task_dict(task)
            sheet.append([
                row.get("task_uid", ""),
                row.get("gap_category") or row.get("gap_type", ""),
                row.get("required_evidence_type") or row.get("missing_evidence_type", ""),
                row.get("target_system", ""),
                row.get("recommended_action", ""),
                ", ".join(row.get("missing_fields") or []),
                row.get("product_title_preview") or sanitize_text(task.product_title)[:80],
                row.get("item_id_masked", ""),
                row.get("sku_code", ""),
                row.get("query_fact_type", ""),
                row.get("failure_type", ""),
                row.get("suggested_fix_area", ""),
                row.get("suggested_owner", ""),
                int(row.get("sample_count") or 0),
                _joined(task.get_latest_buyer_questions()),
                _joined(task.get_latest_agent_replies()),
                _joined(task.get_latest_original_cs_replies()),
                row.get("risk_level", ""),
                row.get("status", ""),
                row.get("priority", ""),
                _context_summary_text(task),
                row.get("summary", ""),
            ])

        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(max(len(str(col[0].value or "")) + 4, 12), 42)
        workbook.save(path)
        return {"output": str(path), "count": len(tasks)}
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Export knowledge gap tasks to Excel.")
    parser.add_argument("--output", default="", help="Excel output path")
    parser.add_argument("--status", default="open", help="Task status filter; empty exports all tasks")
    parser.add_argument("--run-uid", default="", help="Only export tasks related to this eval run")
    args = parser.parse_args()
    result = export_knowledge_gap_tasks(output=args.output or None, status=args.status, run_uid=args.run_uid)
    print(result)


if __name__ == "__main__":
    main()
