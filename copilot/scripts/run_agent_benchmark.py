"""Run active Agent benchmark scenarios."""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def _configure_fixture_database_from_argv(argv: list[str]) -> None:
    """Set the isolated fixture DB before config/app modules are imported."""
    for index, value in enumerate(argv):
        if value == "--fixture" or value.startswith("--fixture="):
            os.environ["COPILOT_BENCHMARK_FIXTURE_MODE"] = "true"
        if value == "--benchmark-db" and index + 1 < len(argv):
            os.environ["COPILOT_KNOWLEDGE_DB_PATH"] = argv[index + 1]
            return
        if value.startswith("--benchmark-db="):
            os.environ["COPILOT_KNOWLEDGE_DB_PATH"] = value.split("=", 1)[1]
            return


def _load_dotenv_safely() -> None:
    try:
        from dotenv import load_dotenv

        load_dotenv(PROJECT_ROOT / ".env")
    except Exception:
        pass


_configure_fixture_database_from_argv(sys.argv[1:])
_load_dotenv_safely()

from app.db import init_db
from app.models.eval_tables import AgentBenchmarkScenario  # noqa: F401 - register table before init_db
from app.services.agent_benchmark_runner_service import AgentBenchmarkRunnerService
from app.services.eval_sanitizer_service import sanitize_obj, sanitize_text


SUMMARY_HEADERS = ("指标", "数值")
DETAIL_HEADERS = [
    "场景 UID",
    "标题",
    "场景类型",
    "问题类型",
    "是否通过",
    "回复状态",
    "可发送",
    "需要人工复核",
    "失败原因",
    "缺少关键点",
    "命中禁用话术",
    "客户完整对话",
    "侧栏上下文",
    "期望回复",
    "Agent 回复",
    "耗时 ms",
]


def _write_rows(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def _group_rows(group: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for key, value in sorted((group or {}).items()):
        if isinstance(value, dict):
            rows.append([
                key,
                value.get("total", 0),
                value.get("passed", 0),
                value.get("failed", 0),
                value.get("pass_rate", 0),
            ])
        else:
            rows.append([key, value])
    return rows


def _write_excel(path: str, payload: dict[str, Any]) -> None:
    from openpyxl import Workbook

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    _write_rows(summary, list(SUMMARY_HEADERS), [
        ["评测批次", payload.get("benchmark_run_uid") or payload.get("run_uid") or ""],
        ["评测状态", payload.get("status", "")],
        ["总样本数", payload.get("total_scenarios", 0)],
        ["通过数", payload.get("passed_count", 0)],
        ["失败数", payload.get("failed_count", 0)],
        ["通过率", payload.get("pass_rate", 0)],
    ])

    by_scenario = workbook.create_sheet("按场景")
    _write_rows(by_scenario, ["场景类型", "总数", "通过数", "失败数", "通过率"], _group_rows(payload.get("by_scenario_type", {})))

    by_fact = workbook.create_sheet("按问题类型")
    _write_rows(by_fact, ["问题类型", "总数", "通过数", "失败数", "通过率"], _group_rows(payload.get("by_query_fact_type", {})))

    failures = workbook.create_sheet("失败原因")
    _write_rows(failures, ["失败原因", "数量"], _group_rows(payload.get("by_failure_reason", {})))

    detail = workbook.create_sheet("逐条结果")
    detail_rows = []
    for item in payload.get("per_scenario_result", []) or []:
        turns = item.get("conversation_turns") or []
        expected = item.get("expected_reply") or {}
        detail_rows.append([
            item.get("scenario_uid", ""),
            item.get("title", ""),
            item.get("scenario_type", ""),
            item.get("query_fact_type", ""),
            "通过" if item.get("passed") else "失败",
            item.get("reply_status", ""),
            bool(item.get("can_send")),
            bool(item.get("requires_human_review")),
            "\n".join(item.get("failure_reasons") or []),
            "\n".join(item.get("missing_key_points") or []),
            "\n".join(item.get("forbidden_claims_hit") or []),
            json.dumps(sanitize_obj(turns), ensure_ascii=False),
            json.dumps(sanitize_obj(item.get("sidecar_context") or {}), ensure_ascii=False),
            sanitize_text(expected.get("expected_reply")),
            sanitize_text(item.get("agent_reply")),
            item.get("latency_ms", 0),
        ])
    _write_rows(detail, DETAIL_HEADERS, detail_rows)
    workbook.save(target)


def _write_json(path: str, payload: dict) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(sanitize_obj(payload), ensure_ascii=True, indent=2), encoding="utf-8")


def run_benchmark_report(
    status: str = "active",
    scenario_type: str = "",
    scenario_uids: list[str] | None = None,
    limit: int | None = None,
    run_uid: str = "",
    json_output: str = "",
    excel_output: str = "",
    fail_on_failure: bool = False,
    include_full_trace: bool = False,
    db_factory=None,
    agent_callable=None,
    dataset_metadata: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], int]:
    result = AgentBenchmarkRunnerService(agent_callable=agent_callable).run_scenarios(
        status=status,
        scenario_type=scenario_type or None,
        scenario_uids=scenario_uids or None,
        limit=limit or None,
        run_uid=run_uid,
        include_full_trace=include_full_trace,
        db_factory=db_factory,
        dataset_metadata=dataset_metadata,
    )
    if json_output:
        _write_json(json_output, result)
    if excel_output:
        _write_excel(excel_output, result)
    if result.get("invalid_run"):
        exit_code = 2
    else:
        exit_code = 1 if fail_on_failure and result.get("failed_count", result.get("failed", 0)) else 0
    return result, exit_code


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Run Agent benchmark scenarios.")
    parser.add_argument("--status", default="active")
    parser.add_argument("--scenario-type", default="")
    parser.add_argument("--scenario-uid", action="append", default=[])
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--run-uid", default="")
    parser.add_argument("--json-output", default="")
    parser.add_argument("--excel-output", default="")
    parser.add_argument("--fail-on-failure", action="store_true")
    parser.add_argument("--include-full-trace", action="store_true")
    parser.add_argument("--fixture", default="")
    parser.add_argument("--fixture-manifest", default="")
    parser.add_argument("--benchmark-db", default="")
    args = parser.parse_args(argv)

    db_factory = None
    dataset_metadata = None
    if args.fixture:
        if not args.benchmark_db:
            parser.error("--fixture requires --benchmark-db")
        from app.services.agent_benchmark_fixture_service import (
            BenchmarkFixtureError,
            fixture_database_metadata,
            fixture_session_factory,
            load_fixture,
        )

        try:
            _payload, manifest = load_fixture(args.fixture, args.fixture_manifest or None)
            metadata = fixture_database_metadata(args.benchmark_db)
            for key in ("dataset_id", "dataset_version", "schema_version", "fixture_sha256"):
                if metadata.get(key) != str(manifest.get(key, "")):
                    raise BenchmarkFixtureError(f"fixture database mismatch: {key}")
            db_factory = fixture_session_factory(args.benchmark_db)
            dataset_metadata = {**metadata, "database_path": str(Path(args.benchmark_db).resolve())}
        except BenchmarkFixtureError as exc:
            parser.error(str(exc))
    elif args.benchmark_db:
        parser.error("--benchmark-db is only supported with --fixture")
    else:
        init_db()
    result, exit_code = run_benchmark_report(
        status=args.status,
        scenario_type=args.scenario_type,
        scenario_uids=args.scenario_uid or None,
        limit=args.limit or None,
        run_uid=args.run_uid,
        json_output=args.json_output,
        excel_output=args.excel_output,
        fail_on_failure=bool(args.fail_on_failure),
        include_full_trace=bool(args.include_full_trace),
        db_factory=db_factory,
        dataset_metadata=dataset_metadata,
    )
    print(json.dumps(sanitize_obj(result), ensure_ascii=False, indent=2))
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
